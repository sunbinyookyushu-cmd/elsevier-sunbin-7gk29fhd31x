# -*- coding: utf-8 -*-
"""Build HSR_OilHedge_results_20260929.xlsx (data build, validation, level DiD, event study, oil-hedge tables).
README / decision text in _xlsx_text_hsr.py."""
import sys
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import LineChart, Reference
sys.path.insert(0, r"..\GACI_FuelShock")
from _xlsx_helpers import st, W, rule, fc, fs, fn, ff, notes, title, widths, style_chart
from _xlsx_text_hsr import README, DECISIONS, SOURCES

OUTF = "HSR_OilHedge_results_20260929.xlsx"
lev = pd.read_csv("_res_hsr_level.csv")
es = pd.read_csv("_res_hsr_es.csv")
hed = pd.read_csv("_res_hsr_hedge.csv")
cov = pd.read_csv("data/station_date_coverage.csv")
val = pd.read_csv("data/validation_km_by_year.csv")
agr = pd.read_csv("data/validation_date_agreement.csv")
unm = pd.read_csv("data/validation_unmatched.csv")
aph = pd.read_csv("data/airport_hsr.csv", parse_dates=["hsr_date_30", "hsr_date_50", "hsr_date_100"])
wb = Workbook()
wb.remove(wb.active)


def text_sheet(name, blocks):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 120
    for r, (kind, txt) in enumerate(blocks, start=1):
        c = ws.cell(row=r, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="맑은 고딕", size=14 if kind == "h1" else 11 if kind == "h2" else 10, bold=kind in ("h1", "h2"))
        if kind == "p":
            ws.row_dimensions[r].height = max(15, 15 * (len(txt) // 95 + txt.count("\n") + 1))


def pick(df, **kw):
    m = df
    for k, v in kw.items():
        m = m[m[k] == v]
    return m.iloc[0] if len(m) else None


text_sheet("README", README)
text_sheet("결정필요_표본정의", DECISIONS)
text_sheet("자료출처", SOURCES)

# ---- H1 validation ----
ws = wb.create_sheet("H1_자료검증")
ncol = 9
r = title(ws, 1, ncol, "Table H1. Global HSR opening panel: coverage and validation")
rule(ws, r, 1, ncol, "top")
hdr = ["Country", "Stations on HSR track", "Dated", "Share dated", "OSM/UIC route km 2005", "2010", "2015", "2019",
       "UIC = OSM year (same ways)"]
for j, h in enumerate(hdr):
    W(ws, r, j + 1, h, wrap=True)
ws.row_dimensions[r].height = 30
rule(ws, r, 1, ncol, "bottom")
r += 1
for x in cov.sort_values("stations_on_hsr", ascending=False).itertuples():
    W(ws, r, 1, x.cc)
    W(ws, r, 2, fn(x.stations_on_hsr))
    W(ws, r, 3, fn(x.dated) if pd.notna(x.dated) else "0")
    W(ws, r, 4, f"{x.share_dated:.2f}" if pd.notna(x.share_dated) else "0.00")
    for j, y in enumerate([2005, 2010, 2015, 2019]):
        v = val[(val.cc == x.cc) & (val.year == y)].ratio_osm_to_uic
        W(ws, r, 5 + j, f"{v.iloc[0]:.2f}" if len(v) and pd.notna(v.iloc[0]) else "")
    a = agr[agr.cc == x.cc]
    W(ws, r, 9, f"{a.same_year_share.iloc[0]:.2f}" if len(a) else "")
    r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 12, 13, ncol)
notes(ws, r, ncol, "Stations: OSM railway=station/halt within 300 m of a highspeed=yes way. Dated: an opening date from UIC, MLIT (Japan) or "
      "OSM start_date. Route km ratio: dated OSM way length / 2 (double track) against the UIC in-operation total up to that year; 1 = full match, "
      "< 1 = undated sections remain. Last column: share of ways carrying both a UIC-path date and an OSM start_date where the years agree "
      "(China: OSM is used first, see decisions).")

# ---- H2 treated airports ----
ws = wb.create_sheet("H2_처치공항")
tr = aph[aph.hsr_date_50.notna()].sort_values("hsr_date_50")
cols = ["airport_iata", "iso3", "hsr_date_50", "hsr_src_50", "n_hsr_50", "km_nearest_hsr", "hsr_date_30", "hsr_date_100", "GACI_96", "seats_96"]
ws.append(["Airport", "Country", "First HSR station within 50 km", "Date source", "Stations within 50 km", "km to nearest HSR station",
           "First within 30 km", "First within 100 km", "GACI 1996", "Seats 1996"])
for x in tr[cols].itertuples(index=False):
    ws.append([v.strftime("%Y-%m-%d") if isinstance(v, pd.Timestamp) and pd.notna(v) else (None if (isinstance(v, float) and np.isnan(v)) else v)
               for v in x])
for c, w_ in zip("ABCDEFGHIJ", [9, 8, 14, 26, 10, 12, 14, 14, 10, 14]):
    ws.column_dimensions[c].width = w_

# ---- H3 level DiD ----
ws = wb.create_sheet("H3_수준효과")
ncol = 8
r = title(ws, 1, ncol, "Table H3. Seats after an HSR station opens within R km (Gardner two-stage DiD, airport x month, 1997-2019)")
rule(ws, r, 1, ncol, "top")
for j, h in enumerate(["Outcome", "Sample", "Radius", "Post (coef.)", "SE (airport cl.)", "SE (country cl.)", "Observations", "Treated airports"]):
    W(ws, r, j + 1, h, wrap=True)
rule(ws, r, 1, ncol, "bottom")
r += 1
OL = {"ln_seats_dom": "ln domestic seats", "ln_seats": "ln seats", "ln_seats_intl": "ln intl. seats", "ln_co2": "ln CO2"}
for x in lev.dropna(subset=["b"]).itertuples():
    W(ws, r, 1, OL.get(x.outcome, x.outcome), align="left")
    W(ws, r, 2, x.sample, align="left")
    W(ws, r, 3, f"{int(x.radius)} km")
    W(ws, r, 4, fc(x.b, x.p_country))
    W(ws, r, 5, fs(x.se_airport))
    W(ws, r, 6, fs(x.se_country, br="[]"))
    W(ws, r, 7, fn(x.n))
    W(ws, r, 8, fn(x.n_treated))
    r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 20, 14, ncol)
ws.column_dimensions["B"].width = 34
notes(ws, r, ncol, "Stage 1: airport and country x year-month FE estimated on untreated observations (never-treated airports; treated airports before "
      "opening). Stage 2: residualised outcome on post. Stars use the country-cluster SE (more conservative). SEs ignore first-stage estimation "
      "error. Airports with a station before 1997 are dropped.")

# ---- H4 event study ----
ws = wb.create_sheet("H4_사건연구")
W(ws, 1, 1, "Event study: years relative to the first HSR station within 50 km (reference = year -1)", bold=True, align="left")
for j, h in enumerate(["Years relative", "Domestic seats", "Lower 95%", "Upper 95%", "All seats", "Lower 95% ", "Upper 95% ",
                       "CO2", "Lower 95%  ", "Upper 95%  "]):
    W(ws, 2, 1 + j, h)
e1 = es[es.outcome == "ln_seats_dom"].set_index("rel_year")
e2 = es[es.outcome == "ln_seats"].set_index("rel_year")
e3 = es[es.outcome == "ln_co2"].set_index("rel_year")
for i, k in enumerate(range(-6, 7)):
    vals = [k]
    for e_ in (e1, e2, e3):
        a = e_.loc[k]
        vals += [a.b, a.b - 1.96 * a.se_country, a.b + 1.96 * a.se_country]
    for j, v in enumerate(vals):
        ws.cell(row=3 + i, column=1 + j, value=float(v))
ch2 = LineChart()
ch2.title = "ln CO2 around the first HSR station within 50 km"
ch2.add_data(Reference(ws, min_col=8, max_col=10, min_row=2, max_row=15), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=3, max_row=15))
style_chart(ch2, ci=True)
ch2.height, ch2.width = 8, 16
ws.add_chart(ch2, "L20")
ch = LineChart()
ch.title = "ln domestic seats around the first HSR station within 50 km"
ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=2, max_row=15), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=15))
style_chart(ch, ci=True)
ch.height, ch.width = 8, 16
ws.add_chart(ch, "L2")

# ---- H5 hedge ----
ws = wb.create_sheet("H5_석유헤지")
EST = [("OLS", "OLS"), ("2SLS-KZ", "2SLS (Känzig)"), ("2SLS-BH", "2SLS (BH)")]
ncol = 1 + 3 * 3
r = title(ws, 1, ncol, "Table H5. Does HSR make air seats more sensitive to jet fuel prices? D12 ln P(t-3) x post-HSR (airport x month, 1997-2019)")
rule(ws, r, 1, ncol, "top")
OUTS = [("d12_ln_seats_dom", "Δ12 ln domestic seats"), ("d12_ln_seats", "Δ12 ln seats"), ("d12_ln_co2", "Δ12 ln CO2")]
for j, (o, lab) in enumerate(OUTS):
    c0 = 2 + 3 * j
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
    W(ws, r, c0, lab)
    rule(ws, r, c0, c0 + 2, "bottom")
r += 1
for j in range(len(OUTS)):
    for k, (_, el) in enumerate(EST):
        W(ws, r, 2 + 3 * j + k, el)
rule(ws, r, 1, ncol, "bottom")
r += 1
for smp in ["all", "ever-treated only", "China", "excl. China", "high-quality dates (CN JP KR TW FR)"]:
    W(ws, r, 1, f"Sample: {smp} (R = 50 km)", italic=True, align="left")
    r += 1
    for lab, key, se_key in [("Δ ln P × post-HSR", "b", "se"), ("Δ ln P × ever-HSR", "b_ever", "se_ever"), ("post-HSR", "b_post", "se_post")]:
        W(ws, r, 1, lab, align="left")
        for j, (o, _) in enumerate(OUTS):
            for k, (e, _) in enumerate(EST):
                x = pick(hed, radius=50, outcome=o, sample=smp, estimator=e)
                if x is None or pd.isna(getattr(x, key)):
                    continue
                pv = x.p if key == "b" else np.nan
                W(ws, r, 2 + 3 * j + k, fc(getattr(x, key), pv))
                W(ws, r + 1, 2 + 3 * j + k, fs(getattr(x, se_key)))
        r += 2
    W(ws, r, 1, "Observations / treated airports / KP F", align="left")
    for j, (o, _) in enumerate(OUTS):
        for k, (e, _) in enumerate(EST):
            x = pick(hed, radius=50, outcome=o, sample=smp, estimator=e)
            if x is not None:
                W(ws, r, 2 + 3 * j + k, f"{fn(x.n)} / {x.n_treated}" + ("" if e == "OLS" else f" / {ff(x.F)}"), size=9)
    r += 1
W(ws, r, 1, "Radius robustness, Δ12 ln domestic seats, all", italic=True, align="left")
r += 1
for R in (30, 100):
    W(ws, r, 1, f"R = {R} km: Δ ln P × post-HSR", align="left")
    for k, (e, _) in enumerate(EST):
        x = pick(hed, radius=R, outcome="d12_ln_seats_dom", sample="all", estimator=e)
        if x is not None:
            W(ws, r, 2 + k, fc(x.b, x.p))
            W(ws, r + 1, 2 + k, fs(x.se))
    r += 2
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 40, 14, ncol)
notes(ws, r, ncol, "Airport and country x year-month FE. post-HSR = 1 from the month the first HSR station within R km opens; ever-HSR = 1 for "
      "airports that get a station at any time (incl. after 2019). A negative Δ ln P × post-HSR coefficient means the airport cuts more seats "
      "when jet fuel gets dearer once HSR is available. Stars (* p<0.1, ** p<0.05, *** p<0.01) use country cluster + Newey-West over months "
      "(L = 12). 2SLS: Känzig news shock or BH supply shock (x -1) 12-month sums, lag 3, x post and x ever.")

for nm, df in [("raw_level", lev), ("raw_eventstudy", es), ("raw_hedge", hed), ("raw_unmatched_uic", unm)]:
    ws = wb.create_sheet(nm)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else (v.strftime("%Y-%m-%d") if isinstance(v, pd.Timestamp) else v)
                   for v in row])
wb.save(OUTF)
print("saved", OUTF, [s.title for s in wb.worksheets])
