# -*- coding: utf-8 -*-
"""Build the interaction-IV heterogeneity table from gaci_tourism_hetero_iv.log
   -> GACI_hetero_interaction.docx  (two panels: remoteness, baseline income)."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.section import WD_ORIENT

def star(p):
    try: p = float(p)
    except Exception: return ''
    return '***' if p < .01 else '**' if p < .05 else '*' if p < .1 else ''

OUT = ['g_int', 'g_vol', 'g_shr', 'lnpc', 'lngdp']
OLAB = {'g_int': 'Trade openness\nln(goods/GDP)', 'g_vol': 'Ln goods\nvolume', 'g_shr': 'Goods share\n(% GDP)',
        'lnpc': 'Ln GDP\nper capita', 'lngdp': 'Ln GDP\n(total)'}
DEC = {'g_int': 3, 'g_vol': 3, 'g_shr': 2, 'lnpc': 3, 'lngdp': 3}
MODLAB = {'remote': 'Panel A. Moderator = remoteness (ln mean sea distance, CERDI; mean-centered)',
          'income': 'Panel B. Moderator = baseline income (1996 ln GDP per capita; mean-centered)'}

d, F, N = {}, {}, {}
for ln in open('gaci_tourism_hetero_iv.log', encoding='utf-8', errors='replace'):
    ln = ln.strip()
    if ln.startswith('HET3|'):
        p = ln.split('|')
        if p[3] == 'FAIL': continue
        _, mod, yv, term, b, se, pv, f, n = p
        d[(mod, yv, term)] = (float(b), float(se), float(pv))
        F[(mod, yv)] = float(f); N[(mod, yv)] = int(float(n))

doc = Document()
s = doc.sections[0]; s.orientation = WD_ORIENT.PORTRAIT
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(s, m, Inches(0.7))
doc.styles['Normal'].font.name = 'Times New Roman'; doc.styles['Normal'].font.size = Pt(9)
doc.add_heading('Heterogeneity by interaction IV: connectivity x moderator (1996-2023)', 1)
p = doc.add_paragraph()
r = p.add_run('Full-sample 2SLS; ln_gaci_cwm and ln_gaci_cwm x moderator both instrumented by tourism_int and '
              'tourism_int x moderator; country + year FE; robust SE. Moderators mean-centered, so "Connectivity '
              '(at mean)" is the effect for the average country and "x moderator" is the slope of that effect in the '
              'moderator. KP = Kleibergen-Paap first-stage F (joint). * p<.10 ** p<.05 *** p<.01.')
r.italic = True; r.font.size = Pt(8)

def panel(mod):
    doc.add_heading(MODLAB[mod], 3)
    t = doc.add_table(rows=0, cols=6); t.style = 'Table Grid'
    h = t.add_row().cells; h[0].text = ''
    for j, o in enumerate(OUT): h[j + 1].text = OLAB[o]
    for c in h:
        for pa in c.paragraphs:
            for rr in pa.runs: rr.bold = True
    rows = [('Connectivity (at mean)', 'main', 'b'), ('   SE', 'main', 'se'),
            ('x moderator (interaction)', 'inter', 'b'), ('   SE', 'inter', 'se'),
            ('KP first-stage F', None, 'F'), ('Observations', None, 'N')]
    for lab, term, kind in rows:
        r = t.add_row().cells; r[0].text = lab
        for j, o in enumerate(OUT):
            dec = DEC[o]
            if kind == 'F':
                if (mod, o) in F: r[j + 1].text = f"{F[(mod,o)]:.1f}" + ("w" if F[(mod,o)] < 10 else "")
            elif kind == 'N':
                if (mod, o) in N: r[j + 1].text = f"{N[(mod,o)]:,}"
            else:
                v = d.get((mod, o, term))
                if not v: continue
                b, se, pv = v
                r[j + 1].text = f"{b:,.{dec}f}{star(pv)}" if kind == 'b' else f"({se:,.{dec}f})"

for mod in ('remote', 'income'):
    panel(mod)
p = doc.add_paragraph()
r = p.add_run('"w" marks KP F < 10. Positive "x remoteness" => connectivity raises trade MORE for remote countries; '
              'positive "x income" => effect larger in richer countries (negative => returns concentrated in poorer).')
r.italic = True; r.font.size = Pt(8)
for tb in doc.tables:
    for row in tb.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs: run.font.size = Pt(7.5)
doc.save('GACI_hetero_interaction.docx')
print('wrote GACI_hetero_interaction.docx')
