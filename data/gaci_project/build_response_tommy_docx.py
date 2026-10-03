# -*- coding: utf-8 -*-
"""build_response_tommy_docx.py -- point-by-point response email to Tommy Cheung (2026-09-01)."""
from docx import Document
from docx.shared import Pt

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(11)

def para(text, bold=False, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold
    p.paragraph_format.space_after = Pt(space_after)
    return p

def item(head, body):
    p = doc.add_paragraph()
    r = p.add_run(head); r.bold = True
    p.paragraph_format.space_after = Pt(2)
    q = doc.add_paragraph(body)
    q.paragraph_format.space_after = Pt(8)

para('Subject: Re: GACI and trade - point-by-point responses to your comments', bold=True, space_after=10)
para('Dear Tommy,')
para("Thank you very much for the careful reading and the detailed comments. They were extremely helpful, particularly on tightening the causal and mechanism claims. We have revised the manuscript accordingly. A small note first: your comments were made on a slightly earlier draft, so a few points had already been revised in the current version; we flag those below. Here are our point-by-point responses, in the order the comments appear.", space_after=10)

items = [
("1. Literature review: more comprehensive coverage of aviation-trade studies, with critical comparison of data, measures, and outcomes.",
"We substantially expanded Section 2.1. We added nine references, including Harrigan (2010, JIE), Hummels and Schaur (2010, JIE), Soderlund (2023, JIE), Micco and Serebrisky (2006, JIE), Umana-Dajud (2019, JDE), Yilmazkuday and Yilmazkuday (2017, RWE), Alderighi and Gaggero (2017, Economics of Transportation), Blonigen and Cristea (2015, JUE), and Fageda (2017, JRS). The aviation-trade strand is now discussed comparatively rather than study by study: for each paper we state the data, the connectivity measure (realised travel volumes, flight-time changes, direct-flight presence, or nonstop frequencies), the outcome, and the identification, and we explain why these bilateral or subnational estimates neither map into one another nor aggregate to the effect of a country's overall network position."),
("2. \"Trade is the flow through which better access becomes higher income\" (not entirely true).",
"Agreed; this sentence had already been rewritten in the current draft. The introduction now presents trade as the flow that responds most robustly and is observed most cleanly among the several flows that distance impedes (people, capital, ideas, goods), and we added a citation linking air services to FDI (Fageda 2017) to acknowledge the other channels."),
("3. \"No global causal estimate\" is too strong.",
"Already softened in the current draft: the claim is now that existing causal evidence is largely confined to bilateral or subnational settings, while global evidence remains primarily descriptive."),
("4. The conceptual framework should connect to the research questions.",
"Section 2.2 now opens by mapping the three channels onto the research questions: Hypothesis 1 addresses the aggregate margin of the first question, Hypotheses 2 and 3 the compositional margins, and cross-country variation in channel strength motivates the heterogeneity analysis of the second question."),
("5. Explain \"excluded by construction.\"",
"The Data section now spells this out: the outcome covers trade in physical goods only, and tourist spending is recorded in the balance of payments as a services export (travel receipts), so it never enters the outcome by definition."),
("6. Why all three outcomes are estimated.",
"We now state the distinct question each answers: the volume elasticity is the total effect, the openness elasticity asks whether a country trades more relative to the size of its economy (integration proper), and the GDP elasticity asks whether the effect operates through scale."),
("7. Mechanism data (BACI): pooled median, capital goods, the 0.80 correlation.",
"All three clarified in an expanded data paragraph. (i) The median unit value is computed once from all product observations pooled over the full sample period, so each product's high/low classification is fixed over time. (ii) Capital goods, which the BEC assigns to an end-use class of their own, are grouped separately and do not enter the intermediate-to-consumption ratio. (iii) We now explain that the 0.80 within-country correlation falls short of unity because BACI reconciles customs records while the WDI series follows balance-of-payments conventions."),
("8. The instrument discussion is too brief; explain the mechanism and the choice of natural heritage.",
"We expanded this considerably. The mechanism is now explicit: airlines deploy marginal aircraft and frequencies where expected passenger demand is highest, and when the global tourism cycle rises, expected leisure demand rises most in destinations with strong natural attractions, so heritage-rich countries receive a disproportionate share of tourism-driven route growth. We also justify natural heritage as the exposure measure: it is rooted in physical geography rather than economic activity, unlike cultural sites, whose listing tracks historical development and urbanisation, and its level is absorbed by the country fixed effects."),
("9. Rationale for the controls; why population is the only time-varying control.",
"Now explained after the main equation: country fixed effects absorb all time-invariant determinants, and the natural time-varying candidates (GDP, income) are themselves outcomes of connectivity, so conditioning on them would control away part of the effect. Population is a standard scale control plausibly unaffected at annual frequency; the control-sensitivity table shows the estimate is insensitive to richer control sets."),
("10. The need for IV comes from endogeneity, not from the continuous nature of GACI.",
"Restructured as suggested: the continuity argument now motivates the measurement choice only, and endogeneity is stated as what necessitates instrumental variables."),
("11. Excluding tourism and services does not establish that tourism cannot indirectly affect merchandise trade (two comments).",
"Agreed. Both the empirical strategy section and the appendix now state that the exclusion removes the mechanical channel from tourist spending but not every indirect effect, which could still operate through income, prices, or the real exchange rate; these are addressed by the control sensitivity, the over-identification and falsification tests, and the Conley bounds."),
("12. Can over-identification establish the exclusion restriction?",
"No, and we now say so where the tests are introduced: over-identification and falsification tests can detect violations but cannot establish the restriction, which remains an identifying assumption."),
("13. Specify the instrument for each endogenous regressor; are the moderators time-varying?",
"Both already addressed in the current draft: the connectivity term and the interaction are instrumented with the corresponding pair (the baseline instrument and its interaction with the moderator), and the three moderators are stated to be country-specific and time-invariant, with the interaction equation written accordingly."),
("14. Significant kappa-1 and beta-2 do not establish a causal mechanism.",
"We agree. The strategy section now states that the exercise is a descriptive decomposition rather than formal causal mediation, that the mediator is itself an outcome of connectivity, and, following your comment, that a significant beta-2 establishes an association between composition and openness rather than a causal effect, since unobserved factors may move both. The mechanism results section and the discussion carry the same reading (\"supporting the proposed mechanisms rather than proving them\")."),
("15. Remoteness should be defined where it is first introduced.",
"Done: the full definition (unweighted mean sea distance to all other sample countries, CERDI database, in logs, time-invariant) now appears in the empirical strategy section at first use."),
("16. Low-cost long-haul and Gulf/Asian hub expansion are possible explanations, not established causes (two comments).",
"Reworded in both places: the regressions establish the temporal difference, and we now describe the industry developments as plausible drivers that coincide with the timing rather than tested causes."),
("17. Composition results support rather than prove the mechanisms.",
"The concluding paragraph of the mechanism section now says exactly this, noting that the composition ratios are equilibrium outcomes that other determinants of openness may also move."),
("18. Is it appropriate to apply the elasticity to pre-2010 connectivity growth?",
"We now flag this directly in the aggregate-contribution section (not only in the discussion): the instrument derives nearly all of its first-stage relevance from the post-2010 period, so the headline figure extrapolates a relationship identified from recent variation to the earlier part of the sample, and we note that a post-2010 counterfactual would be an identification-aligned benchmark."),
("19. Report a sensitivity range for the $7.8 trillion figure.",
"Done, and computed rather than described: applying the endpoints of the openness elasticity's 95 per cent confidence interval to the same country-level calculation yields $0.03-12.4 trillion, and the volume-channel counterpart is $5.1-14.9 trillion. These now appear alongside the cross-measure range ($6.8-7.8 trillion) and the country-level delta-method standard errors in the appendix."),
("20. Is the reverse logic true? (connectivity collapse)",
"The section now opens by stating that symmetry between expansion and contraction is an assumption of the exercise rather than an established result, consistent with the discussion section."),
("21. Why does the implied disruption fall from $7.5 to $1.8 trillion across measures?",
"Now explained: the spread compounds two differences, the smaller openness elasticity of the sum measure (0.659 versus 1.302) and its smaller measured collapse, since summed connectivity is dominated by the count of airports with any scheduled service, which contracted far less in 2020 than capacity-weighted network position."),
("22. Range of estimates, not hard bounds; $7.5 trillion is not a causal loss; \"unambiguous\" is too strong.",
"All three adopted: the figures are now presented as a range of estimates rather than upper and lower bounds; we state explicitly that the implied disruption should not be interpreted as the causal trade loss from reduced connectivity alone; and \"unambiguous\" was replaced with \"consistent across all three measures.\""),
("23. \"The two instruments agree\" and \"the first stage remains adequate (KP F = 8.5).\"",
"Reworded as you suggested: the instruments are now described as \"not statistically inconsistent,\" \"adequate\" was removed, and the weaker first stage (KP F = 8.5) is stated as suggesting weaker first-stage identification, with the over-identified estimate read as corroborating the baseline rather than replacing it."),
("24. Explain the construction of the air/sea-distance instrument and why it is an independent identification path.",
"The appendix paragraph was rewritten with the actual construction: the instrument interacts each country's predetermined 1996 air market access (foreign populations weighted by inverse great-circle distance) with a global aviation-intensity index built from world scheduled seat capacity, while sea market access, constructed identically from CERDI sea distances, enters as a control that absorbs the maritime-geography channel. The residual identifying variation is the aviation-specific gain accruing to air-favoured geographies as world aviation grows, and the exclusion restriction is stated accordingly. We also replaced \"fully independent\" with a precise statement that it shares no variation with tourism or heritage."),
("25. Should some limitations appear in the earlier result sections?",
"Yes; as noted above, the caveats on temporal identification, the mechanism decomposition, the counterfactual extrapolation, and collapse symmetry now appear where the results are presented, not only in the discussion."),
("26. Conclusion too brief.",
"Expanded: it now states the scope of the estimates (most informative about connectivity changes of the kind the instruments induce, in the recent aviation era) and the concrete next steps (bilateral gravity structure and general-equilibrium embedding)."),
]
for h, b in items:
    item(h, b)

para('Thank you again; the paper is much tighter as a result. We would of course be happy to discuss any of these points further.', space_after=10)
para('Best regards,')
para('Sunbin (and Longfei)')

doc.save('Response_to_Tommy_pointbypoint_20260901.docx')
print('saved')
