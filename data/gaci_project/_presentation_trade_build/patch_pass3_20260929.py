# -*- coding: utf-8 -*-
"""Pass 3 (2026-09-29, after the user's hand edit at 14:22): patch the deck IN PLACE.

  * slide 2 (speaker slide the user added): restyled into the deck palette, same content
    (+ Keio joint appointment, which the user's own speaker note already mentions)
  * new slide 4   'Why aviation'            (after the rail -> air origin slide)
  * new slide 16  Northeast Asia hub race   (after the 2024 hub-quality map)
  * new slide 17  country change map         (after the hub race)
  * page numbers renumbered; notes.json regenerated from the deck for build_script_docx.py

Inputs: fun_data.py -> fun_data.json, make_slide_maps_pass3.py -> country_change.json, krjp_extent.json,
assets/map_hubchange_1996_2024.png, assets/map_krjp_about.png.
Do NOT rerun build_deck.py: it would drop the user's slide 2 and this pass."""
import copy, json, math, os, shutil
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE,
                            XL_TICK_LABEL_POSITION, XL_AXIS_CROSSES)
from pptx.oxml.ns import qn
from pptx_helpers import *

HERE = os.path.dirname(os.path.abspath(__file__))
A = lambda f: os.path.join(HERE, 'assets', f)
DECK = os.path.join(HERE, '..', 'GACI_presentation_conference_20260929.pptx')
BACKUP = os.path.join(HERE, '_backup', 'GACI_presentation_conference_20260929_userEdit1422_prePass3.pptx')
F = json.load(open(os.path.join(HERE, 'fun_data.json'), encoding='utf-8'))
CC = json.load(open(os.path.join(HERE, 'country_change.json'), encoding='utf-8'))
KJ = json.load(open(os.path.join(HERE, 'krjp_extent.json'), encoding='utf-8'))
OLD_NOTES = json.load(open(os.path.join(HERE, '_notes_before_pass3.json'), encoding='utf-8'))   # 29-slide build

os.makedirs(os.path.dirname(BACKUP), exist_ok=True)
if not os.path.exists(BACKUP):
    shutil.copy2(DECK, BACKUP)
    print('backup ->', BACKUP)
SRC = BACKUP                     # always patch from the user's 14:22 version, so the script is re-runnable

prs = Presentation(SRC)
assert len(prs.slides) == 30, 'expected the 30-slide user version'
BL = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
LM = 0.7
CW = SW - 2 * LM
MINUS = '−'


# ======================================================================= scaffolding (as build_deck.py)
def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def page_no(s, n=0):
    add_text(s, SW - 1.2, SH - 0.45, 0.6, 0.3, [str(n)], size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def content(tag, title, size=28):
    s = prs.slides.add_slide(BL)
    bg(s, WHITE)
    add_text(s, LM, 0.40, 9, 0.3, [[(tag.upper(), dict(spc=120, b=True))]], size=10.5, color=BRASS)
    add_text(s, LM, 0.66, CW, 1.1, [title], size=size, font=HEAD, color=TEXT, line=0.95)
    page_no(s)
    return s


def footer(s, text, y=SH - 0.47, w=10.6):
    add_text(s, LM, y, w, 0.35, [text], size=10, color=MUTED)


def tag(s, x, y, text, color=MUTED, w=5.0):
    add_text(s, x, y, w, 0.3, [[(text.upper(), dict(spc=100, b=True))]], size=10.5, color=color)


def chart_frame(s, ctype, data, x, y, w, h):
    gf = s.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    chart_base(ch)
    return ch


def plot_layout(chart, x, y, w, h):
    pa = chart._chartSpace.find('.//' + qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = etree.Element(qn('c:layout'))
        pa.insert(0, lay)
    for ch in list(lay):
        lay.remove(ch)
    ml = etree.SubElement(lay, qn('c:manualLayout'))
    for t, val in (('c:layoutTarget', 'inner'), ('c:xMode', 'edge'), ('c:yMode', 'edge'),
                   ('c:x', x), ('c:y', y), ('c:w', w), ('c:h', h)):
        e = etree.SubElement(ml, qn(t))
        e.set('val', str(val))


def cat_skip(chart, n):
    ax = chart._chartSpace.find('.//' + qn('c:catAx'))
    for t in ('c:tickLblSkip', 'c:tickMarkSkip'):
        old = ax.find(qn(t))
        if old is not None:
            ax.remove(old)
    el = etree.Element(qn('c:tickLblSkip')); el.set('val', str(n))
    nm = ax.find(qn('c:noMultiLvlLbl'))
    if nm is not None:
        nm.addprevious(el)
    else:
        ax.append(el)


def korean(box):
    """Hangul runs: East Asian font + ko-KR so PowerPoint does not fall back to a random face."""
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            rPr.set('lang', 'ko-KR'); rPr.set('altLang', 'en-US')
            lat = rPr.find(qn('a:latin'))
            ea = etree.Element(qn('a:ea')); ea.set('typeface', 'Malgun Gothic')
            if lat is not None:
                lat.addnext(ea)
            else:
                rPr.append(ea)


NOTES_TEMPLATE = prs.slides[3].notes_slide     # an existing notes page (the notes master has no placeholders)


def set_notes(s, text):
    ns = s.notes_slide
    if ns.notes_text_frame is None:
        tree = ns.shapes._spTree
        for sh in NOTES_TEMPLATE.shapes:
            if sh.is_placeholder and sh.placeholder_format.type in (101, 2):
                tree.append(copy.deepcopy(sh._element))
    ns.notes_text_frame.text = text.strip()


def move_slide(prs, old_idx, new_idx):
    lst = prs.slides._sldIdLst
    el = list(lst)[old_idx]
    lst.remove(el)
    lst.insert(new_idx, el)


# ======================================================================= slide 2: speaker (restyle)
s2 = prs.slides[1]
user_note = s2.notes_slide.notes_text_frame.text.strip() if s2.has_notes_slide else ''
for sh in list(s2.shapes):
    sh._element.getparent().remove(sh._element)
bg(s2, WHITE)

PW = 4.85                                       # left panel width
rect(s2, 0, 0, PW, SH, fill=INK)
add_text(s2, 0.55, 0.50, 4.0, 0.3, [[('ABOUT ME', dict(spc=120, b=True))]], size=10.5, color=BRASS)
add_text(s2, 0.55, 0.72, 4.2, 0.7, ['Sunbin YOO'], size=34, font=HEAD, color=WHITE)
korean(add_text(s2, 0.57, 1.36, 3.0, 0.4, ['유선빈'], size=17, color=ONDARK))
hline(s2, 0.57, 1.88, 1.17, 1.88, color=BRASS, w=1.5)

tag(s2, 0.55, 2.02, 'Current', color=BRASS, w=4.0)
add_text(s2, 0.55, 2.28, 4.1, 1.3, [
    dict(runs=[('Associate Professor, Department of Energy', dict(b=True))], size=14, color=WHITE, after=0),
    dict(runs='Sungkyunkwan University (SKKU), Korea', size=13, color=ONDARK, after=7),
    dict(runs=[('Project Associate Professor (joint appointment)', dict(b=True))], size=13, color=WHITE, after=0),
    dict(runs='Keio University, Graduate School of Media Design', size=12.5, color=ONDARK, after=0),
], size=13, color=WHITE, line=1.0)
tag(s2, 0.55, 3.62, 'Previously', color=BRASS, w=4.0)
add_text(s2, 0.55, 3.88, 4.1, 0.9, [
    dict(runs=['Kyushu University, ', ('Faculty of Engineering', dict(color=ONDARK))], after=1),
    dict(runs=['The University of Tokyo, ', ('PhD', dict(color=ONDARK))], after=1),
    dict(runs=['Yonsei University, ', ('M.A. in Economics', dict(color=ONDARK))], after=0),
], size=13, color=WHITE, line=1.0)

# Korea-Japan map with the three cities; native labels on the pins
MX, MY, MW = 0.45, 4.78, 4.0
MH = MW * 817 / 1260
s2.shapes.add_picture(A('map_krjp_about.png'), Inches(MX), Inches(MY), Inches(MW), Inches(MH))


def kj(lon, lat):
    return (MX + MW * (lon - KJ['lon0']) / (KJ['lon1'] - KJ['lon0']),
            MY + MH * (KJ['lat1'] - lat) / (KJ['lat1'] - KJ['lat0']))


sx, sy = kj(*KJ['cities']['Seoul'])
tx, ty = kj(*KJ['cities']['Tokyo'])
fx, fy = kj(*KJ['cities']['Fukuoka'])
add_text(s2, sx - 0.55, sy - 0.62, 1.6, 0.45, [dict(runs=[('Seoul', dict(b=True))], after=0),
                                               dict(runs='SKKU · Yonsei', size=10, color=ONDARK)],
         size=11, color=WHITE, line=0.95)
add_text(s2, tx + 0.02, ty + 0.14, 1.25, 0.45, [dict(runs=[('Tokyo', dict(b=True))], after=0),
                                                dict(runs='UTokyo · Keio', size=10, color=ONDARK)],
         size=11, color=WHITE, line=0.95)
add_text(s2, fx - 1.32, fy + 0.02, 1.2, 0.45, [dict(runs=[('Fukuoka', dict(b=True))], after=0),
                                               dict(runs='Kyushu Univ.', size=10, color=ONDARK)],
         size=11, color=WHITE, line=0.95, align=PP_ALIGN.RIGHT)

# right: interests + themes
RX = 5.25
RW = SW - 0.6 - RX
add_text(s2, RX, 0.56, RW, 0.6, ['Research themes'], size=28, font=HEAD, color=TEXT)
add_text(s2, RX, 1.22, RW, 0.62, [
    dict(runs='Energy & environmental economics · Energy policy evaluation · Behavioral economics & choice modeling', after=0),
    dict(runs='Low-carbon transport & infrastructure · AI data centers & clean power systems', after=0)],
    size=12, color=MUTED, line=1.05)
THEMES = [
    ('Transport infrastructure & decarbonization', [
        ('HSR, expressways & energy infra revitalization', 'Transportation Research Part A, 2026'),
        ('Sectoral decarbonization and HSR', 'Energy Economics, 2026'),
        ('Air connectivity and aviation CO₂', 'Under review')]),
    ('AI data centers & power systems', [
        ('Can renewables meet AI power demand? A systematic review', 'RSER, 2026'),
        ('Clean power buildout for AI data centers (Minnesota & Korea)', 'Applied Energy (in prep.)')]),
    ('Energy & mobility behavior', [
        ('Diesel policy abandonment and consumer behavior', 'Energy Economics, 2020'),
        ('Drone vs truck: choice & WTP for last-mile delivery', 'Transportation Research Part E, 2026'),
        ('Ethical dilemma of autonomous vehicles', 'Transportation Research Part A, 2023')]),
    ('Low-carbon cities & solar land use', [
        ('Railway expansion and air pollution in Tokyo over 25 years', 'Sustainable Cities and Society, 2025'),
        ('How much solar can a land-constrained country fit?', 'RSER, under review')]),
]
TX = RX + 0.45         # text column
OX = 10.2              # outlet column
RE = SW - 0.6          # right edge
y = 1.98
hline(s2, RX, y - 0.08, RE, y - 0.08, color=TEXT, w=1.0)
for k, (head, papers) in enumerate(THEMES):
    add_text(s2, RX, y - 0.02, 0.45, 0.45, [str(k + 1)], size=21, font=HEAD, color=BRASS)
    add_text(s2, TX, y + 0.02, 5.0, 0.38, [head], size=16, font=HEAD, color=INK)
    if k == 0:
        add_text(s2, RE - 1.6, y + 0.08, 1.6, 0.3, [[('today’s talk', dict(i=True))]],
                 size=11.5, color=BRASS, align=PP_ALIGN.RIGHT)
    y += 0.39
    for title, outlet in papers:
        add_text(s2, TX, y, OX - TX - 0.1, 0.28, [title], size=12, color=TEXT)
        add_text(s2, OX, y, RE - OX, 0.28, [[(outlet, dict(i=True))]], size=11, color=MUTED)
        y += 0.27
    y += 0.09
    hline(s2, RX, y - 0.03, RE, y - 0.03, color=RULE if k < 3 else TEXT, w=0.75 if k < 3 else 1.0)
    y += 0.09
page_no(s2, 2)
set_notes(s2, (user_note + '\n\n' if user_note else '') + """
[English] I did my M.A. in economics at Yonsei, my PhD at the University of Tokyo, then worked at Kyushu University's Faculty of Engineering before moving to SKKU. Since this September I also hold a joint appointment at Keio's Graduate School of Media Design. My work falls into the four themes on the right; today's paper belongs to the first one, transport infrastructure and decarbonization.
""")
print('slide 2 restyled; y end of themes = %.2f in' % y)


# ======================================================================= new: why aviation
s = content('Why aviation', 'Why aviation? One global network, rebuilt every year and recorded seat by seat')
# (a) world seats, 2019 = 100
tag(s, LM, 1.98, 'Shocks hit everywhere at once')
add_text(s, LM, 2.24, 6.5, 0.3, ['World scheduled seats, 2019 = 100'], size=13, color=MUTED)
yrs = F['asia_years']
cd = CategoryChartData()
cd.categories = [str(v) for v in yrs]
cd.add_series('Seats', F['seats_idx'])
ch = chart_frame(s, XL_CHART_TYPE.LINE, cd, LM - 0.1, 2.5, 6.95, 2.8)
plot_layout(ch, 0.07, 0.05, 0.86, 0.80)
ser = ch.plots[0].series[0]
line_style(ser, INK, 2.75)
ser.marker.style = XL_MARKER_STYLE.NONE
va = ch.value_axis
va.minimum_scale = 0; va.maximum_scale = 120; va.major_unit = 40
style_axis(va, size=12, line=False, grid=True, numfmt='0')
style_axis(ch.category_axis, size=12, line=True)
cat_skip(ch, 4)
i20, i24, i96 = yrs.index(2020), yrs.index(2024), 0
for idx, txt, pos, col in ((i20, '2020: %s%d%%' % (MINUS, round(abs(F['seats_2020_drop']))), XL_LABEL_POSITION.BELOW, BRASS),
                           (i24, '2024: %d' % round(F['seats_idx'][i24]), XL_LABEL_POSITION.ABOVE, INK),
                           (i96, '1996: %d' % round(F['seats_idx'][i96]), XL_LABEL_POSITION.ABOVE, MUTED)):
    pt = ser.points[idx]
    pt.marker.style = XL_MARKER_STYLE.CIRCLE
    pt.marker.size = 8
    pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = col
    pt.marker.format.line.color.rgb = col
    point_label(ser, idx, txt, size=12.5, color=col, pos=pos, b=(idx == i20))

# (b) airports come and go: waffle, 100 squares
X2 = 8.05
tag(s, X2, 1.98, 'Airports come and go')
n_ever, n_all = F['ever'], F['always']
k_all = round(100 * n_all / n_ever)
cell, gap = 0.19, 0.036
for j in range(100):
    col_, row_ = divmod(j, 5)
    rect(s, X2 + col_ * (cell + gap), 2.36 + row_ * (cell + gap), cell, cell,
         fill=INK if j < k_all else BRASS)
wy = 2.36 + 5 * (cell + gap) + 0.16
rect(s, X2, wy + 0.07, 0.16, 0.16, fill=INK)
add_text(s, X2 + 0.26, wy, 4.3, 0.3, [[('{:,}'.format(n_all), dict(b=True)), ' served in all 29 years']], size=13.5, color=TEXT)
rect(s, X2, wy + 0.42, 0.16, 0.16, fill=BRASS)
add_text(s, X2 + 0.26, wy + 0.35, 4.3, 0.3, [[('{:,}'.format(n_ever - n_all), dict(b=True)), ' opened, closed or paused along the way']],
         size=13.5, color=TEXT)
add_text(s, X2, wy + 0.82, 4.55, 0.6,
         ['{:,} airports had scheduled flights in at least one year since 1996. One square is about {} airports.'.format(n_ever, round(n_ever / 100))],
         size=12, color=MUTED, i=True, line=1.02)

# (c) number one, year by year
tag(s, LM, 5.5, 'Number one airport, each year')
add_text(s, SW - LM - 5.5, 5.47, 5.5, 0.3, [[('Six airports have held the top spot', dict(i=True))]],
         size=12.5, color=MUTED, align=PP_ALIGN.RIGHT)
FILL = {'LAX': SAGE, 'FRA': INK, 'CDG': RGBColor(0x8C, 0x87, 0x7F), 'PEK': SLATE,
        'CAN': RGBColor(0x8F, 0xA2, 0xAB), 'DXB': BRASS}
cw_ = CW / 29
for k, (yy, apc, _) in enumerate(F['no1']):
    b = rect(s, LM + k * cw_ + 0.015, 5.82, cw_ - 0.03, 0.46, fill=FILL[apc])
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = apc
    r.font.size = Pt(9.5); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = WHITE
    if yy in (1996, 2000, 2005, 2010, 2015, 2020, 2024):
        add_text(s, LM + k * cw_ - 0.2, 6.31, cw_ + 0.4, 0.25, [str(yy)], size=11, color=MUTED, align=PP_ALIGN.CENTER)
footer(s, 'Seats: all scheduled seats, summed over airports. Airports: IATA codes with scheduled service in the GACI panel, 1996 to 2024. '
          'Number one: highest airport GACI in that year.', w=11.4)
set_notes(s, """
Why aviation? For an economist the air network is an unusually good laboratory, for one main reason: it is a single global network, it is rebuilt every year, and every scheduled seat is on record.
On the left are world scheduled seats, with 2019 set to 100. Seats grew steadily for two decades, then fell by 44 percent in a single year, 2020, and by 2024 were back to 98. A shock like that hits every country at the same time, and we can see it airport by airport.
On the right, airports come and go. 5,428 airports had scheduled flights in at least one year since 1996, but only 2,253 had them in every one of the 29 years.
At the bottom, the top spot changes hands. Los Angeles and Frankfurt traded first place until 2013, Beijing held it from 2015 to 2020, and in 2024 Dubai is number one. Six different airports in 29 years.
Railways and roads change slowly and inside one country. Here the whole world network moves every year. That movement is the variation we use.
""")
why = s


# ======================================================================= new: Northeast Asia hub race
s = content('01  Building the index', 'Northeast Asia’s hub race: Chinese hubs rose to the top, Incheon and Haneda climbed')
RACE = [  # code, label, colour, width, name for table
    ('NRT', 'Narita', GREY, 1.5, 'Tokyo Narita'), ('KIX', 'Kansai', GREY, 1.5, 'Osaka Kansai'),
    ('HKG', 'Hong Kong', GREY, 1.5, 'Hong Kong'), ('TPE', 'Taipei', GREY, 1.5, 'Taipei Taoyuan'),
    ('PEK', 'Beijing', INK, 2.25, 'Beijing Capital'), ('CAN', 'Guangzhou', INK, 2.25, 'Guangzhou'),
    ('PVG', 'Pudong', INK, 2.25, 'Shanghai Pudong'),
    ('HND', 'Haneda', BRASS, 2.75, 'Tokyo Haneda'), ('ICN', 'Incheon', BRASS, 2.75, 'Incheon'),
]
LO, HI = 1, 400
cd = CategoryChartData()
cd.categories = [str(v) for v in yrs]
for code, *_ in RACE:
    cd.add_series(code, F['asia_race'][code])
CX, CY, CWD, CHT = LM, 1.95, 8.35, 4.75
PX, PY, PWf, PHf = 0.075, 0.03, 0.80, 0.86
ch = chart_frame(s, XL_CHART_TYPE.LINE, cd, CX, CY, CWD, CHT)
plot_layout(ch, PX, PY, PWf, PHf)
for ser, (code, lab, col, w, _) in zip(ch.plots[0].series, RACE):
    line_style(ser, col, w)
    ser.marker.style = XL_MARKER_STYLE.NONE
va = ch.value_axis
set_log_reversed(va, LO, HI)
va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
va.format.line.fill.background()
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = RULE
va.major_gridlines.format.line.width = Pt(0.6)
va.crosses = XL_AXIS_CROSSES.MAXIMUM
style_axis(ch.category_axis, size=12, color=TEXT, line=True)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
cat_skip(ch, 4)
py0, ph = CY + PY * CHT, PHf * CHT


def ylog(r):
    return py0 + ph * (math.log10(r) - math.log10(LO)) / (math.log10(HI) - math.log10(LO))


for r in (1, 2, 5, 10, 20, 50, 100, 200):
    add_text(s, CX - 0.12, ylog(r) - 0.12, 0.62, 0.25, [str(r)], size=11, color=MUTED, align=PP_ALIGN.RIGHT)
add_text(s, CX - 0.2, CY - 0.3, 1.8, 0.25, ['World rank'], size=11, color=MUTED)
# end labels, nudged apart
ends = sorted([(ylog(F['asia_race'][c][-1]), lab, col) for c, lab, col, *_ in RACE])
placed = []
for yv, lab, col in ends:
    yy = yv if not placed else max(yv, placed[-1] + 0.2)
    placed.append(yy)
    add_text(s, CX + (PX + PWf) * CWD + 0.06, yy - 0.11, 1.0, 0.22, [lab], size=11,
             color=TEXT if col == GREY else col, b=(col != GREY))
# Incheon opening marker
xi = CX + PX * CWD + PWf * CWD * (yrs.index(2001) + 0.5) / len(yrs)
add_text(s, xi - 1.45, ylog(42) + 0.02, 1.4, 0.3, [[('Incheon opens', dict(i=True))]], size=11, color=BRASS, align=PP_ALIGN.RIGHT)
# right: table
X2 = 9.62
rows = [['Airport', '1996', 'Best', '2024']]
order = sorted(RACE, key=lambda t: F['asia_race'][t[0]][-1])
for code, lab, col, w, name in order:
    ser_ = F['asia_race'][code]
    r96 = ser_[0]
    rows.append([[(name, dict(b=(col == BRASS)))], 'n.a.' if r96 is None else str(r96),
                 str(min(v for v in ser_ if v)), [(str(ser_[-1]), dict(b=True))]])
tag(s, X2, 1.98, 'World rank by airport GACI')
hairline_table(s, X2, 2.3, [1.5, 0.48, 0.5, 0.52], rows, row_h=0.335, size=12.5,
               align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT])
add_text(s, X2, 5.8, 3.0, 0.8, [[('2020: ', dict(b=True, color=INK)),
                                 'Beijing, Guangzhou and Shanghai Pudong ranked 1st, 2nd and 3rd in the world.']],
         size=12.5, color=TEXT, line=1.02)
footer(s, 'Rank among all airports with scheduled service (3,171 to 3,863 a year), log scale. '
          'n.a.: not yet open in 1996 (Pudong opened in 1999, Incheon in 2001).', w=11.4)
set_notes(s, """
A closer look at our own neighbourhood. These are the world ranks of nine Northeast Asian airports, out of roughly 3,800 airports, on a log scale.
The Chinese hubs rose fastest. Guangzhou was 155th in 1996 and Beijing 47th. Beijing was number one in the world from 2015 to 2020. In 2020 Beijing, Guangzhou and Shanghai Pudong were first, second and third, as China's large domestic network restarted early while most international routes were shut.
Incheon entered at 42nd when it opened in 2001 and is 15th in 2024. Haneda climbed from 151st to 31st, most of it after it reopened to regular international flights around 2010. Narita and Kansai moved the other way; Kansai fell from 41st to 88th.
Keep Korea in mind. It comes back on the next slide, and again in the results.
""")
race_slide = s


# ======================================================================= new: country change map
up = round(100 * CC['share_up'])
s = content('01  Building the index', 'Hub quality rose in three of every four economies; Qatar, South Korea and the UAE gained most')
mw = 8.85
s.shapes.add_picture(A('map_hubchange_1996_2024.png'), Inches(0.3), Inches(2.25), Inches(mw))
NAME = {'QA': 'Qatar', 'KR': 'South Korea', 'AE': 'UAE', 'ET': 'Ethiopia', 'TR': 'Türkiye', 'SA': 'Saudi Arabia',
        'UA': 'Ukraine', 'SY': 'Syria'}
X2 = 9.5
tag(s, X2, 1.98, 'Largest gains, 1996 to 2024')
rows = [['Economy', 'Change', 'Rank']]
for c, r96, r24, pct in CC['top'][:6]:
    rows.append([[(NAME[c], dict(b=(c == 'KR')))], '+%d%%' % round(pct), '%d → %d' % (r96, r24)])
hairline_table(s, X2, 2.3, [1.42, 0.72, 0.99], rows, row_h=0.325, size=12.5,
               align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT],
               colors={(2, 0): BRASS})
# South Korea's rank, mini line
tag(s, X2, 4.72, 'South Korea’s rank')
cd = CategoryChartData()
cd.categories = [str(v) for v in CC['kr_years']]
cd.add_series('KR', CC['kr_rank'])
KX, KY, KW, KH = X2 - 0.1, 4.98, 3.3, 1.75
ch = chart_frame(s, XL_CHART_TYPE.LINE, cd, KX, KY, KW, KH)
kpx, kpy, kpw, kph = 0.1, 0.06, 0.8, 0.72
plot_layout(ch, kpx, kpy, kpw, kph)
ser = ch.plots[0].series[0]
line_style(ser, BRASS, 2.5)
ser.marker.style = XL_MARKER_STYLE.NONE
va = ch.value_axis
va.reverse_order = True
va.minimum_scale = 0; va.maximum_scale = 125; va.major_unit = 50
va.crosses = XL_AXIS_CROSSES.MAXIMUM
style_axis(va, size=10.5, line=False, grid=True, numfmt='0')     # native gridlines at 0, 50, 100 (behind the line)
va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
kx0, kpw_in = KX + kpx * KW, kpw * KW
ky0, kph_in = KY + kpy * KH, kph * KH
kyr = lambda r: ky0 + kph_in * r / 125.0
kxi = lambda i: kx0 + kpw_in * (i + 0.5) / len(CC['kr_years'])
for r in (0, 50, 100):
    add_text(s, kx0 - 0.5, kyr(r) - 0.11, 0.42, 0.22, [str(max(r, 1))], size=10.5, color=MUTED, align=PP_ALIGN.RIGHT)
style_axis(ch.category_axis, size=10.5, line=True)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
cat_skip(ch, 7)
ky = CC['kr_years']
i01 = ky.index(2001)
pt = ser.points[i01]
pt.marker.style = XL_MARKER_STYLE.CIRCLE; pt.marker.size = 6
pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = INK
pt.marker.format.line.color.rgb = INK
add_text(s, kxi(i01) + 0.1, kyr(80) - 0.12, 2.4, 0.25,
         [[('2001: Incheon opens, 106th → 42nd', dict(i=True))]], size=10.5, color=INK)
point_label(ser, len(ky) - 1, '12th', size=11, color=BRASS, pos=XL_LABEL_POSITION.ABOVE, b=True)
f1, f2 = CC['falls'][0], CC['falls'][1]
footer(s, 'Hub quality: capacity-weighted mean GACI over each economy’s airports. Ranks among all economies with scheduled service '
          '(%d in 1996, %d in 2024). Largest falls: %s %s%d%%, %s %s%d%%.' % (
              CC['n96'], CC['n24'], NAME[f1[0]], MINUS, round(abs(f1[3])), NAME[f2[0]], MINUS, round(abs(f2[3]))), w=11.4)
set_notes(s, """
The same index at the country level. The map shows how hub quality changed between 1996 and 2024. It rose in three of every four economies, and the darkest areas are the Gulf, Ethiopia, Türkiye and South Korea.
The largest gains are Qatar, up 159 percent, South Korea, up 139 percent, and the UAE, up 121 percent. Ethiopia, Türkiye and Saudi Arabia follow.
South Korea also climbed further in rank than any other economy, from 108th to 12th. Much of that happened in a single year: in 2001, when Incheon opened, Korea went from 106th to 42nd. It dipped during the pandemic and was back to 12th by 2024.
Falls are few. The largest are Ukraine and Syria.
Remember Korea: in the results it has the second-largest trade gain from better connectivity, after China.
""")
ctry_slide = s

# ======================================================================= order, page numbers
# new slides were appended at 30, 31, 32 (0-based); target positions (0-based): why -> 3, race -> 15, country -> 16
move_slide(prs, 30, 3)          # why aviation after the origin slide
move_slide(prs, 31, 15)         # hub race after the 2024 hub map (old 14, now 15 after the insert)
move_slide(prs, 32, 16)
assert len(prs.slides) == 33


def page_boxes(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() \
                and Emu(sh.left).inches > 11.5 and Emu(sh.top).inches > 6.8:
            out.append(sh)
    return out


for i, sl in enumerate(prs.slides, 1):
    for sh in page_boxes(sl):
        runs = sh.text_frame.paragraphs[0].runs
        runs[0].text = str(i)
        for r in runs[1:]:
            r.text = ''
prs.save(DECK)
print('saved', DECK, len(prs.slides), 'slides')

# ======================================================================= notes.json for the script docx
# titles: old build numbering k -> new position (1 -> 1; 2 -> 3; 3..13 -> k+2; 14..29 -> k+4)
old_title = {x['n']: x['title'] for x in OLD_NOTES}
newpos = {1: 1, 2: 3}
newpos.update({k: k + 2 for k in range(3, 14)})
newpos.update({k: k + 4 for k in range(14, 30)})
TIT = {v: old_title[k] for k, v in newpos.items()}
TIT[2] = 'About me'
TIT[4] = 'Why aviation? One global network, rebuilt every year and recorded seat by seat'
TIT[16] = 'Northeast Asia’s hub race: Chinese hubs rose to the top, Incheon and Haneda climbed'
TIT[17] = 'Hub quality rose in three of every four economies; Qatar, South Korea and the UAE gained most'
out = []
for i, sl in enumerate(prs.slides, 1):
    nt = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ''
    out.append(dict(n=i, title=TIT[i], notes=nt))
json.dump(out, open(os.path.join(HERE, 'notes.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
print('notes.json rewritten from the deck:', len(out), 'slides,', sum(len(x['notes'].split()) for x in out), 'words')
