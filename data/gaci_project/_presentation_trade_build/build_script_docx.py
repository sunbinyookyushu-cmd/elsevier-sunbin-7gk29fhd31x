# -*- coding: utf-8 -*-
"""Speaker script (docx) for the GACI trade conference deck. Reads notes.json, which since pass 3
(2026-09-29) is regenerated from the deck itself by patch_pass3_20260929.py (33 slides)."""
import json, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'GACI_presentation_conference_20260929_script.docx')
N = json.load(open(os.path.join(HERE, 'notes.json'), encoding='utf-8'))
WPM = 120
INK = RGBColor(0x1F, 0x3A, 0x32)
BRASS = RGBColor(0xB8, 0x86, 0x2F)
MUTED = RGBColor(0x6F, 0x6A, 0x62)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')   # slide 2 note is partly Korean
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15


def para(text='', size=11, bold=False, italic=False, color=None, font=None, after=6, align=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color is not None:
            r.font.color.rgb = color
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


total_words = sum(len(x['notes'].split()) for x in N)
para('Global Air Connectivity and Trade Openness, 1996–2024', size=18, font='Georgia', color=INK, after=2)
para('Speaker script for the conference talk', size=12, italic=True, color=MUTED, after=2)
para(f'{len(N)} slides, {total_words:,} words, about {round(total_words / WPM)} minutes at {WPM} words per minute '
     '(plus transitions). The same text is in the PowerPoint speaker notes.', size=10, color=MUTED, after=14)

for x in N:
    words = len(x['notes'].split())
    secs = int(round(words / WPM * 60 / 5.0) * 5)
    p = doc.add_paragraph()
    r = p.add_run(f"Slide {x['n']}   ")
    r.bold = True; r.font.color.rgb = BRASS; r.font.size = Pt(11)
    r = p.add_run(x['title'])
    r.bold = True; r.font.color.rgb = INK; r.font.size = Pt(12); r.font.name = 'Georgia'
    r = p.add_run(f"   ~{secs} s")
    r.font.color.rgb = MUTED; r.font.size = Pt(9.5)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    for line in x['notes'].split('\n'):
        if line.strip():
            para(line.strip(), after=5)

# ------------------------------------------------------------------ Q&A preparation
doc.add_page_break()
para('Likely questions', size=16, font='Georgia', color=INK, after=4)
para('Answers use only what the manuscript (TRA_rev20260902_ext2024/main.tex) reports. Items marked "prepare" are not answered in the paper.',
     size=10, italic=True, color=MUTED, after=10)
QA = [
    ('Most trade goes by sea. Is an elasticity above one plausible?',
     'By value about one third of trade flies, and the effect travels through the margins where air operates: connectivity raises the intermediate-to-consumption ratio (1.48) and the high-to-low value-to-weight ratio (1.87), and controlling for these ratios makes the openness coefficient insignificant (1.18 to 0.83). Slides 6, 26, 27.'),
    ('Could heritage-driven tourism raise goods imports directly?',
     'The outcome is merchandise only, so tourist spending never enters it. With GDP controls the openness estimate stays at 1.34 to 1.54. Cultural heritage, the most plausible violator, is rejected by the over-identification test (p = 0.008). An independent air/sea-distance instrument gives 2.30 on volume and agrees with ours (combined 2.26, Hansen J p = 0.69). Conley bounds: volume stays positive for a direct effect up to 36% of the reduced form; openness only up to 4%. Slides 22, 25.'),
    ('The instrument has no power before 2010.',
     'Correct: F = 0.3 before 2010 and 31.6 after. The paper says so explicitly. Re-estimating at every split year from 2002 to 2013 keeps the pre-period weak (F < 4) and the post-period volume interaction positive and significant (0.23 to 0.54), so the pattern is not an artefact of the 2010 cutoff. The estimates describe the recent aviation era. Slide 21.'),
    ('Why hub quality and not the sum or the maximum?',
     'The capacity-weighted mean uses every airport but weights by seats, so a country is not rewarded for many small airfields. The three measures agree in sign and significance; the maximum hub has the strongest first stage (F = 33.0), the sum the weakest (F = 11.0). Slides 13, 24.'),
    ('Is the $10 trillion figure credible?',
     'It is partial equilibrium and applies an elasticity identified from post-2010 variation to changes since 1996. The 95% CI of the elasticity maps into $0.5 to 15.9 trillion; the three measures give $8.0 to 10.0 trillion. Economies without 2024 WDI trade data (31, including the UAE and Qatar) count as zero. Slides 29, 30.'),
    ('Is the mechanism causal?',
     'No claim beyond a decomposition: the composition ratios are equilibrium outcomes themselves, so the paper reads the pattern as supporting the channels, not proving them. Slide 27.'),
    ('Why are the IV estimates much larger than OLS? (prepare)',
     'The manuscript does not discuss the OLS/IV gap. Prepare an answer before the talk.'),
    ('Why robust rather than country-clustered standard errors? (prepare)',
     'The manuscript reports heteroskedasticity-robust SE only. In the June analysis (2023 sample) significance weakened under country clustering. Decide how to answer, and check the 2024 clustered numbers before the talk if you want to quote them.'),
]
for q, a in QA:
    p = para(q, bold=True, color=INK, after=2)
    p.paragraph_format.keep_with_next = True
    para(a, after=10)

# ------------------------------------------------------------------ numbers at a glance
doc.add_page_break()
para('Numbers at a glance', size=16, font='Georgia', color=INK, after=6)
rows = [('Sample', '185 countries, 1996–2024, N = 4,814 (regressions)'),
        ('Network', '29 annual networks, 3,171–3,863 airports per year'),
        ('PCA', 'PC1 explains 63–77% by year (74% pooled); loadings 0.49/0.43/0.39/0.43/0.48'),
        ('First stage', '0.034*** (0.007), KP F = 21.9; pre-2010 F = 0.3, post-2010 F = 31.6'),
        ('2SLS, hub quality', 'openness 1.208** (0.593), volume 2.126*** (0.703), GDP 0.919 (0.624)'),
        ('Plain reading', '10% better hub quality: trade/GDP about +12%, total trade about +22%, GDP about +9% (not significant)'),
        ('Other measures', 'max: 0.897** / 1.580***; sum: 0.635* / 1.119*** (openness / volume)'),
        ('Mechanism', 'interm./consum. 1.483**; high/low VW 1.868***; openness 1.176** falls to 0.825 with both ratios'),
        ('Heterogeneity', 'income Q1 to Q4: 1.52 to 0.71; connectivity Q1 to Q4: 2.44 to 1.49 (KP F about 6 to 10)'),
        ('Aggregate', '$10.0 tn of 2024 goods trade = 21% of world goods trade = 9.2% of world GDP'),
        ('COVID-19', 'hub quality fell in 154 of 175 countries; implied $7.1 tn (cwm), $6.3 tn (max), $1.7 tn (sum)')]
t = doc.add_table(rows=len(rows), cols=2)
t.style = 'Table Grid'
for i, (k, v) in enumerate(rows):
    c0, c1 = t.rows[i].cells
    c0.width, c1.width = Cm(4.0), Cm(12.6)
    c0.text = ''; c1.text = ''
    r = c0.paragraphs[0].add_run(k); r.bold = True; r.font.size = Pt(10)
    r = c1.paragraphs[0].add_run(v); r.font.size = Pt(10)

doc.save(OUT)
print('saved', OUT, total_words, 'words')
