# -*- coding: utf-8 -*-
"""Small python-pptx helpers for the GACI trade deck: text, hairline tables,
native charts with custom error bars, simple network shapes."""
import copy
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.chart import XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ palette
INK = RGBColor(0x1F, 0x3A, 0x32)
BRASS = RGBColor(0xB8, 0x86, 0x2F)
TEXT = RGBColor(0x22, 0x21, 0x1F)
MUTED = RGBColor(0x6F, 0x6A, 0x62)
RULE = RGBColor(0xD9, 0xD5, 0xCC)
TINT = RGBColor(0xF3, 0xF1, 0xEC)
SAGE = RGBColor(0x7F, 0x9A, 0x8E)
GREY = RGBColor(0xBD, 0xB8, 0xAE)
SLATE = RGBColor(0x3F, 0x5A, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ONDARK = RGBColor(0xC9, 0xD6, 0xCF)

HEAD = 'Georgia'
BODY = 'Calibri'
MATH = 'Cambria Math'


def hexs(rgb):
    return str(rgb)


# ------------------------------------------------------------------ text
def _apply_run(run, text, o, base):
    run.text = text
    f = run.font
    f.name = o.get('font', base['font'])
    f.size = Pt(o.get('size', base['size']))
    f.bold = o.get('b', base.get('b', False))
    f.italic = o.get('i', base.get('i', False))
    f.color.rgb = o.get('color', base['color'])
    rPr = run._r.get_or_add_rPr()
    if o.get('sub'):
        rPr.set('baseline', '-25000')
    if o.get('sup'):
        rPr.set('baseline', '30000')
    if o.get('spc'):
        rPr.set('spc', str(o['spc']))
    rPr.set('lang', 'en-US')


def add_text(slide, x, y, w, h, paras, size=16, color=TEXT, font=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, after=6, line=None, margin=0.0, b=False, i=False, wrap=True):
    """paras: list of items; each item is str, or list of runs, or dict(runs=..., size=, color=,
    align=, after=, b=, i=, bullet=True, indent=inches, line=)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = anchor
    base = dict(size=size, color=color, font=font, b=b, i=i)
    for k, item in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        if isinstance(item, dict):
            runs = item.get('runs', item.get('text', ''))
            pb = dict(base)
            for key in ('size', 'color', 'font', 'b', 'i'):
                if key in item:
                    pb[key] = item[key]
            p.alignment = item.get('align', align)
            p.space_after = Pt(item.get('after', after))
            if item.get('before') is not None:
                p.space_before = Pt(item['before'])
            ln = item.get('line', line)
            if ln:
                p.line_spacing = ln
            if item.get('bullet'):
                ind = Inches(item.get('indent', 0.22))
                pPr = p._p.get_or_add_pPr()
                pPr.set('marL', str(int(ind)))
                pPr.set('indent', str(-int(ind)))
                bc = etree.SubElement(pPr, qn('a:buClr'))
                c = etree.SubElement(bc, qn('a:srgbClr'))
                c.set('val', hexs(item.get('bullet_color', BRASS)))
                bf = etree.SubElement(pPr, qn('a:buFont'))
                bf.set('typeface', 'Arial')
                bu = etree.SubElement(pPr, qn('a:buChar'))
                bu.set('char', '•')
        else:
            runs = item
            pb = base
            p.alignment = align
            p.space_after = Pt(after)
            if line:
                p.line_spacing = line
        if isinstance(runs, str):
            runs = [runs]
        for r in runs:
            if isinstance(r, str):
                t, o = r, {}
            else:
                t, o = r
            _apply_run(p.add_run(), t, o, pb)
    return box


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ''
    return s


def hline(slide, x1, y1, x2, y2, color=RULE, w=0.75, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', dash)
    return c


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=1.25, dash=None, both=False):
    c = hline(slide, x1, y1, x2, y2, color, w, dash)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, qn('a:tailEnd'))
    te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    if both:
        he = etree.SubElement(ln, qn('a:headEnd'))
        he.set('type', 'triangle'); he.set('w', 'med'); he.set('len', 'med')
    return c


def node(slide, cx, cy, d, fill, label='', size=12, color=WHITE, line=None, b=True, font=BODY):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    if label:
        _apply_run(p.add_run(), label, {}, dict(size=size, color=color, font=font, b=b))
    return s


# ------------------------------------------------------------------ tables
NOSTYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'


def _ln(tag, color=None, w=0):
    el = etree.Element(qn(tag))
    if color is None:
        el.set('w', '0')
        etree.SubElement(el, qn('a:noFill'))
    else:
        el.set('w', str(int(Pt(w))))
        el.set('cap', 'flat'); el.set('cmpd', 'sng'); el.set('algn', 'ctr')
        sf = etree.SubElement(el, qn('a:solidFill'))
        c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', hexs(color))
        pd = etree.SubElement(el, qn('a:prstDash')); pd.set('val', 'solid')
    return el


def hairline_table(slide, x, y, col_w, rows, row_h=0.36, size=13, header=1, align=None,
                   fills=None, bold_cells=None, colors=None, heavy_after=(), light=True,
                   font=BODY, italic_rows=(), cell_margin=0.06):
    """rows: list of lists of cell specs (str or list of runs). Hairline style:
    heavy rule above first row, below header and at bottom; light rules between body rows."""
    nr, nc = len(rows), len(rows[0])
    tot_h = sum(row_h) if isinstance(row_h, (list, tuple)) else row_h * nr
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(tot_h))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = NOSTYLE
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h if not isinstance(row_h, (list, tuple)) else row_h[r])
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            spec = rows[r][c]
            cell.margin_left = Inches(cell_margin); cell.margin_right = Inches(cell_margin)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            a = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            if align:
                a = align[c]
            p.alignment = a
            bold = r < header
            if bold_cells and (r, c) in bold_cells:
                bold = True
            col = TEXT if r >= header else MUTED
            if colors and (r, c) in colors:
                col = colors[(r, c)]
            base = dict(size=size, color=col, font=font, b=bold, i=(r in italic_rows))
            runs = spec if isinstance(spec, list) else [spec]
            for rr in runs:
                if isinstance(rr, str):
                    t, o = rr, {}
                else:
                    t, o = rr
                _apply_run(p.add_run(), t, o, base)
            # borders + fill
            tcPr = cell._tc.get_or_add_tcPr()
            for ch in list(tcPr):
                tcPr.remove(ch)
            top = bottom = None
            if r == 0:
                top = (TEXT, 1.25)
            if r == header - 1 or r in heavy_after:
                bottom = (TEXT, 0.9)
            elif r == nr - 1:
                bottom = (TEXT, 1.25)
            elif r >= header and light:
                bottom = (RULE, 0.6)
            tcPr.append(_ln('a:lnL'))
            tcPr.append(_ln('a:lnR'))
            tcPr.append(_ln('a:lnT', *top) if top else _ln('a:lnT'))
            tcPr.append(_ln('a:lnB', *bottom) if bottom else _ln('a:lnB'))
            f = fills.get((r, c)) if fills else None
            if f is None and fills:
                f = fills.get((r, None))
            if f is not None:
                sf = etree.SubElement(tcPr, qn('a:solidFill'))
                cc = etree.SubElement(sf, qn('a:srgbClr')); cc.set('val', hexs(f))
            else:
                etree.SubElement(tcPr, qn('a:noFill'))
    return shp


# ------------------------------------------------------------------ charts
def style_axis(ax, size=12, color=MUTED, line=True, grid=False, title=None, title_size=12, numfmt=None):
    ax.tick_labels.font.size = Pt(size)
    ax.tick_labels.font.color.rgb = color
    ax.tick_labels.font.name = BODY
    if numfmt:
        ax.tick_labels.number_format = numfmt
        ax.tick_labels.number_format_is_linked = False
    ax.major_tick_mark = XL_TICK_MARK.NONE
    ax.minor_tick_mark = XL_TICK_MARK.NONE
    if line:
        ax.format.line.color.rgb = GREY
        ax.format.line.width = Pt(0.75)
    else:
        ax.format.line.fill.background()
    ax.has_major_gridlines = grid
    if grid:
        ax.major_gridlines.format.line.color.rgb = RULE
        ax.major_gridlines.format.line.width = Pt(0.6)
    if title:
        ax.has_title = True
        tf = ax.axis_title.text_frame
        tf.text = ''
        _apply_run(tf.paragraphs[0].add_run(), title, {}, dict(size=title_size, color=color, font=BODY, b=False))


def chart_base(chart, legend=False):
    chart.has_title = False
    chart.has_legend = legend
    cs = chart._chartSpace
    # no chart border / fill
    spPr = cs.find(qn('c:spPr'))
    if spPr is None:
        spPr = etree.SubElement(cs, qn('c:spPr'))
        # c:spPr must precede c:txPr / externalData etc.; move to right place
        txPr = cs.find(qn('c:txPr'))
        if txPr is not None:
            txPr.addprevious(spPr)
    for ch in list(spPr):
        spPr.remove(ch)
    etree.SubElement(spPr, qn('a:noFill'))
    ln = etree.SubElement(spPr, qn('a:ln')); etree.SubElement(ln, qn('a:noFill'))
    chart.font.name = BODY
    chart.font.size = Pt(12)
    chart.font.color.rgb = MUTED


def _numlit(vals):
    nl = etree.Element(qn('c:numLit'))
    fc = etree.SubElement(nl, qn('c:formatCode')); fc.text = 'General'
    pc = etree.SubElement(nl, qn('c:ptCount')); pc.set('val', str(len(vals)))
    for i, v in enumerate(vals):
        pt = etree.SubElement(nl, qn('c:pt')); pt.set('idx', str(i))
        vv = etree.SubElement(pt, qn('c:v')); vv.text = repr(float(v))
    return nl


def add_errbars(series, plus, minus=None, color=TEXT, w=1.25, direction=None):
    """Custom +/- error bars on a bar/line/scatter series (native, editable)."""
    if minus is None:
        minus = plus
    ser = series._element
    eb = etree.Element(qn('c:errBars'))
    if direction:
        d = etree.SubElement(eb, qn('c:errDir')); d.set('val', direction)
    t = etree.SubElement(eb, qn('c:errBarType')); t.set('val', 'both')
    vt = etree.SubElement(eb, qn('c:errValType')); vt.set('val', 'cust')
    ne = etree.SubElement(eb, qn('c:noEndCap')); ne.set('val', '0')
    pl = etree.SubElement(eb, qn('c:plus')); pl.append(_numlit(plus))
    mi = etree.SubElement(eb, qn('c:minus')); mi.append(_numlit(minus))
    sp = etree.SubElement(eb, qn('c:spPr'))
    ln = etree.SubElement(sp, qn('a:ln')); ln.set('w', str(int(Pt(w))))
    sf = etree.SubElement(ln, qn('a:solidFill'))
    c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', hexs(color))
    # insert before cat/xVal/val/yVal
    anchor = None
    for tag in ('c:cat', 'c:xVal', 'c:val', 'c:yVal'):
        anchor = ser.find(qn(tag))
        if anchor is not None:
            break
    anchor.addprevious(eb)
    return eb


def set_log_reversed(value_axis, lo, hi):
    sc = value_axis._element.find(qn('c:scaling'))
    lb = etree.Element(qn('c:logBase')); lb.set('val', '10')
    sc.insert(0, lb)
    value_axis.reverse_order = True
    value_axis.minimum_scale = lo
    value_axis.maximum_scale = hi


def point_label(series, idx, text, size=11, color=TEXT, pos=None, b=False):
    pt = series.points[idx]
    dl = pt.data_label
    tf = dl.text_frame
    tf.text = ''
    _apply_run(tf.paragraphs[0].add_run(), text, {}, dict(size=size, color=color, font=BODY, b=b))
    if pos is not None:
        dl.position = pos
    return dl


def no_line(series):
    series.format.line.fill.background()


def line_style(series, color, w=2.0, dash=None):
    series.format.line.color.rgb = color
    series.format.line.width = Pt(w)
    series.smooth = False
    if dash:
        ln = series.format.line._get_or_add_ln()
        pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', dash)


def marker(series, style, size, color, line_color=None):
    series.marker.style = style
    series.marker.size = size
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb = color
    series.marker.format.line.color.rgb = line_color or color


def gap_width(chart, gw=80, overlap=None):
    plot = chart.plots[0]
    plot.gap_width = gw
    if overlap is not None:
        plot.overlap = overlap
