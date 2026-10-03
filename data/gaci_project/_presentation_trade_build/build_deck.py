# -*- coding: utf-8 -*-
"""GACI x trade conference deck (2024 manuscript numbers), 16:9, ~29 slides.
Numbers: TRA_rev20260902_ext2024/main.tex; chart data: chart_data.json (from ext2024).
Run:  python build_deck.py   ->  ../GACI_presentation_conference_20260929.pptx + notes.json"""
import json, math, os
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE,
                            XL_TICK_LABEL_POSITION, XL_AXIS_CROSSES, XL_LEGEND_POSITION)
from pptx.oxml.ns import qn
from pptx_helpers import *

HERE = os.path.dirname(os.path.abspath(__file__))
A = lambda f: os.path.join(HERE, 'assets', f)
D = json.load(open(os.path.join(HERE, 'chart_data.json'), encoding='utf-8'))
OUT = os.path.join(HERE, '..', 'GACI_presentation_conference_20260929.pptx')

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]
NOTES = []
TITLES = {}
SW, SH = 13.333, 7.5
LM = 0.7            # left margin
CW = SW - 2 * LM    # content width


# ======================================================================= scaffolding
def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text.strip()
    NOTES.append(dict(n=len(prs.slides), title=TITLES.get(len(prs.slides), ''), notes=text.strip()))


def page_no(s):
    n = len(prs.slides)
    add_text(s, SW - 1.2, SH - 0.45, 0.6, 0.3, [str(n)], size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def content(tag, title, size=28):
    s = prs.slides.add_slide(BL)
    bg(s, WHITE)
    add_text(s, LM, 0.40, 9, 0.3, [[(tag.upper(), dict(spc=120, b=True))]], size=10.5, color=BRASS)
    add_text(s, LM, 0.66, CW, 1.1, [title], size=size, font=HEAD, color=TEXT, line=0.95)
    TITLES[len(prs.slides)] = title
    page_no(s)
    return s


def footer(s, text, y=SH - 0.47, w=10.6):
    add_text(s, LM, y, w, 0.35, [text], size=10, color=MUTED)


def faint_map(s, alpha=38):
    pic = s.shapes.add_picture(A('title_airports_2024.png'), 0, 0, Inches(SW), Inches(SH))
    blip = pic._element.find('.//' + qn('a:blip'))
    am = etree.SubElement(blip, qn('a:alphaModFix'))
    am.set('amt', str(alpha * 1000))
    return pic


def left_fade(s, w):
    ov = rect(s, 0, 0, w, SH, fill=INK)
    sp = ov._element.spPr
    for ch in list(sp):
        if ch.tag in (qn('a:solidFill'), qn('a:noFill'), qn('a:gradFill')):
            sp.remove(ch)
    gf = etree.Element(qn('a:gradFill')); gf.set('rotWithShape', '1')
    gl = etree.SubElement(gf, qn('a:gsLst'))
    for pos, alpha in ((0, 92), (60, 78), (100, 0)):
        gs = etree.SubElement(gl, qn('a:gs')); gs.set('pos', str(pos * 1000))
        c = etree.SubElement(gs, qn('a:srgbClr')); c.set('val', '1F3A32')
        a = etree.SubElement(c, qn('a:alpha')); a.set('val', str(alpha * 1000))
    lin = etree.SubElement(gf, qn('a:lin')); lin.set('ang', '0'); lin.set('scaled', '0')
    sp.find(qn('a:prstGeom')).addnext(gf)
    return ov


def divider(num, title, question, note):
    s = prs.slides.add_slide(BL)
    bg(s, INK)
    faint_map(s, 45)
    TITLES[len(prs.slides)] = num + '  ' + title
    add_text(s, 0.9, 1.75, 4, 1.6, [num], size=96, font=HEAD, color=BRASS)
    add_text(s, 0.9, 3.45, 8.5, 0.9, [title], size=42, font=HEAD, color=WHITE)
    add_text(s, 0.9, 4.45, 7.2, 1.0, [question], size=19, color=ONDARK, i=True, line=1.05)
    notes(s, note)
    return s


def formula(s, x, y, w, h, runs, size=22, color=TEXT, align=PP_ALIGN.LEFT):
    """runs: list of (text, opts) with sub/sup/i flags; Cambria Math."""
    rr = [(t, dict(o, font=MATH)) for t, o in runs]
    return add_text(s, x, y, w, h, [rr], size=size, color=color, font=MATH, align=align)


def sub(t):
    return (t, dict(sub=True))


def sup(t):
    return (t, dict(sup=True))


def n(t):
    return (t, {})


def it(t):
    return (t, dict(i=True))


def chart_frame(s, ctype, data, x, y, w, h):
    gf = s.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    chart_base(ch)
    return ch


def plot_layout(chart, x, y, w, h):
    """manual inner plot-area layout (fractions of the chart frame)."""
    pa = chart._chartSpace.find('.//' + qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = etree.Element(qn('c:layout'))
        pa.insert(0, lay)
    for ch in list(lay):
        lay.remove(ch)
    ml = etree.SubElement(lay, qn('c:manualLayout'))
    for tag, val in (('c:layoutTarget', 'inner'), ('c:xMode', 'edge'), ('c:yMode', 'edge'),
                     ('c:x', x), ('c:y', y), ('c:w', w), ('c:h', h)):
        e = etree.SubElement(ml, qn(tag))
        e.set('val', str(val))


# ======================================================================= 1 title
s = prs.slides.add_slide(BL)
bg(s, INK)
s.shapes.add_picture(A('title_airports_2024.png'), 0, 0, Inches(SW), Inches(SH))
left_fade(s, 9.6)
TITLES[1] = 'Title'
add_text(s, 0.85, 1.35, 6.2, 0.35, [[('1996–2024   ·   185 COUNTRIES   ·   3,863 AIRPORTS', dict(spc=100, b=True))]],
         size=11, color=BRASS)
add_text(s, 0.85, 1.85, 5.9, 2.3, ['Global Air Connectivity and Trade Openness, 1996–2024'],
         size=38, font=HEAD, color=WHITE, line=0.95)
add_text(s, 0.85, 4.2, 6.0, 0.9, ['Building a 29-year index of the world air network, and estimating what it does to goods trade'],
         size=17, color=ONDARK, i=True, line=1.05)
add_text(s, 0.85, 5.45, 6.2, 0.8, [[('Sunbin YOO', dict(b=True)), ', Junya KUMAGAI, Longfei ZHENG, Chunan WANG, Jinwoo LEE, Anming ZHANG']],
         size=14, color=WHITE, line=1.1)
add_text(s, 0.85, 6.85, 7.5, 0.3, ['Background: every airport with scheduled service in 2024, dot size proportional to its GACI score'],
         size=10, color=ONDARK)
notes(s, """
Thank you for the introduction. I am Sunbin Yoo from SKKU, and this is joint work with Junya Kumagai, Longfei Zheng, Chunan Wang, Jinwoo Lee and Anming Zhang.
The paper does two things. First, it builds a long panel of global air connectivity, every year from 1996 to 2024. The background map is the 2024 network: every airport with scheduled service, sized by its connectivity score. Second, it asks whether a country's position in this network causally raises its goods trade, through which goods, and for whom.
""")

# ======================================================================= 2 origin
s = content('How this project started', 'From railway networks to the world air network')
steps = [
    ('1', 'We simply wanted to graduate from railway research.'),
    ('2', 'We had the chance to meet Professor Anming Zhang (UBC; Co-Editor-in-Chief, Transportation Research Part A).'),
    ('3', 'He suggested that we look at aviation.'),
]
y = 2.15
for num, txt in steps:
    add_text(s, LM, y - 0.12, 0.7, 0.8, [num], size=40, font=HEAD, color=BRASS)
    add_text(s, LM + 0.8, y, 5.4, 1.0, [txt], size=19, color=TEXT, line=1.05)
    y += 1.35
# right: rail line vs air web
X0 = 7.6
add_text(s, X0, 1.95, 4.8, 0.3, [[('RAIL', dict(spc=120, b=True))]], size=10.5, color=MUTED)
hline(s, X0 + 0.1, 2.75, X0 + 4.6, 2.75, color=GREY, w=3)
for k in range(6):
    node(s, X0 + 0.1 + k * 0.9, 2.75, 0.2, WHITE, line=MUTED)
add_text(s, X0, 3.0, 4.8, 0.4, ['Stations along a corridor, within one country'], size=13, color=MUTED)
add_text(s, X0, 3.75, 4.8, 0.3, [[('AIR', dict(spc=120, b=True))]], size=10.5, color=BRASS)
pts = {'a': (X0 + 0.5, 4.6), 'b': (X0 + 1.7, 4.25), 'c': (X0 + 2.9, 4.5), 'd': (X0 + 4.3, 4.3),
       'e': (X0 + 1.0, 5.6), 'f': (X0 + 2.3, 5.45), 'g': (X0 + 3.6, 5.7), 'h': (X0 + 4.5, 5.2)}
edges = ['ab', 'bc', 'cd', 'af', 'bf', 'cf', 'fg', 'cg', 'dh', 'gh', 'ef', 'ae', 'bd', 'fh']
for e in edges:
    (x1, y1), (x2, y2) = pts[e[0]], pts[e[1]]
    hline(s, x1, y1, x2, y2, color=SAGE, w=1.25)
for k, (x, yy) in pts.items():
    big = k in 'cf'
    node(s, x, yy, 0.32 if big else 0.2, BRASS if big else INK)
add_text(s, X0, 6.0, 4.8, 0.4, ['One global network, rewired every year'], size=13, color=MUTED)
notes(s, """
A word on how this project started. Most of our earlier work is on railways and market access. We simply wanted to graduate from railway research. We then had the chance to meet Professor Anming Zhang, and he suggested that we look at aviation.
Aviation turned out to be a good laboratory for the same network questions. A railway is a corridor inside one country. The air network is one global system, it is rewired every year as airlines add and drop routes, and, as 2020 showed, it can contract within months.
""")

# ======================================================================= 3 question
s = content('Motivation', 'Does a country’s place in the air network raise its goods trade?')
cols = [
    ('1', 'Whether, and how', ['Is the effect causal?',
                                'Does it raise openness (trade relative to GDP), or only the size of the economy?',
                                'Which goods respond?']),
    ('2', 'Who gains, and when', ['Richer or poorer, better or less connected economies?',
                                   'What happens when connectivity collapses rather than grows?']),
]
for k, (num, head, items) in enumerate(cols):
    x = LM + k * 6.1
    add_text(s, x, 2.0, 0.9, 1.1, [num], size=54, font=HEAD, color=BRASS)
    add_text(s, x + 0.95, 2.2, 4.9, 0.5, [head], size=22, font=HEAD, color=INK, b=False)
    add_text(s, x + 0.95, 2.85, 4.9, 2.4, [dict(runs=t, bullet=True, after=8) for t in items], size=17, color=TEXT, line=1.05)
rect(s, LM, 5.45, CW, 1.05, fill=TINT)
add_text(s, LM + 0.3, 5.55, CW - 0.6, 0.85,
         [[('Why trade?  ', dict(b=True, color=INK)),
           'Trade is the proximate, observable margin on which connectivity acts. Income sits at the far end of a long and confounded chain.']],
         size=16, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Our question is simple to state: does a country's position in the world air network raise its goods trade?
We split it in two. First, whether and how. Is the effect causal? Does connectivity raise openness, meaning trade relative to GDP, or does it only make the economy bigger? And which goods respond?
Second, who gains and when. Do richer or poorer countries, better or less connected countries, gain more? And what happens when connectivity collapses rather than grows?
We look at trade because it is the most direct measure of integration and the proximate margin on which connectivity acts. Income is at the far end of a long and confounded chain.
""")

# ======================================================================= 4 objection
s = content('Motivation', 'Most trade moves by sea, yet air carries about a third of its value')
add_text(s, LM, 2.0, 5.6, 1.1, ['>80%'], size=64, font=HEAD, color=GREY)
add_text(s, LM, 3.15, 5.4, 0.6, ['of world trade volume travels by sea'], size=18, color=TEXT)
hline(s, LM, 3.95, LM + 5.2, 3.95, color=RULE, w=0.75)
add_text(s, LM, 4.15, 5.6, 1.1, ['~33%'], size=64, font=HEAD, color=INK)
add_text(s, LM, 5.3, 5.4, 0.6, ['of world trade value travels by air'], size=18, color=TEXT)
add_text(s, 7.0, 2.1, 5.6, 1.4,
         ['Our measure is built from passenger networks. They matter even for goods that later go by ship, because they move the people who make trade happen:'],
         size=18, color=TEXT, line=1.08)
add_text(s, 7.0, 3.65, 5.6, 1.6,
         [dict(runs=t, bullet=True, after=6) for t in
          ['buyers meeting sellers', 'managers inspecting suppliers', 'engineers servicing production lines']],
         size=18, color=TEXT)
add_text(s, 7.0, 5.15, 5.6, 0.9,
         [[('Passenger connectivity lowers the cost of the transaction, whichever mode carries the cargo.', dict(i=True, color=INK))]],
         size=17, line=1.05)
footer(s, 'Sources: UNCTAD, Review of Maritime Transport 2021; IATA, The Value of Air Cargo.')
notes(s, """
The first objection we always hear is that trade travels by sea. By weight, that is right: more than 80 percent of world trade volume moves by ship.
By value the picture is different. Roughly one third of world trade by value moves by air.
And our measure is built from passenger networks. Those matter even for goods that are later shipped by sea, because they move the people who make trade happen: buyers meeting sellers, managers inspecting suppliers, engineers servicing production lines. Passenger connectivity lowers the cost of the transaction, whichever mode then carries the cargo.
""")

# ======================================================================= 5 channels
s = content('Motivation', 'Three channels, each with a prediction we can test')
rows = [
    ['Channel', 'What moves by air', 'Prediction', 'Data'],
    [[('Transaction costs', dict(b=True))], 'Buyers, sellers and managers: face-to-face contact that matches partners and builds trust',
     [('H1  ', dict(b=True, color=BRASS)), 'Trade rises relative to GDP (openness)'], 'WDI merchandise trade'],
    [[('Belly-hold cargo', dict(b=True))], 'Time-sensitive, high-value goods; each day in transit costs 0.6–2.1% of value',
     [('H2  ', dict(b=True, color=BRASS)), 'Higher share of high value-to-weight goods'], 'BACI unit values ($/kg)'],
    [[('Value chains', dict(b=True))], 'Engineers, samples and urgent components that keep production networks running',
     [('H3  ', dict(b=True, color=BRASS)), 'More intermediate relative to consumption goods'], 'BACI end use (BEC)'],
]
hairline_table(s, LM, 2.0, [2.3, 4.6, 3.3, 1.9], rows, row_h=[0.45, 1.05, 1.05, 1.05], size=15,
               align=[PP_ALIGN.LEFT] * 4)
add_text(s, LM, 5.95, CW, 0.6,
         ['H2 and H3 concern the composition of trade, which we test with product-level data after the main results.'],
         size=15, color=MUTED, i=True)
footer(s, 'Time cost of transit: Hummels and Schaur (2013, AER).')
notes(s, """
This gives three channels, each with a prediction we can test.
First, transaction costs. Face-to-face contact lowers the cost of finding and trusting partners, so trade should rise relative to GDP. That is hypothesis one.
Second, the belly hold. Much of air cargo flies in passenger aircraft, and for time-sensitive goods each day in transit costs roughly 0.6 to 2.1 percent of their value. So connectivity should raise the share of high value-to-weight goods. Hypothesis two.
Third, value chains. Engineers, samples and urgent components move by air, so intermediate goods should rise relative to consumption goods. Hypothesis three.
Two and three are about composition, and we test them with product-level data after the main results.
""")

# ======================================================================= 6 divider 01
divider('01', 'Building the index', 'How do we measure a country’s position in the world air network?',
        'Part one: how we measure connectivity.')

# ======================================================================= 7 concept
s = content('01  Building the index', 'Connectivity is a position in the network, not a count of routes')


def star(s, cx, cy, strong):
    ang = [200, 330, 90]
    for k, a in enumerate(ang):
        r = 1.3
        hx, hy = cx + r * math.cos(math.radians(a)), cy - r * math.sin(math.radians(a))
        if strong:
            for j in range(4):
                b = a - 55 + j * 37
                sx, sy = hx + 0.62 * math.cos(math.radians(b)), hy - 0.62 * math.sin(math.radians(b))
                hline(s, hx, hy, sx, sy, color=GREY, w=1)
                node(s, sx, sy, 0.13, GREY)
        hline(s, cx, cy, hx, hy, color=INK if strong else GREY, w=3 if strong else 1.5)
        if strong:
            node(s, hx, hy, 0.62, INK, ['DXB', 'FRA', 'SIN'][k], size=11)
        else:
            node(s, hx, hy, 0.26, GREY)
    node(s, cx, cy, 0.46, BRASS, 'A' if strong else 'B', size=14)


star(s, 2.6, 3.75, True)
star(s, 6.9, 3.75, False)
add_text(s, 1.0, 5.6, 3.3, 0.9, [[('Country A', dict(b=True, color=INK))], 'Three routes to global hubs: most of the world within two legs'],
         size=13, color=MUTED, align=PP_ALIGN.CENTER, after=2)
add_text(s, 5.3, 5.6, 3.3, 0.9, [[('Country B', dict(b=True, color=INK))], 'Three routes to small regional fields: same count, far less reach'],
         size=13, color=MUTED, align=PP_ALIGN.CENTER, after=2)
hs = D['half_seats_2024']
add_text(s, 9.35, 2.0, 3.3, 1.0, ['{:,}'.format(hs['n_airports'])], size=54, font=HEAD, color=INK)
add_text(s, 9.35, 3.05, 3.3, 1.1, ['airports ({:.1f}% of {:,}) carried half of all scheduled seats in 2024'.format(100 * hs['share'], hs['n_total'])],
         size=16, color=TEXT, line=1.05)
hline(s, 9.35, 4.3, 12.6, 4.3, color=RULE)
add_text(s, 9.35, 4.45, 3.3, 1.6,
         ['A route to Incheon or Haneda is not a route to Kitakyushu. The measure has to weight capacity and the importance of partners.'],
         size=16, color=TEXT, line=1.05)
notes(s, """
The first design point is that connectivity is not a count.
Look at these two countries. Each has three international routes. Country A connects to Dubai, Frankfurt and Singapore, and through them reaches most of the world within two legs. Country B connects to three small regional fields that lead nowhere. Same count, very different access. A route to Incheon or Haneda is not the same as a route to Kitakyushu.
The network is also extremely concentrated: in 2024 only 106 airports, 2.7 percent of the total, carried half of all scheduled seats. So the measure has to weight by capacity and by the importance of partners.
""")

# ======================================================================= 8 data
s = content('01  Building the index', 'We rebuild the network every year from airline schedules')
add_text(s, LM, 1.95, 6.0, 0.3, [[('OAG SCHEDULE RECORDS', dict(spc=100, b=True))]], size=10.5, color=MUTED)
raw = [['Carrier', 'Origin', 'Country', 'First', 'Business', 'Economy', 'Freq.', 'Seats'],
       ['US', 'LAX', 'US', '0', '0', '175', '1', '175'],
       ['CN', 'CTU', 'CN', '0', '30', '271', '1', '301'],
       ['IN', 'FRA', 'DE', '0', '24', '144', '1', '168'],
       ['QA', 'CKY', 'GN', '0', '12', '147', '1', '159'],
       ['GB', 'MIA', 'US', '0', '20', '214', '1', '234'],
       ['US', 'HND', 'JP', '5', '42', '205', '1', '252']]
hairline_table(s, LM, 2.3, [0.8, 0.75, 0.85, 0.65, 0.9, 0.9, 0.65, 0.7], raw, row_h=0.36, size=12.5)
add_text(s, LM, 4.9, 6.2, 0.6, ['One row per scheduled flight segment: carrier, airports, seats by cabin, frequency'],
         size=12, color=MUTED, i=True)
arrow(s, 7.05, 3.5, 7.55, 3.5, color=BRASS, w=2)
codes = ['ATL', 'LHR', 'DXB', 'HKG', 'SPK']
Am = [['', *codes], ['ATL', '–', '1', '1', '1', '1'], ['LHR', '1', '–', '1', '1', '0'],
      ['DXB', '1', '1', '–', '1', '0'], ['HKG', '1', '1', '1', '–', '0'], ['SPK', '1', '0', '0', '0', '–']]
Wm = [['', *codes], ['ATL', '–', '9.0', '5.0', '4.0', '0.3'], ['LHR', '9.0', '–', '6.0', '5.0', '0'],
      ['DXB', '5.0', '6.0', '–', '7.0', '0'], ['HKG', '4.0', '5.0', '7.0', '–', '0'], ['SPK', '0.3', '0', '0', '0', '–']]
add_text(s, 7.75, 1.95, 2.5, 0.3, [[('A  CONNECTED (0/1)', dict(spc=80, b=True))]], size=10.5, color=MUTED)
add_text(s, 10.3, 1.95, 2.5, 0.3, [[('W  SEAT CAPACITY', dict(spc=80, b=True))]], size=10.5, color=MUTED)
hairline_table(s, 7.75, 2.3, [0.5] + [0.4] * 5, Am, row_h=0.33, size=11.5, cell_margin=0.02)
hairline_table(s, 10.3, 2.3, [0.5] + [0.4] * 5, Wm, row_h=0.33, size=11.5, cell_margin=0.02)
formula(s, 7.75, 4.55, 4.9, 0.5, [it('w'), sub('ij'), n(' = √( '), it('s'), sub('i'), it(' s'), sub('j'), n(' / '), it('d'), sub('i'), it(' d'), sub('j'), n(' )')], size=20, color=INK)
add_text(s, 7.75, 5.1, 4.9, 0.7, ['Link intensity: seat capacity per connection at both ends (s = seats, d = degree)'],
         size=12, color=MUTED, i=True)
rect(s, LM, 5.95, CW, 0.6, fill=TINT)
add_text(s, LM + 0.25, 5.95, CW - 0.5, 0.6,
         [[('29 annual networks, 1996–2024', dict(b=True, color=INK)), '     3,171 to 3,863 active airports per year     scheduled commercial service     links weighted by annual seats']],
         size=14, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
We rebuild the network from airline schedules, OAG route-segment data, every year from 1996 to 2024.
Each row is one scheduled flight segment, with its carrier, airports, seats by cabin and frequency. We aggregate segments into airport pairs and form two matrices for every year: A records which airports are connected, and W how much seat capacity links them. From these we build a link intensity, the geometric mean of capacity per connection at the two ends.
Across the 29 years the network has between 3,171 and 3,863 active airports. Airports enter and leave as service starts and stops, so the network is genuinely rebuilt every year.
""")

# ======================================================================= 9 indicators
s = content('01  Building the index', 'Five indicators describe each airport’s position and flow')
tp = {'ATL': (1.35, 2.75), 'NRT': (3.85, 2.75), 'HND': (2.6, 3.95), 'FUK': (2.6, 5.35)}
for a, b in [('ATL', 'HND'), ('ATL', 'NRT'), ('HND', 'NRT'), ('HND', 'FUK')]:
    hline(s, *tp[a], *tp[b], color=SAGE, w=2.5)
for k, (x, yy) in tp.items():
    node(s, x, yy, 0.78 if k == 'HND' else (0.5 if k == 'FUK' else 0.64), BRASS if k == 'HND' else INK, k, size=12)
add_text(s, LM, 5.85, 3.9, 0.6, ['Toy network: Fukuoka reaches Atlanta only through Haneda'], size=12, color=MUTED, i=True,
         align=PP_ALIGN.CENTER)
rows = [['', 'Indicator', 'What it measures', 'HND', 'FUK'],
        [[('Position', dict(color=MUTED, i=True))], [('Degree', dict(b=True))], 'Number of direct connections', '3', '1'],
        ['', [('Closeness', dict(b=True))], '1 ÷ average legs to every other airport', '1.00', '0.60'],
        ['', [('Eigenvector', dict(b=True))], 'Links to airports that are themselves central', 'high', 'low'],
        [[('Flow', dict(color=MUTED, i=True))], [('Flow betweenness', dict(b=True))], 'Share of passenger flow routed through the airport (flow as electrical current)', 'high', '0'],
        ['', [('Regional importance', dict(b=True))], 'Mean intensity of links within the airport’s own region', '', '']]
hairline_table(s, 5.0, 2.05, [1.05, 2.0, 3.45, 0.55, 0.55], rows, row_h=[0.42, 0.62, 0.62, 0.62, 0.8, 0.8], size=14,
               heavy_after=(3,), align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER])
footer(s, 'Indicator definitions follow Cheung, Wong and Zhang (2020, Transportation Research Part E).')
notes(s, """
On each annual network we compute five centrality indicators, following Cheung, Wong and Zhang in Transportation Research Part E.
Take this toy network. Haneda links to every other airport directly: degree three, closeness one. Fukuoka reaches Atlanta only through Haneda, two legs, so its closeness is 0.60. Eigenvector centrality rewards links to airports that are themselves central.
Those three describe position. The other two describe flow. Flow betweenness treats passengers like current in an electrical circuit and measures how much of it passes through the airport: the transfer-hub role. In the toy network nothing routes through Fukuoka, so it is zero. Regional importance is the average intensity of the airport's links within its own region.
""")

# ======================================================================= 10 PCA
s = content('01  Building the index', 'One index: the first principal component, stable over 29 years')
lab = ['Closeness', 'Eigenvector', 'Flow betweenness', 'Regional importance', 'Degree']
L = D['pca_loadings']
means = [sum(L[k]) / len(L[k]) for k in lab]
sds = [(sum((v - m) ** 2 for v in L[k]) / (len(L[k]) - 1)) ** 0.5 for k, m in zip(lab, means)]
cd = CategoryChartData()
cd.categories = lab
cd.add_series('Mean loading', [round(m, 3) for m in means])
ch = chart_frame(s, XL_CHART_TYPE.BAR_CLUSTERED, cd, LM, 1.95, 6.5, 4.4)
ser = ch.plots[0].series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = INK
gap_width(ch, 55)
add_errbars(ser, sds, color=BRASS, w=1.5)
ch.plots[0].has_data_labels = True
dl = ch.plots[0].data_labels
dl.number_format = '0.00'; dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.INSIDE_END
dl.font.size = Pt(13); dl.font.color.rgb = WHITE; dl.font.bold = True
style_axis(ch.category_axis, size=14, color=TEXT, line=True)
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0.0')
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 0.6; ch.value_axis.major_unit = 0.1
add_text(s, LM, 6.35, 6.5, 0.4, ['Mean PC1 loading across the 29 annual networks; whiskers show ±1 SD across years'],
         size=11, color=MUTED, i=True)
formula(s, 7.7, 2.0, 5.0, 0.6, [n('GACI'), sub('it'), n(' = [ '), it('d̃  f̃  c̃  ẽ  r̃'), n(' ]'), sub('it'), n(' · '), it('α'), sub('t')], size=22, color=INK)
add_text(s, 7.7, 2.6, 5.0, 0.4, ['αₜ = first principal component, re-estimated every year'], size=13, color=MUTED, i=True)
facts = [
    ('63–77%', 'of variance explained by PC1, year by year (74% pooled)'),
    ('≤ 0.03', 'year-to-year SD of each loading; all loadings positive'),
    ('0.74–0.92', 'rank correlations among the five indicators in 2024: related, not redundant'),
    ('0.98', 'correlation with an index from a single pooled PCA'),
]
yy = 3.25
for big, txt in facts:
    add_text(s, 7.7, yy, 1.55, 0.6, [big], size=20, font=HEAD, color=INK)
    add_text(s, 9.3, yy + 0.04, 3.35, 0.7, [txt], size=14, color=TEXT, line=1.0)
    yy += 0.78
notes(s, """
The five indicators are correlated, with rank correlations between 0.74 and 0.92 in 2024, but they are not redundant. We combine them by principal components, re-estimated every year.
The first component explains between 63 and 77 percent of the variance, depending on the year. The loadings are all positive and almost the same size: about 0.49 for degree, 0.48 for regional importance, 0.43 for flow betweenness and eigenvector, and 0.39 for closeness. They barely move over 29 years; the whiskers are one standard deviation across years.
Why re-estimate every year rather than once? Because the network itself changes: airports enter and exit and the scale of the indicators shifts. Re-estimating keeps the weights right for each year's network, and because the weights barely move, nothing is lost in comparability.
So GACI is close to an equal-weighted summary of the five indicators, and a single pooled PCA gives almost the same index, correlation 0.98. That is what makes the index comparable over time.
""")

# ======================================================================= 11 aggregation
s = content('01  Building the index', 'From airports to countries: hub quality is the headline measure')
blocks = [
    (True, 'Hub quality (headline)', [n('GACI'), sub('cwm,c'), n(' = Σ'), sub('i∈c'), it(' s'), sub('i'), n(' GACI'), sub('i'), n('  /  Σ'), sub('i∈c'), it(' s'), sub('i')],
     'Capacity-weighted mean: uses every airport, weights by seats'),
    (False, 'Maximum hub', [n('GACI'), sub('max,c'), n(' = max'), sub('i∈c'), n(' GACI'), sub('i')], 'Quality of the single best gateway'),
    (False, 'Total connectivity', [n('GACI'), sub('sum,c'), n(' = Σ'), sub('i∈c'), n(' GACI'), sub('i')], 'Network footprint; grows with the number of airports'),
]
yy = 2.0
for head, name, fr, desc in blocks:
    if head:
        rect(s, LM, yy - 0.1, 7.1, 1.35, fill=TINT)
    add_text(s, LM + 0.25, yy, 6.6, 0.4, [name], size=17, color=INK if head else TEXT, b=True)
    formula(s, LM + 0.25, yy + 0.38, 6.6, 0.5, fr, size=20, color=TEXT)
    add_text(s, LM + 0.25, yy + 0.83, 6.6, 0.4, [desc], size=14, color=MUTED)
    yy += 1.5
add_text(s, 8.3, 1.95, 4.3, 0.3, [[('ILLUSTRATION', dict(spc=100, b=True))]], size=10.5, color=MUTED)
ill = [['Airport', 'Seats', 'GACI'],
       ['Main hub', '80', '3.0'], ['Second airport', '15', '1.5'], ['Regional field', '5', '0.6'],
       [[('Hub quality', dict(b=True))], '', [('2.66', dict(b=True))]],
       [[('Sum', dict(b=True))], '', [('5.10', dict(b=True))]]]
hairline_table(s, 8.3, 2.3, [2.0, 1.1, 1.1], ill, row_h=0.37, size=14, heavy_after=(3,))
add_text(s, 8.3, 4.7, 4.3, 1.5,
         [[('Add ten small airfields', dict(b=True)), ' (GACI 0.55, one seat unit in total): the sum jumps from 5.10 to 10.60, while hub quality moves from 2.66 to 2.63.']],
         size=14, color=TEXT, line=1.05)
footer(s, 'All three measures give the same pattern of signs and significance; results below use hub quality unless stated.')
notes(s, """
We then move from airports to countries, in three ways.
Our headline is hub quality: the capacity-weighted mean of airport GACI. It uses the entire airport system but puts most weight on the airports that carry the seats.
We also use the maximum, the quality of the single best gateway, and the sum, the total network footprint.
The illustration on the right shows why hub quality is the headline. Adding ten small airfields doubles the sum, but hardly changes hub quality. We do not want to reward a country for opening many tiny airfields.
All three measures give the same pattern of results, so I will show hub quality and mention the others as robustness.
""")

# ======================================================================= 12 bump
s = content('01  Building the index', 'The index tracks the rise of Gulf, Turkish and East Asian hubs')
yrs = [1996, 2000, 2005, 2010, 2015, 2020, 2024]


def tr(r):
    if r is None:
        return None
    return r if r <= 20 else 20 + 6.7 * math.log(r / 20) / math.log(8)


TMIN, TMAX = 0.4, 27.2
cd = CategoryChartData()
cd.categories = [str(y) for y in yrs]
airs = D['rank_bump']['airports']
for a in airs:
    cd.add_series(a['label'], [tr(r) for _, r in a['series']])
cd.add_series('divider', [20.5] * len(yrs))
CX, CY, CWD, CHT = LM, 1.9, 8.2, 4.75
ch = chart_frame(s, XL_CHART_TYPE.LINE_MARKERS, cd, CX, CY, CWD, CHT)
PX, PY, PW, PH = 0.07, 0.03, 0.86, 0.86
plot_layout(ch, PX, PY, PW, PH)
for ser, a in zip(ch.plots[0].series, airs + [None]):
    if a is None:
        line_style(ser, GREY, 0.75, dash='dash')
        ser.marker.style = XL_MARKER_STYLE.NONE
        continue
    col = BRASS if a['riser'] else GREY
    line_style(ser, col, 2.5 if a['riser'] else 1.5)
    marker(ser, XL_MARKER_STYLE.CIRCLE, 6 if a['riser'] else 5, col)
va = ch.value_axis
va.reverse_order = True
va.minimum_scale = TMIN; va.maximum_scale = TMAX
va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
va.has_major_gridlines = False
va.format.line.fill.background()
va.crosses = XL_AXIS_CROSSES.MAXIMUM
style_axis(ch.category_axis, size=13, color=TEXT, line=True)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
# custom rank ticks (left) and airport labels (right) in slide coordinates
py0, ph = CY + PY * CHT, PH * CHT


def ypos(t):
    return py0 + ph * (t - TMIN) / (TMAX - TMIN)


for r in [1, 5, 10, 15, 20, 50, 100, 150]:
    add_text(s, CX - 0.1, ypos(tr(r)) - 0.12, 0.62, 0.25, [str(r)], size=11, color=MUTED, align=PP_ALIGN.RIGHT)
add_text(s, CX - 0.2, CY - 0.28, 1.3, 0.25, ['Rank'], size=11, color=MUTED)
xr = CX + (PX + PW) * CWD + 0.08
for a in airs:
    r24 = a['series'][-1][1]
    add_text(s, xr, ypos(tr(r24)) - 0.11, 0.7, 0.22, [a['label']], size=10.5,
             color=BRASS if a['riser'] else MUTED, b=a['riser'])
footer(s, 'Gold: airports outside the top 20 in 1996. Ranks beyond 20 are compressed below the dashed line.', w=8.5)
# right column: facts + countries
X2 = 9.85
add_text(s, X2, 1.95, 2.8, 0.3, [[('FROM 1996 TO 2024', dict(spc=100, b=True))]], size=10.5, color=MUTED)
rows = [['Dubai', '92 → 1'], ['Istanbul', '98 → 2'], ['Guangzhou', '155 → 10'], ['Beijing', '47 → 9']]
hairline_table(s, X2, 2.3, [1.55, 1.2], rows, row_h=0.36, size=14, header=0, align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
add_text(s, X2, 3.95, 2.8, 0.3, [[('TOP COUNTRIES, 2024', dict(spc=100, b=True))]], size=10.5, color=MUTED)
crow = [[str(c['series'][-1][1]), c['label']] for c in D['rank_bump']['countries'][:6]]
hairline_table(s, X2, 4.3, [0.45, 2.3], crow, row_h=0.33, size=14, header=0, align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT])
add_text(s, X2, 6.35, 2.9, 0.35, ['Country rank: hub quality'], size=11, color=MUTED, i=True)
notes(s, """
Here is what the index shows. These are the rank paths of the fifteen highest-ranked airports of 2024; the gold lines are airports that were outside the top twenty in 1996.
Dubai goes from 92nd to first, Istanbul from 98th to second, Guangzhou from 155th to tenth, and Shanghai Pudong, Beijing, Singapore and Incheon climb into the top fifteen. The large American and European hubs drift down.
At the country level, Singapore, the UAE, the Netherlands, Qatar and Hong Kong lead on hub quality in 2024.
This is also a face-validity check: the index reproduces what anyone in the industry would recognise, the rise of the Gulf carriers, Istanbul and the Chinese hubs.
Notice that most of this reshuffling happens after 2010. That timing will matter for identification.
""")

# ======================================================================= 13 map
s = content('01  Building the index', 'In 2024 the best-connected hubs sit in Singapore, the Gulf and Western Europe')
mw = 11.3
s.shapes.add_picture(A('map_hubquality_2024_fullnetwork.png'), Inches((SW - mw) / 2), Inches(1.75), Inches(mw))
footer(s, 'Hub quality computed from all 2024 airports (212 economies). The regressions use within-country change over time; levels are shown for orientation.')
notes(s, """
This is hub quality in 2024, computed from every airport with scheduled service. The darkest shades are the Gulf states, Western European gateway countries, Turkey and the East and Southeast Asian hubs. Most of Africa and Central Asia is light.
Levels differ a lot across countries, but our design does not use these differences. With country fixed effects, identification comes from changes within a country over time.
""")

# ======================================================================= 14 divider 02
divider('02', 'Identification', 'Airlines follow trade. How do we isolate connectivity that trade did not cause?',
        'Part two: identification.')

# ======================================================================= 15 endogeneity
s = content('02  Identification', 'Airlines add capacity where trade is already growing')


def box(s, x, y, w, h, label, fill, color=WHITE, dash=False, size=15, line=None):
    b = rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    b.adjustments[0] = 0.12
    if dash:
        ln = b.line._get_or_add_ln()
        pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', 'dash')
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    from pptx_helpers import _apply_run
    _apply_run(p.add_run(), label, {}, dict(size=size, color=color, font=BODY, b=True))
    return b


box(s, 0.9, 3.35, 2.7, 0.85, 'Air connectivity (GACI)', INK)
box(s, 5.0, 3.35, 2.4, 0.85, 'Goods trade', INK)
box(s, 2.9, 1.95, 2.6, 0.75, 'Common shocks (booms, crises)', WHITE, color=MUTED, dash=True, size=13, line=GREY)
box(s, 0.9, 5.45, 2.7, 0.95, 'Heritage × world tourism (instrument)', BRASS)
arrow(s, 3.6, 3.62, 5.0, 3.62, color=INK, w=2.25)
arrow(s, 5.0, 3.98, 3.6, 3.98, color=SLATE, w=1.75, dash='dash')
add_text(s, 3.45, 3.12, 1.7, 0.3, ['effect we want'], size=11, color=INK, align=PP_ALIGN.CENTER, b=True)
add_text(s, 3.45, 4.05, 1.7, 0.5, ['route entry follows demand'], size=11, color=SLATE, align=PP_ALIGN.CENTER, line=0.95)
arrow(s, 3.5, 2.7, 2.7, 3.35, color=GREY, w=1.25, dash='dash')
arrow(s, 4.9, 2.7, 5.9, 3.35, color=GREY, w=1.25, dash='dash')
arrow(s, 2.25, 5.45, 2.25, 4.2, color=BRASS, w=2.25)
hline(s, 3.6, 5.8, 5.6, 4.35, color=GREY, w=1.0, dash='sysDot')
add_text(s, 4.35, 5.05, 2.6, 0.5, ['×  no direct path (exclusion)'], size=11, color=MUTED, i=True)
pts_ = ['OLS mixes the effect of connectivity with demand-driven route entry and with shocks that move both.',
        'A difference-in-differences design does not fit either: large airport openings are rare and lumpy, while network position changes every year.',
        'We need variation in connectivity that goods-trade fundamentals did not create.']
add_text(s, 8.1, 2.05, 4.5, 4.4, [dict(runs=t, bullet=True, after=14) for t in pts_[:2]] +
         [dict(runs=[(pts_[2], dict(b=True, color=INK))], bullet=True)], size=17, color=TEXT, line=1.05)
notes(s, """
Why not simply regress trade on connectivity? Because airlines add capacity where trade and income are already growing, and both respond to common shocks such as a commodity boom or a financial crisis. OLS mixes the effect we want with demand-driven route entry.
A difference-in-differences design does not fit either. Large airport openings are rare and lumpy, while a country's network position changes every year through frequencies, partners and re-routing.
So we need variation in connectivity that goods-trade fundamentals did not create. That is the gold box: an instrument that moves connectivity, with no direct path to goods trade.
""")

# ======================================================================= 16 instrument
s = content('02  Identification', 'The instrument: natural heritage × the global tourism cycle')
formula(s, LM, 1.95, 8.0, 0.7, [it('Z'), sub('ct'), n(' = '), it('D̃'), sub('t'), n('  ×  ln(1 + '), it('H'), sub('c'), sup('nat+mix'), n(')')], size=30, color=INK)
add_text(s, LM + 1.0, 2.72, 2.8, 0.5, ['global tourism shift'], size=13, color=BRASS, b=True)
add_text(s, LM + 3.35, 2.72, 3.6, 0.5, ['fixed heritage share'], size=13, color=BRASS, b=True)
ts = D['tour_shift']
cd = CategoryChartData()
ky = sorted(ts, key=int)
cd.categories = ky
cd.add_series('Tourism shift', [ts[k] for k in ky])
ch = chart_frame(s, XL_CHART_TYPE.LINE, cd, LM, 3.35, 6.0, 3.1)
line_style(ch.plots[0].series[0], INK, 2.75)
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0.0')
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 1.0; ch.value_axis.major_unit = 0.2
style_axis(ch.category_axis, size=12, line=True)
ch.category_axis.tick_labels.offset = 50
cax = ch.category_axis._element
for tag, val in (('c:tickLblSkip', '4'), ('c:tickMarkSkip', '4')):
    e = cax.find(qn(tag))
    if e is None:
        e = etree.SubElement(cax, qn(tag))
    e.set('val', val)
add_text(s, LM, 6.45, 6.0, 0.35, ['D̃ₜ: world international arrivals, scaled to [0, 1]'], size=11, color=MUTED, i=True)
items = [
    [('Share: ', dict(b=True, color=INK)), 'natural and mixed UNESCO World Heritage sites (reefs, mountain ranges, national parks). The median country holds one; the richest, nineteen.'],
    [('Shift: ', dict(b=True, color=INK)), 'world tourist arrivals (WDI to 2019, UNWTO for 2020–2024).'],
    [('Why it moves routes: ', dict(b=True, color=INK)), 'airlines place marginal capacity where leisure demand rises most.'],
    [('Comparisons: ', dict(b=True, color=INK)), 'each country is compared with itself over time, so its fixed heritage stock and the world cycle cancel out; only their combination is used.'],
]
add_text(s, 7.3, 2.0, 5.3, 4.6, [dict(runs=t, after=12) for t in items], size=15.5, color=TEXT, line=1.05)
footer(s, 'Related designs: Arezki et al. (2009) use World Heritage listings for tourism; Faber and Gaubert (2019) use coastal attractiveness.')
notes(s, """
Our instrument is a shift-share design.
The share is each country's stock of natural and mixed UNESCO World Heritage sites: reefs, mountain ranges, national parks. These are endowments of geography rather than of economic activity. The median country has one; the richest has nineteen.
The shift is the global cycle in international tourist arrivals, scaled between zero and one. You can see the long rise and the collapse in 2020.
The logic is airline route allocation. When world tourism grows, airlines put marginal aircraft where leisure demand grows most, which is in heritage-rich destinations. Country fixed effects absorb the endowment, year fixed effects absorb the cycle, so only their interaction identifies the effect.
Why natural and mixed sites, and not all heritage? Cultural sites tend to sit in old urban and manufacturing centres, which could affect trade directly. I will come back to this in two slides, because the data themselves reject cultural sites as an instrument.
""")

# ======================================================================= 17 first stage
s = content('02  Identification', 'Strength test passed: heritage-rich countries gained connections as world tourism grew')
fs = D['firststage']
xy = XyChartData()
s1 = xy.add_series('Binned means')
for bx, by in zip(fs['bx'], fs['by']):
    s1.add_data_point(bx, by)
s2 = xy.add_series('Fit')
for xv in (fs['zmin'], fs['zmax']):
    s2.add_data_point(xv, fs['slope'] * xv)
ch = chart_frame(s, XL_CHART_TYPE.XY_SCATTER, xy, LM, 1.95, 7.1, 4.45)
p1, p2 = ch.plots[0].series
no_line(p1)
marker(p1, XL_MARKER_STYLE.CIRCLE, 9, INK)
line_style(p2, BRASS, 2.5)
p2.marker.style = XL_MARKER_STYLE.NONE
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0.00', title='Hub quality (change within country)')
ch.value_axis.minimum_scale = -0.04; ch.value_axis.maximum_scale = 0.03; ch.value_axis.major_unit = 0.01
style_axis(ch.category_axis, size=12, line=True, numfmt='0.0', title='Heritage × tourism instrument (change within country)')
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
ch.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
add_text(s, LM, 6.4, 7.1, 0.35, ['Each dot averages about 240 country-years. Country and year effects and population are removed.'], size=11, color=MUTED, i=True)
rows = [['Effect of instrument on hub quality', [('0.034', dict(b=True)), '***']], ['', '(0.007)'],
        ['Strength test F (rule of thumb: above 10)', [('21.9', dict(b=True))]], ['Country-years', '4,814'],
        ['Strength before 2010', '0.3'], ['Strength from 2010', '31.6']]
hairline_table(s, 8.2, 2.05, [3.1, 1.3], rows, row_h=0.4, size=15, header=0, heavy_after=(3,),
               align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
add_text(s, 8.2, 4.75, 4.4, 1.5,
         ['The instrument only has bite from 2010 onward, when low-cost long-haul flights and Gulf and Asian hubs took off. Our answers describe this recent era.'],
         size=15, color=TEXT, line=1.05)
notes(s, """
Does the instrument actually move connectivity? Yes. Each dot groups about 240 country-years. Where the heritage-times-tourism signal rises, hub quality rises too.
The usual strength test gives 21.9, well above the rule-of-thumb threshold of 10.
One caveat, stated openly. The instrument has almost no bite before 2010, and a lot from 2010 onward. That is when low-cost long-haul flights and the Gulf and Asian hubs took off. So our answers describe the recent era of aviation, and I come back to this at the end.
""")

# ======================================================================= 18 exclusion
s = content('02  Identification', 'Why tourism should affect goods trade only through air links')
rowsx = [
    ('1', 'We only count goods trade', 'Tourist spending is counted as a services export, so it never enters our outcome.'),
    ('2', 'We only use natural and mixed sites', 'Cultural sites sit in old cities and could affect trade directly. When we try them, the data reject them (p = 0.008).'),
    ('3', 'We compare each country with itself over time', 'Country effects absorb how much heritage a country has; year effects absorb the world tourism cycle.'),
]
yy = 2.0
for num, head, txt in rowsx:
    add_text(s, LM, yy - 0.1, 0.8, 1.0, [num], size=44, font=HEAD, color=BRASS)
    add_text(s, LM + 0.85, yy, 6.8, 0.45, [head], size=19, color=INK, b=True)
    add_text(s, LM + 0.85, yy + 0.45, 6.8, 0.9, [txt], size=16, color=TEXT, line=1.05)
    yy += 1.45
rect(s, 8.75, 2.0, 3.88, 4.2, fill=TINT)
add_text(s, 9.0, 2.2, 3.4, 3.9,
         [dict(runs=[('What could still go wrong?', dict(b=True, color=INK))], after=10, size=17),
          dict(runs=['A tourism boom can strengthen the currency and hurt exports. That would make our estimate smaller.'], after=10),
          dict(runs=['Tourism raises income, and income raises imports. Holding GDP fixed, the estimate stays at 1.34–1.54.'], after=10),
          dict(runs=['No test can prove the assumption. These checks make a violation hard to explain.'], i=True, color=MUTED)],
         size=15, color=TEXT, line=1.05)
notes(s, """
The key assumption is that heritage tourism affects goods trade only through air links. Three choices make this plausible.
First, we only count goods trade. Tourist spending is counted as a services export, so it never enters our outcome.
Second, we only use natural and mixed sites. Cultural sites sit in old cities and could affect trade directly. When we try them anyway, the data reject them.
Third, we compare each country with itself over time, so the amount of heritage a country has, and the world tourism cycle, both cancel out.
What could still go wrong? A tourism boom can strengthen the currency and hurt exports, but that would make our estimate smaller. Tourism raises income and income raises imports, but holding GDP fixed the estimate hardly moves. No test can prove the assumption; these checks make a violation hard to explain.
""")

# ======================================================================= 19 divider 03
divider('03', 'Results', 'How much does connectivity move trade, through which goods, and for whom?',
        'Part three: results.')

# ======================================================================= 20 main
s = content('03  Results', '10% better air links raise trade relative to GDP by about 12%')
mut = dict(i=True, color=MUTED, b=False)
main = [
    ['', 'Trade ÷ GDP', 'Total trade', 'GDP'],
    ['', [('(openness)', mut)], [('(volume)', mut)], [('(size)', mut)]],
    ['Simple correlation (OLS)', '0.167***', '1.141***', '0.975***'],
    [[('Causal estimate (IV)', dict(b=True))], [('1.208', dict(b=True)), '**'], [('2.126', dict(b=True)), '***'], [('0.919', dict(b=True))]],
    ['', '(0.593)', '(0.703)', '(0.624)'],
    [[('10% better connections', dict(b=True, color=INK))], [('+12%', dict(b=True, color=BRASS))],
     [('+22%', dict(b=True, color=BRASS))], [('+9%, not sig.', dict(color=MUTED))]],
    [[('Other measures (IV)', dict(i=True, color=MUTED))], '', '', ''],
    ['Best airport only', '0.897**', '1.580***', '0.683'],
    ['All airports summed', '0.635*', '1.119***', '0.483'],
]
hairline_table(s, LM, 1.95, [2.95, 1.45, 1.45, 1.45], main,
               row_h=[0.38, 0.3, 0.4, 0.4, 0.32, 0.46, 0.4, 0.37, 0.37], size=15, header=2,
               fills={(3, None): TINT, (4, None): TINT, (5, None): TINT}, heavy_after=(5,))
add_text(s, LM, 5.6, 7.3, 1.0,
         ['Numbers are elasticities: the % change in each outcome for a 1% rise in hub quality (the seat-weighted connectivity of a country’s airports). '
          'Robust SE in parentheses; * p<0.10, ** p<0.05, *** p<0.01. N = 4,814; strength test F = 21.9.'],
         size=11.5, color=MUTED, line=1.05)
add_text(s, 8.55, 1.95, 4.1, 0.3, [[('TWO PARTS OF THE TOTAL EFFECT', dict(spc=80, b=True))]], size=10.5, color=MUTED)
cd = CategoryChartData()
cd.categories = ['IV']
cd.add_series('More open', [1.208])
cd.add_series('Bigger economy', [0.919])
ch = chart_frame(s, XL_CHART_TYPE.BAR_STACKED, cd, 8.45, 2.35, 4.25, 1.45)
gap_width(ch, 30, overlap=100)
for ser, col in zip(ch.plots[0].series, (INK, SAGE)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
point_label(ch.plots[0].series[0], 0, 'More open 1.21', size=13, color=WHITE, pos=XL_LABEL_POSITION.CENTER, b=True)
point_label(ch.plots[0].series[1], 0, 'Bigger economy 0.92', size=11, color=WHITE, pos=XL_LABEL_POSITION.CENTER, b=True)
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 2.13
ch.category_axis.visible = False
plot_layout(ch, 0.0, 0.05, 1.0, 0.9)
add_text(s, 8.55, 3.8, 4.1, 0.4, [[('= Total trade 2.13', dict(b=True, color=INK))]], size=15, align=PP_ALIGN.RIGHT)
add_text(s, 8.55, 4.45, 4.1, 2.0,
         [dict(runs=[('The headline is the first part: ', dict(b=True, color=INK)), 'better-connected countries trade more for their size.'], after=10),
          dict(runs=['The GDP part is positive but not statistically clear.'], color=MUTED)],
         size=16, color=TEXT, line=1.05)
notes(s, """
Here is the main result. The first row is the simple correlation. The second row is our causal estimate, using the instrument.
The easiest way to read it: if a country's air connections improve by 10 percent, its trade relative to GDP rises by about 12 percent, and its total trade by about 22 percent.
The bar on the right splits the total into two parts: the country becomes more open, and the economy becomes bigger. The bigger-economy part is positive, but not statistically clear. So our headline is the first part: better-connected countries trade more for their size.
The last two rows measure connectivity in other ways, using only the best airport, or summing all airports. The numbers are on different scales, but the story is the same: positive and significant.
""")

# ======================================================================= 21 robustness
s = content('03  Results', 'The answer holds when we change the controls or the instrument')
cats = ['No\ncontrols', 'Population\n(main)', 'GDP', 'Population\n+ GDP', 'Population\n+ GDP per head']
b = [1.516, 1.208, 1.540, 1.505, 1.340]
se = [0.675, 0.593, 0.601, 0.609, 0.645]
cd = CategoryChartData()
cd.categories = cats
cd.add_series('Effect on trade / GDP', b)
ch = chart_frame(s, XL_CHART_TYPE.LINE_MARKERS, cd, LM, 2.25, 6.1, 3.9)
ser = ch.plots[0].series[0]
no_line(ser)
marker(ser, XL_MARKER_STYLE.CIRCLE, 11, INK)
pt = ser.points[1]
pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = BRASS
pt.marker.format.line.color.rgb = BRASS
add_errbars(ser, [1.96 * v for v in se], color=INK, w=1.5)
for k, v in enumerate(b):
    point_label(ser, k, f'{v:.2f}', size=12, color=TEXT, pos=XL_LABEL_POSITION.RIGHT)
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0.0')
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 3.0; ch.value_axis.major_unit = 0.5
style_axis(ch.category_axis, size=11.5, color=TEXT, line=True)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
add_text(s, LM, 1.95, 6.1, 0.3, [[('EFFECT ON TRADE ÷ GDP, WITH DIFFERENT CONTROLS', dict(spc=80, b=True))]], size=10.5, color=MUTED)
add_text(s, LM, 6.2, 6.1, 0.5, ['Whiskers: 95% confidence intervals. Adding GDP is a tough test, since GDP itself responds to connectivity.'], size=11, color=MUTED, i=True)
add_text(s, 7.25, 1.95, 5.4, 0.3, [[('CHECKS ON THE INSTRUMENT (OUTCOME: TOTAL TRADE)', dict(spc=80, b=True))]], size=10.5, color=MUTED)
val = [['Check', 'What we find'],
       ['Split heritage into natural and mixed sites', [('Same answer', dict(b=True)), ' (1.90); the two agree (p = 0.73)']],
       ['Add cultural heritage, which should fail', [('Fails, as expected', dict(b=True)), ' (p = 0.008)']],
       ['A completely different instrument: air vs sea distance', [('Same answer', dict(b=True)), ' (2.30); both together 2.26']],
       ['Let the instrument affect trade directly a little', 'Holds unless that direct effect is over a third (36%) of its total effect; trade ÷ GDP: 4%']]
hairline_table(s, 7.25, 2.3, [2.45, 2.95], val, row_h=[0.4, 0.75, 0.7, 0.85, 0.95], size=13.5,
               align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT])
notes(s, """
Is this result fragile? We tried two kinds of checks.
On the left, we change what else we control for. The effect on trade relative to GDP stays between 1.21 and 1.54 every time. Adding GDP is a tough test, because GDP itself responds to connectivity.
On the right, we check the instrument. Splitting heritage into natural and mixed sites gives the same answer. Cultural heritage, which we expect to fail, does fail, which tells us the test has teeth. A completely different instrument, built from geography, air versus sea distance, gives almost the same number: 2.30, against our 2.13 for total trade.
Finally, if we allow the instrument to affect trade directly a little, the total-trade result survives unless that direct effect is more than a third of its total effect. For trade relative to GDP the margin is thinner, about 4 percent.
""")

# ======================================================================= 22 mechanism 1
s = content('03  Results  ·  Mechanism', 'What changes: better-connected countries trade more parts and more high-value goods')
add_text(s, LM, 2.0, 5.2, 0.3, [[('LOOKING INSIDE TRADE: PRODUCT-LEVEL DATA (BACI)', dict(spc=80, b=True))]], size=10.5, color=MUTED)
add_text(s, LM, 2.35, 5.2, 3.9, [
    dict(runs='About 5,000 products, about 200 countries, exports plus imports', bullet=True, after=10),
    dict(runs=[('Parts ratio: ', dict(b=True, color=INK)), 'intermediate inputs ÷ final consumer goods'], bullet=True, after=10),
    dict(runs=[('Value ratio: ', dict(b=True, color=INK)), 'high-value ÷ low-value goods, by dollars per kilogram'], bullet=True, after=10),
    dict(runs='Ratios cancel out anything that moves all of a country’s trade together', bullet=True, after=10),
], size=16, color=TEXT, line=1.05)
cd = CategoryChartData()
cd.categories = ['Parts ÷ consumer goods', 'High-value ÷ low-value goods']
cd.add_series('IV', [1.483, 1.868])
ch = chart_frame(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, 6.45, 1.95, 6.2, 3.95)
ser = ch.plots[0].series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = INK
gap_width(ch, 90)
add_errbars(ser, [1.96 * 0.653, 1.96 * 0.716], color=BRASS, w=1.75)
for k, t in enumerate(['1.48**', '1.87***']):
    point_label(ser, k, t, size=15, color=WHITE, pos=XL_LABEL_POSITION.INSIDE_BASE, b=True)
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0.0', title='Effect of hub quality (elasticity)')
ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = 3.5; ch.value_axis.major_unit = 0.5
style_axis(ch.category_axis, size=14, color=TEXT, line=True)
rect(s, 6.55, 5.95, 6.05, 0.55, fill=TINT)
add_text(s, 6.75, 5.95, 5.8, 0.55,
         [[('10% better connections: ', dict(b=True, color=INK)), 'parts ratio +15%, value ratio +19%']],
         size=15, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, 6.45, 6.55, 6.2, 0.3, ['IV estimates with 95% CI; N = 4,785; strength test F = 22.1'], size=11, color=MUTED, i=True)
notes(s, """
So how does connectivity raise trade? We look inside trade using product-level data: about 5,000 products for about 200 countries.
We build two simple ratios. The parts ratio: intermediate inputs relative to final consumer goods. The value ratio: high-value relative to low-value goods, measured in dollars per kilogram. These are the goods that tend to fly. Because they are ratios, anything that moves all of a country's trade up or down together cancels out.
Both ratios rise with connectivity. With 10 percent better connections, the parts ratio rises by about 15 percent and the value ratio by about 19 percent. In other words, better-connected countries trade more of the goods that fly.
""")

# ======================================================================= 23 mechanism 2
s = content('03  Results  ·  Mechanism', 'Accounting for the goods mix shrinks the effect by about 30%')
cd = CategoryChartData()
cd.categories = ['Connectivity\nonly', '+ parts\nratio', '+ value\nratio', '+ both\nratios']
vals = [1.176, 0.939, 1.144, 0.825]
ses = [0.588, 0.571, 0.599, 0.583]
cd.add_series('Effect of connectivity on trade / GDP', vals)
ch = chart_frame(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, LM, 2.25, 7.2, 4.1)
ser = ch.plots[0].series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = SAGE
p0 = ser.points[0]
p0.format.fill.solid(); p0.format.fill.fore_color.rgb = INK
gap_width(ch, 70)
add_errbars(ser, [1.96 * v for v in ses], color=TEXT, w=1.25)
for k, t in enumerate(['1.18**', '0.94', '1.14*', '0.83']):
    point_label(ser, k, t, size=14, color=WHITE, pos=XL_LABEL_POSITION.INSIDE_BASE, b=True)
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0.0')
ch.value_axis.minimum_scale = -0.5; ch.value_axis.maximum_scale = 2.5; ch.value_axis.major_unit = 0.5
style_axis(ch.category_axis, size=13, color=TEXT, line=True)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
add_text(s, LM, 1.95, 7.2, 0.3, [[('EFFECT OF CONNECTIVITY ON TRADE ÷ GDP, 95% CI', dict(spc=80, b=True))]], size=10.5, color=MUTED)
add_text(s, 8.4, 2.0, 4.25, 4.5, [
    dict(runs=[('Same countries and years in every bar ', dict(b=True, color=INK)), '(4,785), so the drop comes from the added ratios.'], after=12),
    dict(runs=[('With both ratios, ', dict(b=True, color=INK)), 'the effect falls from 1.18 to 0.83 and is no longer statistically clear.'], after=12),
    dict(runs=[('In plain words: ', dict(b=True, color=INK)), 'the trade gain comes with, and runs through, a shift toward goods that fly.'], after=12),
    dict(runs=['The mix is itself an outcome, so this supports the channel without proving it.'], i=True, color=MUTED),
], size=16, color=TEXT, line=1.05)
notes(s, """
Next: does this change in the mix explain the trade gain? We run the main regression again and add the two ratios, keeping exactly the same countries and years.
The connectivity effect drops from 1.18 to 0.83, about 30 percent smaller, and it is no longer statistically clear. Meanwhile the mix of goods itself strongly predicts how open a country is.
If connectivity simply made all trade a bit cheaper, the mix would not change, and adding the ratios would change nothing. That is not what we see. The trade gain comes with, and runs through, a shift toward goods that fly.
One caution: the mix is itself an outcome, so this supports the channel but does not prove it.
""")

# ======================================================================= 24 who gains
s = content('03  Results', 'Poorer and less-connected countries gain the most')
H = D['het_quartile']
panels = (('income', 'By income in 1996', ['Poorest\n25%', '2nd', '3rd', 'Richest\n25%'],
           [('10% better connections: ', dict(b=True, color=INK)), '+16% trade ÷ GDP for the poorest quarter, +7% for the richest']),
          ('baseconn', 'By air connectivity in 1996', ['Least\nconnected', '2nd', '3rd', 'Best\nconnected'],
           [('10% better connections: ', dict(b=True, color=INK)), '+26% for the least connected quarter, +15% for the best connected']))
for k, (key, head, labels, reading) in enumerate(panels):
    x = LM + k * 6.1
    add_text(s, x, 1.95, 5.8, 0.3, [[(head.upper(), dict(spc=60, b=True))]], size=10.5, color=MUTED)
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series('Effect', [r[0] for r in H[key]])
    ch = chart_frame(s, XL_CHART_TYPE.LINE_MARKERS, cd, x, 2.25, 5.8, 3.3)
    ser = ch.plots[0].series[0]
    no_line(ser)
    marker(ser, XL_MARKER_STYLE.CIRCLE, 12, INK if k == 0 else BRASS)
    add_errbars(ser, [r[1] for r in H[key]], color=INK if k == 0 else BRASS, w=1.5)
    for j, r in enumerate(H[key]):
        point_label(ser, j, f'{r[0]:.2f}', size=13, color=TEXT, pos=XL_LABEL_POSITION.RIGHT, b=True)
    style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0')
    ch.value_axis.minimum_scale = -1; ch.value_axis.maximum_scale = 5; ch.value_axis.major_unit = 1
    style_axis(ch.category_axis, size=12.5, color=TEXT, line=True)
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    add_text(s, x + 0.1, 5.6, 5.6, 0.6, [reading], size=13.5, color=TEXT, line=1.0)
add_text(s, LM, 6.3, CW, 0.5,
         ['Effect of hub quality on trade ÷ GDP by quarter of countries, 95% CI. These estimates are less precise (strength test F ≈ 6–10), so the slope matters more than each number. '
          'A simpler check without this weakness shows the same slope.'],
         size=11, color=MUTED, line=1.0)
notes(s, """
Who gains most? We split countries by how rich and how well connected they were in 1996.
The pattern is the same in both panels: the gain shrinks as we move to richer and better-connected countries. With 10 percent better connections, trade relative to GDP rises about 16 percent in the poorest quarter of countries, and about 7 percent in the richest. By connectivity, it is about 26 percent for the least connected quarter and 15 percent for the best connected.
These estimates are less precise, so the downward slope matters more than each number. A simpler check that avoids this weakness shows the same slope.
The message: an extra air link is worth most where links are scarce.
""")

# ======================================================================= 25 aggregate
s = content('03  Results  ·  Magnitudes', 'Better air links since 1996 add about $10 trillion to 2024 trade')
C = D['continent']
yrs = [int(r[0]) for r in C['Asia']]
cd = CategoryChartData()
cd.categories = [str(y) for y in yrs]
order = ['Africa', 'Europe', 'Latin America', 'North America', 'Pacific/Oceania', 'Middle East', 'Asia']
for c in order:
    cd.add_series(c, [r[1] for r in C[c]])
ch = chart_frame(s, XL_CHART_TYPE.LINE, cd, LM, 2.25, 7.0, 4.0)
for ser, c in zip(ch.plots[0].series, order):
    col = INK if c == 'Asia' else (BRASS if c == 'Middle East' else GREY)
    line_style(ser, col, 2.75 if c in ('Asia', 'Middle East') else 1.5)
last = len(yrs) - 1
point_label(ch.plots[0].series[order.index('Asia')], last, 'Asia', size=12, color=INK, pos=XL_LABEL_POSITION.RIGHT, b=True)
point_label(ch.plots[0].series[order.index('Middle East')], last, 'Middle East', size=12, color=BRASS, pos=XL_LABEL_POSITION.RIGHT, b=True)
point_label(ch.plots[0].series[order.index('Europe')], last, 'Other continents', size=11, color=MUTED, pos=XL_LABEL_POSITION.RIGHT)
style_axis(ch.value_axis, size=12, grid=True, line=False, numfmt='0')
ch.value_axis.minimum_scale = -40; ch.value_axis.maximum_scale = 80; ch.value_axis.major_unit = 20
style_axis(ch.category_axis, size=12, line=True)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
cax = ch.category_axis._element
for tag in ('c:tickLblSkip', 'c:tickMarkSkip'):
    e = cax.find(qn(tag))
    if e is None:
        e = etree.SubElement(cax, qn(tag))
    e.set('val', '4')
plot_layout(ch, 0.06, 0.04, 0.72, 0.84)
add_text(s, LM, 1.95, 7.0, 0.3, [[('EXTRA TRADE ÷ GDP FROM CONNECTIVITY GAINS SINCE 1996 (%)', dict(spc=80, b=True))]], size=10.5, color=MUTED)
add_text(s, LM, 6.25, 7.0, 0.45, ['Each country: 2024 trade × the share implied by its own connectivity gain since 1996 (elasticity 1.21).'],
         size=12, color=MUTED, i=True)
agg = [['Extra 2024 goods trade', [('$10.0 tn', dict(b=True, color=INK))]],
       ['Share of world goods trade', [('21%', dict(b=True, color=INK))]],
       ['Share of world GDP', [('9.2%', dict(b=True, color=INK))]],
       ['Plausible range (95% CI)', '$0.5–15.9 tn'],
       ['Using the other two measures', '$8.0–10.0 tn']]
hairline_table(s, 8.15, 2.3, [3.05, 1.45], agg, row_h=0.55, size=15, header=0, heavy_after=(2,),
               align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
add_text(s, 8.15, 5.25, 4.5, 1.2,
         ['A rough figure: it holds other countries and prices fixed, and applies an effect measured after 2010 back to 1996.'],
         size=13, color=MUTED, i=True, line=1.05)
notes(s, """
How much does this add up to? For each country we take its own improvement in connectivity since 1996 and ask how much of its 2024 trade that improvement accounts for, keeping everything else fixed.
Added up, it comes to about 10 trillion dollars: about a fifth of world goods trade, or about 9 percent of world GDP.
Please treat this as a rough number. The plausible range is wide, from half a trillion to 16 trillion. It holds other countries and prices fixed. And it applies an effect we measure after 2010 back to 1996.
The chart on the left shows where the gains go: mostly to Asia and the Middle East, which pull away from the rest after 2005.
""")

# ======================================================================= 26 contribution map
s = content('03  Results  ·  Magnitudes', 'China gains most, then Korea and Japan; 18 countries lost ground')
mw = 8.6
s.shapes.add_picture(A('map_contrib_trade_2024.png'), Inches(0.45), Inches(1.95), Inches(mw))
top = D['contrib_top10'][:8]
names = {'CHN': 'China', 'KOR': 'Korea', 'JPN': 'Japan', 'USA': 'United States', 'HKG': 'Hong Kong',
         'SGP': 'Singapore', 'VNM': 'Viet Nam', 'TUR': 'Türkiye', 'CAN': 'Canada', 'IND': 'India'}
cd = CategoryChartData()
cd.categories = [names[c] for c, _ in reversed(top)]
cd.add_series('US$ bn', [v for _, v in reversed(top)])
ch = chart_frame(s, XL_CHART_TYPE.BAR_CLUSTERED, cd, 9.2, 2.2, 3.5, 3.9)
ser = ch.plots[0].series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = BRASS
gap_width(ch, 45)
for k, (_, v) in enumerate(reversed(top)):
    point_label(ser, k, f'{v:,.0f}', size=11, color=TEXT, pos=XL_LABEL_POSITION.OUTSIDE_END)
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 3600
style_axis(ch.category_axis, size=12, color=TEXT, line=False)
add_text(s, 9.2, 1.95, 3.5, 0.3, [[('EXTRA 2024 TRADE, US$ BILLION', dict(spc=80, b=True))]], size=10.5, color=MUTED)
add_text(s, LM, 6.05, 11.9, 0.8,
         ['Biggest losses, where connectivity fell: Germany (−$202 bn), Ukraine (−$119 bn), Belgium (−$69 bn). Grey: no 2024 trade data (31 economies, incl. UAE and Qatar), counted as zero.'],
         size=12, color=MUTED, line=1.05)
notes(s, """
Country by country, China gains the most, about 2.7 trillion dollars, followed by Korea, Japan, the United States, Hong Kong and Singapore.
Not everyone gained. In 18 countries connectivity fell, and the biggest implied loss is Germany, about 200 billion dollars, then Ukraine and Belgium.
One data note: countries without 2024 trade data, including the UAE and Qatar, are grey and count as zero.
""")

# ======================================================================= 27 covid
s = content('03  Results  ·  Magnitudes', 'The same logic in reverse: the 2020 collapse')
s.shapes.add_picture(A('map_covid_2019_2020.png'), Inches(0.45), Inches(1.95), Inches(8.6))
add_text(s, 9.3, 2.0, 3.35, 0.3, [[('ROUGH ONE-YEAR TRADE LOSS', dict(spc=80, b=True))]], size=10.5, color=MUTED)
cv = [['Main measure', [('$7.1 tn', dict(b=True, color=INK))]], ['Best airport only', '$6.3 tn'], ['All airports summed', '$1.7 tn']]
hairline_table(s, 9.3, 2.35, [1.95, 1.4], cv, row_h=0.45, size=15, header=0, align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT])
add_text(s, 9.3, 3.9, 3.35, 2.6, [
    dict(runs=['Connections fell in 154 of 175 countries. Biggest losses: United States, Germany, Hong Kong, Korea, Japan.'], after=10),
    dict(runs=[('Treat as an upper range. ', dict(b=True, color=INK)),
               'In 2020 the pandemic cut trade and flights at the same time, so not all of this loss is due to fewer flights.'], after=0),
], size=14, color=TEXT, line=1.05)
footer(s, 'Main measure (hub quality), 2019 to 2020. This assumes the effect works the same way when connectivity falls as when it rises.')
notes(s, """
The same logic should work in reverse. In 2020, connections fell in 154 of 175 countries.
Running the same calculation backwards gives a one-year trade loss of about 7 trillion dollars on our main measure, and less on the others, down to 1.7 trillion.
Please read this as an upper range. Our effect is a long-run effect, and in 2020 the pandemic cut trade and flights at the same time, so not all of this loss is due to fewer flights.
""")

# ======================================================================= 28 conclusion
s = content('Conclusion', 'Better air links raise trade, through goods that fly, most where links are scarce')
finds = [
    ('1', 'Air links cause more trade', '10% better connections: about 12% more trade relative to GDP.'),
    ('2', 'It works through what is traded', 'More parts and more high-value goods, the goods that fly.'),
    ('3', 'Poorer, less-connected countries gain most', 'The payoff falls as income and connectivity rise.'),
    ('4', 'The total is large', 'Roughly $10 tn of 2024 trade since 1996 (a rough estimate).'),
]
yy = 2.05
for num, head, txt in finds:
    add_text(s, LM, yy - 0.1, 0.7, 0.8, [num], size=34, font=HEAD, color=BRASS)
    add_text(s, LM + 0.7, yy, 6.5, 0.4, [head], size=18, color=INK, b=True)
    add_text(s, LM + 0.7, yy + 0.4, 6.5, 0.5, [txt], size=15, color=TEXT)
    yy += 1.05
X3 = 8.15
add_text(s, X3, 2.0, 4.5, 0.3, [[('POLICY', dict(spc=100, b=True))]], size=10.5, color=BRASS)
add_text(s, X3, 2.3, 4.5, 1.2, ['Investing in airports and routes, and opening air markets, are practical ways to raise trade, especially where connections are poor.'],
         size=15, color=TEXT, line=1.05)
hline(s, X3, 3.6, 12.63, 3.6, color=RULE)
add_text(s, X3, 3.75, 4.5, 0.3, [[('LIMITS', dict(spc=100, b=True))]], size=10.5, color=MUTED)
add_text(s, X3, 4.05, 4.5, 1.3, [dict(runs=t, bullet=True, after=3) for t in
                                  ['Totals are rough, one country at a time', 'Evidence comes mainly from 2010 onward', 'Rich-versus-poor comparison is less precise']],
         size=14, color=TEXT)
hline(s, X3, 5.2, 12.63, 5.2, color=RULE)
add_text(s, X3, 5.35, 4.5, 0.3, [[('NEXT', dict(spc=100, b=True))]], size=10.5, color=MUTED)
add_text(s, X3, 5.65, 4.5, 1.0, [dict(runs=t, bullet=True, after=3) for t in
                                  ['Trade between country pairs (gravity)', 'A full model where prices and partners adjust']],
         size=14, color=TEXT)
notes(s, """
Let me sum up in four points.
One: better air links cause more trade. Ten percent better connections mean about 12 percent more trade relative to GDP.
Two: they work through what is traded, more parts and more high-value goods, the goods that fly.
Three: poorer and less-connected countries gain the most.
Four: the total is large, roughly 10 trillion dollars of 2024 trade since 1996.
For policy, investing in airports and routes, and opening air markets, are practical ways to raise trade, especially where connections are poor.
The limits: the totals are rough, the evidence comes mainly from 2010 onward, and the rich-versus-poor comparison is less precise. Next, we want to look at trade between country pairs, and at a full model where prices and partners adjust. Thank you.
""")

# ======================================================================= 29 thanks
s = prs.slides.add_slide(BL)
bg(s, INK)
faint_map(s, 55)
add_text(s, 0.9, 2.3, 8, 1.2, ['Thank you'], size=54, font=HEAD, color=WHITE)
add_text(s, 0.9, 3.55, 7.5, 0.6, ['Global Air Connectivity and Trade Openness, 1996–2024'], size=19, color=ONDARK, i=True)
add_text(s, 0.9, 4.45, 8.2, 0.8, ['Sunbin YOO, Junya KUMAGAI, Longfei ZHENG, Chunan WANG, Jinwoo LEE, Anming ZHANG'], size=14, color=WHITE, line=1.1)
add_text(s, 0.9, 5.5, 7.0, 0.4, ['Comments and questions are very welcome.'], size=15, color=BRASS)
TITLES[len(prs.slides)] = 'Thank you'
notes(s, 'Thank you. I look forward to your comments and questions.')

prs.save(OUT)
json.dump(NOTES, open(os.path.join(HERE, 'notes.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('saved', OUT, len(prs.slides), 'slides; notes words =', sum(len(t['notes'].split()) for t in NOTES))
