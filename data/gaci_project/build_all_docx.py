# -*- coding: utf-8 -*-
"""Combine main (incl GDP) + heterogeneity + COVID/Russia shocks into one Word file."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.section import WD_ORIENT

def star(p):
    try: p = float(p)
    except Exception: return ''
    return '***' if p < .01 else '**' if p < .05 else '*' if p < .1 else ''

OUT = ['g_vol', 'g_int', 'g_shr', 'lnpc', 'lngdp']
OLAB = {'g_vol': 'Ln goods\nvolume', 'g_int': 'Goods openness\nln(goods/GDP)', 'g_shr': 'Goods share\n(% GDP)',
        'lnpc': 'Ln GDP\nper capita', 'lngdp': 'Ln GDP\n(total)'}
DEC = {'g_vol': 3, 'g_int': 3, 'g_shr': 2, 'lnpc': 3, 'lngdp': 3}

main, mF, mN = {}, {}, {}
for ln in open('gaci_tourism_main.log', encoding='utf-8', errors='replace'):
    ln = ln.strip()
    if ln.startswith('MAINP|'):
        _, tr, yv, b, sr, pr, sc, pc, F, ob, op, N = ln.split('|')
        main[(tr, yv)] = dict(b=float(b), sr=float(sr), pr=float(pr), sc=float(sc), pc=float(pc), ob=float(ob), op=float(op))
        mF[tr] = float(F); mN[tr] = int(float(N))

HOUT = ['g_vol', 'g_int', 'g_shr', 'lnpc', 'lngdp']
het, hF = {}, {}
for ln in open('gaci_tourism_hetero2.log', encoding='utf-8', errors='replace'):
    ln = ln.strip()
    if ln.startswith('RES2|'):
        p = ln.split('|')
        if p[3] == 'FAIL': continue
        het[(p[2], p[1])] = (float(p[3]), float(p[4]), float(p[5])); hF[p[2]] = float(p[6])
GLAB = {'all': 'All (pooled)', 'gaci_below': '1996 GACI < median', 'gaci_above': '1996 GACI > median',
        'gdp_below': '1996 GDP < median', 'gdp_above': '1996 GDP > median', 'region_AF': 'Africa', 'region_AS': 'Asia',
        'region_EU': 'Europe', 'region_LA': 'Latin America', 'region_ME': 'Middle East', 'region_NA': 'North America', 'region_SW': 'Pacific/SW'}
GORD = ['all', 'gaci_below', 'gaci_above', 'gdp_below', 'gdp_above', 'region_AF', 'region_AS', 'region_EU', 'region_LA', 'region_ME', 'region_NA', 'region_SW']

sh = {}
for ln in open('gaci_shock_export2.log', encoding='utf-8', errors='replace'):
    ln = ln.strip()
    if ln.startswith('SHOCK2|'):
        _, yv, v, b, se, p, F, N = ln.split('|')
        sh[(yv, v)] = dict(b=float(b), se=float(se), p=float(p), F=F, N=int(float(N)))

doc = Document()
s = doc.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Inches(11), Inches(8.5)
doc.styles['Normal'].font.name = 'Times New Roman'; doc.styles['Normal'].font.size = Pt(9)

def ital(t, sz=8):
    pp = doc.add_paragraph(); r = pp.add_run(t); r.italic = True; r.font.size = Pt(sz)

def bold_hdr(cells):
    for c in cells:
        for pa in c.paragraphs:
            for r in pa.runs: r.bold = True

doc.add_heading('Air Connectivity, Goods Trade, and Income: Tourism-Heritage IV', level=1)
ital('Two-stage least squares. Air connectivity instrumented by tourism_int = (world tourism demand, t) x '
     'ln(1 + natural/mixed UNESCO heritage endowment). Country and year fixed effects. Goods = merchandise trade '
     '(services/tourism excluded). Robust SE in (parentheses), country-clustered SE in [brackets]. '
     '* p<0.10, ** p<0.05, *** p<0.01 by robust SE; "+" significant at 5% under clustering.')

doc.add_heading('Table 1. Main results: OLS vs. IV (sample 1996-2019)', level=2)
t = doc.add_table(rows=0, cols=6); t.style = 'Table Grid'
h = t.add_row().cells; h[0].text = ''
for j, o in enumerate(OUT): h[j + 1].text = OLAB[o]
bold_hdr(h)

def panel(tr, title):
    rr = t.add_row().cells; rr[0].merge(rr[5]); run = rr[0].paragraphs[0].add_run(title); run.bold = True
    for lab, kind in [('Air connectivity (IV)', 'iv'), ('   Robust SE', 'sr'), ('   Clustered SE', 'sc'),
                      ('OLS benchmark', 'ols'), ('First-stage KP F', 'F'), ('Observations', 'N')]:
        r = t.add_row().cells; r[0].text = lab
        for j, o in enumerate(OUT):
            dd = main.get((tr, o)); dec = DEC[o]
            if dd is None: continue
            if kind == 'iv': r[j + 1].text = f"{dd['b']:,.{dec}f}{star(dd['pr'])}"
            elif kind == 'sr': r[j + 1].text = f"({dd['sr']:,.{dec}f})"
            elif kind == 'sc': r[j + 1].text = f"[{dd['sc']:,.{dec}f}]" + ('+' if dd['pc'] < .05 else '')
            elif kind == 'ols': r[j + 1].text = f"{dd['ob']:,.{dec}f}{star(dd['op'])}"
            elif kind == 'F': r[j + 1].text = f"{mF[tr]:.1f}"
            elif kind == 'N': r[j + 1].text = f"{mN[tr]:,}"

panel('lng', 'Panel A. Treatment = total connectivity (log GACI sum)')
panel('ln_gaci_cwm', 'Panel B. Treatment = hub quality (log capacity-weighted-mean GACI)')

doc.add_heading('Table 2. Heterogeneity by 1996 GACI, 1996 GDP, and region (sample 1996-2019)', level=2)
ncol2 = len(HOUT) + 2
t2 = doc.add_table(rows=0, cols=ncol2); t2.style = 'Table Grid'
h2 = t2.add_row().cells; h2[0].text = 'Subsample'
for j, o in enumerate(HOUT): h2[j + 1].text = OLAB[o]
h2[ncol2 - 1].text = 'First-stage\nKP F'; bold_hdr(h2)
for g in GORD:
    if g not in hF: continue
    r = t2.add_row().cells; run = r[0].paragraphs[0].add_run(GLAB[g]); run.bold = (g == 'all')
    for j, o in enumerate(HOUT):
        if (g, o) in het:
            b, se, p = het[(g, o)]; r[j + 1].text = f"{b:,.{DEC[o]}f}{star(p)}\n({se:,.{DEC[o]}f})"
    fv = hF[g]; r[ncol2 - 1].text = f"{fv:.1f}" + (" (weak)" if fv < 10 else "")
ital('Treatment = log connectivity (GACI sum). Cells with first-stage F < 10 are weakly identified.')

doc.add_heading('Table 3. COVID and Russia-war shock interactions, all outcomes (sample 1996-2023)', level=2)
ital('IV (2SLS) with both shocks: connectivity, connectivity x COVID, and connectivity x Russia are jointly '
     'instrumented by tourism_int and its interactions. Tourism demand extended to 2023 (UNWTO recovery ratios; '
     'the 2020 COVID collapse enters the instrument). COVID = 2020-2021, Russia = 2022-2023; year FE absorbs the '
     'main shock dummies. Robust SE in parentheses.')
ncol3 = len(OUT) + 1
t3 = doc.add_table(rows=0, cols=ncol3); t3.style = 'Table Grid'
h3 = t3.add_row().cells; h3[0].text = ''
for j, o in enumerate(OUT): h3[j + 1].text = OLAB[o]
bold_hdr(h3)
for term, tl in [('lng', 'GACI (baseline)'), ('gaci_c', 'GACI x COVID'), ('gaci_r', 'GACI x Russia')]:
    r = t3.add_row().cells; r[0].text = tl
    for j, o in enumerate(OUT):
        d = sh.get((o, term))
        if d is None: continue
        r[j + 1].text = f"{d['b']:,.{DEC[o]}f}{star(d['p'])}\n({d['se']:,.{DEC[o]}f})"
rF = t3.add_row().cells; rF[0].text = 'First-stage KP F'
rN = t3.add_row().cells; rN[0].text = 'Observations'
for j, o in enumerate(OUT):
    a = sh.get((o, 'lng'))
    rF[j + 1].text = f"{float(a['F']):.1f}" if a else ''
    rN[j + 1].text = f"{a['N']:,}" if a else ''
ital('All GACI-shock interactions are statistically insignificant across every outcome: the connectivity-trade '
     '(and connectivity-income) relationship does not detectably change during COVID or the Russia war. Note the '
     'shock-window identification is limited because the tourism instrument collapses to ~0 during the 2020 tourism shutdown.')

doc.save('GACI_tourism_results.docx')
print('wrote GACI_tourism_results.docx (main=%d, hetero=%d, shock=%d)' % (len(main), len(hF), len(sh)))
