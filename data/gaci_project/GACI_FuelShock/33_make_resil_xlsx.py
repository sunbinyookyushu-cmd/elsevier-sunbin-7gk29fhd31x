# -*- coding: utf-8 -*-
"""Build GACI_FuelShock_Resilience_20260929.xlsx: resilience index (Zhang, Cheung & Zhang 2027 method)
x jet-fuel price shocks. Text sheets come from _xlsx_text_resil.py."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, ScatterChart, Reference, Series
from _xlsx_helpers import st, W, rule, fc, fs, fn, ff, notes, title, widths, style_chart
from _xlsx_text_resil import README, DECISIONS, DESIGN

OUTF = "GACI_FuelShock_Resilience_20260929.xlsx"
rs = pd.read_csv("_res_resil.csv")
rv = pd.read_csv("_res_resil_valid.csv")
re_ = pd.read_csv("_res_resil_events.csv")
res_ = pd.read_csv("_res_resil_es.csv")
rep = pd.read_csv("_res_resil_replication.csv")
rob = pd.read_csv("_res_resil_events_rob.csv")
shk = pd.read_csv("_res_resil_shrink.csv")
esr = pd.read_csv("_res_resil_es_rob.csv")
iu = pd.read_csv("_res_resil_invU.csv")
Ridx = pd.read_csv("resilience_airport.csv")
EST = {"OLS": "OLS", "2SLS-KZ": "2SLS (Känzig)", "2SLS-BH": "2SLS (BH)"}
LAB = {"d12_ln_seats": "Δ12 ln seats", "d12_ln_flights": "Δ12 ln departures", "d12_ln_skm": "Δ12 ln seat-km", "d12_ln_co2": "Δ12 ln CO2"}
NOTE_DK = ("SE in parentheses: country cluster + Newey-West over months (L = 12; Thompson 2011), floored at the larger one-way "
           "component; stars (* p<0.1, ** p<0.05, *** p<0.01) refer to these. Brackets: country cluster only.")

def pick(df, **kw):
    m = df
    for k, v in kw.items():
        m = m[m[k] == v]
    return m.iloc[0] if len(m) else None

wb = Workbook()
wb.remove(wb.active)

def text_sheet(name, blocks):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 120
    for r, (kind, txt) in enumerate(blocks, start=1):
        c = ws.cell(row=r, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="맑은 고딕", size=14 if kind == "h1" else 11 if kind == "h2" else 10, bold=kind in ("h1", "h2"))
        if kind == "p":
            ws.row_dimensions[r].height = max(15, 15 * (len(txt) // 95 + txt.count("\n") + 1))

text_sheet("README", README)
text_sheet("결정필요_표본정의", DECISIONS)
text_sheet("연구설계", DESIGN)

# ---- R1 replication ----
ws = wb.create_sheet("R1_재현검증")
ncol = 13
r = title(ws, 1, ncol, "Table R1. Replication of the resilience index against Zhang, Cheung & Zhang (2027, TR-E 217, 105192)")
rule(ws, r, 1, ncol, "top")
hdr = ["Group", "N", "Mean", "SD", "P10", "P50", "P90", "N (paper)", "Mean (paper)", "SD (paper)", "P10 (paper)", "P50 (paper)", "P90 (paper)"]
for j, h in enumerate(hdr):
    W(ws, r, j + 1, h, wrap=True)
rule(ws, r, 1, ncol, "bottom")
r += 1
W(ws, r, 1, "Panel A. Table 5: reliability-adjusted resilience index", italic=True, align="left")
r += 1
for x in rep[rep.section == "Table 5"].itertuples():
    vals = [x.group, fn(x.N), f"{x.mean:.3f}", f"{x.sd:.3f}", f"{x.p10:.3f}", f"{x.p50:.3f}", f"{x.p90:.3f}", fn(x.N_paper),
            f"{x.mean_paper:.3f}", f"{x.sd_paper:.3f}", f"{x.p10_paper:.3f}", f"{x.p50_paper:.3f}", f"{x.p90_paper:.3f}"]
    for j, v in enumerate(vals):
        W(ws, r, j + 1, v, align="left" if j == 0 else "center")
    r += 1
W(ws, r, 1, "Panel B. Fig. 10: R on GACI 2024 and its square (airports with ≥ 5 observed years)", italic=True, align="left")
r += 1
for x in rep[rep.section != "Table 5"].itertuples():
    W(ws, r, 1, x.group, align="left")
    W(ws, r, 2, fn(x.N))
    W(ws, r, 3, f"{x.mean:.3f}")
    W(ws, r, 4, fs(x.sd, 3))
    W(ws, r, 8, fn(x.N_paper))
    W(ws, r, 9, "" if pd.isna(x.mean_paper) else f"{x.mean_paper:.2f}")
    r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 24, 10, ncol)
notes(ws, r, ncol, "Index: 8 crisis windows of the paper's Table 3; Depth, Speed (H = 6, τ = 2), Adaptability (Theil-Sen rank-percentile slope "
      "and k = 5 post-crisis upgrade), equal weights, empirical-Bayes shrinkage toward 0.5 with λ = 2 for airports with 1-2 valid crises. "
      "Unstated choices fixed here: the post-crisis upgrade is min-max scaled to [0,1]; airports absent in a crisis year are not imputed. "
      "Panel B: coefficient (HC1 SE); peak = −a1/(2 a2) with delta-method SE. Paper: peak ≈ 1.85 (range 1.81-1.89 across its robustness checks).")

# ---- R2 main ----
ws = wb.create_sheet("R2_본분석")
outs = ["d12_ln_seats", "d12_ln_flights", "d12_ln_skm", "d12_ln_co2"]
ests = ["OLS", "2SLS-KZ", "2SLS-BH"]
ncol = 1 + len(outs) * len(ests)
r = title(ws, 1, ncol, "Table R2. Resilience and the response to jet fuel price shocks: airport × month (12-month differences)")
rule(ws, r, 1, ncol, "top")
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
W(ws, r, 2, "Dependent variable")
r += 1
for j, o in enumerate(outs):
    c0 = 2 + 3 * j
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
    W(ws, r, c0, LAB[o])
    rule(ws, r, c0, c0 + 2, "bottom")
r += 1
for j in range(len(outs)):
    for k, e in enumerate(ests):
        W(ws, r, 2 + 3 * j + k, EST[e])
r += 1
for c in range(2, ncol + 1):
    W(ws, r, c, f"({c - 1})")
rule(ws, r, 1, ncol, "bottom")
r += 1
PL = {"pre2004": "Panel A. Resilience built from 1996–2003 (AFC, 9/11, SARS); test 2004–2019",
      "pre2008": "Panel B. Resilience built from 1996–2007 (AFC, 9/11, SARS); test 2008–2019",
      "realtime": "Panel C. Real-time resilience R(i, y−1), rebuilt every year; test 2004–2019",
      "full": "Panel D. Full-sample resilience (8 crises, 1996–2024), overlaps the fuel episodes; test 1997–2019"}
R2 = rs[(rs.table == "RS1")]
for v, pl in PL.items():
    W(ws, r, 1, pl, italic=True, align="left")
    r += 1
    W(ws, r, 1, "Δ ln P(jet) × resilience (z)", align="left")
    for j, o in enumerate(outs):
        for k, e in enumerate(ests):
            x = pick(R2, panel=v, outcome=o, col=e, term="x")
            c = 2 + 3 * j + k
            W(ws, r, c, fc(x.b, x.p))
            W(ws, r + 1, c, fs(x.se))
            W(ws, r + 2, c, fs(x.se_cl, br="[]"), size=9)
    r += 3
    for lab, f_ in [("Observations", lambda x, e: fn(x.n)), ("Airports", lambda x, e: fn(x.n_air)),
                    ("KP F (HAC)", lambda x, e: ff(x.F) if e != "OLS" else "")]:
        W(ws, r, 1, lab, align="left")
        for j, o in enumerate(outs):
            for k, e in enumerate(ests):
                x = pick(R2, panel=v, outcome=o, col=e, term="x")
                W(ws, r, 2 + 3 * j + k, f_(x, e))
        r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 44, 13, ncol)
notes(ws, r, ncol, "Δ12 ln y_it = a_i + d_t + b (Δ12 ln P_(t−3) × R_i) + e_it, airport and year-month FE; R_i standardized over the airports of each "
      "panel (Panel C: within each year, with R(i, y−1) also entered on its own). b > 0: more resilient airports cut capacity less when fuel prices "
      "rise. 2SLS instruments: 12-month sums of the Känzig news shock or BH supply shock (× −1), lagged three months, × R_i. " + NOTE_DK)

# ---- R3 horse race ----
ws = wb.create_sheet("R3_위치와경쟁")
TL = {"x": None, "x_G": "Δ ln P × GACI (z)", "x_G2": "Δ ln P × GACI squared", "x_size_z": "Δ ln P × ln seats 1996 (z)",
      "x_intl_z": "Δ ln P × intl. share 1996 (z)", "x_stage_z": "Δ ln P × ln stage length 1996 (z)"}
R3 = rs[rs.table == "RS2"]
cols = list(dict.fromkeys(R3.col))
ncol = 1 + len(cols)
r = 1
for v, gy in [("pre2004", 2003), ("pre2008", 2007)]:
    r = title(ws, r, ncol, f"Table R3{'ab'[['pre2004', 'pre2008'].index(v)]}. Resilience versus network position, Δ12 ln seats ({PL[v].split('; ')[1]}; GACI measured in {gy})")
    rule(ws, r, 1, ncol, "top")
    for k, c in enumerate(cols):
        W(ws, r, 2 + k, c, wrap=True)
    ws.row_dimensions[r].height = 30
    rule(ws, r, 1, ncol, "bottom")
    r += 1
    for plab, est, nsel in [("OLS", "OLS", "common sample"), ("2SLS (BH)", "2SLS", "BH"), ("2SLS (Känzig)", "2SLS", "KZ")]:
        W(ws, r, 1, "Panel: " + plab, italic=True, align="left")
        r += 1
        d = R3[(R3.panel == v) & (R3.estimator == est) & (R3.note.str.endswith(nsel) if est != "OLS" else R3.note.eq("common sample"))]
        for term in ["x", "x_G", "x_G2", "x_size_z", "x_intl_z", "x_stage_z"]:
            if term == "x":
                W(ws, r, 1, "Δ ln P × resilience (z)  [col. (2): × GACI]", align="left")
            else:
                W(ws, r, 1, TL[term], align="left")
            anyv = False
            for k, c in enumerate(cols):
                x = pick(d, col=c, term=term)
                if x is None:
                    continue
                anyv = True
                W(ws, r, 2 + k, fc(x.b, x.p))
                W(ws, r + 1, 2 + k, fs(x.se))
            r += 2 if anyv else 0
        W(ws, r, 1, "Observations", align="left")
        for k, c in enumerate(cols):
            x = pick(d, col=c, term="x")
            if x is not None:
                W(ws, r, 2 + k, fn(x.n))
        r += 1
    rule(ws, r - 1, 1, ncol, "bottom")
    r += 1
widths(ws, 44, 15, ncol)
notes(ws, r, ncol, "Common sample of airports with resilience, GACI in the build-end year and 1996 characteristics. Column (2) replaces resilience by GACI. "
      "Column (6) replaces year-month FE by country × year-month FE. In 2SLS every interaction is instrumented by the shock × the same variable. "
      + NOTE_DK.split(" Brackets")[0] + ".")

# ---- R4 components, terciles, predictive validity ----
ws = wb.create_sheet("R4_구성요소_예측력")
ncol = 7
r = title(ws, 1, ncol, "Table R4. Which component matters, and does past resilience predict later crises?")
rule(ws, r, 1, ncol, "top")
for j, h in enumerate(["", "Built 1996–2003: OLS", "2SLS (BH)", "2SLS (Känzig)", "Built 1996–2007: OLS", "2SLS (BH)", "2SLS (Känzig)"]):
    W(ws, r, j + 1, h, wrap=True)
ws.row_dimensions[r].height = 30
rule(ws, r, 1, ncol, "bottom")
r += 1
W(ws, r, 1, "Panel A. Δ12 ln seats on Δ ln P × component (z), one component at a time", italic=True, align="left")
r += 1
R4 = rs[rs.table == "RS3"]
for comp, lab in [("depth_z", "Depth (shallow losses)"), ("speed_z", "Speed (fast recovery)"), ("adapt_z", "Adaptability (upgrading)")]:
    W(ws, r, 1, lab, align="left")
    for j, v in enumerate(["pre2004", "pre2008"]):
        for k, (est, nt) in enumerate([("OLS", ""), ("2SLS", "BH"), ("2SLS", "KZ")]):
            d = R4[(R4.panel == v) & (R4.col == comp) & (R4.estimator == est) & (R4.term == "x")]
            if est == "2SLS":
                d = d[d.note == nt]
            if len(d):
                x = d.iloc[0]
                W(ws, r, 2 + 3 * j + k, fc(x.b, x.p))
                W(ws, r + 1, 2 + 3 * j + k, fs(x.se))
    r += 2
W(ws, r, 1, "Panel B. Terciles of resilience (bottom = reference), OLS", italic=True, align="left")
r += 1
for term, lab in [("x", "Middle tercile × Δ ln P"), ("x_T_top", "Top tercile × Δ ln P")]:
    W(ws, r, 1, lab, align="left")
    for j, v in enumerate(["pre2004", "pre2008"]):
        x = pick(R4, panel=v, col="terciles", term=term)
        if x is not None:
            W(ws, r, 2 + 3 * j, fc(x.b, x.p))
            W(ws, r + 1, 2 + 3 * j, fs(x.se))
    r += 2
rule(ws, r - 1, 1, ncol, "bottom")
r += 1
OL = {"depth_GFC+H1N1": "Depth, GFC 2008–09", "speed_GFC+H1N1": "Speed, GFC 2008–09", "depth_Ash+Tohoku": "Depth, ash/Tōhoku 2010–11",
      "depth_Ebola+MERS": "Depth, Ebola/MERS 2014–15", "depth_COVID": "Depth, COVID 2020–21", "speed_COVID": "Speed, COVID 2020–21",
      "upgrade_COVID": "Upgrade, COVID (k = 3)"}
r = title(ws, r, ncol, "Panel C. Predictive validity: later crisis performance on earlier resilience (z), airport cross-section")
rule(ws, r, 1, ncol, "top")
for j, h in enumerate(["Outcome", "R built 1996–2007", "", "R built 1996–2007, + ln GACI 2007", "", "N", "Mean of outcome"]):
    W(ws, r, j + 1, h, wrap=True)
rule(ws, r, 1, ncol, "bottom")
r += 1
for o, lab in OL.items():
    a = pick(rv, resilience="R_pre2008", outcome=o, control="none")
    b = pick(rv, resilience="R_pre2008", outcome=o, control="ln GACI at build end")
    W(ws, r, 1, lab, align="left")
    W(ws, r, 2, fc(a.b, a.p))
    W(ws, r, 3, fs(a.se))
    W(ws, r, 4, fc(b.b, b.p))
    W(ws, r, 5, fs(b.se))
    W(ws, r, 6, fn(a.n))
    W(ws, r, 7, f"{a.ymean:.3f}")
    r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 40, 16, ncol)
notes(ws, r, ncol, "Panels A-B: specification of Table R2 with the composite replaced by one component (or tercile dummies). Panel C: OLS across airports, "
      "SE clustered by country; outcomes are the crisis-level components of the full-sample index for crises after 2007.")

# ---- R5 2022 spike ----
ws = wb.create_sheet("R5_2022급등")
outs = ["ln_seats", "ln_seats_dom", "rec", "rec_dom"]
OLb = {"ln_seats": "ln seats", "ln_seats_dom": "ln domestic seats", "rec": "ln seats − ln seats (same month 2019)",
       "rec_dom": "ln dom. seats − same month 2019"}
specs = list(dict.fromkeys(re_.spec))
ncol = 1 + len(outs) * len(specs)
r = title(ws, 1, ncol, "Table R5. 2022 fuel price spike: seats by pre-Covid resilience (built 1996–2019, 6 crises), airport × month 2021-01..2024-06")
rule(ws, r, 1, ncol, "top")
for j, o in enumerate(outs):
    c0 = 2 + len(specs) * j
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + len(specs) - 1)
    W(ws, r, c0, OLb[o], wrap=True)
    rule(ws, r, c0, c0 + len(specs) - 1, "bottom")
ws.row_dimensions[r].height = 28
r += 1
for j in range(len(outs)):
    for k, s in enumerate(specs):
        W(ws, r, 2 + len(specs) * j + k, s, wrap=True, size=9)
ws.row_dimensions[r].height = 28
rule(ws, r, 1, ncol, "bottom")
r += 1
for tm, lab in [("x_spike", "Resilience (z) × spike 2022-03..12"), ("x_after", "Resilience (z) × after 2023-01..2024-06")]:
    W(ws, r, 1, lab, align="left")
    for j, o in enumerate(outs):
        for k, s in enumerate(specs):
            x = pick(re_, outcome=o, spec=s, term=tm)
            if x is None:
                continue
            c = 2 + len(specs) * j + k
            W(ws, r, c, fc(x.b, x.p))
            W(ws, r + 1, c, fs(x.se))
    r += 2
W(ws, r, 1, "Observations", align="left")
for j, o in enumerate(outs):
    for k, s in enumerate(specs):
        x = pick(re_, outcome=o, spec=s, term="x_spike")
        if x is not None:
            W(ws, r, 2 + len(specs) * j + k, fn(x.n))
r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 40, 13, ncol)
r = notes(ws, r, ncol, "Airport × calendar-month FE (airport FE for the recovery outcomes) and country × year-month FE; reference period 2021-01..2022-02. "
          "SE: country cluster + Newey-West over months (L = 6). The monthly coefficients below show whether the gap opens with the spike or earlier.")
W(ws, r, 1, "Event study (recovery outcome, reference 2022-02)", bold=True, align="left")
r += 1
for j, h in enumerate(["Year-month", "Coefficient", "Lower 95%", "Upper 95%"]):
    W(ws, r, 1 + j, h)
r0 = r
for i, x in enumerate(res_.itertuples()):
    ws.cell(row=r + 1 + i, column=1, value=x.ym)
    ws.cell(row=r + 1 + i, column=2, value=float(x.b))
    ws.cell(row=r + 1 + i, column=3, value=float(x.b - 1.96 * x.se))
    ws.cell(row=r + 1 + i, column=4, value=float(x.b + 1.96 * x.se))
ch = LineChart()
ch.title = "Pre-Covid resilience × month, recovery vs same month 2019 (reference 2022-02)"
ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=r0, max_row=r0 + len(res_)), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + len(res_)))
style_chart(ch, ci=True, skip=6)
ch.height, ch.width = 8, 18
ws.add_chart(ch, f"F{r0}")


# ---- R6 2022 robustness and shrinkage check ----
ws = wb.create_sheet("R6_2022강건성_축소")
ncol = 7
r = title(ws, 1, ncol, "Table R6. 2022 spike: robustness, and resilience versus number of past crises (recovery outcome vs same month 2019)")
rule(ws, r, 1, ncol, "top")
for j, h in enumerate(["Check", "Outcome", "GACI 2019 × period", "Spike 2022-03..12", "After 2023-01..2024-06", "Observations", "Airports"]):
    W(ws, r, j + 1, h, wrap=True)
ws.row_dimensions[r].height = 30
rule(ws, r, 1, ncol, "bottom")
r += 1
W(ws, r, 1, "Panel A. Robustness of the resilience (shrunk, pre-Covid) × period coefficients", italic=True, align="left")
r += 1
OUTL = {"rec": "all seats", "rec_dom": "domestic seats"}
for (chk, y, gc), d in rob.groupby(["check", "outcome", "gaci_control"], sort=False):
    W(ws, r, 1, chk, align="left")
    W(ws, r, 2, OUTL[y])
    W(ws, r, 3, "yes" if gc else "no")
    for tm, c_ in [("spike", 4), ("after", 5), ("placebo", 4), ("lnP x R", 4)]:
        x = pick(d, term=tm)
        if x is not None:
            W(ws, r, c_, fc(x.b, x.p))
            W(ws, r + 1, c_, fs(x.se))
    x0 = d.iloc[0]
    W(ws, r, 6, fn(x0.n))
    W(ws, r, 7, fn(x0.n_air))
    r += 2
W(ws, r, 1, "Panel B. Which part of the index carries the effect? (all seats, excl. mainland China, GACI 2019 × period)", italic=True, align="left")
r += 1
for (mod), d in shk.groupby("moderator", sort=False):
    for x in d.itertuples():
        W(ws, r, 1, f"{mod}: {x.term}", align="left")
        W(ws, r, 2, "all seats")
        W(ws, r, 3, "yes")
        W(ws, r, 4, fc(x.b, x.p))
        W(ws, r + 1, 4, fs(x.se))
        W(ws, r, 6, fn(x.n))
        W(ws, r, 7, fn(x.n_air))
        r += 2
rule(ws, r - 1, 1, ncol, "bottom")
ws.column_dimensions["A"].width = 58
for c_ in "BCDEFG":
    ws.column_dimensions[c_].width = 15
r = notes(ws, r, ncol, "Airport FE + country × year-month FE; reference 2021-01..2022-02. Placebo: 2021-07..12 against 2021-01..06 within the pre-period. "
          "'raw R' = composite before the empirical-Bayes shrinkage; 'n valid crises' = number of crises with a >= 1% dip and observed recovery; "
          "n2, n3 = at least 2 / 3 valid crises. SE: country cluster + Newey-West over months (L = 6).")
W(ws, r, 1, "Event study, excl. mainland China, GACI 2019 × month partialled out (reference 2022-02)", bold=True, align="left")
r += 1
for j, h in enumerate(["Year-month", "Coefficient", "Lower 95%", "Upper 95%"]):
    W(ws, r, 1 + j, h)
r0 = r
for i, x in enumerate(esr.itertuples()):
    ws.cell(row=r + 1 + i, column=1, value=x.ym)
    ws.cell(row=r + 1 + i, column=2, value=float(x.b))
    ws.cell(row=r + 1 + i, column=3, value=float(x.b - 1.96 * x.se))
    ws.cell(row=r + 1 + i, column=4, value=float(x.b + 1.96 * x.se))
ch = LineChart()
ch.title = "Pre-Covid resilience × month, excl. mainland China, GACI × month controlled"
ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=r0, max_row=r0 + len(esr)), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + len(esr)))
style_chart(ch, ci=True, skip=6)
ch.height, ch.width = 8, 18
ws.add_chart(ch, f"F{r0}")

# ---- R7 inverse-U decomposition ----
ws = wb.create_sheet("R7_역U분해")
ncol = 8
r = title(ws, 1, ncol, "Table R7. The inverse-U between resilience and GACI 2024: which part of the index produces it?")
rule(ws, r, 1, ncol, "top")
for j, h in enumerate(["Outcome", "Crisis-count dummies", "GACI 2024", "", "GACI 2024 squared", "", "Peak (GACI)", "N"]):
    W(ws, r, j + 1, h)
rule(ws, r, 1, ncol, "bottom")
r += 1
OLx = {"R_full": "Resilience (paper definition, shrunk)", "R_raw_full": "Resilience before shrinkage", "n_valid_full": "Number of valid crises",
       "depth_full": "Depth component", "speed_full": "Speed component", "adapt_full": "Adaptability component"}
for x in iu.itertuples():
    W(ws, r, 1, OLx[x.outcome], align="left")
    W(ws, r, 2, "yes" if x.n_valid_dummies else "no")
    W(ws, r, 3, f"{x.a1:.4f}")
    W(ws, r, 4, fs(x.se1))
    W(ws, r, 5, fc(x.a2, x.p2))
    W(ws, r, 6, fs(x.se2))
    W(ws, r, 7, f"{x.peak:.2f}" if x.a2 < 0 else "no maximum")
    W(ws, r, 8, fn(x.n))
    r += 1
rule(ws, r - 1, 1, ncol, "bottom")
widths(ws, 40, 13, ncol)
notes(ws, r, ncol, "OLS across airports with >= 5 observed years and a 2024 GACI value, HC1 SE. Peak = −a1/(2 a2), reported when a2 < 0.")

# ---- Fig: resilience vs GACI 2024 (replication of Fig. 10) ----
ws = wb.create_sheet("Fig_레질리언스_GACI")
g = pd.read_csv(r"..\GACI1996_2024_new_panel_data.csv")
g24 = g[g.Year == 2024].set_index("Airport").GACI
d = Ridx.set_index("airport_iata")[["R_full", "n_years_full"]].join(g24.rename("g24")).dropna()
d = d[d.n_years_full >= 5].sort_values("g24")
c = np.polyfit(d.g24, d.R_full, 2)
d["fit"] = np.polyval(c, d.g24)
for j, h in enumerate(["Airport", "GACI 2024", "Resilience (full)", "Quadratic fit"]):
    W(ws, 1, 1 + j, h)
for i, (a, x) in enumerate(d.iterrows()):
    ws.cell(row=2 + i, column=1, value=a)
    ws.cell(row=2 + i, column=2, value=float(x.g24))
    ws.cell(row=2 + i, column=3, value=float(x.R_full))
    ws.cell(row=2 + i, column=4, value=float(x.fit))
n = len(d)
sc = ScatterChart()
sc.title = "Resilience vs GACI 2024 (replication of Zhang et al. 2027, Fig. 10)"
xs = Reference(ws, min_col=2, min_row=2, max_row=n + 1)
s1 = Series(Reference(ws, min_col=3, min_row=1, max_row=n + 1), xs, title_from_data=True)
s1.marker.symbol = "circle"
s1.marker.size = 3
s1.graphicalProperties.line.noFill = True
s1.marker.graphicalProperties.solidFill = "8C8C8C"
s1.marker.graphicalProperties.line.noFill = True
s2 = Series(Reference(ws, min_col=4, min_row=1, max_row=n + 1), xs, title_from_data=True)
s2.marker.symbol = "none"
s2.graphicalProperties.line.solidFill = "137C7C"
s2.graphicalProperties.line.width = 28000
sc.series += [s1, s2]
sc.x_axis.title = "GACI 2024"
sc.y_axis.title = "Resilience index"
sc.x_axis.delete = False
sc.y_axis.delete = False
sc.legend.position = "r"
sc.y_axis.scaling.min = 0.45
sc.y_axis.scaling.max = 0.9
sc.x_axis.majorGridlines = None
sc.height, sc.width = 10, 18
ws.add_chart(sc, "F2")

# ---- index values ----
ws = wb.create_sheet("레질리언스_지수값")
iso = pd.read_csv("airport_base.csv").set_index("airport_iata").iso3
out = Ridx.copy()
out.insert(1, "iso3", out.airport_iata.map(iso))
out.insert(2, "GACI_2024", out.airport_iata.map(g24))
ws.append(list(out.columns))
for row in out.itertuples(index=False):
    ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else v for v in row])

for nm, df in [("raw_resil", rs), ("raw_valid", rv), ("raw_events", re_), ("raw_eventstudy", res_), ("raw_replication", rep),
               ("raw_2022_robust", rob), ("raw_shrinkage", shk), ("raw_eventstudy_robust", esr), ("raw_invU", iu)]:
    ws = wb.create_sheet(nm)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else v for v in row])
wb.save(OUTF)
print("saved", OUTF, [s.title for s in wb.worksheets])
