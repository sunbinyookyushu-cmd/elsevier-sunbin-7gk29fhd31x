# -*- coding: utf-8 -*-
"""Build GACI_FuelShock_Designs_20260929.xlsx: alternative fuel-shock designs (40_fuel_designs.py,
41_local_currency.py). Text sheets from _xlsx_text_designs.py."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from _xlsx_helpers import st, W, rule, fc, fs, fn, ff, notes, title, widths
from _xlsx_text_designs import README, DECISIONS, MENU

OUTF = "GACI_FuelShock_Designs_20260929.xlsx"
rd = pd.read_csv("_res_designs.csv")
ts = pd.read_csv("_res_designs_ts.csv")
lf = pd.read_csv("_res_localfx.csv")
NOTE = ("Airport × month, Δ12 differences (price lagged 3 months), 1997–2019, airports with a 1996 GACI value; airport + year-month FE unless stated. "
        "SE in parentheses: country cluster + Newey-West over months (L = 12), floored at the larger one-way component; stars (* p<0.1, ** p<0.05, "
        "*** p<0.01) refer to these. Brackets: country cluster only.")

def pick(df, **kw):
    m = df
    for k, v in kw.items():
        m = m[m[k] == v]
    return m.iloc[0] if len(m) else None

wb = Workbook()
wb.remove(wb.active)

def text_sheet(name, blocks):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 120
    for r, (kind, txt) in enumerate(blocks, start=1):
        c = ws.cell(row=r, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="맑은 고딕", size=14 if kind == "h1" else 11 if kind == "h2" else 10, bold=kind in ("h1", "h2"))
        if kind == "p":
            ws.row_dimensions[r].height = max(15, 15 * (len(txt) // 95 + txt.count("\n") + 1))

text_sheet("README", README)
text_sheet("결정필요_표본정의", DECISIONS)
text_sheet("설계메뉴", MENU)

def matrix(ws, r, ttl, df, rows_, cols_, note, first=46, colw=14, brackets=True, extra=None):
    """rows_: list of ('panel', text) or (label, filter); cols_: list of (header, filter). Cell = filter_row & filter_col."""
    ncol = 1 + len(cols_)
    r = title(ws, r, ncol, ttl)
    rule(ws, r, 1, ncol, "top")
    for k, (h, _) in enumerate(cols_):
        W(ws, r, 2 + k, h, wrap=True)
    ws.row_dimensions[r].height = 42
    rule(ws, r, 1, ncol, "bottom")
    r += 1
    for item in rows_:
        if item[0] == "panel":
            W(ws, r, 1, item[1], italic=True, align="left")
            r += 1
            continue
        lab, fr = item
        W(ws, r, 1, lab, align="left", wrap=True)
        for k, (_, fcol) in enumerate(cols_):
            x = pick(df, **{**fr, **fcol})
            if x is None:
                continue
            W(ws, r, 2 + k, fc(x.b, x.p))
            W(ws, r + 1, 2 + k, fs(x.se))
            if brackets:
                W(ws, r + 2, 2 + k, fs(x.se_cl, br="[]"), size=9)
        r += 3 if brackets else 2
    if extra:
        for lab, fnc in extra:
            W(ws, r, 1, lab, align="left")
            for k, (_, fcol) in enumerate(cols_):
                v = fnc(fcol)
                if v is not None:
                    W(ws, r, 2 + k, v)
            r += 1
    rule(ws, r - 1, 1, ncol, "bottom")
    widths(ws, first, colw, ncol)
    return notes(ws, r, ncol, note)

# ---- D1 exposure ----
ws = wb.create_sheet("D1_연료비노출")
OUTL = {"d12_ln_seats": "Δ12 ln seats", "d12_ln_co2": "Δ12 ln CO2", "d12_ln_int": "Δ12 ln CO2 per seat-km", "d12_ln_gauge": "Δ12 ln seats per departure"}
cols_ = [(f"{OUTL[o]}: {e}", dict(outcome=o, col=c)) for o in OUTL for e, c in [("OLS", "OLS"), ("2SLS Känzig", "2SLS-KZ"), ("2SLS BH", "2SLS-BH")]]
D1 = rd[rd.table == "D1"]
rows_ = [(lab, dict(panel=H, term="x")) for H, lab in [("z_fuelseat", "Δ ln P × fuel burn per seat 1996 (z)"), ("z_int", "Δ ln P × CO2 per seat-km 1996 (z)"),
                                                        ("z_stage", "Δ ln P × stage length 1996 (z)")]]
nobs = lambda fcol: fn(pick(D1, panel="z_fuelseat", term="x", **fcol).n) if pick(D1, panel="z_fuelseat", term="x", **fcol) is not None else None
Fk = lambda fcol: ff(pick(D1, panel="z_fuelseat", term="x", **fcol).F) if fcol["col"] != "OLS" and pick(D1, panel="z_fuelseat", term="x", **fcol) is not None else None
r = matrix(ws, 1, "Table D1. Who is exposed: jet fuel price × 1996 fuel-cost exposure (one exposure per regression)", D1, rows_, cols_,
           "Fuel burn per seat = CO2 on departing flights / departing seats in 1996 (CO2 = 3.16 × fuel), a direct measure of the fuel cost per "
           "passenger trip. CO2 per seat-km measures fleet fuel efficiency; stage length = seat-km / seats. Correlations across airports: fuel/seat and "
           "stage 0.87, CO2 per seat-km and stage −0.74. " + NOTE, first=40, colw=12,
           extra=[("Observations (fuel/seat row)", nobs), ("KP F (HAC, fuel/seat row)", Fk)])
D1j = rd[rd.table == "D1j"]
TL = {"x_z_fuelseat": "Δ ln P × fuel burn per seat (z)", "x_H_g": "Δ ln P × GACI 1996 (z)", "x_z_int": "Δ ln P × CO2 per seat-km (z)",
      "x_z_stage": "Δ ln P × stage length (z)", "x_z_size": "Δ ln P × ln seats 1996 (z)", "x_z_intl": "Δ ln P × intl. share 1996 (z)"}
for o in ["d12_ln_seats", "d12_ln_co2"]:
    specs = list(dict.fromkeys(D1j.panel))
    cols_ = [(f"{s_} | {e}", dict(panel=s_, col=c, outcome=o)) for s_ in specs for e, c in [("OLS", "OLS"), ("BH", "2SLS-BH")]]
    rows_ = [(TL[t_], dict(term=t_)) for t_ in TL]
    r = matrix(ws, r, f"Table D1 (cont.). Fuel-cost exposure against network position, {OUTL[o]}", D1j, rows_, cols_,
               "Each column enters all listed interactions jointly; 2SLS instruments each interaction by the BH shock × the same exposure. "
               "The country × month column compares airports within a country and month. " + NOTE.split(" SE in")[0] + ".", first=40, colw=13, brackets=False)

# ---- D2 decomposition ----
ws = wb.create_sheet("D2_공급vs수요충격")
D2 = rd[rd.table == "D2"]
cols_ = [(f"{hl}: {OUTL[o]}", dict(panel=H, outcome=o)) for H, hl in [("H_g", "× GACI 1996"), ("z_fuelseat", "× fuel/seat 1996")] for o in ["d12_ln_seats", "d12_ln_co2"]]
labs = ["Supply (BH, x -1)", "Economic activity (BH)", "Oil consumption demand (BH)", "Oil inventory demand (BH)"]
rows_ = [(l_, dict(label=l_)) for l_ in labs]
tsd = ts[ts.design == "D2 BH decomposition"].reset_index(drop=True)
r = matrix(ws, 1, "Table D2. Supply-driven versus demand-driven oil price movements (reduced form, all four shocks jointly, per 1 SD of each 12-month sum)",
           D2, rows_, cols_, "Shocks: Baumeister and Hamilton (2019) structural oil market shocks, 12-month sums lagged three months, each scaled by its SD "
           "over 1997–2019 and interacted with the moderator. The panel below gives each shock's effect on Δ12 ln jet fuel price (time series, Newey-West "
           "L = 12), for reading the reduced forms in price units. " + NOTE, first=36, colw=18)
W(ws, r, 1, "Time-series: Δ12 ln jet price (t−3) on the four 12-month shock sums", bold=True, align="left")
r += 1
for j, h in enumerate(["Shock", "Effect of 1 SD on Δ12 ln P", "SE", "R² (joint)"]):
    W(ws, r, 1 + j, h)
r += 1
for i, x in enumerate(tsd.itertuples()):
    W(ws, r, 1, labs[i], align="left")
    W(ws, r, 2, fc(x.b * x.sd_x, x.p))
    W(ws, r, 3, fs(x.se * x.sd_x))
    W(ws, r, 4, f"{x.r2:.2f}")
    r += 1

# ---- D3 persistent vs transitory ----
ws = wb.create_sheet("D3_지속vs일시충격")
D3 = rd[rd.table == "D3"]
rows_ = []
for pnl in list(dict.fromkeys(D3.col)):
    rows_.append(("panel", "Panel " + pnl))
    for l_ in list(dict.fromkeys(D3[D3.col == pnl].label)):
        rows_.append((l_, dict(col=pnl, label=l_)))
cols_ = [(f"{hl}: {OUTL[o]}", dict(panel=H, outcome=o)) for H, hl in [("H_g", "× GACI 1996"), ("z_fuelseat", "× fuel/seat 1996")] for o in ["d12_ln_seats", "d12_ln_co2"]]
r = matrix(ws, 1, "Table D3. Persistent versus transitory fuel price movements (components entered jointly)", D3, rows_, cols_,
           "(a) Expected price: Baumeister (2022) 12-month-ahead WTI expectation, CPI-deflated; transitory = spot WTI minus the expectation. "
           "(b) Persistent = 12-month trailing mean of ln real jet price; transitory = deviation from it. (c) Känzig (2021) OPEC-announcement news shock "
           "(about future supply) against the realized BH supply shock, reduced form per SD. Coefficients in (a) and (b) are differential elasticities per "
           "SD of the moderator. " + NOTE, first=52, colw=18)
W(ws, r, 1, "Time-series: how each component moves the jet price (Δ12 ln jet, t−3)", bold=True, align="left")
r += 1
for j, h in enumerate(["Design", "Component", "Coefficient", "SE", "R²"]):
    W(ws, r, 1 + j, h)
r += 1
for x in ts[ts.design.str.startswith("D3")].itertuples():
    W(ws, r, 1, x.design, align="left")
    W(ws, r, 2, x.term)
    W(ws, r, 3, fc(x.b, x.p))
    W(ws, r, 4, fs(x.se))
    W(ws, r, 5, f"{x.r2:.2f}")
    r += 1

# ---- D4-D5 ----
ws = wb.create_sheet("D4_D5_제트고유_비대칭")
D45 = rd[rd.table.isin(["D4", "D5"])]
rows_ = []
for pnl, lab in [("crude vs jet-specific", "Panel A. Crude oil versus jet-specific (refining) component"), ("asymmetry", "Panel B. Price rises versus falls"),
                 ("regime", "Panel C. High-price versus low-price regime")]:
    rows_.append(("panel", lab))
    for l_ in list(dict.fromkeys(D45[D45.col == pnl].label)):
        rows_.append((l_, dict(col=pnl, label=l_)))
matrix(ws, 1, "Table D4–D5. Jet-specific price component, asymmetry and price regimes", D45, rows_, cols_,
       "Crack spread = ln(jet price / Brent per gallon). Regime: real jet price three months earlier above or below USD 2.5 per gallon (2019 dollars). "
       + NOTE, first=46, colw=18)

# ---- L local currency ----
ws = wb.create_sheet("L_현지통화연료가격")
OL = {"d12_ln_seats": "Δ12 ln seats", "d12_ln_co2": "Δ12 ln CO2", "d12_ln_seats_dom": "Δ12 ln domestic seats", "d12_ln_seats_intl": "Δ12 ln intl. seats"}
cols_ = [(OL[o], dict(outcome=o)) for o in OL]
rows_ = [("panel", "Panel A. Local real jet price, airport + year-month FE (identified from real exchange rates; includes their demand effects)"),
         ("Δ12 ln local real jet price (t−3)", dict(design="L1 local real price, airport + month FE", term="d12_lnP_loc_l3"))]
for H, hl in [("fuel burn per seat 1996 (z)", "fuel/seat"), ("GACI 1996 (z)", "GACI")]:
    for k_, smp in enumerate(["all countries", "excl. |D12 ln q| > 0.5"], start=1):
        rows_.append(("panel", f"Panel {'B' if hl == 'fuel/seat' else 'C'}{k_}. Real exchange rate × {hl}, airport + country × year-month FE; {smp}"))
        rows_.append(("Δ12 ln real exchange rate × exposure", dict(col=f"{smp}; exposure = {H}", term="x_q")))
        rows_.append(("Δ12 ln USD jet price × exposure (control)", dict(col=f"{smp}; exposure = {H}", term="x_p")))
matrix(ws, 1, "Table L. Country-specific fuel cost variation from real exchange rates", lf, rows_, cols_,
       "Real exchange rate q = local currency per USD (BIS monthly averages) × US CPI / local CPI (WDI annual CPI, interpolated log-linearly). "
       "Local real jet price = real USD jet price × q. Panels B-C: within a country and month (country × year-month FE absorb the demand effects of the "
       "exchange rate), are airports with higher fuel cost per seat hit harder when the local real cost of fuel rises? 'excl. |D12 ln q| > 0.5' drops "
       "342 of 47,056 country-months with real exchange-rate swings above 50% in a year (currency crises: COD, AGO, VEN, SSD, SDN, IRN, BLR, SRB, "
       "IDN 1998, RUS 1998/2015, ...). " + NOTE, first=52, colw=16)

for nm, df in [("raw_designs", rd), ("raw_timeseries", ts), ("raw_localfx", lf)]:
    ws = wb.create_sheet(nm)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else v for v in row])
wb.save(OUTF)
print("saved", OUTF, [s.title for s in wb.worksheets])
