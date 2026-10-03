# -*- coding: utf-8 -*-
"""Shared cell / table helpers for the result workbooks (Times New Roman, three-line tables, coef + (SE))."""
import numpy as np
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Times New Roman"
thin, thick = Side(style="thin"), Side(style="medium")


def st(p):
    if p is None or not np.isfinite(p):
        return ""
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


def W(ws, r, c, v, bold=False, italic=False, align="center", size=10, wrap=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, size=size, bold=bold, italic=italic)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return cell


def rule(ws, r, c1, c2, side="top", s=thin):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=r, column=c)
        b = cell.border
        cell.border = Border(top=s if side == "top" else b.top, bottom=s if side == "bottom" else b.bottom,
                             left=b.left, right=b.right)


def fc(b, p, d=4):
    return "" if b is None or not np.isfinite(b) else f"{b:.{d}f}{st(p)}"


def fs(se, d=4, br="()"):
    return "" if se is None or not np.isfinite(se) else f"{br[0]}{se:.{d}f}{br[1]}"


def fn(n):
    return "" if n is None or not np.isfinite(n) else f"{int(n):,}"


def ff(F):
    return "" if F is None or not np.isfinite(F) else (f"{F:,.1f}" if F < 1e5 else f"{F:.2e}")


def notes(ws, r, ncol, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    W(ws, r, 1, text, align="left", size=9, wrap=True)
    ws.row_dimensions[r].height = max(30, 13 * (len(text) // (12 * ncol) + 2))
    return r + 2


def title(ws, r, ncol, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    W(ws, r, 1, text, italic=True, align="center", size=11)
    return r + 1


def widths(ws, first=40, rest=13, ncol=20):
    ws.column_dimensions["A"].width = first
    for c in range(2, ncol + 1):
        ws.column_dimensions[get_column_letter(c)].width = rest


def style_chart(ch, ci=True, skip=None, colors=("137C7C", "8C8C8C", "8C8C8C")):
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
    ch.x_axis.tickLblPos = "low"
    ch.legend.position = "r"
    if skip:
        ch.x_axis.tickLblSkip = skip
        ch.x_axis.tickMarkSkip = skip
    for i, s_ in enumerate(ch.series):
        s_.smooth = False
        s_.marker.symbol = "none"
        s_.graphicalProperties.line.solidFill = colors[i] if i < len(colors) else "8C8C8C"
        s_.graphicalProperties.line.width = 22000 if i == 0 or not ci else 12700
        if ci and i > 0:
            s_.graphicalProperties.line.dashStyle = "dash"
