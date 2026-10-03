"""Template for the carrier-level data request to Fangyu (2026-09-30).

Example rows = flights from Fangyu's August sample (Emission_Data_Sample.csv),
grouped exactly as the requested file would be. Check numbers come from the
August delivery (airport_month_emissions.csv).
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent
FANGYU = ROOT.parent / "GACI_CO2" / "data_from_fangyu"
OUT_DIR = ROOT / "data_request_20260930"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "Airport_Carrier_Month_Template.xlsx"

KEYS = ["airport_iata", "year", "month", "dom_intl",
        "marketing_carrier", "operating_carrier"]
VALUES = ["n_dep_flights", "dep_seats"]

# ---- example rows from the August flight-level sample ----------------------
s = pd.read_csv(FANGYU / "Emission_Data_Sample.csv")
date = pd.to_datetime(s["Time series"], format="%Y/%m/%d")
s = s.assign(
    airport_iata=s["Origin"],
    year=date.dt.year,
    month=date.dt.month,
    dom_intl=s["InternationalDomestic"],
    marketing_carrier=s["Marketing Carrier"],
    operating_carrier=s["Operating Carrier"],
)
agg = (s.groupby(KEYS, as_index=False)
        .agg(n_dep_flights=("Seats", "size"), dep_seats=("Seats", "sum")))

# mainline next to its regional partner, then international and non-US rows
order = [("ATL", "DL", "DL"), ("SLC", "DL", "OO"),
         ("EWR", "UA", "UA"), ("ORF", "UA", "YX"),
         ("ROK", "QF", "SSQ"), ("MEX", "Y4", "Y4")]
rows = []
for ap, mk, op in order:
    r = agg[(agg.airport_iata == ap) & (agg.marketing_carrier == mk)
            & (agg.operating_carrier == op)]
    assert len(r) == 1, (ap, mk, op)
    rows.append(r.iloc[0])
ex = pd.DataFrame(rows)[KEYS + VALUES]

# ---- check numbers from the August delivery ---------------------------------
aug = pd.read_csv(FANGYU / "delivery_20260818" / "airport_month_emissions.csv",
                  usecols=["airport_iata", "year", "month", "dom_intl",
                           "n_dep_flights", "dep_seats"])
chk = aug[(aug.airport_iata == "SLC") & (aug.year == 2019)
          & (aug.month == 7) & (aug.dom_intl == "Domestic")]
assert len(chk) == 1
chk_f, chk_s = int(chk.n_dep_flights.iloc[0]), int(chk.dep_seats.iloc[0])

# ---- workbook ----------------------------------------------------------------
GREY_H, GREY_C = "D9D9D9", "F2F2F2"
YEL_H, YEL_C = "FFD966", "FFF2CC"
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
bold = Font(name="Calibri", size=11, bold=True)
norm = Font(name="Calibri", size=11)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


wb = Workbook()

# Sheet 1: the file itself
ws = wb.active
ws.title = "Template"
cols = KEYS + VALUES
for j, c in enumerate(cols, start=1):
    cell = ws.cell(row=1, column=j, value=c)
    cell.font = bold
    cell.fill = fill(GREY_H if c in KEYS else YEL_H)
    cell.border = box
    cell.alignment = Alignment(horizontal="center")
for i, rec in enumerate(ex.itertuples(index=False), start=2):
    for j, c in enumerate(cols, start=1):
        v = getattr(rec, c)
        v = int(v) if c in ("year", "month") + tuple(VALUES) else v
        cell = ws.cell(row=i, column=j, value=v)
        cell.font = norm
        cell.fill = fill(GREY_C if c in KEYS else YEL_C)
        cell.border = box
        if c in VALUES:
            cell.number_format = "#,##0"
        cell.alignment = Alignment(horizontal="right" if c in VALUES
                                   or c in ("year", "month") else "left")
widths = {"airport_iata": 13, "year": 7, "month": 7, "dom_intl": 14,
          "marketing_carrier": 19, "operating_carrier": 19,
          "n_dep_flights": 15, "dep_seats": 12}
for j, c in enumerate(cols, start=1):
    ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = widths[c]
ws.freeze_panes = "A2"

# legend to the right of the table
ws.column_dimensions["I"].width = 3
ws.column_dimensions["J"].width = 4
ws.column_dimensions["K"].width = 58
legend = [
    (GREY_H, "Grey columns: group the flights by these."),
    (YEL_H, "Yellow columns: the two numbers we need."),
    (None, "The example rows are flights from your August sample (2023),"),
    (None, "so each has one flight. In the real file each row is a monthly total."),
]
for k, (col, text) in enumerate(legend, start=1):
    if col:
        c = ws.cell(row=k, column=10, value="")
        c.fill = fill(col)
        c.border = box
    t = ws.cell(row=k, column=11, value=text)
    t.font = norm

# Sheet 2: column notes
wd = wb.create_sheet("Columns")
head = ["column", "group by / count / sum", "description",
        "field in the flight data", "example"]
spec = [
    ("airport_iata", "group by", "Departure airport, IATA code (as in the August file)", "Origin", "SLC"),
    ("year", "group by", "Calendar year", "Time series", 2023),
    ("month", "group by", "Calendar month (1 to 12)", "Time series", 9),
    ("dom_intl", "group by", "Domestic or International (as in the August file)", "InternationalDomestic", "Domestic"),
    ("marketing_carrier", "group by", "Airline that sells the flight", "Marketing Carrier", "DL"),
    ("operating_carrier", "group by", "Airline that flies the flight", "Operating Carrier", "OO"),
    ("n_dep_flights", "count", "Number of departing flights", "number of flights", 1),
    ("dep_seats", "sum", "Total seats on these flights", "Seats", 70),
]
for j, h in enumerate(head, start=1):
    c = wd.cell(row=1, column=j, value=h)
    c.font = bold
    c.fill = fill(GREY_H)
    c.border = box
for i, rec in enumerate(spec, start=2):
    for j, v in enumerate(rec, start=1):
        c = wd.cell(row=i, column=j, value=v)
        c.font = norm
        c.border = box
        c.fill = fill(YEL_C if rec[0] in VALUES else GREY_C)
        c.alignment = Alignment(horizontal="left", vertical="top")

notes_row = len(spec) + 3
wd.cell(row=notes_row, column=1, value="Notes").font = bold
notes = [
    "1. One row per airport, year, month, dom_intl, marketing carrier and operating carrier. Departing flights only.",
    "2. Years 1996 to 2019.",
    "3. Check: adding up the rows of one airport, month and dom_intl should give the August numbers.",
    f"    For example, SLC, 2019, 7, Domestic has n_dep_flights = {chk_f:,} and dep_seats = {chk_s:,} in the August file.",
    "4. Any format is fine: csv, csv.gz, parquet, or one file per year.",
]
for k, t in enumerate(notes, start=1):
    c = wd.cell(row=notes_row + k, column=1, value=t)
    c.font = norm
for letter, w in zip("ABCDE", (20, 22, 52, 25, 11)):
    wd.column_dimensions[letter].width = w
wd.freeze_panes = "A2"

# print setup so a PDF check is readable
for sh in (ws, wd):
    sh.page_setup.orientation = "landscape"
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0

wb.save(OUT)
print(ex.to_string(index=False))
print("check SLC 2019-07 Domestic:", chk_f, chk_s)
print("saved", OUT)
