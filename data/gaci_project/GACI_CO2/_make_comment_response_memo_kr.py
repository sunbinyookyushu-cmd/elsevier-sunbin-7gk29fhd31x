# -*- coding: utf-8 -*-
"""코멘트별 대응 설명 메모(국문) -> 코멘트대응_설명메모_20260903.docx"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Malgun Gothic")
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2.0)

def H(text, size=13):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
def B(label, text):
    p = doc.add_paragraph(); r = p.add_run(label + " "); r.bold = True; p.add_run(text); p.paragraph_format.space_after = Pt(3)
def P(text):
    p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph(); r = p.add_run("GACI x 항공 CO2: 공저자 코멘트별 대응 설명 메모"); r.bold = True; r.font.size = Pt(15)
P("작성 2026-09-03 (Sunbin). 기준 원고: co2_overleaf_20260903_cl\\main_co2_nature_20260903.tex (전 국가-연도 표 국가 클러스터 SE, 공항 표는 공항 클러스터). 각 항목은 요구 → 한 일 → 결과 → 원고 위치 → 남은 것 순.")

H("전체적으로 바꾼 것", 14)
B("1. 추론을 국가 클러스터로 통일.", "논문의 모든 국가-연도 회귀(Table 1, 분해, 이질성, 시기분할, 스필오버, 배제제약 진단, 함수형·그래디언트 표, 매개분석 Panel A)가 국가 클러스터(184개) 표준오차와 클러스터-robust 1단계 통계량을 보고. 8월판은 전부 heteroskedasticity-robust였음. 공항 회귀는 공항 클러스터. robust 결과는 SI Table 1에만 비교용으로 보존. 결과: 5.67의 se 0.41 → 1.12, 1단계 F 154 → 18.4; 효율 항과 기재·스테이지 성분은 개별 유의성 상실; 분할표본 1단계가 약해져 이질성 그래디언트는 pooled 상호작용이 담당.")
B("2. Results를 롱페이 제안 5-layer로 재구성,", "Methods에 항등식 분해, 공항 설계, 스필오버 설계, 배제제약 진단, 함수형 소절 신설.")
B("3. 스필오버 헤드라인 교체.", "8월판 역거리 이웃효과(7.45)는 참고행으로만 남기고, own·이웃 동시도구화 인접국 스펙 + Anderson-Rubin 구간 + 순열 플라시보를 본명으로.")
B("4. 8월판 이질성 수치 정정.", "국제선 CO2로 추정된 값이 총CO2로 표기돼 있던 것을 총CO2 추정치로 교체하고 coefplot 재생성.")
B("5. 해석 축소.", "탄력성을 네트워크 이륙기의 배출 반응으로 제시하고, 수준 스펙·연속 그래디언트·귀속 범위 35~43%를 42.5% 단일값 대신 병기.")
B("6. 배제제약을 2단으로 제시.", "통과하는 핵심 3검정을 Methods·서론·Discussion에 쓰고, 나머지 6검정과 캐비앗 2개는 ED Table 7·SI Table 1을 가리키는 한 문단으로. 도구는 그대로(Feyrer 상호작용; 대안은 1단계 없음).")
B("7. 디스플레이 추가.", "본문 표 4·그림 9, ED 표 10·그림 7, SI 표 3. Nature 한도 초과라 투고 시 삭감 필요.")
H("I. 롱페이(LZ) 코멘트", 14)
H("1. 5.67이 어디서 오는가: 항등식 분해")
B("요구:", "CO2 = 편수 x 기재 x 로드팩터 x 거리 x 집약도 형태로 성분별 탄력성을 추정해 합이 5.67이 되게 보이고, 스케일 효과와 효율 효과를 정량적으로 구분.")
B("한 일:", "데이터가 스케줄 기반이라 로드팩터가 상수로 소거됨을 명시하고, CO2 = 편수 x (좌석/편) x (km/편) x (CO2/좌석-km) 항등식으로 동일 표본에서 네 성분을 각각 2SLS로 추정. 소득 tercile, 연결성 tercile, 국제/국내 세그먼트별로 같은 항등식 적용. 워터폴 그림 신설. 기존 매개분석은 Extended Data로 강등.")
B("결과:", "5.67 = 편수 6.98(se 1.38)*** - 기재 0.30(0.44) - 스테이지 0.60(0.58) - 집약도 0.40(0.30). 스케일 소계(좌석-km) 6.07 = 107%, 효율 항 -7%는 점추정이며 비유의. 집약도 95% 구간 하한(-0.99)을 잡아도 스케일의 16%만 상쇄. 국제선 5.69 = 6.03 + 0.58 - 0.40 - 0.52, 국내선 -0.64(ns). 소득별 분해는 tercile 내 1단계가 약해(F 1.7~5.7) indicative로 표기.")
B("원고:", "Layer 2 소절 'Where the elasticity comes from', Table 2(tab:decomp), Fig 워터폴, ED Table 4, Methods 식(2)(3).")
B("남은 것:", "기단연비 within-type vs type-mix 분리는 Fangyu의 기종별 집계가 있어야 가능. 효율 마진은 '작고 부정확, 어떤 경우에도 스케일을 상쇄 못 함'으로 서술.")

H("2. 공항 레벨: 허브 집중인가")
B("요구:", "6,000+ 공항 데이터로 허브/소형, 연결성, 국내/국제 지향, 기초 규모, 지역/소득별 이질성과 상위 1/5/10% 공항의 기여 집중도.")
B("한 일:", "공항 FE + 국가x연 FE(공항 클러스터)로 같은 나라 같은 해의 공항끼리 비교. 1996 기준 글로벌 상위 1%/5%, 연결성 및 트래픽 tercile, 국제선 비중, 지역, 호스트국 소득별 분할과 허브 상호작용. GACI의 capacity 성분 우려에 대비해 토폴로지 성분(eigen, closeness, betweenness, degree)만으로 재추정. 공항별 귀속배출 = CO2_2023 x [1 - exp(-b x dln GACI)]로 랭킹, 로렌츠 곡선과 Gini를 배출 자체의 집중도와 비교. 공항 Feyrer 쉬프터 IV 파일럿.")
B("결과:", "국가 내 탄력성 3.72. 허브(상위 5%) 2.11 vs 비허브 4.05, 상호작용 -2.56*** → 한계효과는 시스템적. 그러나 귀속배출은 배출보다 더 집중(상위 1% 45.8% vs 41.0%, Gini 0.945 vs 0.916)이고 1996 허브의 몫은 68.5%로 배출 몫 75.3%보다 작음 → 신흥허브(DXB 19.4 Mt, DOH, PVG, IST, ICN) 주도. 공항 합 345 Mt은 국가 356 Mt과 정합. IV 파일럿은 1단계 F 0.4~0.7로 실패.")
B("원고:", "Layer 1 소절 'Airport-level evidence', Fig 집중도, Table 3(tab:airport_conc), ED Table 5, SI Table 2, Layer 5의 '상위 5% 공항 커버 시 85%' 문장.")
B("남은 것:", "공항 인과 추정 불가(공개). 국가 최대공항은 국가x연 FE에 흡수돼 상호작용으로만 식별.")

H("3. 스필오버 정식화")
B("요구:", "(a) own GACI 통제 하 1차 이웃효과, (b) 거리감쇠(인접국, 500km, 1,000km, 2,000km 밴드 또는 연속), (c) 지역 스필오버, 메커니즘과의 연결.")
B("한 일:", "W 8종 구축(인접국, 5최근접, 거리밴드 5, 커널 5, 하위지역/항공블록 in-out, 기존 역거리). own과 이웃 동시도구화(Sanderson-Windmeijer 조건부 F), 이웃 단독도구화 + own 쉬프터 통제, 국가 클러스터 Anderson-Rubin 구간(격자 반전), 하위지역x연 FE, 순열 플라시보(국가 라벨 치환 500회), own vs 이웃 마진 표.")
B("결과:", "8월판 역거리 7.45는 국가 클러스터 t 1.3, 순열 p 0.15, 5,000km 초과 밴드 -7.0으로 글로벌 트렌드 오염 판정. 본명은 인접국: 단독도구화 2.05(0.61) F 175, AR [0.9, 3.2]; 동시도구화 own 3.48(1.29)/이웃 1.93(0.69), SW F 6.7/13.2, subset AR [0.1, 3.1], 순열 p 0.004; 하위지역x연 FE 2.07(0.83). 거리감쇠는 매끄럽지 않고 인접국에 집중. 마진: 이웃 연결성이 좌석-km 2.56***, 집약도 -0.63***, 국제선 CO2 2.23***, 편수 1.22와 스테이지 1.04는 비유의.")
B("원고:", "Layer 4 전면 재작성, Table 4(tab:spillover), ED Table 6, ED Fig 4(거리 프로파일), ED Fig 5(순열), Methods 'Spillover design'(AR 설명 포함).")
B("남은 것:", "Conley 공간 HAC SE, 하위지역(22개) wild cluster bootstrap. 메커니즘이 W 정의에 따라 달라지는 점과 동시도구화에서 own 집약도 +0.32(ns)는 본문에 명시.")

H("4. 배제제약 강화")
B("요구:", "도구가 항공 연결성 경로로만 작동한다는 검정을 추가.")
B("한 일:", "배제제약 직접 검정 3개: (1) 플라시보 아웃컴, 비항공 배출이 도구에 반응하는지; (2) 플라시보 지리, 같은 항공 사이클을 항공 대신 해운 시장접근에 곱해도 항공 CO2를 설명하는지; (3) zero-first-stage, 도구가 연결성을 못 움직이는 나라에서 배출이 도구에 반응하는지. 보조 검정 6개(개발 통제, 대안 기술지수, 리드, Hansen J, 세계 사이클 호스레이스, SE 변형)는 ED Table 7·SI Table 1에만.")
B("결과:", "3개 모두 통과. (1) 총 비항공 CO2 RF 0.05(se 0.12) vs 항공 0.73(0.16). (2) 동시 투입 시 해운 0.08(0.30) vs 항공 0.64(0.22). (3) 상위 연결성 tercile: 1단계 F 0.2, RF 0.10(0.19) vs 하위 2.22, 0.41. 보조 검정의 캐비앗 2개는 Methods에 명시: 석탄·시멘트 CO2가 도구와 함께 움직임(0.90, 0.67), 항공 사이클이 세계GDP·무역 사이클과 공선(0.82, 0.84). 대안 도구 구성은 1단계 없음(F 1.0, 3.0)이라 Feyrer 단일 도구 유지.")
B("원고 제시 방식:", "2단 구조. Methods는 핵심 3검정(비항공 플라시보, 항공 vs 해운 지리, zero-first-stage) 통과를 한 문단으로 서술하고, 나머지 6검정은 ED Table 7·SI Table 1을 가리키는 둘째 문단에서 요약하며 캐비앗 2개(석탄·시멘트 공동 움직임, 사이클 공선성)를 명시. 서론과 Discussion은 3검정만 언급. 도구는 Feyrer 단일 도구 유지(대안 구성은 1단계 F 1.0, 3.0으로 무용).")
B("남은 것:", "선택 사항: 중국·인도 제외와 bad-control 바운드로 석탄·시멘트 공동 움직임이 항공 추정을 오염시키지 않음을 보이기.")

H("5. 5-layer 구조")
P("Results를 인과효과(국가+공항) → 분해 → 이질성 → 스필오버 → 정책 순으로 재배치. Methods에 분해식, 공항 설계와 파일럿, 배제제약 진단, 스필오버 설계, 함수형과 연속 이질성 소절 신설. 매개분석, 시기분할, 측정치 강건성은 Extended Data. 08-26판 이질성 수치(11.43/5.91/0.35)가 국제선 CO2를 총CO2로 잘못 표기했던 것을 총CO2(10.40/6.14/1.83)로 정정.")

H("II. 유이푸(Yifu) 코멘트", 14)
H("1. 로그 저기초 문제: semi-log와 수준 스펙")
B("한 일:", "ln CO2 ~ GACI 수준(semi-log), CO2 Mt ~ GACI 수준(수준-수준), CO2 Mt ~ ln GACI, 기초연결성 tercile별 semi-log와 수준.")
B("결과:", "semi-log 4.09(1.01) log point/GACI 단위(평균에서 탄력성 4.4), 수준 4.2(4.0) Mt/단위(ns), tercile semi-log 10.3(2.3)/4.6(2.6)/1.1(1.7) → 절대 변화 기준에서도 저기초 국가에 집중. 수준 스펙은 대형 배출국에 좌우돼 부정확.")
B("원고:", "Layer 3 소절 'What the elasticity represents: the take-off stage', ED Table 8.")

H("2. post-2010 약한 1단계, 이륙 단계 해석, 사고 IV")
B("한 일:", "헤드라인을 '네트워크 이륙기의 배출 반응'으로 해석 변경(초록, Results, Discussion). Aviation Safety Network 사고 IV 파일럿 결과를 SI로 공개.")
B("결과:", "1996~2007 5.61(1.10) F 25 vs 2010~2023(COVID 제외) 3.01(1.19) F 11, 비제한 후기 F 3.2. 사고 IV는 1단계 F 1~4로 단독 불가, 지연 치명사고 IV 6.2(2.4), Feyrer와 과식별 Hansen p 0.70/0.95, 공동 5.4.")
B("원고:", "위 소절, Methods 'Instruments and validity' 사고 IV 문장, SI Table 3.")
B("남은 것:", "사고 IV 표는 robust 값(클러스터 재추정 미실시, 표 노트에 명시).")

H("3. 공통 탄력성 귀속의 민감도")
B("한 일:", "연결성 tercile, 소득 tercile, 연속 상호작용(선형, 2차), semi-log 규칙으로 국가별 탄력성을 달리해 2023 귀속 재계산.")
B("결과:", "294~356 Mt(35~43%). 공통 5.67이 상한. 하향 조정 최대 USA -10, CHN -8, JPN -3, ESP -3 Mt; DEU -16 → -11.")
B("원고:", "Layer 5 귀속 문단의 범위 문장, ED Table 10, Discussion 한계.")

H("4. 기초연결성 연속 이질성")
B("한 일:", "1996 GACI 분위 분할, ln GACI x 중심화 ln 기초GACI(선형, 2차) 상호작용(도구도 동일 상호작용), 백분위별 fitted 탄력성과 델타법 SE, 그래디언트 그림.")
B("결과:", "Q1 8.5(1.7), Q2 8.5(2.9), Q3 3.3(2.3), Q4 미식별(F 0.0), Q5 4.1(2.5). 상호작용 -3.99(1.00), 결합 1단계 F 9. p10 7.1 → p50 6.0 → p90 4.1, 2차항은 바닥 형태. 계단형 하락이며 최하위만의 현상이 아님.")
B("원고:", "ED Table 9, Fig gradient, 소절 본문.")

H("III. 공통: 추론 규약과 남은 취약점", 14)
P("국가-연도 표 전부를 국가 클러스터(184개)로 통일했고 robust 결과는 SI Table 1에만 남김. 그 결과 Table 1은 5.67(1.12), t 5.1, F 18.4; 효율 항과 기재, 스테이지 성분은 비유의; 분할표본 1단계가 약해져 이질성 그래디언트의 근거를 pooled 상호작용(소득 -3.33(0.57), 연결성 -2.33(0.69))으로 옮김; Hansen J는 5% 기각.")
P("취약점 목록(회신 docx의 Open issues와 동일): (1) 역거리 스필오버 불안정, 인접국 결과도 동시도구화 1단계 약함(AR로 방어), (2) 배제제약의 섹터 재구성과 사이클 공선성, (3) 확장기 및 저기초 국가에 집중된 식별(상위 tercile F 0.2, Q4 미식별), (4) 08월판 이질성 라벨 오류 정정 필요 공지, (5) 공항 레벨 인과 부재, (6) 수준 스펙 부정확, (7) 국내선 결과 부정확, (8) 추론 규약 변경, (9) 분량 초과와 로컬 미컴파일.")
P("관련 파일: Response_LZ_results_20260903.docx, Response_Yifu_results_20260903.docx(영문 point-by-point), CO2_results_summary_20260903_LZ.xlsx(31시트), co2_overleaf_20260903_cl.zip(Overleaf 업로드용).")
out = os.path.join(HERE, "코멘트대응_설명메모_20260903.docx")
try:
    doc.save(out)
except PermissionError:
    out = out.replace(".docx", "_v2.docx"); doc.save(out); print("target locked; saved as", out)
txt = "\n".join(p.text for p in doc.paragraphs)
print("saved", out, "| em-dash:", txt.count(chr(8212)), "| chars:", len(txt))
