# -*- coding: utf-8 -*-
"""Add regression-model, estimation, results-table and ranking slides to
GACI_CO2_presentation.pptx IN PLACE (2026-09-03).

New slides (positions refer to the 62-slide deck before insertion):
  after 15  CO2 country ranking 2023 (top 15 emitters)
  after 33  Airport GACI ranking 2023 (top 15 + Korean airports)
  after 36  Country GACI ranking 2023 (hub quality vs network size)
  after 39  MODEL: the estimating equation (KAIST-deck style)
  after 41  ESTIMATION: 2SLS step by step, with first stage / RF / 2SLS numbers
  after 44  RESULTS TABLE 1: main results (OLS, 2SLS, aggregation rules)
  after 49  RESULTS TABLE 2: heterogeneity + temporal split
  after 51  RESULTS TABLE 3: spillovers (neighbour connectivity)
Numbers: co2_tables_20260826.tex (Nature-version tables); first stage and
reduced form re-estimated with pyfixest (identical 2SLS 5.669 / 0.412).
Rankings: GACI1996_2024_new_panel_data.csv, gaci_co2_panel.csv,
co2_country_year.csv (2023).
Footers of ALL slides are renumbered to 'i / N' after insertion.
A timestamped backup of the pptx is written first. Asserts the deck has not
been patched yet (no 'ESTIMATING EQUATION' kicker), so it runs once.
"""
import os, re, shutil, datetime
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

GACI = r"C:\Users\sunbi\managi-lab Dropbox\Sunbin Yoo\Research Box ^-^\2026\GACI"
CO2 = os.path.join(GACI, "GACI_CO2")
PPTX = os.path.join(GACI, "GACI_CO2_presentation.pptx")

NAVY = RGBColor(0x23, 0x2A, 0x55); GOLD = RGBColor(0xE2, 0xA8, 0x5A)
INK = RGBColor(0x1E, 0x1E, 0x24); GRAY = RGBColor(0x6E, 0x6E, 0x6E)
LAV = RGBColor(0xBC, 0xC5, 0xDE); PUNCH = RGBColor(0xF2, 0xE9, 0xF7)
RED = RGBColor(0xB0, 0x41, 0x3E); GREEN = RGBColor(0x2E, 0x7D, 0x5B)
RED_T = RGBColor(0xF8, 0xE9, 0xE8); GREEN_T = RGBColor(0xE6, 0xF2, 0xEC)
GOLD_T = RGBColor(0xFA, 0xF0, 0xE1); NAVY_T = RGBColor(0xE9, 0xEB, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); LIGHT = RGBColor(0xF5, 0xF6, 0xFA)
FOOTER = "Air Connectivity and Aviation CO2, 1996\u20132023"
EQF = "Cambria"

# ----------------------------------------------------------------- helpers
def tb(s, x, y, w, h, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = wrap
    return box


def endsz(p, size):
    """Set the end-of-paragraph font size so empty/short table cells do not
    inherit the 18pt default and inflate the row height."""
    from pptx.oxml.ns import qn
    pPr = p._p
    epr = pPr.find(qn("a:endParaRPr"))
    if epr is None:
        from lxml import etree
        epr = etree.SubElement(pPr, qn("a:endParaRPr"))
    epr.set("sz", str(int(size * 100)))


def setp(p, text, size, bold=False, color=INK, font="Calibri", align=None,
         italic=False, sub=False, sup=False):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.name = font; r.font.color.rgb = color
    if sub:
        r.font._element.set("baseline", "-25000")
    if sup:
        r.font._element.set("baseline", "30000")
    if align is not None:
        p.alignment = align
    return r


def rect(s, x, y, w, h, fill, line=None, rounded=False, radius=None):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if rounded and radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def new_slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    return s


def kicker(s, text):
    box = tb(s, 0.60, 0.32, 12.0, 0.5)
    setp(box.text_frame.paragraphs[0], " ".join(text.upper()), 12, True, NAVY)
    rect(s, 0.62, 0.78, 1.40, 0.04, GOLD)


def title(s, text, size=28):
    box = tb(s, 0.58, 0.92, 12.10, 1.0)
    setp(box.text_frame.paragraphs[0], text, size, True, INK, "Verdana")


def footer(s):
    box = tb(s, 12.20, 7.05, 1.0, 0.35)
    setp(box.text_frame.paragraphs[0], "0 / 0", 10, False, GRAY,
         align=PP_ALIGN.RIGHT)
    rect(s, 0.62, 7.12, 0.55, 0.05, NAVY)
    box2 = tb(s, 1.25, 7.00, 8.0, 0.35)
    setp(box2.text_frame.paragraphs[0], FOOTER, 10, False, GRAY)


def punchline(s, text, fill=PUNCH, color=NAVY, y=6.02, x=1.40, w=10.50,
              h=0.78, size=17):
    shp = rect(s, x, y, w, h, fill, rounded=True, radius=0.5)
    tf = shp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    setp(tf.paragraphs[0], text, size, True, color, align=PP_ALIGN.CENTER)
    return shp


def note(s, text, y=6.45, x=0.62, w=12.1, size=10.5):
    box = tb(s, x, y, w, 0.5)
    setp(box.text_frame.paragraphs[0], text, size, False, GRAY, italic=True)
    return box


def panel(s, x, y, w, h, fill=LIGHT, accent=NAVY, accent_w=0.10):
    rect(s, x, y, w, h, fill)
    rect(s, x, y, accent_w, h, accent)


def eqline(p, parts, size=20, color=INK):
    """parts: list of (text, style); style chars: i italic, b bold, s subscript,
    r red, g gray, n normal. e.g. ('ct','is') italic subscript."""
    for text, st in parts:
        setp(p, text, size * (0.72 if "s" in st else 1.0),
             bold="b" in st, italic="i" in st, sub="s" in st,
             color=RED if "r" in st else (GRAY if "g" in st else color),
             font=EQF)


def circ(s, x, y, d, num, fill):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                             Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0)
    setp(tf.paragraphs[0], num, 18, True, WHITE, "Verdana",
         align=PP_ALIGN.CENTER)


# ---- regression-table builder (KAIST Table style, native editable table)
RH = {"header": 0.34, "panel": 0.26, "coef": 0.30, "se": 0.23, "stat": 0.27,
      "row": 0.25}


def reg_table(s, x, y, colw, rows, fs=11):
    """rows: list of (kind, cells). kind in header/panel/coef/se/stat/row.
    cell may be str or (str, color, bold)."""
    ncol = len(colw)
    heights = [RH[k] for k, _ in rows]
    gfx = s.shapes.add_table(len(rows), ncol, Inches(x), Inches(y),
                             Inches(sum(colw)), Inches(sum(heights)))
    t = gfx.table
    t.first_row = False; t.horz_banding = False
    for j, w in enumerate(colw):
        t.columns[j].width = Inches(w)
    for i, (kind, cells) in enumerate(rows):
        t.rows[i].height = Inches(heights[i])
        if kind == "panel":
            t.cell(i, 0).merge(t.cell(i, ncol - 1))
        for j in range(ncol):
            c = t.cell(i, j)
            c.margin_left = Inches(0.07); c.margin_right = Inches(0.05)
            c.margin_top = Inches(0.0); c.margin_bottom = Inches(0.0)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = {"header": NAVY, "panel": NAVY_T}.get(
                kind, WHITE)
            if kind == "panel" and j > 0:
                continue
            val = cells[j] if j < len(cells) else ""
            color, bold = None, False
            if isinstance(val, tuple):
                val, color, bold = val
            if kind == "header":
                color = color or WHITE; bold = True; size = fs
            elif kind == "panel":
                color = color or NAVY; bold = True; size = fs - 0.5
            elif kind == "se":
                color = color or GRAY; size = fs - 2
            elif kind == "stat":
                color = color or INK; size = fs - 1
            else:
                color = color or INK; size = fs
            tf = c.text_frame; tf.word_wrap = True
            setp(tf.paragraphs[0], str(val), size, bold, color,
                 italic=(kind == "panel"),
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
            endsz(tf.paragraphs[0], size)
    return gfx


def rank_table(s, x, y, colw, header, rows, fs=11, hl=None, rh=0.285,
               align_first=(0, 1, 2)):
    """Simple ranking table; hl = set of row indices (0-based, data rows)
    to highlight (Korea)."""
    n = len(rows) + 1
    gfx = s.shapes.add_table(n, len(colw), Inches(x), Inches(y),
                             Inches(sum(colw)), Inches(rh * n))
    t = gfx.table; t.first_row = False; t.horz_banding = False
    for j, w in enumerate(colw):
        t.columns[j].width = Inches(w)
    for i in range(n):
        t.rows[i].height = Inches(rh)
        cells = header if i == 0 else rows[i - 1]
        is_hl = hl is not None and (i - 1) in hl
        for j, val in enumerate(cells):
            c = t.cell(i, j)
            c.margin_left = Inches(0.06); c.margin_right = Inches(0.05)
            c.margin_top = Inches(0.0); c.margin_bottom = Inches(0.0)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = (NAVY if i == 0 else
                                     GOLD_T if is_hl else
                                     (WHITE if i % 2 else LIGHT))
            color = WHITE if i == 0 else (NAVY if is_hl else INK)
            bold = (i == 0) or is_hl or (j == 1)
            tf = c.text_frame; tf.word_wrap = False
            setp(tf.paragraphs[0], str(val), fs, bold, color,
                 align=PP_ALIGN.LEFT if j in align_first else PP_ALIGN.CENTER)
            endsz(tf.paragraphs[0], fs)
    return gfx


NAME = {"USA": "United States", "CHN": "China", "GBR": "United Kingdom",
        "JPN": "Japan", "ARE": "UAE", "IND": "India", "DEU": "Germany",
        "ESP": "Spain", "AUS": "Australia", "FRA": "France", "CAN": "Canada",
        "TUR": "Turkey", "BRA": "Brazil", "RUS": "Russia", "KOR": "Korea",
        "MEX": "Mexico", "ITA": "Italy", "QAT": "Qatar", "SGP": "Singapore",
        "THA": "Thailand", "NLD": "Netherlands", "HKG": "Hong Kong",
        "CHE": "Switzerland", "SAU": "Saudi Arabia", "ISR": "Israel",
        "ETH": "Ethiopia", "AUT": "Austria", "IRL": "Ireland",
        "DNK": "Denmark", "BEL": "Belgium", "IDN": "Indonesia",
        "MYS": "Malaysia"}


def cname(iso):
    return NAME.get(iso, iso)


def fmt(v, d=3):
    return f"{v:.{d}f}"


# ================================================================ data
def load_rankings():
    a = pd.read_csv(os.path.join(GACI, "GACI1996_2024_new_panel_data.csv"))
    ap = pd.read_csv(os.path.join(CO2, "airport_co2_panel.csv"))
    iso = ap.drop_duplicates("airport_iata").set_index("airport_iata")["iso3"]
    t = a[a.Year == 2023].sort_values("GACI", ascending=False).reset_index(drop=True)
    t["iso3"] = t.Airport.map(iso); t["rank"] = t.index + 1
    r96 = a[a.Year == 1996].sort_values("GACI", ascending=False).reset_index(drop=True)
    r96["r96"] = r96.index + 1
    t = t.merge(r96[["Airport", "r96"]], on="Airport", how="left")
    n_ap = len(t)
    air_top = t.head(15)
    kor = t[t.iso3 == "KOR"].head(6)

    g = pd.read_csv(os.path.join(CO2, "gaci_co2_panel.csv"))
    c = g[g.y == 2023].copy()
    cw = c.sort_values("gaci_cwmean", ascending=False).reset_index(drop=True)
    cw["rank"] = cw.index + 1
    n_c = len(cw)
    ssum = (ap[ap.year == 2023].groupby("iso3")
            .agg(gsum=("GACI", "sum"), n=("GACI", "size"))
            .sort_values("gsum", ascending=False).reset_index())
    ssum["rank"] = ssum.index + 1

    co = pd.read_csv(os.path.join(CO2, "co2_country_year.csv"))
    e = co[co.year == 2023].sort_values("co2_bunker", ascending=False).reset_index(drop=True)
    e["rank"] = e.index + 1; e["Mt"] = e.co2_bunker / 1e9
    e["Mt_intl"] = e.co2_bunker_intl / 1e9; e["Mt_lto"] = e.co2_lto / 1e9
    e["share"] = e.Mt / e.Mt.sum() * 100
    e96 = co[co.year == 1996].sort_values("co2_bunker", ascending=False).reset_index(drop=True)
    e96["r96"] = e96.index + 1; e96["Mt96"] = e96.co2_bunker / 1e9
    e = e.merge(e96[["iso3", "r96", "Mt96"]], on="iso3", how="left")
    return dict(air_top=air_top, kor=kor, n_ap=n_ap, cw=cw, n_c=n_c,
                ssum=ssum, emit=e, world=e.Mt.sum())


# ================================================================ slides
def slide_co2_ranking(prs, D):
    s = new_slide(prs)
    kicker(s, "Data")
    title(s, "The emitters: top 15 countries, 2023")
    e = D["emit"]
    rows = []
    for _, r in e.head(15).iterrows():
        rows.append((int(r["rank"]), cname(r.iso3), r.iso3, f"{r.Mt:.1f}",
                     f"{r.share:.1f}%", f"{r.Mt_intl / r.Mt * 100:.0f}%",
                     f"{r.Mt_lto:.1f}", f"{int(r.r96)}", f"{r.Mt96:.1f}"))
    hl = {i for i, r in enumerate(rows) if r[2] == "KOR"}
    rank_table(s, 0.62, 1.95,
               [0.55, 1.75, 0.65, 1.05, 0.95, 1.05, 1.05, 1.0, 1.05],
               ("#", "Country", "ISO", "CO2 2023 (Mt)", "World share",
                "Intl share", "LTO (Mt)", "Rank 1996", "CO2 1996 (Mt)"),
               rows, fs=10.5, hl=hl, rh=0.27)
    # side panel
    panel(s, 10.05, 1.95, 2.68, 4.32, NAVY_T, GOLD)
    box = tb(s, 10.30, 2.05, 2.35, 4.2)
    tf = box.text_frame
    setp(tf.paragraphs[0], "Reading the list", 13, True, NAVY)
    items = [
        (f"{D['world']:.0f} Mt", " of world aviation CO2 in 2023 under the bunker convention."),
        ("US + China", " = 35% of the world total. China: 8th in 1996, 2nd today."),
        ("Movers", ": UAE 20th to 5th, Turkey 32nd to 12th, India 19th to 6th, Qatar into the top 20."),
        ("Korea", ": 15th, 14.4 Mt, 90% international."),
        ("Intl share", " = international share of bunker CO2; LTO = physical landing-and-take-off CO2 at own airports."),
    ]
    for k, v in items:
        p = tf.add_paragraph(); p.space_before = Pt(5)
        setp(p, k, 11, True, NAVY); setp(p, v, 11, False, INK)
    punchline(s, "Aviation CO2 is concentrated, and its geography has shifted "
                 "toward the newly connecting world since 1996.",
              y=6.35, h=0.60, size=15)
    footer(s)
    return s


def slide_airport_ranking(prs, D):
    s = new_slide(prs)
    kicker(s, "Building GACI")
    title(s, f"Airport ranking, 2023: the top 15 (of {D['n_ap']:,} scored)")
    rows = []
    for _, r in D["air_top"].iterrows():
        rows.append((int(r["rank"]), r.Airport, cname(r.iso3), fmt(r.GACI, 2),
                     f"{r.TotalCapacity / 1e6:.0f}",
                     "new" if pd.isna(r.r96) else f"{int(r.r96)}"))
    rank_table(s, 0.62, 1.95, [0.5, 0.95, 2.0, 1.15, 1.3, 1.15],
               ("#", "Airport", "Country", "GACI 2023", "Seats (M)", "Rank 1996"),
               rows, fs=11, rh=0.27, align_first=(1, 2))
    # Korea panel
    panel(s, 7.95, 1.95, 4.78, 4.32, NAVY_T, GOLD)
    box = tb(s, 8.20, 2.03, 4.4, 0.4)
    setp(box.text_frame.paragraphs[0], "Korean airports in the same ranking",
         13, True, NAVY)
    krows = []
    for _, r in D["kor"].iterrows():
        krows.append((int(r["rank"]), r.Airport, fmt(r.GACI, 2),
                      f"{r.TotalCapacity / 1e6:.1f}",
                      "new" if pd.isna(r.r96) else f"{int(r.r96)}"))
    rank_table(s, 8.20, 2.45, [0.65, 0.95, 1.0, 1.0, 0.85],
               ("#", "Airport", "GACI", "Seats (M)", "1996"),
               krows, fs=10.5, rh=0.255, align_first=(1,))
    box = tb(s, 8.20, 4.35, 4.4, 1.9)
    tf = box.text_frame
    items = [
        ("ICN", " sits in the world top 25: Korea's one global gateway, opened in 2001 and already above NRT and HND in reach."),
        ("GMP and CJU", " carry the world's busiest route yet rank 300-500: enormous domestic volume, little network reach."),
        ("The top of the table", " is FRA, IST, DXB: Istanbul and Dubai were ranked in the 90s in 1996."),
    ]
    first = True
    for k, v in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_before = Pt(4)
        setp(p, k, 10.5, True, NAVY); setp(p, v, 10.5, False, INK)
    punchline(s, "GACI ranks network position: the Gulf and Turkey climbed "
                 "from the 90s into the top 3 in one generation.",
              y=6.35, h=0.60, size=15)
    footer(s)
    return s


def slide_country_ranking(prs, D):
    s = new_slide(prs)
    kicker(s, "Building GACI")
    title(s, "Country ranking, 2023: hub quality vs network size")
    cw = D["cw"]; ss = D["ssum"]
    # left: hub quality (cwm)
    box = tb(s, 0.62, 1.90, 6.0, 0.35)
    setp(box.text_frame.paragraphs[0],
         f"Hub quality: capacity-weighted mean GACI (the regressor), {D['n_c']} countries",
         12, True, NAVY)
    rows = []
    for _, r in cw.head(15).iterrows():
        rows.append((int(r["rank"]), cname(r.c), fmt(r.gaci_cwmean, 2),
                     fmt(r.gaci_max, 2), f"{int(r.n_air)}"))
    kr = cw[cw.c == "KOR"].iloc[0]
    rows.append((int(kr["rank"]), "Korea", fmt(kr.gaci_cwmean, 2),
                 fmt(kr.gaci_max, 2), f"{int(kr.n_air)}"))
    rank_table(s, 0.62, 2.25, [0.5, 1.9, 1.25, 1.25, 1.05],
               ("#", "Country", "Hub quality", "Best airport", "Airports"),
               rows, fs=10.5, rh=0.235, hl={15}, align_first=(1,))
    # right: network size (sum)
    box = tb(s, 7.0, 1.90, 6.0, 0.35)
    setp(box.text_frame.paragraphs[0],
         "Network size: sum of airport GACI (an alternative aggregation)",
         12, True, NAVY)
    rows2 = []
    for _, r in ss.head(15).iterrows():
        rows2.append((int(r["rank"]), cname(r.iso3), f"{r.gsum:.0f}", f"{int(r.n)}"))
    ks = ss[ss.iso3 == "KOR"].iloc[0]
    rows2.append((int(ks["rank"]), "Korea", f"{ks.gsum:.0f}", f"{int(ks.n)}"))
    rank_table(s, 7.0, 2.25, [0.5, 1.9, 1.5, 1.4],
               ("#", "Country", "GACI sum", "Airports"),
               rows2, fs=10.5, rh=0.235, hl={15}, align_first=(1,))
    punchline(s, "The regression uses hub quality: small hub states (UAE, Netherlands, "
                 "Singapore) lead; the size list is a different world. Korea: 19th vs 44th.",
              y=6.38, h=0.58, size=13.5)
    footer(s)
    return s


def slide_model(prs):
    s = new_slide(prs)
    kicker(s, "Identification  \u00b7  Estimating equation")
    title(s, "The model: a log-log panel with two-way fixed effects")
    # equation box
    rect(s, 0.62, 1.95, 12.1, 1.55, LIGHT, line=LAV)
    box = tb(s, 0.85, 2.05, 11.7, 0.7)
    p = box.text_frame.paragraphs[0]
    eqline(p, [("ln CO", "n"), ("2,", "s"), ("ct", "is"), ("  =  ", "n"),
               ("\u03b2", "br"), (" ln GACI", "n"), ("ct", "is"),
               ("  +  \u03b3 ln pop", "n"), ("ct", "is"),
               ("  +  \u03b4 ln seaMA", "n"), ("ct", "is"),
               ("  +  \u03b1", "n"), ("c", "is"), ("  +  \u03c4", "n"),
               ("t", "is"), ("  +  \u03b5", "n"), ("ct", "is"),
               ("                    (1)", "g")], size=22)
    box = tb(s, 0.85, 2.80, 11.7, 0.6)
    p = box.text_frame.paragraphs[0]
    setp(p, "\u03b2", 15, True, RED, EQF)
    setp(p, "  = elasticity of national aviation CO2 with respect to hub quality: "
            "the % change in CO2 per 1% change in GACI.   Benchmark: ", 13, False, INK)
    setp(p, "\u03b2 = 1", 13, True, NAVY)
    setp(p, " (emissions grow in proportion to connectivity).", 13, False, INK)
    # variables panel
    panel(s, 0.62, 3.72, 8.15, 3.05, LIGHT, NAVY)
    box = tb(s, 0.90, 3.78, 7.8, 3.0)
    tf = box.text_frame
    setp(tf.paragraphs[0], "Variables", 13, True, NAVY)
    vars_ = [
        ("CO2", "ct", "national aviation CO2 in year t; headline = bunker rule. "
                      "Also LTO-only, 50/50 split, international only, seat-km, CO2 per seat-km."),
        ("GACI", "ct", "hub quality: capacity-weighted mean of airport GACI scores (alternatives: max, mean, sum)."),
        ("pop", "ct", "population (log)."),
        ("seaMA", "ct", "sea-distance market access (log): the surface-shipping alternative, held fixed."),
        ("\u03b1", "c", "country fixed effects: geography, size, institutions."),
        ("\u03c4", "t", "year fixed effects: fuel prices, world traffic cycle, technology."),
        ("\u03b5", "ct", "error; robust standard errors."),
    ]
    for v, sub, d in vars_:
        p = tf.add_paragraph(); p.space_before = Pt(2.5)
        setp(p, v, 11.5, True, NAVY, EQF, italic=True)
        setp(p, sub, 8.5, True, NAVY, EQF, italic=True, sub=True)
        setp(p, "   " + d, 11, False, INK)
    # identification panel
    panel(s, 9.0, 3.72, 3.73, 3.05, NAVY, GOLD)
    box = tb(s, 9.25, 3.78, 3.4, 3.0)
    tf = box.text_frame
    setp(tf.paragraphs[0], "Why OLS is not enough", 13, True, WHITE)
    for k, v in [
        ("Reverse causality", ": airlines add capacity where demand is already growing."),
        ("Measurement", ": the index is an estimate; noise attenuates OLS."),
        ("Sample", ": 184 countries \u00d7 1996\u20132023, N = 4,634 country-years."),
        ("Answer", ": instrument ln GACI with fixed 1996 geography \u00d7 the world aviation cycle (next slides). "
                   "OLS gives 3.59; 2SLS gives 5.67."),
    ]:
        p = tf.add_paragraph(); p.space_before = Pt(5)
        setp(p, k, 11, True, GOLD); setp(p, v, 11, False, LAV)
    footer(s)
    return s


def slide_estimation(prs):
    s = new_slide(prs)
    kicker(s, "Identification  \u00b7  Estimation")
    title(s, "Two-stage least squares, step by step")
    # step 1
    panel(s, 0.62, 1.95, 6.0, 4.05, LIGHT, NAVY)
    circ(s, 0.90, 2.12, 0.55, "1", NAVY)
    box = tb(s, 1.55, 2.12, 5.0, 0.5)
    setp(box.text_frame.paragraphs[0], "First stage: predict connectivity from geography",
         14, True, INK)
    box = tb(s, 0.90, 2.75, 5.6, 0.5)
    p = box.text_frame.paragraphs[0]
    eqline(p, [("ln GACI", "n"), ("ct", "is"), (" = ", "n"), ("\u03c0", "b"),
               (" Z", "n"), ("ct", "is"), (" + \u03b3", "n"), ("1", "s"),
               (" ln pop", "n"), ("ct", "is"), (" + \u03b4", "n"), ("1", "s"),
               (" ln seaMA", "n"), ("ct", "is"), (" + \u03b1", "n"), ("c", "is"),
               (" + \u03c4", "n"), ("t", "is"), (" + u", "n"), ("ct", "is")],
           size=15)
    box = tb(s, 0.90, 3.25, 5.6, 0.5)
    p = box.text_frame.paragraphs[0]
    eqline(p, [("Z", "n"), ("ct", "is"), (" = a", "n"), ("t", "is"),
               (" \u00d7 ln airMA", "n"), ("c,1996", "is")], size=15)
    setp(p, "     (Feyrer-type shifter)", 11, False, GRAY, italic=True)
    box = tb(s, 0.90, 3.75, 5.6, 1.2)
    tf = box.text_frame
    for k, sub, v in [
        ("a", "t", ": world aviation capacity index, min-max scaled 0 (1996) to 1: the global cycle nobody steers."),
        ("airMA", "c,1996", ": air market access on fixed 1996 geography, the differential exposure to that cycle."),
    ]:
        p = tf.paragraphs[0] if k == "a" else tf.add_paragraph()
        p.space_before = Pt(3)
        setp(p, k, 11, True, NAVY, EQF, italic=True)
        setp(p, sub, 8, True, NAVY, EQF, italic=True, sub=True)
        setp(p, v, 11, False, INK)
    # first stage result chips
    rect(s, 0.90, 5.05, 5.45, 0.80, NAVY_T, rounded=True, radius=0.15)
    box = tb(s, 1.05, 5.10, 5.2, 0.75)
    tf = box.text_frame
    p = tf.paragraphs[0]
    setp(p, "\u03c0\u0302 = 0.128", 15, True, NAVY, EQF)
    setp(p, "  (s.e. 0.010),  t = 12.4,  Kleibergen\u2013Paap F = 153.6", 12.5, False, INK)
    p = tf.add_paragraph()
    setp(p, "Reduced form: ln CO2 on Z directly = 0.726 (s.e. 0.057).", 12, False, INK)
    # step 2
    panel(s, 6.95, 1.95, 5.78, 4.05, NAVY, GOLD)
    circ(s, 7.23, 2.12, 0.55, "2", GOLD)
    box = tb(s, 7.88, 2.12, 4.8, 0.5)
    setp(box.text_frame.paragraphs[0], "Second stage: plug the prediction into (1)",
         14, True, WHITE)
    box = tb(s, 7.23, 2.75, 5.4, 0.6)
    p = box.text_frame.paragraphs[0]
    eqline(p, [("ln CO", "n"), ("2,", "s"), ("ct", "is"), (" = ", "n"),
               ("\u03b2", "b"), (" ln G\u0302ACI", "n"), ("ct", "is"),
               (" + \u03b3 ln pop", "n"), ("ct", "is"), (" + \u03b4 ln seaMA", "n"),
               ("ct", "is"), (" + \u03b1", "n"), ("c", "is"), (" + \u03c4", "n"),
               ("t", "is"), (" + \u03b5", "n"), ("ct", "is")], size=15,
           color=WHITE)
    box = tb(s, 7.23, 3.40, 5.4, 1.5)
    tf = box.text_frame
    for k, v in [
        ("Exactly identified", ": one instrument, one endogenous regressor, so \u03b2 is a ratio of two reduced forms: 0.726 / 0.128 = 5.67."),
        ("Only the geography-driven part", " of connectivity growth is used; demand-driven capacity is purged."),
        ("Inference", ": robust SEs; Conley (spatial) SEs and plausibly-exogenous bounds as checks."),
    ]:
        p = tf.paragraphs[0] if k.startswith("Exactly") else tf.add_paragraph()
        p.space_before = Pt(3)
        setp(p, k, 11, True, GOLD); setp(p, v, 11, False, LAV)
    rect(s, 7.23, 5.05, 5.25, 0.80, GOLD, rounded=True, radius=0.15)
    box = tb(s, 7.38, 5.10, 5.0, 0.75)
    tf = box.text_frame
    p = tf.paragraphs[0]
    setp(p, "\u03b2\u0302 = 5.669", 17, True, NAVY, EQF)
    setp(p, "  (s.e. 0.412),  p < 0.001", 12.5, True, NAVY)
    p = tf.add_paragraph()
    setp(p, "OLS on the same sample: 3.585 (0.175). N = 4,634.", 12, False, NAVY)
    punchline(s, "A 1% rise in hub quality raises national aviation CO2 by 5.7%: "
                 "far from the proportional benchmark of one.",
              y=6.22, h=0.62, size=15)
    footer(s)
    return s


def slide_table_main(prs):
    s = new_slide(prs)
    kicker(s, "Results  \u00b7  Table 1")
    title(s, "Main results: OLS versus 2SLS across six outcomes")
    R = RED
    rows = [
        ("header", ("Dependent variable (log)", "Bunker CO2", "LTO CO2", "50/50 CO2",
                    "Intl. CO2", "Seat-km", "CO2 / seat-km")),
        ("panel", ("Panel A. OLS, country and year FE",)),
        ("coef", ("ln GACI (cwm)", "3.585***", "3.043***", "3.566***", "3.874***",
                  "3.941***", "\u22120.355***")),
        ("se", ("", "(0.175)", "(0.177)", "(0.175)", "(0.198)", "(0.183)", "(0.043)")),
        ("panel", ("Panel B. 2SLS, Feyrer instrument, capacity-weighted mean (headline)",)),
        ("coef", ("ln GACI (cwm)", ("5.669***", R, True), "5.921***", "5.620***",
                  "5.690***", ("6.068***", R, True), ("\u22120.399***", R, True))),
        ("se", ("", "(0.412)", "(0.448)", "(0.409)", "(0.415)", "(0.443)", "(0.118)")),
        ("stat", ("Kleibergen\u2013Paap F", "153.6", "153.6", "153.6", "152.7", "153.6", "153.6")),
        ("panel", ("Panel C. 2SLS, alternative aggregations of airport GACI to the country",)),
        ("coef", ("ln GACI (sum)", "2.283***", "2.384***", "2.263***", "2.291***",
                  "2.443***", "\u22120.161***")),
        ("se", ("KP F = 90.3", "(0.224)", "(0.219)", "(0.222)", "(0.237)", "(0.250)", "(0.052)")),
        ("coef", ("ln GACI (max)", "4.822***", "5.036***", "4.780***", "4.836***",
                  "5.161***", "\u22120.340***")),
        ("se", ("KP F = 172.0", "(0.329)", "(0.348)", "(0.327)", "(0.343)", "(0.361)", "(0.101)")),
        ("coef", ("ln GACI (mean)", "5.897***", "6.159***", "5.846***", "5.902***",
                  "6.313***", "\u22120.415***")),
        ("se", ("KP F = 136.2", "(0.591)", "(0.623)", "(0.587)", "(0.593)", "(0.624)", "(0.122)")),
        ("stat", ("N (country-years)", "4,634", "4,634", "4,634", "4,633", "4,634", "4,634")),
    ]
    reg_table(s, 0.62, 1.90, [2.35, 1.2, 1.15, 1.15, 1.15, 1.15, 1.25], rows, fs=11)
    # headline panel
    panel(s, 10.30, 1.90, 2.43, 4.35, LIGHT, GOLD)
    box = tb(s, 10.52, 1.98, 2.15, 4.3)
    tf = box.text_frame
    setp(tf.paragraphs[0], "Headline", 13, True, NAVY)
    p = tf.add_paragraph(); p.space_before = Pt(4)
    setp(p, "\u03b2 = 5.67 (0.41)", 13, True, RED)
    p = tf.add_paragraph()
    setp(p, "+1% hub quality = +5.7% national aviation CO2. Benchmark of one rejected by a wide margin.", 10.5, False, INK)
    p = tf.add_paragraph(); p.space_before = Pt(5)
    setp(p, "OLS understates", 11, True, NAVY)
    p = tf.add_paragraph()
    setp(p, "3.59 vs 5.67: the 2SLS estimate exceeds OLS, so the net OLS bias is downward.", 10.5, False, INK)
    p = tf.add_paragraph(); p.space_before = Pt(5)
    setp(p, "Allocation rule irrelevant", 11, True, NAVY)
    p = tf.add_paragraph()
    setp(p, "5.62 to 5.92 across territorial, bunker and split conventions.", 10.5, False, INK)
    p = tf.add_paragraph(); p.space_before = Pt(5)
    setp(p, "Scale vs technique", 11, True, NAVY)
    p = tf.add_paragraph()
    setp(p, "Seat-km +6.07 and intensity \u22120.40: volume, not efficiency.", 10.5, False, INK)
    note(s, "Notes: country-year panel, 184 countries, 1996\u20132023. All columns include country and year fixed "
            "effects and control for log population and log sea market access. Instrument: world aviation cycle \u00d7 "
            "1996 air market access. Robust SE in parentheses. ***, **, *: 1, 5, 10% significance.",
         y=6.45, w=9.6, size=10)
    footer(s)
    return s


def slide_table_hetero(prs):
    s = new_slide(prs)
    kicker(s, "Results  \u00b7  Table 2")
    title(s, "Who and when: split-sample 2SLS")
    R = RED
    # left: heterogeneity
    box = tb(s, 0.62, 1.88, 6.2, 0.35)
    setp(box.text_frame.paragraphs[0], "A. Heterogeneity (total bunker CO2; terciles on 1996 values)", 12, True, NAVY)
    rows = [
        ("header", ("Subsample", "\u03b2", "s.e.", "KP F", "N")),
        ("panel", ("By baseline income",)),
        ("row", ("Low income", ("11.431***", R, True), "(1.338)", "35.2", "1,595")),
        ("row", ("Middle income", "5.906***", "(0.782)", "48.2", "1,547")),
        ("row", ("High income", ("0.350", GRAY, False), "(1.106)", "10.1", "1,491")),
        ("panel", ("By baseline connectivity",)),
        ("row", ("Low connectivity", ("9.018***", R, True), "(0.656)", "117.4", "1,452")),
        ("row", ("Middle connectivity", "5.642***", "(0.908)", "29.9", "1,526")),
        ("row", ("High connectivity", ("0.452", GRAY, False), "(2.201)", "1.8", "1,655")),
        ("panel", ("By region",)),
        ("row", ("Europe", ("0.236", GRAY, False), "(0.616)", "55.5", "1,189")),
        ("row", ("Asia\u2013Pacific", "5.424***", "(0.667)", "64.4", "1,110")),
        ("row", ("Africa", ("8.385***", R, True), "(1.688)", "20.2", "1,230")),
        ("row", ("Latin America", "3.229***", "(0.900)", "14.1", "683")),
    ]
    reg_table(s, 0.62, 2.22, [2.3, 1.15, 0.9, 0.85, 0.9], rows, fs=10.5)
    note(s, "Middle East and North America: not identified (KP F < 2). Robust SE in parentheses; "
            "***, **, *: 1, 5, 10%.", y=5.95, w=6.2, size=9.5)
    # right: temporal
    box = tb(s, 6.95, 1.88, 6.0, 0.35)
    setp(box.text_frame.paragraphs[0], "B. Temporal split (2SLS, Feyrer instrument)", 12, True, NAVY)
    rows2 = [
        ("header", ("Sample", "Bunker CO2", "Intl. CO2", "Seat-km", "Intensity", "KP F")),
        ("coef", ("Full, 1996\u20132023", "5.669***", "5.690***", "6.068***", "\u22120.399***", "153.6")),
        ("se", ("N = 4,634", "(0.412)", "(0.415)", "(0.443)", "(0.118)", "")),
        ("coef", ("Expansion era, 1996\u20132007", ("5.613***", R, True), "5.809***", "4.837***", "+0.776***", "75.0")),
        ("se", ("N = 1,883", "(0.665)", "(0.704)", "(0.626)", "(0.249)", "")),
        ("coef", ("Mature era, 2010\u201323 ex-COVID", ("3.005***", R, True), "2.965***", "3.984***", "\u22120.979***", "36.7")),
        ("se", ("N = 2,065", "(0.635)", "(0.610)", "(0.668)", "(0.248)", "")),
        ("coef", ("Unrestricted 1996\u20132009", "6.109***", "6.262***", "5.537***", "+0.571***", "127.7")),
        ("se", ("N = 2,218", "(0.565)", "(0.590)", "(0.534)", "(0.192)", "")),
        ("coef", ("Unrestricted 2010\u20132023", ("1.213", GRAY, False), ("1.051", GRAY, False),
                  ("1.829", GRAY, False), ("\u22120.617", GRAY, False), ("7.0", GRAY, False))),
        ("se", ("N = 2,415 (weak IV)", "(1.394)", "(1.463)", "(1.598)", "(0.739)", "")),
    ]
    reg_table(s, 6.95, 2.22, [2.0, 0.85, 0.8, 0.8, 0.85, 0.6], rows2, fs=10.5)
    rect(s, 6.95, 5.30, 5.78, 0.95, NAVY_T, rounded=True, radius=0.12)
    box = tb(s, 7.10, 5.33, 5.5, 0.9)
    tf = box.text_frame
    p = tf.paragraphs[0]
    setp(p, "Pre vs post difference: ", 11, True, NAVY)
    setp(p, "2.61 (p = 0.005). The elasticity halves as the network matures but stays far above one. "
            "The weak unrestricted post-2010 first stage is a COVID artefact: dropping 2020\u201321 restores F to 36.7.",
         11, False, INK)
    punchline(s, "The elasticity is the sound of a network being built: newly connecting, "
                 "low-income countries carry it; mature networks are near carbon-neutral at the margin.",
              y=6.40, h=0.58, size=14)
    footer(s)
    return s


def slide_table_spill(prs):
    s = new_slide(prs)
    kicker(s, "Results  \u00b7  Table 3")
    title(s, "Spillovers: neighbours' connectivity, own outcomes")
    R = RED
    box = tb(s, 0.62, 1.88, 9.0, 0.35)
    setp(box.text_frame.paragraphs[0],
         "Panel A. 2SLS: neighbour connectivity instrumented by neighbours' Feyrer shifters (inverse-distance weighted)",
         12, True, NAVY)
    rows = [
        ("header", ("Dependent variable (log)", "Bunker CO2", "Intl. CO2", "CO2 / seat-km", "Trade volume")),
        ("coef", ("ln neighbour GACI", ("7.454***", R, True), ("10.237***", R, True), "1.312**", ("2.072 (ns)", GRAY, False))),
        ("se", ("", "(1.994)", "(2.092)", "(0.534)", "(2.011)")),
        ("coef", ("Own Feyrer shifter (reduced form)", "0.254**", "0.076", "\u22120.134***", "0.157")),
        ("se", ("", "(0.122)", "(0.128)", "(0.037)", "(0.124)")),
        ("stat", ("Kleibergen\u2013Paap F", "134.1", "134.0", "134.1", "134.1")),
        ("stat", ("N", "4,623", "4,622", "4,623", "4,623")),
    ]
    reg_table(s, 0.62, 2.22, [2.9, 1.5, 1.5, 1.5, 1.5], rows, fs=11)
    box = tb(s, 0.62, 4.30, 9.0, 0.35)
    setp(box.text_frame.paragraphs[0],
         "Panel B. Margins of the own response to neighbour connectivity (same specification)",
         12, True, NAVY)
    rows2 = [
        ("header", ("Dependent variable (log)", "Flights", "Seat-km", "Aircraft size", "Stage length", "Intl. share")),
        ("coef", ("ln neighbour GACI", ("2.730 (ns)", GRAY, False), "6.141***", ("5.437***", R, True),
                  ("\u22122.026**", R, True), "1.046***")),
        ("se", ("", "(2.216)", "(2.083)", "(1.170)", "(1.008)", "(0.281)")),
    ]
    reg_table(s, 0.62, 4.64, [2.9, 1.2, 1.2, 1.2, 1.2, 1.2], rows2, fs=11)
    # side panel
    panel(s, 9.75, 1.88, 2.98, 4.05, NAVY, GOLD)
    box = tb(s, 9.98, 1.95, 2.7, 4.0)
    tf = box.text_frame
    setp(tf.paragraphs[0], "How to read it", 13, True, WHITE)
    for k, v in [
        ("Neighbour GACI", ": distance-weighted average of all other countries' log hub quality."),
        ("Why reduced form for own", ": own and neighbour shifters are collinear within country, so only one can be instrumented."),
        ("Result", ": +1% in neighbours' hubs raises own CO2 by 7.5% and international CO2 by 10%, with no trade gain."),
        ("Margins", ": not more flights, but bigger aircraft, shorter stages, more international share: feeding the foreign hub."),
    ]:
        p = tf.add_paragraph(); p.space_before = Pt(5)
        setp(p, k, 10.5, True, GOLD); setp(p, v, 10.5, False, LAV)
    note(s, "Notes: country and year FE; controls log population and log sea market access; robust SE in parentheses. "
            "***, **, *: 1, 5, 10% significance. ns = not significant.", y=5.62, w=9.0, size=10)
    punchline(s, "Feeding foreign hubs: the CO2 lands at home, the trade gain does not.",
              y=6.30, h=0.60, size=15)
    footer(s)
    return s


# ================================================================ main
def move_slide(prs, old_index, new_index):
    lst = prs.slides._sldIdLst
    el = list(lst)[old_index]
    lst.remove(el)
    lst.insert(new_index, el)


def renumber_footers(prs):
    n = len(prs.slides)
    pat = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
    cnt = 0
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and pat.match(sh.text_frame.text or ""):
                p = sh.text_frame.paragraphs[0]
                runs = list(p.runs)
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
                runs[0].text = f"{i} / {n}"
                cnt += 1
    return cnt, n


if __name__ == "__main__":
    prs = Presentation(PPTX)
    assert not any(sh.has_text_frame and "E S T I M A T I N G" in sh.text_frame.text
                   for s in prs.slides for sh in s.shapes), "already patched"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    bak = PPTX.replace(".pptx", f"_backup_{stamp}_preRegSlides.pptx")
    shutil.copy2(PPTX, bak); print("backup:", bak)
    n0 = len(prs.slides); print("slides before:", n0)
    D = load_rankings()
    # (anchor slide number in the ORIGINAL deck, builder)
    plan = [
        (15, lambda: slide_co2_ranking(prs, D)),
        (33, lambda: slide_airport_ranking(prs, D)),
        (36, lambda: slide_country_ranking(prs, D)),
        (39, lambda: slide_model(prs)),
        (41, lambda: slide_estimation(prs)),
        (44, lambda: slide_table_main(prs)),
        (49, lambda: slide_table_hetero(prs)),
        (51, lambda: slide_table_spill(prs)),
    ]
    inserted = 0
    for anchor, build in plan:          # ascending anchors: offset grows
        build()
        target = anchor + inserted      # 0-based index = position after anchor
        move_slide(prs, len(prs.slides) - 1, target)
        inserted += 1
        print(f"inserted after original slide {anchor} -> now slide {target + 1}")
    cnt, n = renumber_footers(prs)
    print(f"footers renumbered: {cnt} of {n} slides")
    prs.save(PPTX)
    print("saved", PPTX)
