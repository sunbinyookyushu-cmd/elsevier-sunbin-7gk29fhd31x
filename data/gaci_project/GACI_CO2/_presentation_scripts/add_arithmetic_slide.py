# -*- coding: utf-8 -*-
"""Add ONE slide, 'The arithmetic: how much GACI rose, how much CO2 that explains',
right after the '06 The carbon bill' divider of GACI_CO2_presentation.pptx (2026-09-03).
Re-uses the helpers of add_regression_slides.py. Numbers computed here from
_attribution_scc.csv (dln_cwm, att_tot_t; beta 5.669), gaci_co2_panel.csv and
co2_country_year.csv. Backup first; asserts the slide is not already present.
"""
import os, shutil, datetime
import numpy as np, pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from add_regression_slides import (PPTX, CO2, NAVY, GOLD, INK, GRAY, LAV, RED,
                                   GREEN, NAVY_T, GOLD_T, GREEN_T, RED_T, WHITE, LIGHT, EQF,
                                   tb, setp, rect, new_slide, kicker, title, footer,
                                   punchline, note, panel, eqline, rank_table,
                                   move_slide, renumber_footers, cname)

BETA = 5.669


def data():
    g = pd.read_csv(os.path.join(CO2, "gaci_co2_panel.csv"))
    co = pd.read_csv(os.path.join(CO2, "co2_country_year.csv")).rename(columns={"iso3": "c", "year": "y"})
    att = pd.read_csv(os.path.join(CO2, "_attribution_scc.csv")).set_index("c")
    g23 = g[g.y == 2023].set_index("c").gaci_cwmean
    e96 = co[co.y == 1996].set_index("c").co2_bunker / 1e9
    e23 = co[co.y == 2023].set_index("c").co2_bunker / 1e9
    d = att.join(g23.rename("g23")).join(e96.rename("e96")).join(e23.rename("e23"))
    d["g96"] = d.g23 / np.exp(d.dln_cwm)
    d["gpct"] = (np.exp(d.dln_cwm) - 1) * 100
    d["epct"] = (d.e23 / d.e96 - 1) * 100
    d["att"] = d.att_tot_t / 1e6
    d["attsh"] = d.att / d.e23 * 100
    world = dict(e96=e96.sum(), e23=e23.sum(), att=d.att.sum(),
                 med=d.gpct.median(),
                 wmean=((np.exp(d.dln_cwm) - 1) * d.e23).sum() / d.e23.sum() * 100,
                 nneg=int((d.dln_cwm <= 0).sum()), n=len(d))
    return d, world


def build(prs, d, W):
    s = new_slide(prs)
    kicker(s, "The carbon bill  ·  The arithmetic")
    title(s, "How much did GACI rise, and how much CO2 is that?")
    # ---- three-step flow
    xs = [0.62, 4.30, 8.20]; ws = [3.45, 3.65, 4.53]; y = 1.90; h = 1.32
    fills = [NAVY_T, GOLD_T, RED_T]; acc = [NAVY, GOLD, RED]
    heads = ["1  Connectivity actually rose", "2  × the elasticity", "3  = the carbon bill"]
    bodies = [
        [("Median country: ", f"+{W['med']:.0f}%"), (" GACI 1996–2023; emissions-weighted ", f"+{W['wmean']:.0f}%"),
         (f". {W['nneg']} of {W['n']} countries fell. Korea ", "+104%"), (" (0.88 → 1.79), UAE +76%, Turkey +73%, China +39%, US +4%.", "")],
        [("β = 5.67: ", "each +1% GACI = +5.67% CO2"), (". Counterfactual with 1996 connectivity: ", ""),
         ("E₂₀₂₃ × exp(−β × Δln GACI)", ""), (". Attributed = E₂₀₂₃ × [1 − exp(−β Δln GACI)].", "")],
        [("World aviation CO2 ", f"{W['e96']:.0f} → {W['e23']:.0f} Mt"), (f" (+{(W['e23']/W['e96']-1)*100:.0f}%) over 1996–2023. Attributed to connectivity growth: ", ""),
         (f"{W['att']:.0f} Mt = {W['att']/W['e23']*100:.1f}%", ""), (f" of 2023 emissions, about {W['att']/(W['e23']-W['e96'])*100:.0f}% of the net increase.", "")],
    ]
    for i in range(3):
        rect(s, xs[i], y, ws[i], h, fills[i], rounded=True, radius=0.08)
        rect(s, xs[i], y, 0.08, h, acc[i])
        box = tb(s, xs[i] + 0.18, y + 0.06, ws[i] - 0.3, h - 0.1)
        tf = box.text_frame
        setp(tf.paragraphs[0], heads[i], 12.5, True, acc[i])
        p = tf.add_paragraph(); p.space_before = Pt(3)
        for plain, bold in bodies[i]:
            setp(p, plain, 10.5, False, INK)
            if bold:
                setp(p, bold, 10.5, True, acc[i])
        if i < 2:
            ar = s.shapes.add_shape(13, Inches(xs[i] + ws[i] + 0.02), Inches(y + h / 2 - 0.13),
                                    Inches(0.19), Inches(0.26))   # 13 = RIGHT_ARROW
            ar.fill.solid(); ar.fill.fore_color.rgb = GOLD; ar.line.fill.background()
    # ---- table
    rows = []
    order = ["KOR", "CHN", "ARE", "TUR", "IND", "JPN", "USA", "GBR", "DEU"]
    for c in order:
        r = d.loc[c]
        rows.append((cname(c), f"{r.g96:.2f} → {r.g23:.2f}", f"{r.gpct:+.0f}%",
                     f"{r.e96:.1f} → {r.e23:.1f}", f"{r.epct:+.0f}%",
                     f"{r.att:+.1f}", f"{r.attsh:.0f}%"))
    rows.append(("World (184 countries)", f"median {W['med']:+.0f}%", f"wtd {W['wmean']:+.0f}%",
                 f"{W['e96']:.0f} → {W['e23']:.0f}", f"{(W['e23']/W['e96']-1)*100:+.0f}%",
                 f"{W['att']:+.0f}", f"{W['att']/W['e23']*100:.0f}%"))
    rank_table(s, 0.62, 3.40, [2.05, 1.35, 0.95, 1.45, 0.95, 1.2, 1.2],
               ("Country", "GACI 1996 → 2023", "Δ GACI", "CO2 1996 → 2023 (Mt)",
                "Δ CO2", "Attributed (Mt)", "Share of 2023"),
               rows, fs=10.5, rh=0.245, hl={0, 9}, align_first=(0,))
    # ---- caveat panel
    panel(s, 10.05, 3.40, 2.68, 2.75, LIGHT, GOLD)
    box = tb(s, 10.25, 3.46, 2.42, 2.65)
    tf = box.text_frame
    setp(tf.paragraphs[0], "Read with care", 12, True, NAVY)
    for k, v in [
        ("Marginal elasticity", ": estimated from year-to-year variation, then applied to 28-year changes."),
        ("Saturation", ": the attributed share tends to 100% for big movers (Korea 98%), so country rows are upper bounds."),
        ("Negative rows", ": Germany's GACI fell 9%, so the formula credits it with avoided CO2."),
    ]:
        p = tf.add_paragraph(); p.space_before = Pt(3)
        setp(p, k, 9.5, True, NAVY); setp(p, v, 9.5, False, INK)
    punchline(s, "Most countries moved little; the bill is a story of a few dozen fast connectors. "
                 "Korea doubled its hub quality and doubled its aviation CO2.",
              y=6.40, h=0.56, size=13.5)
    footer(s)
    return s


if __name__ == "__main__":
    prs = Presentation(PPTX)
    assert not any(sh.has_text_frame and "T H E   A R I T H M E T I C" in sh.text_frame.text
                   for s in prs.slides for sh in s.shapes), "already patched"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    bak = PPTX.replace(".pptx", f"_backup_{stamp}_preArith.pptx")
    shutil.copy2(PPTX, bak); print("backup:", bak)
    anchor = None
    for i, s in enumerate(prs.slides):
        txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if "The carbon bill" in txt and "Attribution, who pays" in txt:
            anchor = i; break
    assert anchor is not None, "divider not found"
    d, W = data()
    print({k: (round(v, 1) if isinstance(v, float) else v) for k, v in W.items()})
    build(prs, d, W)
    move_slide(prs, len(prs.slides) - 1, anchor + 1)
    cnt, n = renumber_footers(prs)
    print(f"inserted at slide {anchor + 2}; footers {cnt}/{n}")
    prs.save(PPTX); print("saved")
