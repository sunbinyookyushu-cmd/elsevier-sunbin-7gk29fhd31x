# -*- coding: utf-8 -*-
"""Build GACI_CO2_presentation.pptx in the visual language of GACI_presentation.pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

GACI = r"C:\Users\sunbi\managi-lab Dropbox\Sunbin Yoo\Research Box ^-^\2026\GACI"
CO2 = os.path.join(GACI, "GACI_CO2")
MEDIA = os.path.join(CO2, "_presentation_scripts", "media_from_gaci_pptx")
OUT = os.path.join(GACI, "GACI_CO2_presentation.pptx")

NAVY = RGBColor(0x23, 0x2A, 0x55)
GOLD = RGBColor(0xE2, 0xA8, 0x5A)
INK = RGBColor(0x1E, 0x1E, 0x24)
GRAY = RGBColor(0x6E, 0x6E, 0x6E)
LAV = RGBColor(0xBC, 0xC5, 0xDE)
PUNCH = RGBColor(0xF2, 0xE9, 0xF7)
RED = RGBColor(0xB0, 0x41, 0x3E)
GREEN = RGBColor(0x2E, 0x7D, 0x5B)
RED_T = RGBColor(0xF8, 0xE9, 0xE8)
GREEN_T = RGBColor(0xE6, 0xF2, 0xEC)
GOLD_T = RGBColor(0xFA, 0xF0, 0xE1)
NAVY_T = RGBColor(0xE9, 0xEB, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIDLINE = RGBColor(0x44, 0x4D, 0x77)

prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

FOOTER = "Air Connectivity and Aviation CO2, 1996\u20132023"
TOTAL = 61
_pageno = [0]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    _pageno[0] += 1
    return s


def tb(s, x, y, w, h, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = wrap
    return box


def setp(p, text, size, bold=False, color=INK, font="Calibri", align=None, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return r


def rect(s, x, y, w, h, fill, line=None, rounded=False, radius=None):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if rounded and radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def kicker(s, text, dark=False):
    box = tb(s, 0.60, 0.32, 12.0, 0.5)
    setp(box.text_frame.paragraphs[0], " ".join(text.upper()), 12, True,
         LAV if dark else NAVY)
    rect(s, 0.62, 0.78, 1.40, 0.04, GOLD)


def title(s, text, size=28, color=INK):
    box = tb(s, 0.58, 0.92, 12.10, 1.0)
    setp(box.text_frame.paragraphs[0], text, size, True, color, "Verdana")


def footer(s, dark=False):
    c = LAV if dark else GRAY
    box = tb(s, 12.20, 7.05, 1.0, 0.35)
    setp(box.text_frame.paragraphs[0], f"{_pageno[0]} / {TOTAL}", 10, False, c,
         align=PP_ALIGN.RIGHT)
    if not dark:
        rect(s, 0.62, 7.12, 0.55, 0.05, NAVY)
        box2 = tb(s, 1.25, 7.00, 8.0, 0.35)
        setp(box2.text_frame.paragraphs[0], FOOTER, 10, False, c)


def bullets(s, items, x=0.70, y=2.05, w=12.0, h=4.5, size=19, gap=8, color=INK):
    box = tb(s, x, y, w, h)
    tf = box.text_frame
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        if isinstance(it, tuple):
            plain, boldpart = it
            setp(p, "\u2022  " + plain, size, False, color)
            setp(p, boldpart, size, True, color)
        else:
            setp(p, "\u2022  " + it, size, False, color)
    return box


def punchline(s, text, fill=PUNCH, color=NAVY, y=6.02, x=1.40, w=10.50, h=0.78, size=17):
    shp = rect(s, x, y, w, h, fill, rounded=True, radius=0.5)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    setp(tf.paragraphs[0], text, size, True, color, align=PP_ALIGN.CENTER)
    return shp


def divider(num, name, sub, icon=None):
    s = new_slide(NAVY)
    rect(s, 0.0, 3.05, 0.14, 1.20, GOLD)
    box = tb(s, 0.70, 2.35, 11.5, 2.6)
    tf = box.text_frame
    setp(tf.paragraphs[0], num, 54, True, GOLD, "Verdana")
    p2 = tf.add_paragraph()
    setp(p2, name, 32, True, WHITE, "Verdana")
    p3 = tf.add_paragraph()
    p3.space_before = Pt(6)
    setp(p3, sub, 16, False, LAV)
    if icon:
        ib = tb(s, 9.90, 2.50, 3.0, 2.0, wrap=False)
        setp(ib.text_frame.paragraphs[0], icon, 96, False, GOLD,
             align=PP_ALIGN.CENTER)
    footer(s, dark=True)
    return s


def picture(s, path, y=1.95, max_w=11.6, max_h=4.55, caption=None, cap_y=None,
            x=None):
    from PIL import Image
    im = Image.open(path)
    ar = im.size[0] / im.size[1]
    w = max_w
    h = w / ar
    if h > max_h:
        h = max_h
        w = h * ar
    if x is None:
        x = (13.3333 - w) / 2
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if caption:
        cy = cap_y if cap_y is not None else y + h + 0.08
        box = tb(s, 1.0, cy, 11.33, 0.7)
        setp(box.text_frame.paragraphs[0], caption, 12.5, False, GRAY,
             align=PP_ALIGN.CENTER)
    return w, h


def stat_tile(s, x, y, w, h, big, label, big_color=NAVY, fill=NAVY_T, big_size=30,
              label_size=12.5, label_color=None):
    shp = rect(s, x, y, w, h, fill, rounded=True, radius=0.12)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    setp(tf.paragraphs[0], big, big_size, True, big_color, "Verdana",
         align=PP_ALIGN.CENTER)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    setp(p2, label, label_size, False, label_color or INK, align=PP_ALIGN.CENTER)
    return shp


def chip(s, x, y, w, h, text, fill=NAVY_T, color=NAVY, size=13, bold=True):
    shp = rect(s, x, y, w, h, fill, rounded=True, radius=0.5)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    setp(tf.paragraphs[0], text, size, bold, color, align=PP_ALIGN.CENTER)
    return shp


def make_table(s, x, y, w, h, data, fontsize=12, header_fill=NAVY,
               first_col_bold=True):
    rows, cols = len(data), len(data[0])
    gfx = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w),
                             Inches(h))
    table = gfx.table
    table.first_row = False
    table.horz_banding = False
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            header = (r == 0)
            cell.fill.fore_color.rgb = header_fill if header else WHITE
            if not header and c == 0:
                cell.fill.fore_color.rgb = NAVY_T
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            val = data[r][c]
            color = WHITE if header else INK
            bold = header or (first_col_bold and c == 0)
            if isinstance(val, tuple):
                val, color, bold = val
            setp(tf.paragraphs[0], str(val), fontsize, bold, color,
                 align=PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT)
    return gfx


def arrow_right(s, x, y, w=0.5, h=0.28, color=GOLD):
    shp = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


# ============================= 1. TITLE ==================================
s = new_slide(NAVY)
rect(s, 0.0, 0.0, 0.16, 7.5, GOLD)
box = tb(s, 0.65, 0.95, 4.0, 0.30)
setp(box.text_frame.paragraphs[0], "Research Seminar", 12, True, LAV)
rect(s, 0.67, 1.35, 1.50, 0.05, GOLD)
box = tb(s, 0.62, 1.70, 12.0, 1.8)
tf = box.text_frame
tf.word_wrap = True
setp(tf.paragraphs[0], "How Much Does Air Connectivity", 33, True, WHITE, "Verdana")
p = tf.add_paragraph()
setp(p, "Increase Aviation CO2?  \u2708", 33, True, WHITE, "Verdana")
box = tb(s, 0.65, 3.55, 11.8, 0.6)
setp(box.text_frame.paragraphs[0],
     "Flight-stage emissions for 6,000+ airports, a Feyrer-type instrument, "
     "and the carbon bill of connecting the world, 1996\u20132023",
     18, False, LAV)
rect(s, 0.67, 4.55, 11.5, 0.02, MIDLINE)
box = tb(s, 0.65, 4.75, 11.5, 1.6)
tf = box.text_frame
setp(tf.paragraphs[0],
     "The GACI team  \u00b7  Yifu \u00b7 Sunbin \u00b7 Ray \u00b7 Lisa \u00b7 Jinwoo \u00b7 "
     "Chunan \u00b7 Fangyu \u00b7 Junya", 13, False, LAV)
p = tf.add_paragraph()
setp(p, "(author order TBD)", 11, False, LAV)
footer(s, dark=True)

# ==================== 2. WHAT IS TRANSPORTATION ECONOMICS ================
s = new_slide()
kicker(s, "BACKGROUND")
title(s, "First: what is transportation economics?")
bullets(s, [
    "The economics of moving people and goods: how mobility is demanded, "
    "supplied, priced, and built.",
    "Core questions: which roads, rails, and airports are worth building; "
    "how to price them (congestion charges, fares, fuel taxes); and what "
    "they cost society beyond the ticket (congestion, accidents, noise, "
    "emissions).",
    "Its central insight: transport is a DERIVED demand. Nobody flies for "
    "the pleasure of sitting in seat 43E; people buy access, and access "
    "shapes trade, jobs, and where activity locates.",
], size=18)
box = tb(s, 0.5, 4.45, 12.33, 0.9)
setp(box.text_frame.paragraphs[0],
     "\U0001F697   \U0001F68C   \U0001F684   ✈️   \U0001F6A2",
     36, False, INK, align=PP_ALIGN.CENTER)
punchline(s, "Transport economics asks what a connection is worth, and who "
             "pays its full cost.")
footer(s)

# ============ 3. TRANSPORTATION AND ENERGY ECONOMICS =====================
s = new_slide()
kicker(s, "BACKGROUND")
title(s, "And transportation & energy economics?")
bullets(s, [
    "The intersection of the two fields: transport runs almost entirely on "
    "energy, and mostly on oil.",
    "It studies how transport choices and infrastructure drive energy use "
    "and emissions, and how to decarbonise mobility (fuel taxes, EVs, SAF, "
    "modal shift) without losing the access benefits.",
], y=1.95, h=2.0, size=18, gap=8)
chip(s, 0.85, 4.15, 3.7, 1.30, "\U0001F30D Transport ≈ 1/4 of global "
     "energy-related CO2", NAVY_T, NAVY, 14)
chip(s, 4.85, 4.15, 3.7, 1.30, "✈️ Aviation ≈ 2–3% of global CO2, "
     "and growing", RED_T, RED, 14)
chip(s, 8.85, 4.15, 3.7, 1.30, "\U0001F50B And aviation has no near-term "
     "electrification path", GOLD_T, NAVY, 14)
punchline(s, "Today's talk sits exactly at this intersection: a network "
             "investment (connectivity) meets an energy outcome (jet fuel "
             "burned).", y=5.90)
footer(s)

# ============================= 4. HOOK ===================================
s = new_slide(NAVY)
box = tb(s, 0.5, 1.35, 12.33, 2.6)
setp(box.text_frame.paragraphs[0], "42.5%", 130, True, GOLD, "Verdana",
     align=PP_ALIGN.CENTER)
box = tb(s, 1.2, 4.10, 10.93, 0.9)
setp(box.text_frame.paragraphs[0],
     "of the world's 2023 aviation CO2 traces to connectivity growth since 1996.",
     24, True, WHITE, align=PP_ALIGN.CENTER)
box = tb(s, 1.7, 5.15, 9.93, 1.1)
tf = box.text_frame
tf.word_wrap = True
setp(tf.paragraphs[0],
     "That is 356 Mt of CO2, an external cost of $18\u201368 billion every year.",
     17, False, LAV, align=PP_ALIGN.CENTER)
p = tf.add_paragraph()
p.space_before = Pt(6)
setp(p, "This talk is about where that number comes from, and what could change it.",
     17, False, LAV, align=PP_ALIGN.CENTER)
footer(s, dark=True)

# ==================== 3. WHAT THIS PAPER DOES ============================
s = new_slide()
kicker(s, "TODAY'S TALK")
title(s, "What this paper does")
bullets(s, [
    "We compute CO2 for every scheduled flight in the world, stage by stage, "
    "for 6,000+ airports over 1996\u20132023.",
    "We then ask a causal question: when a country becomes better connected "
    "by air, how much does its aviation CO2 rise?",
    "The answer is: much more than proportionally. A 1% gain in connectivity "
    "raises national aviation CO2 by 5.7%.",
    "We finish with the bill: who owes it, where it is booked, and what "
    "(only fuel switching) can shrink it.",
])
punchline(s, "A data contribution (the emissions panel) and a causal "
             "contribution (the elasticity).")
footer(s)

# ==================== 4. SEQUEL: BENEFIT VS BILL =========================
s = new_slide()
kicker(s, "TODAY'S TALK")
title(s, "One index, two ledgers: the benefit and the bill")
b1 = rect(s, 0.85, 2.05, 5.6, 3.6, NAVY_T, rounded=True, radius=0.08)
tf = b1.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3); tf.margin_top = Inches(0.25)
setp(tf.paragraphs[0], "THE BENEFIT  (our companion trade paper)", 15, True, NAVY)
p = tf.add_paragraph(); p.space_before = Pt(12)
setp(p, "+1.30%", 40, True, NAVY, "Verdana")
p = tf.add_paragraph()
setp(p, "trade openness per 1% rise in hub quality", 14, False, INK)
p = tf.add_paragraph(); p.space_before = Pt(10)
setp(p, "\u2248 $7.8 trillion of 2023 trade, 17% of the world total", 14, False, INK)
b2 = rect(s, 6.90, 2.05, 5.6, 3.6, RED_T, rounded=True, radius=0.08)
tf = b2.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3); tf.margin_top = Inches(0.25)
setp(tf.paragraphs[0], "THE BILL  (this paper)", 15, True, RED)
p = tf.add_paragraph(); p.space_before = Pt(12)
setp(p, "+5.67%", 40, True, RED, "Verdana")
p = tf.add_paragraph()
setp(p, "aviation CO2 per 1% rise in hub quality", 14, False, INK)
p = tf.add_paragraph(); p.space_before = Pt(10)
setp(p, "356 Mt of 2023 CO2, $18\u201368 billion per year", 14, False, INK)
punchline(s, "Same countries, same connectivity index, same IV logic: "
             "the trade paper found the benefit, this paper prices the bill.")
footer(s)

# ============================= 5. ROADMAP ================================
s = new_slide()
kicker(s, "ROADMAP")
title(s, "Seven parts")
parts = [
    ("01", "Motivation", "Policy premise meets carbon", "💡"),
    ("02", "Data", "CO2 for every scheduled flight", "🛰️"),
    ("03", "Building GACI", "Five centralities, PCA, one index", "🕸️"),
    ("04", "Identification", "A geography instrument", "🔍"),
    ("05", "Results", "The elasticity and its anatomy", "📈"),
    ("06", "The carbon bill", "Attribution, mismatch, SAF", "💸"),
    ("07", "Conclusion", "What it implies", "🎯"),
]
for i, (n, t, d, ic) in enumerate(parts):
    if i < 4:
        x = 0.72 + i * 3.05
        y = 2.15
    else:
        x = 2.24 + (i - 4) * 3.05
        y = 4.30
    shp = rect(s, x, y, 2.85, 1.85, NAVY_T if i < 4 else GOLD_T,
               rounded=True, radius=0.10)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18); tf.margin_top = Inches(0.14)
    setp(tf.paragraphs[0], n + "  ", 20, True, GOLD, "Verdana")
    setp(tf.paragraphs[0], ic, 15, False, GOLD)
    p = tf.add_paragraph()
    setp(p, t, 14, True, NAVY, "Verdana")
    p = tf.add_paragraph(); p.space_before = Pt(4)
    setp(p, d, 11.5, False, INK)
footer(s)

# ======================== 6. DIVIDER 01 ==================================
divider("01", "Motivation", "Policy premise meets carbon reality", "💡")

# ======================== 7. MOTIVATION ==================================
s = new_slide()
kicker(s, "MOTIVATION")
title(s, "Everyone wants connectivity; nobody priced the carbon")
bullets(s, [
    "Governments subsidise routes, expand airports, and sign air-service "
    "agreements on the premise that connectivity raises trade, tourism, and "
    "productivity (it does: our companion paper finds +1.3% trade openness "
    "per 1% of connectivity).",
    "But aviation is among the hardest transport sectors to decarbonise: no "
    "near-term electrification path, and emissions booked under separate "
    "bunker conventions, outside national inventories.",
    "The direction is obvious (more connections = more flying). The magnitude "
    "and structure are not, and no causal estimate exists at the global level.",
])
punchline(s, "The carbon consequence of connectivity policy is a first-order "
             "question with no causal answer. Until now.")
footer(s)

# ======================== 8. THREE QUESTIONS =============================
s = new_slide()
kicker(s, "MOTIVATION")
title(s, "Three empirical questions")
qs = [
    ("01", "\U0001F4CF", "Proportionality", "Does a 1% gain in connectivity "
     "raise emissions by more or less than 1%? Hubs consolidate traffic into "
     "fewer, fuller, more efficient movements, so the answer is not "
     "obvious.", NAVY_T, NAVY),
    ("02", "\U0001F5FA️", "Incidence", "Whose emissions respond? Poor or "
     "rich countries? And does the response spill across borders through "
     "the network?", GOLD_T, GOLD),
    ("03", "\U0001F4D2", "Accounting", "Are emissions booked where they are "
     "generated? International aviation follows bunker conventions, not "
     "territory.", RED_T, RED),
]
for i, (n, ic, t, d, fill, accent) in enumerate(qs):
    x = 0.85 + i * 4.0
    shp = rect(s, x, 2.15, 3.7, 3.55, fill, rounded=True, radius=0.08)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.20)
    tf.margin_top = Inches(0.20)
    setp(tf.paragraphs[0], n + "  ", 26, True, accent, "Verdana")
    setp(tf.paragraphs[0], ic, 22, False, accent)
    p = tf.add_paragraph()
    setp(p, t, 18, True, NAVY, "Verdana")
    p = tf.add_paragraph(); p.space_before = Pt(8)
    setp(p, d, 13.5, False, INK)
punchline(s, "Magnitude, geography, and bookkeeping: all three turn out to matter.")
footer(s)

# ======================== 9. DIVIDER 02 ==================================
divider("02", "Data", "CO2 for every scheduled flight on Earth, 1996\u20132023", "🛰️")

# ======================== 10. STAGE DIAGRAM ==============================
s = new_slide()
kicker(s, "DATA")
title(s, "Carbon for every flight, stage by stage")
bullets(s, [
    "Source: OAG schedules, every scheduled commercial flight, Jan 1996 to "
    "Jun 2024.",
    "EEA/EMEP Guidebook: fuel burn by engine type \u00d7 stage length, "
    "\u00d7 3.15 kg CO2 per kg of Jet A.",
], y=1.90, h=1.3, size=16, gap=4)
# stage strip
stages = [
    ("Taxi-out", 1.30, GOLD_T, GOLD),
    ("Take-off", 1.30, GOLD_T, GOLD),
    ("Climb", 1.30, GOLD_T, GOLD),
    ("CRUISE", 3.30, RED_T, RED),
    ("Approach", 1.30, GOLD_T, GOLD),
    ("Taxi-in", 1.30, GOLD_T, GOLD),
]
x = 0.95
y = 3.90
plane = tb(s, 0.0, 3.28, 13.3333, 0.55)
setp(plane.text_frame.paragraphs[0],
     "\u2708" + " " * 130, 22, False, NAVY, align=PP_ALIGN.CENTER)
for name, w, fill, accent in stages:
    shp = rect(s, x, y, w, 0.85, fill, rounded=True, radius=0.25)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    setp(tf.paragraphs[0], name, 13 if w < 2 else 16, True,
         accent if name == "CRUISE" else NAVY, align=PP_ALIGN.CENTER)
    x += w + 0.18
    if name != "Taxi-in":
        arrow_right(s, x - 0.175, y + 0.30, 0.16, 0.24, LAV)
# brackets
chip(s, 0.95, 5.05, 4.06, 0.5, "\U0001F6EB\U0001F6EC LTO phases: 11.8% of "
     "CO2", GOLD_T, NAVY, 13)
chip(s, 5.21, 5.05, 3.30, 0.5, "\U0001F4A8 Cruise: ~88% of CO2", RED_T, RED,
     13)
chip(s, 8.71, 5.05, 3.55, 0.5, "Departure + arrival sides kept separate", NAVY_T,
     NAVY, 12)
punchline(s, "Stage-resolved once, aggregated under any allocation rule, "
             "with no double counting.", y=5.95)
footer(s)

# ======================== 11. VALIDATION TILES ===========================
s = new_slide()
kicker(s, "DATA")
title(s, "Does it add up? Yes, to the tonne (almost) ✅")
stat_tile(s, 0.85, 2.10, 3.7, 1.75, "0.919 Gt", "our 2019 world total vs "
          "ICCT's 0.92 Gt", NAVY)
stat_tile(s, 4.85, 2.10, 3.7, 1.75, "r = 0.99", "log-log match with OWID/ICCT "
          "country series, every year", NAVY)
stat_tile(s, 8.85, 2.10, 3.7, 1.75, "11.8%", "LTO share of fuel burn, right in "
          "the standard range", NAVY)
stat_tile(s, 0.85, 4.10, 3.7, 1.75, "6,000+", "airports with stage-resolved "
          "monthly CO2", GOLD, GOLD_T)
stat_tile(s, 4.85, 4.10, 3.7, 1.75, "184", "countries in the estimation panel",
          GOLD, GOLD_T)
stat_tile(s, 8.85, 4.10, 3.7, 1.75, "4,634", "country-years, 1996\u20132023",
          GOLD, GOLD_T)
punchline(s, "Schedule-based capacity, not load factors: the carbon content "
             "of the network itself.", y=6.15)
footer(s)

# ======================== 12. ALLOCATION RULES ===========================
s = new_slide()
kicker(s, "DATA")
title(s, "Whose emissions are they? Three answers, one dataset")
rules = [
    ("\U0001F6EC Territorial (LTO)", "Each landing-and-take-off phase booked "
     "at its own airport. Physical, local.", NAVY_T, NAVY, ""),
    ("\U0001F6EB Bunker (headline)", "Cruise assigned to the departure "
     "country: the IEA fuel-uplift convention.", GOLD_T, NAVY,
     "\u2605 our headline"),
    ("\u2696\ufe0f 50/50 split", "Cruise shared equally between the endpoint "
     "countries.", RED_T, NAVY, ""),
]
for i, (t, d, fill, accent, tag) in enumerate(rules):
    x = 0.85 + i * 4.0
    shp = rect(s, x, 2.15, 3.7, 3.0, fill, rounded=True, radius=0.08)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.22)
    setp(tf.paragraphs[0], t, 17, True, NAVY, "Verdana")
    if tag:
        p = tf.add_paragraph()
        setp(p, tag, 13, True, GOLD)
    p = tf.add_paragraph(); p.space_before = Pt(8)
    setp(p, d, 14, False, INK)
punchline(s, "Spoiler: the elasticity barely moves across all three "
             "(5.62\u20135.92). The result is not an accounting artefact.",
          y=5.65)
footer(s)

# ======================== 13. LEVELS MAP =================================
s = new_slide()
kicker(s, "DATA")
title(s, "Where aviation CO2 lives, 2023")
picture(s, os.path.join(CO2, "CO2_map_levels_2023.png"), y=2.05, max_w=11.9,
        max_h=4.35,
        caption="National aviation CO2 under the bunker convention, 2023. "
                "This is the outcome variable.")
footer(s)

# ======================== 14. DIVIDER 03 =================================
divider("03", "Building GACI", "What the index is, and exactly how we compute it", "🕸️")

# ======================== 15. CONCEPT ====================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, 'What is "air connectivity", really?')
bullets(s, [
    "It is not how many airports or routes a country has.",
    "It is a country's POSITION in the world air network: how easily, "
    "through how many and how important partners, it can reach everywhere "
    "else.",
    "Two airports with 'international' in the name are not equal: Incheon "
    "plugs into the world's hub network; Yangyang has a handful of "
    "point-to-point routes.",
])
punchline(s, "Connectivity = network position, not a simple count.")
footer(s)

# ======================== 16. CONCEPT: THE INDEX =========================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "How we measure it: the Global Air Connectivity Index")
bullets(s, [
    "GACI summarises each airport's position in the weighted world air "
    "network as a single, time-comparable number (Cheung, Wong & Zhang, "
    "2020).",
    "It blends three ideas:",
], y=1.95, h=1.7, size=18, gap=6)
chip(s, 0.85, 3.85, 3.7, 1.35, "REACH: how many, and how well-placed, are "
     "its connections", NAVY_T, NAVY, 13.5)
chip(s, 4.85, 3.85, 3.7, 1.35, "FLOW: how much passenger traffic moves "
     "through it", GOLD_T, NAVY, 13.5)
chip(s, 8.85, 3.85, 3.7, 1.35, "STANDING: how important are the airports "
     "it connects to", RED_T, RED, 13.5)
punchline(s, "One number that captures reach, flow, and standing in the "
             "global network.", y=5.85)
footer(s)

# ======================== 17. PIPELINE ===================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "How the index is built: the pipeline")
box = tb(s, 0.70, 2.00, 12.0, 0.5)
setp(box.text_frame.paragraphs[0],
     "From raw schedules to one number per airport per year:", 16, False, INK)
picture(s, os.path.join(MEDIA, "image1.png"), y=3.30, max_w=10.2, max_h=3.05,
        caption="Each step is computed for every year, then aggregated to "
                "countries.")
chip(s, 9.42, 5.42, 2.30, 0.40, "today: USE FOR CO2 \u2708", GOLD_T, RED, 11.5)
footer(s)

# ======================== 18. RAW DATA ===================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "The raw data: OAG airline schedules")
picture(s, os.path.join(MEDIA, "image2.png"), y=2.11, max_w=8.93, max_h=4.47,
        x=0.32)
shp = rect(s, 9.56, 2.25, 3.45, 1.45, NAVY_T, rounded=True, radius=0.12)
tf = shp.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.15)
setp(tf.paragraphs[0], "Connected dummy", 14, True, NAVY)
p = tf.add_paragraph(); p.space_before = Pt(6)
setp(p, "Total seat capacity (the link weight)", 14, True, NAVY)
box = tb(s, 9.56, 3.95, 3.45, 1.6)
setp(box.text_frame.paragraphs[0],
     "Route-segment level, every scheduled commercial flight, "
     "1996\u20132024H1.", 12.5, False, GRAY)
footer(s)

# ======================== 19. STEP 0 =====================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 0: building a clean weighted network")
bullets(s, [
    "Source: global airline schedule data at the route-segment level (OAG; "
    "it is big).",
    "Keep only scheduled commercial service; exclude military, charter, "
    "seaplanes, heliports, and airports that closed in-window.",
    "Aggregate segments per airport pair per year into one link; the link "
    "weight is passenger seat capacity.",
    "Two matrices per year: adjacency A (who connects) and weight W (how "
    "much).",
], y=1.90, h=2.3, size=16, gap=5)
AP = ["", "ICN", "GMP", "CJU", "PUS", "YNY"]
adj = [
    AP,
    ["ICN", "–", "0", "1", "0", "0"],
    ["GMP", "0", "–", "1", "1", "0"],
    ["CJU", "1", "1", "–", "1", "1"],
    ["PUS", "0", "1", "1", "–", "0"],
    ["YNY", "0", "0", "1", "0", "–"],
]
wgt = [
    AP,
    ["ICN", "–", "0", "0.5", "0", "0"],
    ["GMP", "0", "–", "17.0", "4.0", "0"],
    ["CJU", "0.5", "17.0", "–", "2.5", "0.1"],
    ["PUS", "0", "4.0", "2.5", "–", "0"],
    ["YNY", "0", "0", "0.1", "0", "–"],
]
box = tb(s, 1.55, 4.15, 4.4, 0.35)
setp(box.text_frame.paragraphs[0], "Adjacency A: who connects", 13, True, NAVY)
make_table(s, 1.55, 4.50, 4.4, 1.95, adj, fontsize=11)
box = tb(s, 7.15, 4.15, 4.7, 0.35)
setp(box.text_frame.paragraphs[0], "Weight W: seat capacity (millions, "
     "illustrative)", 13, True, NAVY)
make_table(s, 7.15, 4.50, 4.7, 1.95, wgt, fontsize=11)
box = tb(s, 1.55, 6.60, 10.5, 0.4)
setp(box.text_frame.paragraphs[0],
     "The Korean corner of the matrix: Gimpo–Jeju alone is the world's "
     "busiest air route.", 12.5, False, GRAY)
footer(s)

# ======================== 20. STEP 1A ====================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 1a: topological indicators (position)")
bullets(s, [
    "Degree: the number of direct connections (raw reach).",
    "Closeness: inverse of the average number of transfers needed to reach "
    "all other airports (accessibility).",
], y=1.60, h=1.1, size=16, gap=4)
box = tb(s, 0.60, 2.62, 5.9, 0.35)
setp(box.text_frame.paragraphs[0],
     "Toy domestic network: GMP–CJU, GMP–PUS, GMP–YNY, CJU–PUS", 13, True,
     NAVY)
hop = [
    ["Hops", "GMP", "CJU", "PUS", "YNY", "Avg"],
    ["GMP", "–", "1", "1", "1", "1.00"],
    ["CJU", "1", "–", "1", "2", "1.33"],
    ["PUS", "1", "1", "–", "2", "1.33"],
    ["YNY", "1", "2", "2", "–", "1.67"],
]
make_table(s, 0.60, 3.00, 5.55, 1.70, hop, fontsize=11)
clo = [
    ["Airport", "Closeness", "Reading"],
    ["GMP", "1.00", "Hub: everyone in one hop"],
    ["CJU", "0.75", "Intermediate"],
    ["PUS", "0.75", "Intermediate"],
    ["YNY", "0.60", "Spoke: always via Gimpo"],
]
make_table(s, 0.60, 4.90, 5.55, 1.80, clo, fontsize=11)
box = tb(s, 6.75, 2.62, 6.0, 0.7)
setp(box.text_frame.paragraphs[0],
     '+ Eigenvector: how much am I connected to "BIG" airports?', 15, True,
     NAVY)
eig = [
    ["Airport", "Degree (# friends)", "Eigenvector (rank of friends)"],
    ["GMP", "3", ("High", NAVY, True)],
    ["CJU", "2", ("High", NAVY, True)],
    ["PUS", "2", ("High", NAVY, True)],
    ["YNY", "1", ("Low", GRAY, True)],
]
make_table(s, 6.75, 3.35, 5.9, 2.05, eig, fontsize=11.5)
box = tb(s, 6.75, 5.60, 5.9, 1.0)
setp(box.text_frame.paragraphs[0],
     "Yangyang's one friend is a big hub, but one friend is still one "
     "friend.", 13, False, GRAY)
footer(s)

# ======================== 21. STEP 1B ====================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 1b: volumetric indicators (passenger flow)")
bullets(s, [
    "Flow betweenness: the share of passengers transferring through the "
    "airport; the hub / transfer role (Gimpo < Incheon: transfers "
    "concentrate at ICN).",
    "Regional importance: average connection intensity within its own "
    "region (Gimpo > Incheon: the domestic backbone).",
    "These two flow measures make GACI flow-aware, not purely topological.",
])
punchline(s, "Position says where you sit; flow says how much actually "
             "moves through you.")
footer(s)

# ================= 21b. FIVE INDICATORS, KOREAN AIRPORTS =================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "The five indicators, read through Korean airports")
H = (NAVY, True)
M = (INK, False)
L = (GRAY, False)


def lv(txt, style):
    return (txt, style[0], style[1])


kor = [
    ["Airport", "Degree", "Closeness", "Eigenvector", "Flow betw.",
     "Regional imp."],
    ["ICN  Incheon", lv("High", H), lv("High", H), lv("High", H),
     lv("High", H), lv("Mid", M)],
    ["GMP  Gimpo", lv("Mid", M), lv("Mid", M), lv("Mid", M), lv("Low", L),
     lv("High", H)],
    ["CJU  Jeju", lv("Mid", M), lv("Mid", M), lv("Low", L), lv("Low", L),
     lv("High", H)],
    ["PUS  Gimhae", lv("Mid", M), lv("Mid", M), lv("Mid", M), lv("Low", L),
     lv("Mid", M)],
    ["YNY  Yangyang", lv("Low", L), lv("Low", L), lv("Low", L), lv("Low", L),
     lv("Low", L)],
]
make_table(s, 0.85, 2.05, 11.6, 2.85, kor, fontsize=13)
box = tb(s, 0.85, 5.10, 11.6, 0.9)
tf = box.text_frame
tf.word_wrap = True
setp(tf.paragraphs[0],
     "ICN: global gateway, strong on every margin.  GMP and CJU: enormous "
     "traffic (the world's busiest route) but little global position: "
     "volume is not connectivity.  YNY: connected on paper only.",
     13.5, False, INK)
punchline(s, "Five airports, five different network stories: no single "
             "indicator captures them all, so we combine all five.",
          y=6.15, h=0.70)
footer(s)

# ======================== 22. FLOW BETWEENNESS ===========================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Flow betweenness: passengers as electrical current")
picture(s, os.path.join(MEDIA, "image7.jpg"), y=2.05, max_w=9.6, max_h=4.20,
        caption="Each route is a resistor with resistance inversely "
                "proportional to its passenger intensity; betweenness = "
                "current through the node. Incheon plays this HKG-style "
                "bridge role between North America and Southeast Asia. "
                "Source: Cheung et al. (2020).")
footer(s)

# ======================== 23. STEP 1 FORMULAS ============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 1: the exact formulas")
picture(s, os.path.join(MEDIA, "image8.png"), y=1.95, max_w=9.0, max_h=4.75)
footer(s)

# ======================== 24. CORRELATIONS ===============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "The five indicators are related but not redundant")
picture(s, os.path.join(MEDIA, "image9.jpg"), y=1.85, max_w=8.0, max_h=4.45,
        caption="Pairwise rank correlations 0.71\u20130.87: each adds "
                "information, so we combine rather than pick one. Source: "
                "Cheung et al. (2020).")
footer(s)

# ======================== 25. STEP 2 PCA =================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 2: combine the five into one index (PCA)")
bullets(s, [
    "Five indicators, different scales, partly correlated.",
    "GACI = the first principal component: the single weighted combination "
    "capturing the most variance across airports.",
    "On our 1996\u20132023 panel, PC1 explains 74% of total variance, so "
    "one number retains most of the information.",
])
punchline(s, "One index, five ingredients, no arbitrary weights.")
footer(s)

# ======================== 26. STEP 2 FORMULA =============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 2: the aggregation formula")
picture(s, os.path.join(MEDIA, "image10.png"), y=2.00, max_w=9.0, max_h=4.60)
footer(s)

# ======================== 27. PCA WEIGHTS ================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "PCA weights, computed on our panel")
picture(s, os.path.join(MEDIA, "image11.png"), y=1.95, max_w=9.0, max_h=4.30,
        caption="First-component loadings on our 1996\u20132023 data: all "
                "positive and similar; PC1 explains 74% of variance.")
footer(s)

# ======================== 28. TIME COMPARABILITY =========================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Making GACI comparable over time")
bullets(s, [
    "A naive yearly PCA would re-scale every year, so scores could not be "
    "compared across years.",
    "We standardise so that rankings WITHIN a year and relative scores "
    "ACROSS years are both preserved.",
    "The result: one index that tracks an airport or country up and down "
    "the global hierarchy from 1996 to 2023.",
])
punchline(s, "Essential here: the CO2 question is about changes over 28 "
             "years.")
footer(s)

# ======================== 29. FACE VALIDITY ==============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Face validity: high GACI = the world's major hubs")
picture(s, os.path.join(MEDIA, "image12.jpg"), y=1.85, max_w=8.0, max_h=4.45,
        caption="GACI tracks passenger volume; the top of the scale is ATL, "
                "PEK, DXB, LHR, HKG, SIN. Source: Cheung et al. (2020).")
footer(s)

# ======================== 30. STEP 3 =====================================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 3: from airports to countries")
bullets(s, [
    "Hub quality = capacity-weighted MEAN (headline): average gateway "
    "quality, not inflated by having many tiny airports.",
    "Total connectivity = SUM: the overall network footprint, larger for "
    "big countries.",
    "In this paper the headline treatment is hub quality; sum, max, and "
    "unweighted mean all reappear as robustness checks (spoiler: they "
    "agree).",
])
punchline(s, "ln GACI (capacity-weighted mean) is the right-hand side of "
             "everything that follows.")
footer(s)

# ======================== 31. STEP 3 FORMULAS ============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "Step 3: the aggregation formulas")
picture(s, os.path.join(MEDIA, "image13.png"), y=2.00, max_w=9.0, max_h=4.60)
footer(s)

# ======================== 32. GACI MAP 2023 ==============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "The built index, mapped: connectivity by country, 2023")
picture(s, os.path.join(MEDIA, "image14.png"), y=2.05, max_w=11.9,
        max_h=4.35,
        caption="Total connectivity (country sum), 2023. Yellow = highly "
                "connected.")
footer(s)

# ======================== 33. TREATMENT MAP ==============================
s = new_slide()
kicker(s, "BUILDING GACI")
title(s, "The treatment: who got connected, 1996\u20132023")
picture(s, os.path.join(CO2, "CO2_map_dlngaci.png"), y=2.05, max_w=11.9,
        max_h=4.35,
        caption="Change in log connectivity, 1996\u20132023. The Gulf, "
                "Turkey, China, and Korea connect; parts of Europe quietly "
                "disconnect.")
footer(s)

# ======================== 34. DIVIDER 04 =================================
divider("04", "Identification", "Breaking reverse causality with geography", "🔍")

# ======================== 17. ENDOGENEITY ================================
s = new_slide()
kicker(s, "IDENTIFICATION")
title(s, "The problem: airlines follow the economy")
bullets(s, [
    "Airlines add capacity where income, trade, and travel demand are already "
    "growing: reverse causality, upward and downward biases at once.",
    "OLS would also be attenuated by measurement error in the index (it is: "
    "OLS = 3.59 vs 2SLS = 5.67).",
    "We need variation in connectivity that has nothing to do with a "
    "country's economic trajectory.",
])
punchline(s, "Our answer: geography that was fixed before the boom, "
             "interacted with a global technology cycle.")
footer(s)

# ======================== 18. IV DIAGRAM =================================
s = new_slide()
kicker(s, "IDENTIFICATION")
title(s, "A Feyrer-type air-vs-sea geography instrument")


def ivbox(x, y, w, h, head, body, fill=NAVY_T, accent=NAVY):
    shp = rect(s, x, y, w, h, fill, rounded=True, radius=0.10)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.14); tf.margin_right = Inches(0.14)
    setp(tf.paragraphs[0], head, 14.5, True, accent, align=PP_ALIGN.CENTER)
    if body:
        p = tf.add_paragraph()
        setp(p, body, 11.5, False, INK, align=PP_ALIGN.CENTER)
    return shp


ivbox(0.75, 2.20, 3.10, 1.30, "\U0001F5FA️ 1996 air-vs-sea market-access "
      "geography", "fixed before the sample", GOLD_T, NAVY)
box = tb(s, 3.90, 2.45, 0.55, 0.7)
setp(box.text_frame.paragraphs[0], "\u00d7", 30, True, GRAY, align=PP_ALIGN.CENTER)
ivbox(4.45, 2.20, 3.10, 1.30, "\U0001F310 World aviation technology cycle",
      "global, no country steers it", GOLD_T, NAVY)
arrow_right(s, 7.68, 2.68, 0.55, 0.34)
ivbox(8.35, 2.20, 2.10, 1.30, "GACI", "first stage", NAVY_T, NAVY)
arrow_right(s, 10.55, 2.68, 0.55, 0.34)
ivbox(11.20, 2.20, 1.75, 1.30, "\U0001F4A8 Aviation CO2", "", RED_T, RED)
box = tb(s, 0.75, 3.75, 12.2, 0.5)
setp(box.text_frame.paragraphs[0],
     "Controls: log population, log sea market access, country and year "
     "fixed effects.", 13, False, GRAY)
chip(s, 0.85, 4.45, 3.7, 1.05, "Kleibergen\u2013Paap F = 153.6", NAVY_T, NAVY, 15)
chip(s, 4.85, 4.45, 3.7, 1.05, "Stress test (cycle \u00d7 baseline size): "
     "F 120 \u2192 106. Tourism IV: 14 \u2192 0, dies.", GOLD_T, NAVY, 12.5)
chip(s, 8.85, 4.45, 3.7, 1.05, "Conley bounds: tolerates direct effects up to "
     "85% of the reduced form", GREEN_T, GREEN, 12.5)
punchline(s, "The only instrument we tried that survives its own execution "
             "squad.", y=5.90)
footer(s)

# ==================== EXCLUSION RESTRICTION ==============================
s = new_slide()
kicker(s, "IDENTIFICATION")
title(s, "The fine print: the exclusion restriction")
box = tb(s, 0.70, 1.95, 12.0, 0.45)
setp(box.text_frame.paragraphs[0],
     "The instrument may affect CO2 through ONE door only: connectivity.",
     16, True, NAVY)


def exbox(x, y, w, h, head, fill, accent):
    shp = rect(s, x, y, w, h, fill, rounded=True, radius=0.12)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
    setp(tf.paragraphs[0], head, 14.5, True, accent, align=PP_ALIGN.CENTER)
    return shp


exbox(0.85, 2.60, 3.30, 1.00, "\U0001F50D Instrument (geography × cycle)",
      GOLD_T, NAVY)
arrow_right(s, 4.35, 2.92, 0.75, 0.36, GREEN)
box = tb(s, 4.28, 2.42, 0.9, 0.45, wrap=False)
setp(box.text_frame.paragraphs[0], "✓", 22, True, GREEN,
     align=PP_ALIGN.CENTER)
exbox(5.30, 2.60, 2.60, 1.00, "✈️ Connectivity (GACI)", NAVY_T, NAVY)
arrow_right(s, 8.10, 2.92, 0.75, 0.36, GREEN)
box = tb(s, 8.03, 2.42, 0.9, 0.45, wrap=False)
setp(box.text_frame.paragraphs[0], "✓", 22, True, GREEN,
     align=PP_ALIGN.CENTER)
exbox(9.05, 2.60, 3.30, 1.00, "\U0001F4A8 Aviation CO2", RED_T, RED)
# forbidden direct path underneath
rect(s, 2.50, 4.05, 7.60, 0.06, RED)
tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(10.10), Inches(3.93),
                         Inches(0.45), Inches(0.30))
tri.fill.solid(); tri.fill.fore_color.rgb = RED; tri.line.fill.background()
tri.shadow.inherit = False
xc = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.95), Inches(3.72),
                        Inches(0.70), Inches(0.70))
xc.fill.solid(); xc.fill.fore_color.rgb = WHITE
xc.line.color.rgb = RED; xc.line.width = Pt(2.5)
xc.shadow.inherit = False
tf = xc.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
setp(tf.paragraphs[0], "✗", 24, True, RED, align=PP_ALIGN.CENTER)
box = tb(s, 2.50, 4.42, 7.60, 0.4)
setp(box.text_frame.paragraphs[0],
     "no direct path to CO2 allowed (untestable by assumption, so we "
     "attack it three ways)", 13, False, RED, align=PP_ALIGN.CENTER)
threats = [
    ("\U0001F6A2 Geography also moves SEA trade",
     "controlled directly: ln sea market access in every regression"),
    ("\U0001F4B0 Big, rich countries ride the global cycle anyway",
     "stress test: add cycle × 1996 pop, GDP, capacity. F 120 → 106, "
     "β stable (the test that killed our other instruments)"),
    ("\U0001F52C Some direct effect might still remain",
     "Conley bounds: results survive direct effects up to 85% of the "
     "reduced form"),
]
for i, (thr, ans) in enumerate(threats):
    y = 5.00 + i * 0.62
    shp = rect(s, 0.85, y, 5.55, 0.54, RGBColor(0xEF, 0xEF, 0xEF),
               rounded=True, radius=0.3)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.14)
    setp(tf.paragraphs[0], thr, 11.5, True, GRAY)
    arrow_right(s, 6.50, y + 0.14, 0.30, 0.26, GOLD)
    shp = rect(s, 6.90, y, 5.60, 0.54, GREEN_T, rounded=True, radius=0.3)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.14)
    setp(tf.paragraphs[0], "✓ " + ans, 10.5, True, GREEN)
footer(s)

# ======================== 19. RF QUINTILE ================================
s = new_slide()
kicker(s, "IDENTIFICATION")
title(s, "The raw first look: shifter up, emissions up")
picture(s, os.path.join(CO2, "CO2_rf_quintile.png"), y=2.30, max_w=11.9,
        max_h=3.8,
        caption="Reduced form by instrument quintile: countries with stronger "
                "geography-driven connectivity growth emit more.")
footer(s)

# ======================== 20. DIVIDER 04 =================================
divider("05", "Results", "The elasticity, and everything inside it", "📈")

# ======================== 21. HERO 5.7 ===================================
s = new_slide(NAVY)
box = tb(s, 0.5, 1.10, 12.33, 2.3)
setp(box.text_frame.paragraphs[0], "+1%  \u2192  +5.7%", 92, True, GOLD,
     "Verdana", align=PP_ALIGN.CENTER)
box = tb(s, 1.0, 3.65, 11.33, 0.8)
setp(box.text_frame.paragraphs[0],
     "connectivity up one percent, national aviation CO2 up 5.67 percent "
     "(s.e. 0.41)", 22, True, WHITE, align=PP_ALIGN.CENTER)
box = tb(s, 1.6, 4.75, 10.13, 1.3)
tf = box.text_frame
tf.word_wrap = True
setp(tf.paragraphs[0],
     "The benchmark is an elasticity of one: hubs could in principle absorb "
     "growth efficiently.", 16, False, LAV, align=PP_ALIGN.CENTER)
p = tf.add_paragraph(); p.space_before = Pt(6)
setp(p, "An elasticity of one is rejected at any conventional level.",
     16, True, LAV, align=PP_ALIGN.CENTER)
footer(s, dark=True)

# ======================== 22. SIX OUTCOME TILES ==========================
s = new_slide()
kicker(s, "RESULTS")
title(s, "However you count it: far above one")
tiles = [
    ("+5.67***", "Bunker CO2 (headline)", RED, RED_T),
    ("+5.92***", "LTO-only (territorial)", RED, RED_T),
    ("+5.62***", "50/50 cruise split", RED, RED_T),
    ("+5.69***", "International CO2", RED, RED_T),
    ("+6.07***", "Seat-kilometres", NAVY, NAVY_T),
    ("\u22120.40***", "CO2 per seat-km", GREEN, GREEN_T),
]
for i, (big, lab, c, f) in enumerate(tiles):
    col = i % 3
    row = i // 3
    stat_tile(s, 0.85 + col * 4.0, 2.05 + row * 1.90, 3.7, 1.65, big, lab, c, f,
              big_size=27)
box = tb(s, 0.85, 5.80, 11.7, 0.45)
setp(box.text_frame.paragraphs[0],
     "Aggregation checks: GACI sum +2.28***, max +4.82***, mean +5.90*** "
     "(KP F 90\u2013172). 2SLS, Feyrer IV, N = 4,634.", 13, False, GRAY)
punchline(s, "Allocation rule, aggregation, instrument: nothing pushes the "
             "elasticity anywhere near one.", y=6.30, h=0.70)
footer(s)

# ======================== 23. SCALE VS TECHNIQUE =========================
s = new_slide()
kicker(s, "RESULTS")
title(s, "Efficiency is real. It is also hopelessly outgunned.")
# bar graphic: scale +6.07 vs technique -0.40
base_x = 3.6
maxbar = 8.2
box = tb(s, 0.7, 2.35, 2.8, 0.6)
setp(box.text_frame.paragraphs[0], "Scale (seat-km)", 16, True, NAVY)
bar = rect(s, base_x, 2.30, maxbar * 6.07 / 6.07, 0.70, RED, rounded=True,
           radius=0.3)
tf = bar.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
setp(tf.paragraphs[0], "+6.07%", 18, True, WHITE, align=PP_ALIGN.CENTER)
box = tb(s, 0.7, 3.55, 2.8, 0.6)
setp(box.text_frame.paragraphs[0], "Technique (CO2 per seat-km)", 14, True, NAVY)
bar = rect(s, base_x, 3.50, 0.60, 0.70, GREEN, rounded=True, radius=0.3)
box = tb(s, base_x + 0.72, 3.58, 2.5, 0.55, wrap=False)
setp(box.text_frame.paragraphs[0], "\u22120.40%", 17, True, GREEN)
bullets(s, [
    "The same 1% of connectivity raises seat-kilometres by 6.07%, faster "
    "than emissions rise.",
    "So emissions per seat-km fall by 0.40%: hubbing does deliver larger, "
    "fuller, longer-range operations.",
    ("But the efficiency margin offsets ", "less than one tenth of the scale "
     "response."),
], y=4.55, size=16, gap=6)
punchline(s, "The emissions consequence of connectivity is a volume story, "
             "not an efficiency story.", y=6.25, h=0.70)
footer(s)

# ======================== 24. EFFICIENCY CURVE ===========================
s = new_slide()
kicker(s, "RESULTS")
title(s, "Zoom to the airport: the efficiency curve")
picture(s, os.path.join(CO2, "CO2_efficiency_curve.png"), y=1.90, max_w=10.5,
        max_h=3.70,
        caption="2023 cross-section by connectivity ventile: median intensity "
                "falls 284 \u2192 77 kg per 1,000 seat-km while the "
                "international share climbs from 1% to 60%.", cap_y=5.80)
punchline(s, "Efficiency is real at the node, and dominated by scale at the "
             "aggregate.", y=6.50, h=0.60, size=15)
footer(s)

# ======================== 25. HETEROGENEITY ==============================
s = new_slide()
kicker(s, "RESULTS")
title(s, "Who drives it: the newly connecting world")
picture(s, os.path.join(CO2, "CO2_hetero_coefplot.png"), y=1.85, max_w=6.6,
        max_h=4.45, x=0.60)
box = tb(s, 7.55, 2.15, 5.2, 4.1)
tf = box.text_frame
tf.word_wrap = True
hpts = [
    ("By 1996 income:  ", "11.4 (bottom) \u00b7 5.9 (middle) \u00b7 0.35 ns "
     "(top)"),
    ("By region:  ", "Africa 8.4 \u00b7 Asia\u2013Pacific 5.4 \u00b7 Latin "
     "America 3.2 \u00b7 Europe 0.24 ns"),
    ("Efficiency gains  ", "also appear only in the bottom income tercile."),
]
first = True
for lead, rest in hpts:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(14)
    setp(p, lead, 16, True, NAVY)
    setp(p, rest, 16, False, INK)
punchline(s, "For mature networks, connectivity growth is roughly "
             "carbon-neutral at the margin. The elasticity is the sound of "
             "countries entering the network.", y=6.45, h=0.72, size=14)
footer(s)

# ======================== 26. TEMPORAL ===================================
s = new_slide()
kicker(s, "RESULTS")
title(s, "And it is fading as the network matures")
picture(s, os.path.join(CO2, "CO2_temporal_coefplot.png"), y=1.95, max_w=7.6,
        max_h=4.15, x=0.60)
box = tb(s, 8.45, 2.25, 4.3, 4.0)
tf = box.text_frame
tf.word_wrap = True
tpts = [
    ("5.61 \u2192 3.01  ", "excluding crisis years: the elasticity halves from "
     "1996\u20132007 to 2010\u20132023 (p = 0.005)."),
    ("+0.78 \u2192 \u22120.98  ", "the efficiency margin only appears in the mature "
     "era."),
    ("COVID alone  ", "breaks the late-period first stage (F 36.7 \u2192 7.0 with "
     "2020\u201321 included)."),
]
first = True
for lead, rest in tpts:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(14)
    setp(p, lead, 15, True, NAVY)
    setp(p, rest, 14, False, INK)
punchline(s, "Still far above one in both eras: maturity slows the meter, it "
             "does not stop it.", y=6.45, h=0.65, size=15)
footer(s)

# ======================== 27. COMPOSITION ================================
s = new_slide()
kicker(s, "RESULTS")
title(s, "Anatomy: volume carries it, hub quality recomposes it")
box = tb(s, 0.85, 1.95, 6.0, 0.5)
setp(box.text_frame.paragraphs[0], "What mediates the effect?", 16, True, NAVY)
stat_tile(s, 0.85, 2.50, 3.55, 1.55, "86.5%", "of the effect runs through "
          "seat-kilometres", NAVY, NAVY_T, big_size=26)
stat_tile(s, 4.60, 2.50, 3.55, 1.55, "83.9%", "through flight frequency",
          NAVY, NAVY_T, big_size=26)
stat_tile(s, 8.35, 2.50, 3.55, 1.55, "~0%", "through the international share",
          GRAY, RGBColor(0xEF, 0xEF, 0xEF), big_size=26)
box = tb(s, 0.85, 4.30, 11.5, 0.5)
setp(box.text_frame.paragraphs[0],
     "A different lever, hub quality (secondary tourism IV), moves the mix "
     "instead:", 16, True, NAVY)
chip(s, 0.85, 4.85, 3.55, 0.85, "Total CO2: no response", NAVY_T, NAVY, 14)
chip(s, 4.60, 4.85, 3.55, 0.85, "International CO2: +3.16***", RED_T, RED, 14)
chip(s, 8.35, 4.85, 3.55, 0.85, "CO2 per seat-km: +1.23***", RED_T, RED, 14)
punchline(s, "Market access scales flying; hub quality tilts it toward "
             "international long-haul.", y=6.15, h=0.70)
footer(s)

# ======================== 28. SPILLOVER ==================================
s = new_slide()
kicker(s, "RESULTS")
title(s, "Your neighbour's hub, your emissions")
# diagram: neighbours -> own country
c1 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.1), Inches(2.35), Inches(2.5),
                        Inches(1.7))
c1.fill.solid(); c1.fill.fore_color.rgb = GOLD_T; c1.line.color.rgb = GOLD
c1.shadow.inherit = False
tf = c1.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
setp(tf.paragraphs[0], "✈️ Neighbours' connectivity +1%", 14, True, NAVY,
     align=PP_ALIGN.CENTER)
arrow_right(s, 3.80, 2.98, 1.0, 0.44, GOLD)
c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.95), Inches(2.35), Inches(2.5),
                        Inches(1.7))
c2.fill.solid(); c2.fill.fore_color.rgb = RED_T; c2.line.color.rgb = RED
c2.shadow.inherit = False
tf = c2.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
setp(tf.paragraphs[0], "\U0001F4A8 Own aviation CO2 +7.45***", 15, True, RED,
     align=PP_ALIGN.CENTER)
chips = [
    ("International CO2 +10.2***", RED_T, RED),
    ("Trade volume: +2.1, ns", RGBColor(0xEF, 0xEF, 0xEF), GRAY),
    ("KP F = 134", NAVY_T, NAVY),
]
for i, (t, f, c) in enumerate(chips):
    chip(s, 8.00, 2.30 + i * 0.65, 4.45, 0.55, t, f, c, 13)
box = tb(s, 0.85, 4.45, 11.5, 0.45)
setp(box.text_frame.paragraphs[0],
     "The margins are the converse of the own-network response:", 15, True, NAVY)
chip(s, 0.85, 4.95, 2.80, 0.80, "Flights: ns", RGBColor(0xEF, 0xEF, 0xEF),
     GRAY, 13)
chip(s, 3.85, 4.95, 2.80, 0.80, "Aircraft size +5.4***", GOLD_T, NAVY, 13)
chip(s, 6.85, 4.95, 2.80, 0.80, "Intl share +1.0***", GOLD_T, NAVY, 13)
chip(s, 9.85, 4.95, 2.80, 0.80, "Stage length \u22122.0**", GOLD_T, NAVY, 13)
punchline(s, "Feeding foreign hubs: the CO2 lands at home, the trade gains "
             "do not. National accounting misses a regional externality.",
          y=6.15, h=0.72, size=15)
footer(s)

# ======================== 29. DIVIDER 05 =================================
divider("06", "The carbon bill", "Attribution, who pays, and the only real lever", "💸")

# ======================== 30. ATTRIBUTED MAP =============================
s = new_slide()
kicker(s, "THE CARBON BILL")
title(s, "356 Mt: the CO2 the connectivity boom built")
picture(s, os.path.join(CO2, "CO2_map_attributed_2023.png"), y=2.05, max_w=11.9,
        max_h=4.35,
        caption="2023 aviation CO2 attributed to post-1996 connectivity "
                "growth: 356 Mt, 42.5% of the world total. White = zero.")
footer(s)

# ==================== SOCIAL COST OF CARBON EXPLAINER ====================
s = new_slide()
kicker(s, "THE CARBON BILL")
title(s, "The price tag: the social cost of carbon")
box = tb(s, 0.70, 1.95, 12.0, 0.5)
setp(box.text_frame.paragraphs[0],
     "SCC = the present value of all future damage caused by one extra "
     "tonne of CO2.", 16, True, NAVY)


def sccbox(x, w, head, body, fill, accent):
    shp = rect(s, x, 2.60, w, 1.55, fill, rounded=True, radius=0.10)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
    setp(tf.paragraphs[0], head, 14.5, True, accent, align=PP_ALIGN.CENTER)
    p = tf.add_paragraph()
    setp(p, body, 11.5, False, INK, align=PP_ALIGN.CENTER)


sccbox(0.85, 3.40, "\U0001F3ED 1 tonne of CO2, emitted today",
       "one flight's worth at a time", NAVY_T, NAVY)
arrow_right(s, 4.45, 3.22, 0.55, 0.34)
sccbox(5.15, 3.60, "\U0001F321️ A century of damages",
       "heat, crops, storms, sea level, health, productivity", RED_T, RED)
arrow_right(s, 8.95, 3.22, 0.55, 0.34)
sccbox(9.65, 2.85, "\U0001F4B5 Discounted back to today",
       "= one price per tonne", GREEN_T, GREEN)
box = tb(s, 0.70, 4.40, 12.0, 0.4)
setp(box.text_frame.paragraphs[0],
     "Three standard price tags (2020 USD per tonne):", 15, True, NAVY)
vals = [
    ("$51", "US Interagency Working Group (2021)", "3% discount rate",
     NAVY_T, NAVY),
    ("$185", "Rennert et al. (2022, Nature)", "2% near-term discount rate",
     GOLD_T, NAVY),
    ("$190", "US EPA (2023) central value", "2% Ramsey discounting",
     RED_T, RED),
]
for i, (v, src, dr, fill, accent) in enumerate(vals):
    shp = rect(s, 0.85 + i * 4.0, 4.85, 3.7, 1.30, fill, rounded=True,
               radius=0.12)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
    setp(tf.paragraphs[0], v + "  ", 24, True, accent, "Verdana",
         align=PP_ALIGN.CENTER)
    setp(tf.paragraphs[0], "per tonne", 12, False, INK)
    p = tf.add_paragraph()
    setp(p, src + " · " + dr, 11, False, INK, align=PP_ALIGN.CENTER)
box = tb(s, 0.85, 6.20, 11.7, 0.35)
setp(box.text_frame.paragraphs[0],
     "The gap is mostly the discount rate, not the climate science: how "
     "much do we care about damages 50 years out?", 12.5, False, GRAY)
punchline(s, "Attach this price tag to our 356 Mt: an $18–68 billion "
             "bill, recurring every year the traffic persists.", y=6.62,
          h=0.60, size=14)
footer(s)

# ======================== 31. ATTRIBUTION BARS + SCC =====================
s = new_slide()
kicker(s, "THE CARBON BILL")
title(s, "Who owes what: the country bill, priced \U0001F4B8")
picture(s, os.path.join(CO2, "CO2_attribution_bars.png"), y=1.90, max_w=7.4,
        max_h=4.55, x=0.55)
box = tb(s, 8.30, 2.05, 4.3, 4.3)
tf = box.text_frame
tf.word_wrap = True
rows = [
    ("China", "+88 Mt", "$16.7 bn/yr"),
    ("United States", "+36 Mt", "$6.9 bn/yr"),
    ("UAE", "+25 Mt", "$4.7 bn/yr"),
    ("Japan / Turkey", "+17 Mt each", ""),
    ("Korea", "+14 Mt", ""),
    ("Germany", "\u221216 Mt", "\u2212$3.0 bn/yr"),
]
first = True
for name, mt, cost in rows:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(7)
    setp(p, name + "  ", 15, True, NAVY)
    setp(p, mt + ("  \u00b7  " + cost if cost else ""), 15, False,
         GREEN if mt.startswith("\u2212") else RED)
p = tf.add_paragraph(); p.space_before = Pt(8)
setp(p, "World: $18.2 / 65.9 / 67.7 bn per year at SCC $51 / $185 / $190.",
     13, False, GRAY)
punchline(s, "A climate bill in the tens of billions of dollars, recurring "
             "annually while the traffic persists.", y=6.55, h=0.62, size=15)
footer(s)

# ======================== 32. MISMATCH MAP ===============================
s = new_slide()
kicker(s, "THE CARBON BILL")
title(s, "Generated here, booked there")
picture(s, os.path.join(CO2, "CO2_map_mismatch_2023.png"), y=1.90, max_w=11.5,
        max_h=3.85,
        caption="Bunker-attributed share minus physical LTO share, 2023: "
                "+1.4 pp for the US and the UAE, \u22123.5 pp for China. "
                "Positive at every international mega-hub.", cap_y=5.90)
punchline(s, "Any scheme that allocates responsibility by booked emissions, "
             "CORSIA included, embeds a transfer across these lines.",
          y=6.55, h=0.62, size=14)
footer(s)

# ======================== 33. SAF SCENARIOS ==============================
# 2026-09-02: QR code / live simulator replaced by a plain static version
# (levels figure + "how it is computed" box + native table). The right-hand
# panel and caption are imported from patch_saf_slide.py so that this rebuild
# and the in-place patch of the shipped 62-slide deck stay identical.
from patch_saf_slide import saf_right_panel, CAPTION as SAF_CAPTION
s = new_slide()
kicker(s, "THE CARBON BILL")
title(s, "The only lever that moves the level: change the fuel")
picture(s, os.path.join(CO2, "CO2_saf_levels.png"), y=1.90, max_w=7.0,
        max_h=3.80, x=0.55)
box = tb(s, 0.55, 5.84, 12.2, 0.7)
setp(box.text_frame.paragraphs[0], SAF_CAPTION, 12.5, False, GRAY,
     align=PP_ALIGN.CENTER)
saf_right_panel(s)
punchline(s, "The full 2050 path (≥65% life-cycle saving) offsets about "
             "as much CO2 as the past generation of connectivity growth "
             "created.", y=6.62, h=0.60, size=14)
footer(s)

# ======================== 34. DIVIDER 06 =================================
divider("07", "Conclusion", "Three myths, three implications, one bill", "🎯")

# ======================== 35. MYTH VS FACT ===============================
s = new_slide()
kicker(s, "CONCLUSION")
title(s, "Three comfortable myths, audited")
myths = [
    ("\u201cHub efficiency will absorb the traffic growth.\u201d",
     "It offsets less than one tenth of the scale response."),
    ("\u201cEmissions are booked where they are generated.\u201d",
     "Hub accounting shifts up to 3.5 pp of world emissions across borders."),
    ("\u201cYour emissions depend on your own network.\u201d",
     "+1% in neighbours' connectivity raises your CO2 by 7.5%."),
]
hx = tb(s, 0.85, 1.95, 5.8, 0.4)
setp(hx.text_frame.paragraphs[0], "MYTH \U0001F9DA", 14, True, GRAY)
hx = tb(s, 6.95, 1.95, 5.6, 0.4)
setp(hx.text_frame.paragraphs[0], "WHAT THE DATA SAY \U0001F52C", 14, True,
     RED)
for i, (m, f) in enumerate(myths):
    y = 2.45 + i * 1.30
    shp = rect(s, 0.85, y, 5.8, 1.10, RGBColor(0xEF, 0xEF, 0xEF), rounded=True,
               radius=0.12)
    tf = shp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.15)
    setp(tf.paragraphs[0], m, 14, False, GRAY, italic=True)
    arrow_right(s, 6.72, y + 0.42, 0.35, 0.26, GOLD)
    shp = rect(s, 7.15, y, 5.35, 1.10, RED_T, rounded=True, radius=0.12)
    tf = shp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.15)
    setp(tf.paragraphs[0], f, 14, True, RED)
punchline(s, "Connectivity policy needs a carbon line in its cost-benefit "
             "table. Here is the number to put there.", y=6.30, h=0.70)
footer(s)

# ======================== 36. SUMMARY & POLICY ===========================
s = new_slide()
kicker(s, "CONCLUSION")
title(s, "Summary and policy")
bullets(s, [
    "New data: flight-stage CO2 for 6,000+ airports, 1996\u20132023, valid "
    "under any allocation rule.",
    "New estimate: the connectivity-emissions elasticity is 5.67, far above "
    "one, concentrated in newly connecting countries, spilling across "
    "borders.",
    "The bill: 42.5% of 2023 aviation CO2 (356 Mt, $18\u201368 bn per year) "
    "traces to the post-1996 connectivity boom.",
    "Policy: evaluate connectivity investments jointly with fuel and "
    "demand-side instruments; efficiency gains are not a decarbonisation "
    "strategy, and booked-emissions schemes embed cross-border transfers.",
], size=17, gap=8)
punchline(s, "Connect if it pays, but put the carbon in the ledger, and bring "
             "the fuel.", y=6.30, h=0.70)
footer(s)

# ======================== 37. DISCUSSION =================================
s = new_slide()
kicker(s, "CONCLUSION")
title(s, "Points to discuss (a.k.a. our limitations)")
bullets(s, [
    "Schedule-based emissions: carbon content of scheduled capacity, not of "
    "realised load factors.",
    "The instrument is strongest in the expansion era; the late period is "
    "identified only excluding COVID years.",
    "Airport-level results are descriptive: no valid airport-level "
    "instrument survived our checks.",
    "Attribution and SAF scenarios are partial-equilibrium accounting; "
    "non-CO2 climate effects (contrails, NOx) are out of scope.",
    "Where next: load factors? Airline-level behaviour? Non-CO2 forcing? "
    "Your ideas welcome.",
], size=17, gap=8)
footer(s)

# ======================== 38. THANK YOU ==================================
s = new_slide(NAVY)
rect(s, 0.0, 0.0, 0.16, 7.5, GOLD)
box = tb(s, 0.65, 2.60, 12.0, 1.2)
setp(box.text_frame.paragraphs[0], "Thank you  \u2708", 44, True, WHITE,
     "Verdana")
box = tb(s, 0.65, 3.90, 11.8, 0.9)
tf = box.text_frame
tf.word_wrap = True
setp(tf.paragraphs[0],
     "How Much Does Air Connectivity Increase Aviation CO2?", 18, False, LAV)
p = tf.add_paragraph()
setp(p, "+1% connectivity \u2192 +5.7% CO2  \u00b7  356 Mt since 1996  \u00b7  "
        "only fuel switching bends the level", 15, False, GOLD)
footer(s, dark=True)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
