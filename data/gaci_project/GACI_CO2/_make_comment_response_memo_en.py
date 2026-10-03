# -*- coding: utf-8 -*-
"""Comment-by-comment response memo (English) -> Comment_response_memo_20260903.docx"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2.0)

def H(text, size=13):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
def B(label, text):
    p = doc.add_paragraph(); r = p.add_run(label + " "); r.bold = True; p.add_run(text); p.paragraph_format.space_after = Pt(3)
def P(text):
    p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph(); r = p.add_run("GACI and aviation CO2: how each coauthor comment was addressed"); r.bold = True; r.font.size = Pt(15)
P("Prepared 3 September 2026 by Sunbin Yoo. Reference manuscript: co2_overleaf_20260903_cl\\main_co2_nature_20260903.tex (all country-year tables with standard errors clustered by country; airport tables clustered by airport). Each item follows the order: request, what was done, results, where in the manuscript, what remains.")

H("Overall changes made while addressing the comments", 14)
B("1. Inference switched to country-level clustering.", "Every country-year regression in the paper (Table 1, the decomposition, heterogeneity, temporal split, spillovers, exclusion diagnostics, functional-form and gradient tables, mediation Panel A) now reports standard errors clustered by country (184 clusters) and cluster-robust first-stage statistics. The August draft used heteroskedasticity-robust errors throughout. Airport regressions cluster by airport. Robust results are kept only in Supplementary Table 1 for comparison. Headline consequence: 5.67 (s.e. 1.12 instead of 0.41), first-stage F 18.4 instead of 154; the efficiency term and the gauge and stage components are no longer individually significant; split-sample first stages are weak, so the heterogeneity gradient is carried by the pooled interactions.")
B("2. Results restructured into the five layers proposed by Prof. Zheng,", "with new Methods subsections on the identity decomposition, the airport design, the spillover design, the exclusion diagnostics, and functional form.")
B("3. Spillover headline replaced.", "The inverse-distance neighbour effect of the August draft (7.45) is retained only as a reference row; the paper now reports the contiguity specification with own and neighbour connectivity jointly instrumented, plus Anderson-Rubin sets and a permutation placebo.")
B("4. Correction of the August heterogeneity numbers,", "which had been estimated on international CO2 but labelled as total CO2; the total-CO2 estimates replace them and the coefficient plot was regenerated.")
B("5. Interpretation narrowed.", "The elasticity is presented as the emissions response during the take-off stage of network development, with level specifications, a continuous baseline gradient, and an attribution range of 35 to 43 percent in place of the single 42.5 percent.")
B("6. Exclusion restriction presented in two tiers.", "Three direct tests that pass are stated in Methods, the introduction and the Discussion; the remaining six checks and two caveats are in one paragraph pointing to ED Table 7 and SI Table 1. The instrument is unchanged (Feyrer interaction; alternatives have no first stage).")
B("7. New displays.", "Four main tables and nine main figures, ten Extended Data tables, seven Extended Data figures and three Supplementary tables; this exceeds Nature-family limits and will be cut at submission.")
H("I. Comments from Prof. Zheng (LZ)", 14)
H("1. Where the 5.67 elasticity comes from: identity decomposition")
B("Request:", "Express CO2 as flights x seats per flight x load factor x distance x carbon intensity, estimate the elasticity of each component, show that they add up to 5.67, and quantify scale versus efficiency.")
B("What was done:", "Because emissions are built from scheduled flights, load factor is a constant and cancels; this is stated. The identity CO2 = flights x (seats per flight) x (km per flight) x (CO2 per seat-km) is estimated component by component by 2SLS on the identical sample. The same identity is applied within income and connectivity terciles and separately to international and domestic departures. A waterfall figure was added. The earlier mediation analysis moves to Extended Data.")
B("Results:", "5.67 = flights 6.98 (s.e. 1.38)*** - gauge 0.30 (0.44) - stage 0.60 (0.58) - intensity 0.40 (0.30). The scale subtotal (seat-km) is 6.07, or 107 percent; the efficiency term is minus 7 percent in point estimate and not significant. Even at the lower end of its 95 percent interval (-0.99) the intensity term offsets only 16 percent of scale. International: 5.69 = 6.03 + 0.58 - 0.40 - 0.52; domestic -0.64 (ns). Within-tercile decompositions rest on weak first stages (F 1.7 to 5.7) and are labelled indicative.")
B("In the manuscript:", "Layer 2 subsection 'Where the elasticity comes from', Table 2 (tab:decomp), waterfall figure, ED Table 4, Methods equations (2) and (3).")
B("Remaining:", "The split of carbon intensity into within-type fuel efficiency and fleet mix needs an aircraft-type aggregate from Fangyu. The efficiency margin is now described as small and imprecise, and in any case unable to offset scale.")

H("2. Airport level: is the effect hub-driven?")
B("Request:", "Use the 6,000+ airport panel for heterogeneity by hub status, connectivity, domestic versus international orientation, baseline size, region and income; rank airports by connectivity-induced emissions and report the shares of the top 1, 5 and 10 percent.")
B("What was done:", "Airport and country-by-year fixed effects (clustered by airport), so airports are compared with other airports of the same country in the same year. Splits by 1996 global top 1 and 5 percent, connectivity and traffic terciles, international share, region and host-country income, plus hub interactions. Topology-only components (eigenvector, closeness, betweenness, degree) as a safeguard against the capacity term in airport GACI. Airport-level attributed emissions = CO2 in 2023 x [1 - exp(-b x change in log GACI)], ranked, with Lorenz curves and Gini benchmarked against the concentration of emissions themselves. An airport Feyrer-shifter instrument was piloted.")
B("Results:", "Within-country elasticity 3.72. Hubs (top 5 percent) 2.11 versus 4.05 elsewhere, interaction -2.56***: the marginal response is systemic. But attributed emissions are more concentrated than emissions (top 1 percent 45.8 versus 41.0 percent; Gini 0.945 versus 0.916) and the 1996 hubs hold 68.5 percent of the attributed total against 75.3 percent of emissions, so the new hubs (Dubai 19.4 Mt, Doha, Shanghai Pudong, Istanbul, Incheon) lead. The airport sum (345 Mt) is consistent with the country total (356 Mt). The identification pilot has no first stage (F 0.4 to 0.7).")
B("In the manuscript:", "Layer 1 subsection 'Airport-level evidence', concentration figure, Table 3 (tab:airport_conc), ED Table 5, SI Table 2, and a Layer 5 sentence: covering the top 5 percent of airports covers 85 percent of connectivity-driven emissions.")
B("Remaining:", "No causal airport estimate is available (disclosed). Each country's largest airport is absorbed by country-by-year effects and enters only through the interaction.")

H("3. Formal spillover analysis")
B("Request:", "(a) first-order neighbour effect controlling for own GACI, (b) distance decay (contiguous, 500 km, 1,000 km, 2,000 km bands or a continuous function), (c) regional spillovers, and a link to the mechanisms.")
B("What was done:", "Eight weighting schemes (contiguity, five nearest, five distance bands, five kernels, within/outside sub-region and aviation bloc, and the earlier inverse distance). Own and neighbour connectivity instrumented jointly (Sanderson-Windmeijer conditional F), neighbour instrumented alone with the own shifter controlled, country-clustered Anderson-Rubin sets by grid inversion, sub-region-by-year fixed effects, a permutation placebo (500 random relabellings of countries in the weight matrix), and an own-versus-neighbour margins table.")
B("Results:", "The August inverse-distance estimate (7.45) has a clustered t of 1.3, a permutation p of 0.15, and the band beyond 5,000 km carries -7.0, so wide exposures track global trends. The headline is contiguity: neighbour instrumented alone 2.05 (0.61), F 175, AR set [0.9, 3.2]; joint 2SLS own 3.48 (1.29) and neighbour 1.93 (0.69), SW F 6.7/13.2, subset AR [0.1, 3.1], permutation p 0.004; with sub-region-by-year effects 2.07 (0.83). There is no smooth decay; the spillover is a land-neighbour phenomenon. Margins: neighbours' connectivity raises seat-km by 2.56*** and lowers intensity by 0.63***, loads on international CO2 (2.23***); flights (1.22) and stage (1.04) are positive but imprecise.")
B("In the manuscript:", "Layer 4 rewritten; Table 4 (tab:spillover), ED Table 6, ED Fig. 4 (distance profile), ED Fig. 5 (permutation), Methods 'Spillover design' including the Anderson-Rubin construction.")
B("Remaining:", "Conley spatial HAC errors and a wild cluster bootstrap for the 22 sub-regions. The text states that the mechanism depends on the weighting and that the own intensity coefficient in the joint specification (+0.32, ns) differs from the single-equation estimate.")

H("4. Strengthening the exclusion restriction")
B("Request:", "More tests that the instrument affects aviation CO2 only through air connectivity.")
B("What was done:", "Three direct tests of the exclusion restriction: (1) placebo outcomes, whether non-aviation emissions respond to the instrument; (2) placebo geography, whether the same aviation cycle interacted with sea market access instead of air market access explains aviation CO2; (3) zero first stage, whether emissions respond to the instrument in countries where it does not move connectivity. Six secondary checks (development controls, alternative technology series, leads, Hansen J, rival global cycles, standard-error variants) are kept in ED Table 7 and SI Table 1.")
B("Results:", "All three pass. (1) Total non-aviation CO2 reduced form 0.05 (s.e. 0.12) against 0.73 (0.16) for aviation CO2. (2) With both interactions entered, sea 0.08 (0.30) versus air 0.64 (0.22). (3) Top baseline-connectivity tercile: first-stage F 0.2, reduced form 0.10 (0.19), against 2.22 and 0.41 in the lower terciles. Two caveats from the secondary checks are stated in Methods: coal and cement CO2 co-move with the instrument (0.90, 0.67), and the aviation cycle is collinear with world GDP and trade cycles (0.82, 0.84). Alternative instrument constructions have no first stage (F 1.0, 3.0), so the Feyrer interaction stays as the single instrument.")
B("How it is presented:", "Two tiers. Methods states the three direct tests (non-aviation placebo, air-versus-sea geography, zero first stage), all passing, in one paragraph; the six further checks are summarised in a second paragraph that points to ED Table 7 and SI Table 1 and states the two caveats (coal and cement co-movement; cycle collinearity). The introduction and Discussion refer to the three tests only. The Feyrer interaction remains the single instrument: the alternative constructions (air-minus-sea advantage x cycle; air interaction with the sea interaction as a control) have no usable first stage (F 1.0 and 3.0).")
B("Remaining:", "Optional: exclusion of China and India and a bad-control bound to show that the coal and cement co-movement does not contaminate the aviation estimate.")

H("5. Five-layer structure")
P("Results reordered as causal effect (country and airport), decomposition, heterogeneity, spillovers, policy. Methods gains subsections on the identity, the airport design and pilot, the exclusion diagnostics, the spillover design, and functional form and continuous heterogeneity. Mediation, the temporal split and measure robustness are in Extended Data. The August heterogeneity numbers (11.43, 5.91, 0.35), which were estimated on international CO2 but labelled as total, are corrected to total CO2 (10.40, 6.14, 1.83).")

H("II. Comments from Yifu", 14)
H("1. The low-base problem: semi-log and level specifications")
B("What was done:", "log CO2 on GACI in levels (semi-log), CO2 in Mt on GACI in levels, CO2 in Mt on log GACI, and semi-log and level specifications within baseline-connectivity terciles.")
B("Results:", "Semi-log 4.09 (1.01) log points per GACI unit (an elasticity of 4.4 at the sample mean); level 4.2 (4.0) Mt per unit, not significant; tercile semi-log 10.3 (2.3), 4.6 (2.6), 1.1 (1.7). The gradient survives in absolute terms; level specifications are dominated by large emitters and imprecise.")
B("In the manuscript:", "Layer 3 subsection 'What the elasticity represents: the take-off stage', ED Table 8.")

H("2. Weak post-2010 first stage, the take-off interpretation, the accident instrument")
B("What was done:", "The headline is now interpreted as the emissions response during the take-off stage of network development (abstract, Results, Discussion). The Aviation Safety Network accident-instrument pilot is reported in the SI.")
B("Results:", "1996 to 2007: 5.61 (1.10), F 25; 2010 to 2023 excluding COVID: 3.01 (1.19), F 11; unrestricted later period F 3.2. Accident instruments have first-stage F of 1 to 4 and cannot stand alone; lagged fatal accidents give 6.2 (2.4); the Hansen test does not reject joint validity with the Feyrer instrument (p 0.70 and 0.95); the joint estimate is 5.4.")
B("In the manuscript:", "The subsection above, Methods 'Instruments and validity', SI Table 3.")
B("Remaining:", "The accident table uses heteroskedasticity-robust errors (stated in the note); a clustered re-estimation has not been run.")

H("3. Sensitivity of the attribution to a common elasticity")
B("What was done:", "The 2023 attribution is recomputed with country-specific elasticities by connectivity tercile, by income tercile, from the continuous interaction (linear and quadratic), and from the semi-log slope applied to the absolute change in GACI.")
B("Results:", "294 to 356 Mt (35 to 43 percent); the common elasticity is the upper end. Largest downward revisions: United States -10, China -8, Japan -3, Spain -3 Mt; Germany's negative attribution moves from -16 to -11.")
B("In the manuscript:", "Range sentence in the Layer 5 attribution paragraph, ED Table 10, Discussion limitations.")

H("4. Continuous heterogeneity by baseline connectivity")
B("What was done:", "Quintile splits on 1996 GACI; pooled 2SLS with log GACI interacted with centred log baseline GACI (linear and quadratic), instruments interacted identically; fitted elasticities and delta-method standard errors at baseline percentiles; a gradient figure.")
B("Results:", "Q1 8.5 (1.7), Q2 8.5 (2.9), Q3 3.3 (2.3), Q4 unidentified (F 0.0), Q5 4.1 (2.5). Interaction -3.99 (1.00), joint first-stage F 9. Fitted elasticity 7.1 at the tenth percentile, 6.0 at the median, 4.1 at the ninetieth, with the quadratic term indicating a floor. The decline is a step and is not confined to the bottom of the distribution.")
B("In the manuscript:", "ED Table 9, gradient figure, subsection text.")

H("III. Inference convention and remaining weak points", 14)
P("All country-year tables now use standard errors clustered by country (184 clusters); heteroskedasticity-robust results are kept only in SI Table 1. Consequences: Table 1 is 5.67 (1.12), t 5.1, F 18.4; the efficiency, gauge and stage components are not individually significant; split-sample first stages are weak, so the heterogeneity gradient rests on the pooled interactions (income -3.33 (0.57), connectivity -2.33 (0.69)); the Hansen test against the tourism instrument rejects at 5 percent.")
P("Weak points (identical to the Open issues section of the response documents): (1) the inverse-distance spillover is not robust and the contiguity joint specification has a weak first stage, defended by Anderson-Rubin sets; (2) sectoral recomposition and cycle collinearity in the exclusion diagnostics; (3) identification concentrated in the expansion era and in low-baseline countries (top tercile F 0.2, fourth quintile unidentified); (4) the August heterogeneity mislabel must be communicated; (5) no causal airport estimate; (6) level specifications imprecise; (7) domestic-traffic results imprecise; (8) inference convention changed; (9) length above Nature limits and no local compilation (Overleaf required).")
P("Related files: Response_LZ_results_20260903.docx and Response_Yifu_results_20260903.docx (point-by-point responses for the coauthors), CO2_results_summary_20260903_LZ.xlsx (31 sheets), co2_overleaf_20260903_cl.zip (for Overleaf).")
out = os.path.join(HERE, "Comment_response_memo_20260903.docx")
try:
    doc.save(out)
except PermissionError:
    out = out.replace(".docx", "_v2.docx"); doc.save(out); print("target locked; saved as", out)
txt = "\n".join(p.text for p in doc.paragraphs)
print("saved", out, "| em-dash:", txt.count(chr(8212)), "| words:", len(txt.split()))
