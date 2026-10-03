# -*- coding: utf-8 -*-
"""Point-by-point response (Japanese) to Junya's comments on the Estimation
part of the GACI-CO2 manuscript (2026-09-25), with before/after manuscript text.
BEFORE text is cut from base_overleaf_main_co2_nature_20260908.tex (the Overleaf
version holding Junya's Methods); AFTER text is cut from paste_blocks_20260925.tex,
so the response shows exactly what will be pasted. Style follows _make_response_lz.py."""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Response_Junya_comments_20260925.docx")
BASE = open(os.path.join(HERE, "base_overleaf_main_co2_nature_20260908.tex"), encoding="utf-8").read()
NEW = open(os.path.join(HERE, "paste_blocks_20260925.tex"), encoding="utf-8").read()
JP = "Yu Gothic"
NAVY = RGBColor(0x23, 0x2A, 0x55)
ORANGE = RGBColor(0xB0, 0x6A, 0x1E)


# ---------------------------------------------------------------- text helpers
def grab(src, start, end, after=None):
    """Substring from `start` through `end` (inclusive); search begins at `after`."""
    i0 = src.index(after) if after else 0
    i = src.index(start, i0)
    j = src.index(end, i) + len(end)
    return src[i:j]


def mathplain(m):
    if m.strip() == "-":
        return "−"
    m = m.replace("^{***}", "***").replace("^{**}", "**").replace("^{*}", "*")
    m = re.sub(r"\\sum_\{([^}]*)\}", lambda k: "Σ⟨" + k.group(1) + "⟩ ", m)
    m = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", m)
    for a, b in [("\\,", ""), ("\\times", "×"), ("\\approx", "≈"), ("\\neq ", "≠"), ("\\neq", "≠"), ("\\ln", "ln "),
                 ("\\theta", "θ"), ("\\exp", "exp"), ("\\lambda", "λ"), ("\\rho", "ρ"), ("\\alpha", "α"),
                 ("\\tau", "τ"), ("\\varepsilon", "ε"), ("\\beta", "β")]:
        m = m.replace(a, b)
    m = re.sub(r"_\{([^}]*)\}", r"_\1", m)
    m = re.sub(r"\^\{([^}]*)\}", r"^\1", m)
    m = m.replace("{", "").replace("}", "").replace("⟨", "_{").replace("⟩", "}")
    m = re.sub(r"ln\s+", "ln ", m)
    return m.strip()


EQ = {
    "eq:seama": "airMA_ct = Σ_{j≠c} Pop_jt / d^air_cj ,    seaMA_ct = Σ_{j≠c} Pop_jt / d^sea_cj        (eq:seama)",
    "eq:feyrer": "Z_ct = a_t × ln airMA_c,1996        (eq:feyrer)",
}


def plain(s):
    """LaTeX -> readable plain text; returns a list of paragraphs."""
    def eqrep(k):
        lab = re.search(r"\\label\{([^}]*)\}", k.group(0)).group(1)
        return "\n\n" + EQ[lab] + "\n\n"
    s = re.sub(r"\\begin\{equation\}.*?\\end\{equation\}", eqrep, s, flags=re.S)
    s = s.replace("CO$_2$", "CO2").replace("\\citet{feyrer2019}", "Feyrer (2019)")
    s = s.replace("\\citep{bertoli2016}", "(Bertoli et al., 2016)")
    s = re.sub(r"~?\\ref\{([^}]*)\}", lambda k: " [" + k.group(1) + "]", s)
    s = re.sub(r"\$([^$]*)\$", lambda k: mathplain(k.group(1)), s)
    s = s.replace("\\item ", "").replace("\\%", "%").replace("\\ ", " ").replace("\\$", "$")
    s = s.replace("\\quad ", "").replace("--", "–").replace("~", " ")
    paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", s)]
    return [p for p in paras if p]


# ---------------------------------------------------------------- document helpers
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), JP)
for sname in ["List Bullet", "Heading 1", "Heading 2", "Heading 3", "Title"]:
    s = doc.styles[sname]
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), JP)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.0)


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcol)
    tcPr.append(shd)


def para(text, bold=False, size=None, color=None, space_after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Calibri"; r.font.color.rgb = NAVY
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), JP)
    return h


def comment_box(text, who="Junyaさんのコメント"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, "EEF1F7")
    c.paragraphs[0].text = ""
    lab = c.paragraphs[0].add_run(who + "： ")
    lab.bold = True; lab.font.color.rgb = NAVY
    lines = text.strip().split("\n")
    c.paragraphs[0].add_run(lines[0])
    for line in lines[1:]:
        p = c.add_paragraph(); p.add_run(line)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def response(items, label="回答"):
    if label:
        p = doc.add_paragraph()
        r = p.add_run(label + "："); r.bold = True; r.font.color.rgb = ORANGE
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
    for kind, txt in items:
        if kind == "p":
            q = doc.add_paragraph(txt); q.paragraph_format.space_after = Pt(4)
        elif kind == "b":
            q = doc.add_paragraph(txt, style="List Bullet"); q.paragraph_format.space_after = Pt(2)


CHG = [0]


def change(title, before, after):
    """Numbered before/after box. before/after: list of paragraphs, or a str (Japanese note)."""
    CHG[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f"修正{CHG[0]}　"); r.bold = True; r.font.color.rgb = ORANGE; r.font.size = Pt(9.5)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, (lab, content, fill) in enumerate([("修正前", before, "F3F3F3"), ("修正後", after, "FBF4EA")]):
        c0, c1 = t.rows[k].cells
        c0.width = Cm(1.7); c1.width = Cm(14.9)
        shade(c0, fill); shade(c1, fill)
        c0.paragraphs[0].text = ""
        rr = c0.paragraphs[0].add_run(lab); rr.bold = True; rr.font.size = Pt(8.5)
        paras = [content] if isinstance(content, str) else content
        c1.paragraphs[0].text = ""
        for n, ptxt in enumerate(paras):
            q = c1.paragraphs[0] if n == 0 else c1.add_paragraph()
            rr = q.add_run(ptxt); rr.font.size = Pt(8.5)
            if ptxt.startswith(("airMA_ct", "Z_ct")):
                q.paragraph_format.left_indent = Cm(1.0)
            q.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return CHG[0]


def table(headers, rows, widths=None, fs=9, panel_rows=False):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(fs)
        shade(c, "D9DEEA")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(fs)
            if panel_rows and str(row[0]).startswith("Panel"):
                rr.bold = True; shade(cells[i], "F2F3F7")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------------------------------------------------------- manuscript pieces
B = {  # BEFORE (Overleaf, Junya's version)
    "gaci_n": grab(BASE, "The estimation sample covers 184 countries", "country-year observations."),
    "est_sea": grab(BASE, "$\\mathrm{pop}_{ct}$ and $\\mathrm{seaMA}_{ct}$ denote", "idiosyncratic error term."),
    "est_cl": grab(BASE, "Standard errors are clustered at the country level (184 clusters)", "cluster-robust."),
    "iv2": grab(BASE, "To address this endogeneity, we instrument connectivity using a Feyrer-type shifter", "from 120.1 to 105.8."),
    "six": grab(BASE, "Six further checks reinforce our findings", "($F$ of 1.0 and 5.6)."),
    "other": grab(BASE, "Finally, we consider other potential instruments", "post-2010 first stage."),
    "airport": grab(BASE, "We initially explored an instrumental variable strategy", "strictly causal effects."),
    "ar": grab(BASE, "Finally, we implement robust inference procedures", "exceeds the true empirical estimate."),
    "spill_note": grab(BASE, "\\item Notes: Neighbour connectivity is the leave-out average", "10 percent.", after="\\label{tab:spillover}"),
    "ex2_note": grab(BASE, "\\item Specification as in Table~\\ref{tab:main}, column 1. Panel A replaces", "$p<0.1$.", after="\\label{tab:exclusion2}"),
    "acc_note": grab(BASE, "\\item Notes: Country-year accident counts", "10 percent.", after="\\label{tab:accident_iv}"),
    "res_range": grab(BASE, "It stays\nbetween 4.4 and 4.6", "in disguise"),
}
A = {  # AFTER (paste blocks)
    "gaci_n": grab(NEW, "The estimation sample covers 182 countries", "used in the analysis."),
    "iv2": grab(NEW, "To address this endogeneity, we instrument connectivity with a shift-share", "from 120.1 to 105.8."),
    "six": grab(NEW, "Further checks support these results", "($F$ of 1.0 and 5.6)."),
    "other": grab(NEW, "Finally, an aviation-accident instrument", "tab:accident_iv})."),
    "airport": grab(NEW, "We also piloted an instrumental-variable strategy", "descriptive evidence."),
    "ar": grab(NEW, "Finally, the joint first stage is weak", "exceeds the actual one."),
    "spill_note": grab(NEW, "\\item Notes: Neighbour connectivity is the leave-out average", "10 percent.", after="BLOCK 8"),
    "ex2_note": grab(NEW, "\\item Specification as in Table~\\ref{tab:main}, column 1. Panel A replaces", "$p<0.1$.", after="BLOCK 9"),
    "acc_note": grab(NEW, "\\item Notes: Country-year accident counts", "10 percent.", after="BLOCK 10"),
    "exE_note": grab(NEW, "Panel D: terciles of 1996 capacity-weighted GACI. Panel E:", "treated as missing.", after="BLOCK 13"),
}
P = {k: plain(v) for k, v in B.items()}
Q = {k: plain(v) for k, v in A.items()}
est_sea_after = [P["est_sea"][0].replace("sea market access, respectively", "sea market access (equation [eq:seama]), respectively")]
assert est_sea_after[0] != P["est_sea"][0]

# ================================================================ title
para("GACI CO2 論文：Estimation 部分へのコメントへの対応", bold=True, size=15, color="232A55")
para("2026年9月25日。対象原稿：Overleaf の main_co2_nature_20260908.tex（Junyaさんの Methods が入った版）。"
     "定義はすべて推定コードで確認し、必要な再推定も済ませています。", size=9.5, color="555555")

para("Junyaさん")
para("Estimation 部分の執筆とコメントをありがとうございました。質問5点と提案3点について、コードで定義を確認したうえで対応を決めました。"
     "以下、コメントごとに回答と原稿の修正（修正前・修正後の文面）を示します。確認の過程で見つかったデータ上の問題3点についても、"
     "再推定まで済ませています（第3節）。修正はすべて差し替え文として用意してあり、Overleaf にはこれから反映します。")
para("修正文中の [tab:…] [eq:…] は LaTeX のラベルです。表の数値行の差し替えを含む LaTeX 全文は paste_blocks_20260925.tex にあります。",
     size=9, color="555555")

heading("対応一覧", 1)
table(["#", "コメント", "対応", "修正"],
      [["Q1", "air market access の計算方法", "式を追加（国レベルで直接計算）", "修正1"],
       ["Q2", "sea market access の計算方法", "式とデータ出典を追加", "修正1・2・3"],
       ["Q3", "tourism-heritage IV の出典・計算方法", "P1 に従い削除", "修正6・7・9"],
       ["Q4", "a_t の定義・データ出典", "本文と Supp. Table の注に明記", "修正1・9"],
       ["Q5", "Summary statistics がない", "Supplementary Table を新規追加", "修正4・5"],
       ["P1", "heritage IV を削除", "採用", "修正6・7・9"],
       ["P2", "leads を削除", "採用", "修正8・9"],
       ["P3", "ED Table 8 の下段を削除", "一部採用（AR 集合のみ注に残す）", "修正10・11"],
       ["追加", "確認中に見つかった点3件", "再推定済み", "修正4・12〜17"]],
      widths=[1.1, 5.4, 6.2, 3.9])

# ================================================================ 1. questions
heading("1. 質問への回答", 1)

heading("Q1. air market access の計算方法", 2)
comment_box("""air market access の計算方法
→airport-level のところには "air market access, defined as the population-weighted inverse distance to all foreign countries' aviation centroids." と書いてあるが、国レベル分析ではこれを国ごとに平均したということで正しいか。""")
response([
    ("p", "空港レベルの値を国ごとに平均したものではなく、国レベルで直接計算しています（GACI/build_feyrer_iv.py）。"),
    ("b", "airMA_c,1996 = Σ_{j≠c} Pop_j,1996 / d_cj"),
    ("b", "d_cj は c 国と j 国の「航空重心」間の大圏距離（km）です。航空重心は、ネットワーク内にあるその国の空港の緯度・経度の単純平均です。"),
    ("b", "Pop は WDI の総人口です。自国人口は除外し、距離減衰パラメータは θ = 1 です。相手国の集合は sea market access と同じです（CERDI の海上距離がある国）。"),
    ("b", "操作変数は Z_ct = a_t × ln airMA_c,1996 です。1996年の人口で計算するため、1996年以降にネットワークに参入した32か国にも基準値があります。"),
    ("p", "\"population-weighted inverse distance\" は加重平均ではなく和なので、原稿では式で定義し直しました。"
          "空港レベルの値は、国の航空重心の代わりに空港の座標を使った同じ式です（修正7）。"),
])
response([("p", "Instruments and validity の第2段落を、air・sea market access と操作変数の定義式を含む2段落に差し替えます。"
                "Q2（sea MA）と Q4（a_t）への対応もこの修正に含まれます。")], label="原稿の修正")
change("Methods > Instruments and validity、第2段落（差し替え）", P["iv2"], Q["iv2"])

heading("Q2. sea market access の計算方法", 2)
comment_box("sea market access の計算方法")
response([
    ("b", "seaMA_ct = Σ_{j≠c} Pop_jt / dsea_cj"),
    ("b", "dsea は CERDI-seadistance database（Bertoli, Goujon and Santoni 2016, CERDI Working Paper 2016/07）の港湾間の海上距離です。内陸国には通過国の港が割り当てられています。"),
    ("b", "コントロールとして入っているのは、当年の人口で計算した時変の ln seaMA_ct で、1996年固定の値ではありません。"
          "1996年固定の値は、ED Table（exclusion diagnostics）Panel C のプラセボ操作変数 a_t × ln seaMA_c,1996 だけで使っています。"),
])
response([("p", "定義は修正1に入っています。加えて Estimation 節から定義式を参照し、参考文献を2件追加します。")], label="原稿の修正")
change("Methods > Estimation、式 (main) の説明文", P["est_sea"], est_sea_after)
change("refs_co2.bib（参考文献の追加）", "（なし）",
       ["Feyrer, J. (2019). Trade and Income: Exploiting Time Series in Geography. American Economic Journal: Applied Economics, 11(4), 1–35. doi:10.1257/app.20170616",
        "Bertoli, S., Goujon, M., and Santoni, O. (2016). The CERDI-seadistance database. CERDI Working Paper 2016/7."])

heading("Q3. tourism-heritage instrument のデータソースと計算方法", 2)
comment_box("""a tourism-heritage instrument -> データソース、具体的な計算方法
→個人的には、tourism-heritage instrument を使った robustness check はすべて抜いてよいと思っています（結果があまりよくなさそうなので）""")
response([
    ("p", "参考までに定義は次のとおりです。Z^H_ct = T_t × ln(1 + H_c)。"
          "T_t は世界の国際観光客到着数（WDI ST.INT.ARVL の世界計、1996–2019年。2020–23年は UNWTO の対2019年比で延長）を min-max 正規化したもの、"
          "H_c は UNESCO 世界遺産のうち自然遺産と複合遺産の数（2023年時点）です。"),
    ("p", "ご提案どおり削除するので、原稿には定義を載せません。削除の修正は P1 にまとめています（修正6・7・9）。"),
])

heading("Q4. a_t（air technology）の計算方法とデータソース", 2)
comment_box("a_t (air technology) の計算方法。World seat capacityと書いてあるが、seat capacity の total value ってことで合っているか。world seat capacity, world flights, world seat-km などのデータソースはどこか。")
response([
    ("p", "はい、座席容量の合計です。"),
    ("b", "OAG スケジュールから作った空港別の年間座席容量を、パネル対象国の全空港で合計し、1996–2023年で [0,1] に min-max 正規化しています。ICAO などの外部統計ではなく、自前の OAG データです。"),
    ("b", "パネル対象国の合計はネットワーク全体の容量の約95–98%で、全空港で合計した場合との相関は 0.999 です。"),
    ("b", "Supp. Table の代替系列もすべて自前のデータです。world flights と world seat-km は排出インベントリの出発便数と出発座席キロの世界合計、"
          "world GACI sum は空港 GACI の合計、fuel-efficiency index は −ln（世界 CO2 / 世界座席キロ）で、いずれも同じく min-max 正規化しています。"),
    ("b", "ED Table（exclusion diagnostics）Panel B のライバル循環は、WDI の世界 GDP（実質）と世界輸出（実質）、Brent 原油価格（FRED）です。"),
])
response([("p", "a_t の定義は修正1の本文に、代替系列の作り方は Supp. Table（旧1）の注（修正9）に入れました。")], label="原稿の修正")

heading("Q5. Summary statistics", 2)
comment_box("Summary statistics が入っていない。追加分析も含めると多くの変数があるが、すべて含めた summary statistics table を入れた方が良い")
response([
    ("p", "ご指摘のとおりです。推定サンプル（182か国、1996–2023年、4,634 国・年）について、追加分析の変数も含めた要約統計量の表を Supplementary Table として追加します。"
          "6つのパネル（航空アウトカム、GACI、操作変数とコントロール、開発コントロール、非航空プラセボ、空間変数）に40変数を収めています。"),
])
response([("p", "GACI 節の最後の段落から新しい表を参照し、表は Supplementary Tables の先頭に置きます（表の中身は付録）。"
                "この文では推定サンプルの国数も直しています（第3節 (1)）。")], label="原稿の修正")
change("Methods > Global Air Connectivity Index、最後の段落", P["gaci_n"], Q["gaci_n"])
change("Supplementary Tables（新規の表を先頭に追加）", "（なし）",
       "Supplementary Table [tab:sumstat]「Summary statistics, estimation sample」を追加（6パネル、40変数、注付き）。内容は付録を参照。")

# ================================================================ 2. proposals
heading("2. 提案への対応", 1)

heading("P1. heritage の操作変数を削除（採用）", 2)
comment_box("heritage の操作変数は消してよいのでは。Feyrer shifter と異なる除外制約の説明をする必要があるので、メインでない、かつ結果の良くない操作変数は消してよい気がする（supplementary table 1 panel C）。")
response([
    ("p", "採用します。"),
    ("b", "単独推定は 1.39 (1.59) で有意でなく、サイズ×サイクルのストレステストでは F が 13.8 から 1 未満に落ちます。"
          "さらに Feyrer との同時推定の Hansen J は p = 0.063 で、10%水準では過剰識別制約が棄却されるため、残すとかえって弱点になります。"),
    ("b", "過剰識別の検定は事故IV（Hansen p = 0.82, 0.96）で示せます。事故IVの数値は国クラスターで再推定した値です（第3節 (3)）。"),
], label="対応")
response([("p", "Methods の2か所と Supp. Table（旧1）から heritage の記述をすべて外します（表の修正は修正9）。")], label="原稿の修正")
change("Methods > Instruments and validity、最終段落（heritage の文を削除し、事故IVの段落に）", P["other"], Q["other"])
change("Methods > Airport-Level Analysis、第2段落（heritage-proximity の文を削除、空港 MA を式で定義）", P["airport"], Q["airport"])

heading("P2. 操作変数の leads を削除（採用）", 2)
comment_box("""Supplementary Table 1 (Leads of the instrument are uninformative because the technology index is a smooth trend).
→なら補足結果から抜いてよいのではないか""")
response([
    ("p", "採用します。"),
    ("b", "国固定効果の下では、lead と当期値の差は (a_{t+k} − a_t) × 地理だけです。a_t はほぼ線形に伸びるのでこの差は国固定効果にほぼ吸収され、"
          "識別は2001–03年、2009年、2020年の落ち込みからしか来ません。"),
    ("b", "一方で lead の係数は有意（0.39, 0.30）なので、「情報がない」と注記して残すと、プレトレンドの問題に見えるリスクの方が大きいと判断しました。"),
    ("b", "査読でプレトレンドを問われた場合は、ED Table（exclusion diagnostics）Panel D（第一段階のない上位三分位では誘導形も 0.09 で有意でない）と、"
          "Panel C（海の地理を使ったプラセボ）で答えられます。"),
], label="対応")
response([("p", "Methods の段落を書き直し、Supp. Table（旧1）を再構成します。"
                "この段落の \"horse race ... supports our specification\" は、ED Table Panel B（世界 GDP・輸出のサイクルを入れると第一段階の F が 0.2 と 0.0 に落ちる）と合わないため、"
                "caveat として書き直しました。開発コントロールの範囲は再推定後の値（4.4–4.7）です（第3節 (2)）。")], label="原稿の修正")
change("Methods > Instruments and validity、「Six further checks」の段落", P["six"], Q["six"])
change("Supplementary Table（旧1）[tab:exclusion2]：表の構成と注",
       ["Caption: Alternative instrument constructions, leads, overidentification and inference",
        "Panel A. Technology series in the interaction（8行）",
        "Panel B. Leads of the instrument（Three-year lead、Five-year lead）",
        "Panel C. Overidentification and inference（Feyrer and tourism-heritage instruments jointly、Hansen J、Tourism-heritage instrument alone、Heteroskedasticity-robust、Clustered by country、Two-way clustered）",
        "Note: " + P["ex2_note"][0]],
       ["Caption: Alternative technology series, distance decay and inference",
        "Panel A. Technology series in the interaction（8行、数値は変更なし）",
        "旧 Panel B（leads）は削除",
        "Panel B. Inference（Heteroskedasticity-robust、Clustered by country (baseline)、Two-way clustered）。旧 Panel C の heritage 関連3行は削除",
        "Note: " + Q["ex2_note"][0]])

heading("P3. ED Table 8 の下段を削除（一部採用）", 2)
comment_box("""Extended Data Table 8: Connectivity spillovers: own and neighbouring countries' connectivity, jointly instrumented Panel A lower part
"Neighbour connectivity instrumented, own shifter in reduced form; s.e. clustered by country; AR = Anderson–Rubin 95% set" の部分
→spatial model の robustness check としては結構細かいことをやっているので、robustness check の結果が今かなり多いので、抜いてもよいのではないか。ほとんど解釈もされていないので問題ない気がする""")
response([
    ("p", "下段ブロックは削除しますが、Anderson–Rubin 集合を1つだけ表の注に残します。"),
    ("b", "sub-region × year FE の行は ED Table（spatial structure of the spillover）に同じ推定があり、500 km・5近傍・逆距離の行は上段と重複しているので、削除しても失うものはありません。"),
    ("b", "ただし contiguity の Anderson–Rubin 95% 集合 [0.2, 3.0] は表の注に移して残します。own と neighbour を同時に操作する joint 2SLS の KP F は 3.3"
          "（SW F は 6.7 / 13.0）と弱く、隣国係数 1.85 を弱操作変数に頑健な形で支える根拠はこの集合だけだからです。"),
], label="対応")
response([("p", "表の下段と注、Methods の空間節の AR の段落を修正します。")], label="原稿の修正")
change("Extended Data Table 8 [tab:spillover]：Panel A 下段と注",
       ["Panel A 下段（\"Neighbour connectivity instrumented, own shifter in reduced form ...\"）の見出しと8行："
        "Contiguous countries / Contiguous, sub-region × year FE / Countries within 500 km / Five nearest countries / "
        "Contiguous, joint 2SLS / Within 500 km, joint 2SLS / Inverse distance, own shifter in reduced form / s.e. clustered by country",
        P["spill_note"][0]],
       ["Panel A は上段の5行のみ（数値は変更なし）。Panel B は変更なし。",
        Q["spill_note"][0]])
change("Methods > Spatial Econometric Specifications、最終段落", P["ar"], Q["ar"])

# ================================================================ 3. found issues
heading("3. 確認中に見つかった点（再推定済み）", 1)

heading("(1) 推定サンプルの国数", 2)
response([
    ("p", "原稿の「184 countries」と「184 clusters」は誤りで、正しくは 182 か国です。Stata ログのクラスター数は 182 で、"
          "シングルトン1件（ニューカレドニア 2015年）を除いて 4,634 観測になります。184 は帰属計算（SCC 表）の国数なので、そちらはそのままです。"
          "GACI 節の文は修正4で直しています。"),
], label=None)
change("Methods > Estimation、標準誤差の文", P["est_cl"], [P["est_cl"][0].replace("(184 clusters)", "(182 clusters)")])

heading("(2) ベネズエラの trade/GDP", 2)
response([
    ("p", "WDI がベネズエラの trade/GDP を1991–2011年について欠損ではなく 0 で返していました。1996–2011年の16観測を欠損に直し、"
          "ED Table（exclusion diagnostics）Panel E を再推定しました。変わるのは Panel E だけで、すべての行が1%水準で有意のままです。"),
], label=None)
change("ED Table [tab:exclusion] Panel E の数値",
       ["+ trade/GDP: 4.43*** (1.32), KP F 9.6, N = 4,618",
        "+ urban share: 4.35*** (1.30), KP F 10.4, N = 4,618",
        "+ FDI/GDP: 4.54*** (1.20), KP F 12.0, N = 4,432",
        "+ log tourist arrivals: 4.64*** (1.40), KP F 9.8, N = 3,460",
        "Baseline on the common sample: 5.62*** (1.25), KP F 18.0, N = 3,460"],
       ["+ trade/GDP: 4.45*** (1.40), KP F 8.9, N = 4,602",
        "+ urban share: 4.38*** (1.38), KP F 9.6, N = 4,602",
        "+ FDI/GDP: 4.57*** (1.27), KP F 11.1, N = 4,416",
        "+ log tourist arrivals: 4.72*** (1.50), KP F 8.9, N = 3,444",
        "Baseline on the common sample: 5.69*** (1.32), KP F 16.9, N = 3,444",
        "（Baseline と + log GDP の行は変更なし）"])
change("ED Table [tab:exclusion] の注（Panel E の説明を追加）",
       "Panel D: terciles of 1996 capacity-weighted GACI.", Q["exE_note"])
change("Results > 最初の小節（Table 1 の段落）",
       plain(B["res_range"]), plain(B["res_range"].replace("4.4 and 4.6", "4.4 and 4.7")))

heading("(3) 事故IV表の標準誤差", 2)
response([
    ("p", "表の注は「国クラスター」でしたが、実際の推定は heteroskedasticity-robust でした。heritage IV を削除すると過剰識別の検定はこの表だけになるので、"
          "国クラスターで再推定しました。係数は変わらず、標準誤差が大きくなります。Hansen 検定の結論は変わりません。"
          "死者数単独の 2SLS と同時推定の intensity は有意でなくなりますが、どちらも本文で引用していない数値です。"
          "Methods の事故IVの文は修正6で新しい値に合わせています。"),
], label=None)
change("Supplementary Table [tab:accident_iv] の数値",
       ["2SLS, lagged fatal accidents: 5.59** (2.25) / 4.57** (1.80) / −0.92 (0.65); KP F 4.4",
        "2SLS, lagged accident deaths: 5.26* (3.12) / 4.24 (2.80) / −0.28 (0.99); KP F 2.0",
        "2SLS, lagged any-fatal-accident indicator: 6.05 (4.31) / 5.34 (3.98) / −1.02 (1.49); KP F 1.2",
        "Feyrer + fatal accidents: 5.09*** (0.40) / 4.90*** (0.39) / −0.34*** (0.12); KP F 66.8; Hansen p 0.81 / 0.86 / 0.27",
        "Feyrer + accident deaths: 5.08*** (0.40) / 4.89*** (0.39) / −0.33*** (0.12); KP F 66.3; Hansen p 0.95 / 0.82 / 0.96"],
       ["2SLS, lagged fatal accidents: 5.59** (2.24) / 4.57** (1.82) / −0.92 (0.69); KP F 4.2",
        "2SLS, lagged accident deaths: 5.26 (3.41) / 4.24 (3.02) / −0.28 (0.94); KP F 1.5",
        "2SLS, lagged any-fatal-accident indicator: 6.05 (4.83) / 5.34 (4.28) / −1.02 (1.30); KP F 1.1",
        "Feyrer + fatal accidents: 5.09*** (1.07) / 4.90*** (1.03) / −0.34 (0.30); KP F 9.2; Hansen p 0.82 / 0.87 / 0.37",
        "Feyrer + accident deaths: 5.08*** (1.07) / 4.89*** (1.03) / −0.33 (0.30); KP F 8.4; Hansen p 0.96 / 0.84 / 0.96",
        "（列は Bunker CO2 / Intl. CO2 / Intensity）"])
change("Supplementary Table [tab:accident_iv] の注", P["acc_note"], Q["acc_note"])

# ================================================================ 4. next
heading("4. 今後の進め方", 1)
response([
    ("b", "上記の修正は、こちらで Overleaf の原稿（Junyaさんの Methods が入った版）に反映します。9月23日に差し替えた Longfei さんの統合図もあわせて反映します。"),
    ("b", "反映後、修正1・6・8・11（Instruments and validity と空間節）の英文を確認していただけると助かります。"),
], label=None)
para("よろしくお願いします。")

# ================================================================ appendix: sumstat table
doc.add_page_break()
heading("付録　追加する Supplementary Table（要約統計量）", 1)
rows_tex = open(os.path.join(HERE, "sumstat_rows.tex"), encoding="utf-8").read().splitlines()


def untex(s):
    s = s.replace("$<$", "<").replace("$-$", "−").replace("CO$_2$", "CO2").replace("\\%", "%").replace("\\$", "$")
    s = s.replace("excl.\\ ", "excl. ").replace("$Z_{ct}=a_t\\times\\ln \\mathrm{airMA}_{c,1996}$", "Z = a_t × ln airMA_1996")
    s = s.replace("$W\\ln\\mathrm{GACI}$", "W ln GACI").replace("$WZ$", "WZ").replace("$a_t$", "a_t").replace("0--1", "0–1")
    return s.strip()


rows = []
for ln in rows_tex:
    ln = ln.strip().rstrip("\\").strip()
    m = re.match(r"\\multicolumn\{6\}\{l\}\{\\textit\{(.*)\}\}", ln)
    if m:
        rows.append([untex(m.group(1)), "", "", "", "", ""])
        continue
    rows.append([untex(c) for c in ln.split("&")])
table(["Variable", "N", "Mean", "s.d.", "Min", "Max"], rows, widths=[7.4, 1.6, 1.8, 1.8, 1.7, 2.0], fs=8, panel_rows=True)
para("推定サンプル（182か国、1996–2023年、4,634 国・年）。Domestic CO2 は国内排出が正の国・年、隣国変数は陸上の隣国がある 3,725 国・年について集計。"
     "ベネズエラの trade/GDP（1996–2011年）は欠損として扱っています。作成スクリプト：sumstat_co2.py。", size=8.5, color="555555")

doc.save(OUT)
print("saved", OUT, "| changes:", CHG[0])
