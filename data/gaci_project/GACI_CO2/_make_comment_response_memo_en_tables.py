# -*- coding: utf-8 -*-
"""Comment-by-comment response memo (English, with tables built from the clustered
result files) -> Comment_response_memo_20260903.docx (or _v2 if locked).
Supersedes _make_comment_response_memo_en.py."""
import os, sys, math
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Cm
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
R = lambda fn: pd.read_csv(os.path.join(HERE, fn))
def star(p):
    return "" if (p is None or (isinstance(p, float) and math.isnan(p))) else ("***" if p < .01 else ("**" if p < .05 else ("*" if p < .1 else "")))
def cs(b, se, p=None, nd=2):
    if p is None:
        p = math.erfc(abs(b / se) / 2 ** 0.5) if se and not math.isnan(se) else float("nan")
    return f"{b:.{nd}f}{star(p)} ({se:.{nd}f})"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
def H(text, size=13):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
def B(label, text):
    p = doc.add_paragraph(); r = p.add_run(label + " "); r.bold = True; p.add_run(text); p.paragraph_format.space_after = Pt(3)
def P(text):
    p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(4)
def T(headers, rows, fs=9, note=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(fs)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(fs)
    p = doc.add_paragraph(note or ""); p.paragraph_format.space_after = Pt(4)
    if note:
        p.runs[0].font.size = Pt(8.5); p.runs[0].italic = True

SE = "Standard errors clustered by country in parentheses; *** p<0.01, ** p<0.05, * p<0.1."

p = doc.add_paragraph(); r = p.add_run("GACI and aviation CO2: how each coauthor comment was addressed"); r.bold = True; r.font.size = Pt(15)
P("Prepared 3 September 2026 by Sunbin Yoo. Reference manuscript: co2_overleaf_20260903_cl\\main_co2_nature_20260903.tex (all country-year tables with standard errors clustered by country; airport tables clustered by airport). Each item follows the order: request, what was done, results, where in the manuscript, what remains. Tables are generated from the result files behind the manuscript.")

H("Overall changes made while addressing the comments", 14)
B("1. Inference switched to country-level clustering.", "Every country-year regression (Table 1, decomposition, heterogeneity, temporal split, spillovers, exclusion diagnostics, functional-form and gradient tables, mediation Panel A) now reports standard errors clustered by country (184 clusters) and cluster-robust first-stage statistics. The August draft used heteroskedasticity-robust errors throughout. Airport regressions cluster by airport. Robust results are kept only in Supplementary Table 1. Headline consequence: 5.67 (s.e. 1.12 instead of 0.41), first-stage F 18.4 instead of 154; the efficiency term and the gauge and stage components are no longer individually significant; split-sample first stages are weak, so the heterogeneity gradient is carried by the pooled interactions.")
B("2. Results restructured into the five layers proposed by Prof. Zheng,", "with new Methods subsections on the identity decomposition, the airport design, the spillover design, the exclusion diagnostics, and functional form.")
B("3. Spillover headline replaced.", "The inverse-distance neighbour effect of the August draft (7.45) is retained only as a reference row; the paper now reports the contiguity specification, with Anderson-Rubin sets and a permutation placebo.")
B("4. Correction of the August heterogeneity numbers,", "which had been estimated on international CO2 but labelled as total CO2; the total-CO2 estimates replace them and the coefficient plot was regenerated.")
B("5. Interpretation narrowed.", "The elasticity is presented as the emissions response during the take-off stage of network development, with level specifications, a continuous baseline gradient, and an attribution range of 35 to 43 percent in place of the single 42.5 percent.")
B("6. Exclusion restriction presented in two tiers.", "Three direct tests that pass are stated in Methods, the introduction and the Discussion; the remaining six checks and two caveats are in one paragraph pointing to ED Table 7 and SI Table 1. The instrument is unchanged (Feyrer interaction; alternatives have no first stage).")
B("7. New displays.", "Four main tables and nine main figures, ten Extended Data tables, seven Extended Data figures and three Supplementary tables; this exceeds Nature-family limits and will be cut at submission.")

# ============================================================ LZ 1
H("I. Comments from Prof. Zheng (LZ)", 14)
H("1. Where the 5.67 elasticity comes from: identity decomposition")
B("Request:", "Express CO2 as flights x seats per flight x load factor x distance x carbon intensity, estimate the elasticity of each component, show that they add up to 5.67, and quantify scale versus efficiency.")
B("What was done:", "Load factor is constant by construction (scheduled capacity) and cancels; this is stated. The identity CO2 = flights x (seats per flight) x (km per flight) x (CO2 per seat-km) is estimated component by component by 2SLS on the identical sample, and also within income and connectivity terciles and separately for international and domestic departures. A waterfall figure was added; the mediation analysis moves to Extended Data.")
m = R("_feyrer_mechanism_cl.csv").set_index("outc")
tot = m.loc["ln_co2_tot", "b"]
rows = []
for k, lbl in [("ln_flights", "Flights (frequency)"), ("ln_gauge", "Seats per flight (gauge)"), ("ln_stage", "Km per flight (stage length)"), ("ln_skm", "Scale subtotal (= seat-km)"), ("ln_intensity", "CO2 per seat-km (intensity)"), ("ln_co2_tot", "Total CO2 (Table 1)")]:
    rows.append([lbl, cs(m.loc[k, "b"], m.loc[k, "se"]), f"{100*m.loc[k,'b']/tot:.0f}%"])
T(["Component", "Elasticity to GACI (s.e.)", "Share of total"], rows, note="Same Feyrer instrument, controls and fixed effects as Table 1; N = 4,634; first-stage F = 18.4. " + SE)
d = R("_decomp_hetero_cl.csv")
def dget(blk, grp, seg, outc):
    return d[(d.block == blk) & (d.grp == grp) & (d.seg == seg) & (d.outc == outc)].iloc[0]
rows = []
for blk, grp, seg, lbl in [("A_income", "inc_low", "tot", "Low income"), ("A_income", "inc_mid", "tot", "Middle income"), ("A_income", "inc_high", "tot", "High income"),
                           ("C_segment", "all", "intl", "International departures"), ("C_segment", "all", "dom", "Domestic departures")]:
    cells = [lbl] + [cs(dget(blk, grp, seg, o).b, dget(blk, grp, seg, o).se) for o in ["ln_co2", "ln_flights", "ln_gauge", "ln_stage", "ln_intensity"]] + [f"{dget(blk, grp, seg, 'ln_co2').kpf:.1f}"]
    rows.append(cells)
T(["Group / segment", "Total CO2", "Flights", "Seats per flight", "Km per flight", "CO2 per seat-km", "First-stage F"], rows, fs=8.5,
  note="Components add to the total within each row. Within-tercile first stages are weak (F below 6), so these rows are indicative. " + SE)
B("Reading:", "The response is a frequency response. Flights carry the whole effect; gauge, stage length and intensity are not individually distinguishable from zero under clustering. In point estimates scale is 107 percent of the total and efficiency offsets 7 percent; even at the lower end of its 95 percent interval (-0.99) the intensity term would offset only 16 percent of scale. Domestic emissions do not respond.")
B("In the manuscript:", "Layer 2 subsection, Table 2 (tab:decomp), waterfall figure, ED Table 4, Methods equations (2) and (3).")
B("Remaining:", "The split of intensity into within-type fuel efficiency and fleet mix needs an aircraft-type aggregate from Fangyu.")

# ============================================================ LZ 2
H("2. Airport level: is the effect hub-driven?")
B("Request:", "Use the 6,000+ airport panel for heterogeneity by hub status, connectivity, orientation, baseline size, region and income; rank airports by connectivity-induced emissions and report the shares of the top 1, 5 and 10 percent.")
B("What was done:", "Airport and country-by-year fixed effects (clustered by airport), so airports are compared within the same country and year; splits by 1996 hub status, connectivity and traffic terciles, international share, region and host-country income; topology-only regressors as a safeguard; airport-level attribution ranked with Lorenz curves and Gini against the concentration of emissions themselves; an airport Feyrer-shifter pilot.")
h = pd.read_csv(os.path.join(HERE, "_airport_hetero.csv"), keep_default_na=False)  # region code "NA" must not become NaN
h["p"] = pd.to_numeric(h["p"], errors="coerce"); h["b"] = pd.to_numeric(h["b"]); h["se"] = pd.to_numeric(h["se"]); h["nap"] = pd.to_numeric(h["nap"])
def ah(panel, grp, outc="ln_co2"):
    return h[(h.panel == panel) & (h.grp == grp) & (h.outc == outc)].iloc[0]
rows = []
for panel, grp, lbl in [("ALL", "all", "All airports"), ("HUB", "top1", "Global top 1 percent (1996 seat-km)"), ("HUB", "not_top1", "Others"), ("HUB", "top5", "Global top 5 percent"), ("HUB", "not_top5", "Others"),
                        ("GAC", "gac1", "Baseline connectivity: low tercile"), ("GAC", "gac2", "middle"), ("GAC", "gac3", "high"),
                        ("ORI", "domestic_only", "Domestic-only airports"), ("ORI", "international", "International share above one half"),
                        ("REG", "AF", "Africa"), ("REG", "EU", "Europe"), ("REG", "AS", "Asia"), ("REG", "NA", "North America"), ("TOPO", "close", "Topology only: closeness centrality")]:
    r = ah(panel, grp); ri = ah(panel, grp, "ln_intensity")
    rows.append([lbl, cs(r.b, r.se, r.p), cs(ri.b, ri.se, ri.p), f"{int(r.nap):,}"])
it = h[(h.panel == "INT") & (h.grp == "top5_x_gaci") & (h.outc == "ln_co2")].iloc[0]
rows.append(["Pooled interaction: top 5 percent x ln GACI (CO2)", cs(it.b, it.se, it.p), "", f"{int(it.nap):,}"])
T(["Airport group", "log CO2 on log airport GACI", "log CO2/seat-km on log GACI", "Airports"], rows, fs=8.5,
  note="Airport and country-by-year fixed effects; standard errors clustered by airport. Each country's largest airport is absorbed by country-by-year effects and enters only through the interaction.")
c = R("_airport_concentration.csv")
T(["Measure (3,833 airports, 2023)", "Top 1% share", "Top 5%", "Top 10%", "Top 25%", "Gini"], [[r.measure, f"{100*r.top1:.1f}", f"{100*r.top5:.1f}", f"{100*r.top10:.1f}", f"{100*r.top25:.1f}", f"{r.gini:.3f}"] for r in c.itertuples()],
  note="Attributed = CO2 in 2023 x [1 - exp(-b x change in log GACI 1996-2023)]; airport sum 345 Mt versus country total 356 Mt. The 260 airports in the 1996 global top 5 percent hold 68.5 percent of the attributed total against 75.3 percent of emissions.")
tp = R("_airport_top20.csv").head(8)
T(["Airport", "Country", "CO2 2023 (Mt)", "Attributed (Mt)", "Share of world attributed"], [[r.airport_iata, r.iso3, f"{r.co2_mt:.1f}", f"{r.att_country:.1f}", f"{100*r.share_country:.1f}%"] for r in tp.itertuples()])
B("Reading:", "The marginal response is systemic (hubs have the lower elasticity), but the emissions it generates are more concentrated than emissions themselves and are led by the new hubs rather than the 1996 incumbents. The airport identification pilot has no first stage (F 0.4 to 0.7), so airport results stay descriptive.")
B("In the manuscript:", "Layer 1 subsection 'Airport-level evidence', concentration figure, Table 3, ED Table 5, SI Table 2; Layer 5: covering the top 5 percent of airports covers 85 percent of connectivity-driven emissions.")

# ============================================================ LZ 3
H("3. Formal spillover analysis")
B("Request:", "(a) first-order neighbour effect controlling for own GACI, (b) distance decay, (c) regional spillovers, and a link to the mechanisms.")
B("What was done:", "Eight weighting schemes; own and neighbour connectivity instrumented jointly (Sanderson-Windmeijer F), neighbour alone with the own shifter controlled, Anderson-Rubin sets by grid inversion, sub-region-by-year fixed effects, a permutation placebo (500 relabellings), and an own-versus-neighbour margins table.")
ar = R("_spill_cluster_ar.csv"); sb = R("_spill_bands_cl.csv"); sp = R("_spill_supp.csv")
def ga(item, var):
    r = ar[(ar.item == item) & (ar["var"] == var)]; return r.iloc[0] if len(r) else None
def gb(panel, item, var):
    r = sb[(sb.panel == panel) & (sb.item == item) & (sb["var"] == var)]; return r.iloc[0] if len(r) else None
rows = []
r = ga("single_contig_isocode", "nbr_g_contig"); a = ga("single_contig_AR", "nbr_g_contig")
rows.append(["Contiguous countries; neighbour instrumented, own shifter controlled", "controlled", cs(r.b, r.se, r.p), f"[{a.ar_lo:.1f}, {a.ar_hi:.1f}]", f"{r.kpf:.1f}"])
r = ga("single_contig_isocode_rxy", "nbr_g_contig")
rows.append(["  same, with sub-region x year fixed effects", "controlled", cs(r.b, r.se, r.p), "--", f"{r.kpf:.1f}"])
r = ga("joint_contig_isocode", "nbr_g_contig"); o = ga("joint_contig_isocode", "ln_gaci_cwm")
rows.append(["Contiguous; own and neighbour jointly instrumented", cs(o.b, o.se, o.p), cs(r.b, r.se, r.p), f"[{r.ar_lo:.1f}, {r.ar_hi:.1f}] (subset AR)", f"{r.kpf:.1f} (SW {r.swf_own:.1f} / {r.swf_nbr:.1f})"])
for W, lbl in [("knn5", "Five nearest countries; jointly instrumented"), ("b1", "Within 500 km; jointly instrumented")]:
    r = ga(f"joint_{W}_isocode", f"nbr_g_{W}"); o = ga(f"joint_{W}_isocode", "ln_gaci_cwm")
    rows.append([lbl, cs(o.b, o.se, o.p), cs(r.b, r.se, r.p), f"[{r.ar_lo:.1f}, {r.ar_hi:.1f}]", f"{r.kpf:.1f}"])
hy = gb("A", "hybrid_rf_control", "nbr_g_inv")
rows.append(["Inverse distance, all countries (August draft)", "RF-controlled", cs(hy.b, hy.se, hy.p), "--", f"{hy.kpf:.1f}"])
T(["Specification (log bunker CO2)", "Own ln GACI", "Neighbour ln GACI", "AR 95% set (neighbour)", "KP F"], rows, fs=8.5,
  note="Permutation placebo (countries relabelled in the weight matrix, 500 draws): p = 0.004 for contiguity, 0.010 for five nearest, 0.146 for inverse distance. " + SE)
rows = []
for v, lbl in [("nbr_g_contig", "Contiguous"), ("nbr_g_b1", "0-500 km"), ("nbr_g_b2", "500-1,000 km"), ("nbr_g_b3", "1,000-2,000 km"), ("nbr_g_b4", "2,000-5,000 km"), ("nbr_g_b5", "Beyond 5,000 km")]:
    r = gb("B", "band_single", v); rows.append([lbl, cs(r.b, r.se, r.p), f"{r.kpf:.1f}"])
T(["One neighbour definition at a time", "Neighbour ln GACI", "KP F"], rows, note="No smooth decay: the effect sits on land neighbours; the far band is large, negative and imprecise, a sign that wide exposures track global trends.")
rows = []
for yv, lbl in [("ln_co2_tot", "Bunker CO2"), ("ln_co2_intl", "International CO2"), ("ln_co2_dom", "Domestic CO2"), ("ln_skm", "Seat-km"), ("ln_flights", "Flights"), ("ln_gauge", "Seats per flight"), ("ln_stage", "Km per flight"), ("ln_intensity", "CO2 per seat-km")]:
    o = sp[(sp.panel == "MECH") & (sp.item == f"{yv}_country") & (sp["var"] == "ln_gaci_cwm")].iloc[0]; n = sp[(sp.panel == "MECH") & (sp.item == f"{yv}_country") & (sp["var"] == "nbr_g_contig")].iloc[0]
    rows.append([lbl, cs(o.b, o.se, o.p), cs(n.b, n.se, n.p)])
T(["Outcome (contiguity, joint 2SLS)", "Own ln GACI", "Neighbour ln GACI"], rows, note="Neighbours' connectivity raises own seat-km and lowers CO2 per seat-km; it loads on international rather than domestic emissions. " + SE)
B("Reading:", "The August inverse-distance estimate does not survive clustering or the permutation placebo. The contiguity estimate does, in both the single-endogenous and the joint form (Anderson-Rubin sets exclude zero), and it survives sub-region-by-year fixed effects. The spillover is therefore reported as a land-neighbour effect of about 2.")
B("In the manuscript:", "Layer 4, Table 4, ED Table 6, ED Figs 4 and 5, Methods 'Spillover design'.")
B("Remaining:", "Conley spatial HAC errors; wild cluster bootstrap for the 22 sub-regions.")

# ============================================================ LZ 4
H("4. Strengthening the exclusion restriction")
B("Request:", "More tests that the instrument affects aviation CO2 only through air connectivity.")
B("What was done:", "Three direct tests: (1) placebo outcomes, whether non-aviation emissions respond to the instrument; (2) placebo geography, whether the same aviation cycle interacted with sea instead of air market access explains aviation CO2; (3) zero first stage, whether emissions respond where the instrument does not move connectivity. Six secondary checks are kept in ED Table 7 and SI Table 1.")
e = R("_exclusion_suite_cl.csv")
def ge(panel, item, stat):
    return e[(e.panel == panel) & (e.item == item) & (e.stat == stat)].iloc[0]
rows = [["(1) Reduced form on aviation CO2 (reference)", cs(ge("A", "ln_co2_tot", "RF").b, ge("A", "ln_co2_tot", "RF").se), "--"],
        ["(1) Reduced form on total fossil CO2 excl. domestic aviation", cs(ge("A", "ln_co2_exav", "RF").b, ge("A", "ln_co2_exav", "RF").se), "Passes: no general development channel"],
        ["(2) Sea-geography interaction, entered with the air interaction", cs(ge("C", "sea_with_air", "RF").b, ge("C", "sea_with_air", "RF").se), "Passes: sea geography carries nothing"],
        ["(2) Air-geography interaction, entered with the sea interaction", cs(ge("C", "air_with_sea", "RF").b, ge("C", "air_with_sea", "RF").se), "Air geography carries the effect (first stage F 3.0 in this joint form)"],
        ["(3) Top baseline-connectivity tercile: first stage", cs(ge("D", "con3", "FS").b, ge("D", "con3", "FS").se) + f", F = {ge('D','con3','FS').f:.1f}", "Instrument does not move connectivity here"],
        ["(3) Top tercile: reduced form on CO2", cs(ge("D", "con3", "RF").b, ge("D", "con3", "RF").se), "Passes: no direct effect where the channel is absent"],
        ["(3) Lower terciles: reduced form on CO2", cs(ge("D", "con1", "RF").b, ge("D", "con1", "RF").se) + " / " + cs(ge("D", "con2", "RF").b, ge("D", "con2", "RF").se), "for comparison"]]
T(["Direct test", "Coefficient (s.e.)", "Reading"], rows, note=SE)
alt = [ge("F", k, "IV").b for k in ["fey_fl", "fey_skm", "fey_gaci", "fey_eff", "fey_th05", "fey_th15"]]
rows = [["Development controls added (GDP, trade, urban, FDI, arrivals)", f"{min(ge('E',k,'IV').b for k in ['c1_gdp','c2_trade','c3_urban','c4_fdi','c5_arrivals']):.2f} to {max(ge('E',k,'IV').b for k in ['c1_gdp','c2_trade','c3_urban','c4_fdi','c5_arrivals']):.2f}; F 9.6 to 12", "stable"],
        ["Alternative technology series; decay exponents 0.5 / 1.5", f"{min(alt):.2f} to {max(alt):.2f}; F 13 to 19", "stable"],
        ["Leads of the instrument (t+3, t+5)", f"{ge('G','lead3_lead','RF').b:.2f}*** / {ge('G','lead5_lead','RF').b:.2f}***", "uninformative: the index is a smooth trend"],
        ["Hansen J against the tourism-heritage instrument", f"J = {ge('H','hansen_j','J').b:.1f}, p = {ge('H','hansen_j','J').p:.2f}", "rejects at 5 percent; expected, that instrument fails the stress test"],
        ["Rival cycles x geography (world GDP, trade)", f"first-stage F {ge('B','plus_gdp','FS').f:.1f} / {ge('B','plus_trade','FS').f:.1f} when added", "caveat: not separable from the aviation cycle (r = 0.82, 0.84)"],
        ["Coal and cement CO2 reduced forms", f"{cs(ge('A','ln_coal','RF').b, ge('A','ln_coal','RF').se)} / {cs(ge('A','ln_cement','RF').b, ge('A','ln_cement','RF').se)}", "caveat: sectoral recomposition co-moves with the instrument"],
        ["Robust versus clustered s.e. of the headline", f"0.41 versus {ge('I','cluster_country','IV').se:.2f}", "inference convention changed to clustering"]]
T(["Secondary check (ED Table 7, SI Table 1)", "Result", "Reading"], rows, fs=8.5)
B("Reading:", "The three direct tests pass. Two caveats remain and are stated rather than resolved. Alternative instrument constructions that would avoid them (air-minus-sea advantage x cycle, F = 1.0; air interaction with the sea interaction as a control, F = 3.0) have no usable first stage, so the Feyrer interaction remains the single instrument.")
B("In the manuscript:", "Methods 'Exclusion-restriction diagnostics' (two paragraphs), ED Table 7, SI Table 1, ED Fig. 6; introduction and Discussion refer to the three tests.")

# ============================================================ LZ 5
H("5. Five-layer structure")
P("Results reordered as causal effect (country and airport), decomposition, heterogeneity, spillovers, policy. Methods gains subsections on the identity, the airport design and pilot, the exclusion diagnostics, the spillover design, and functional form and continuous heterogeneity. Mediation, the temporal split and measure robustness are in Extended Data. The August heterogeneity numbers (11.43, 5.91, 0.35), which were estimated on international CO2 but labelled as total, are corrected (see II.4 below).")

# ============================================================ Yifu
H("II. Comments from Yifu", 14)
H("1. The low-base problem: semi-log and level specifications")
B("What was done:", "log CO2 on GACI in levels (semi-log), CO2 in Mt on GACI in levels, CO2 in Mt on log GACI, and the semi-log and level slopes within baseline-connectivity terciles.")
y = R("_yifu_suite_cl.csv")
def gy(panel, item, var=None):
    r = y[(y.panel == panel) & (y.item == item)]
    if var: r = r[r["var"] == var]
    return r.iloc[0]
mg = gy("C", "moment_gaci").b; mc = gy("C", "moment_co2_mt").b
rows = [["Log CO2 on log GACI (reference)", cs(gy("A","loglog").b, gy("A","loglog").se, gy("A","loglog").p), f"{gy('A','loglog').b:.2f}", f"{gy('A','loglog').kpf:.1f}"],
        ["Log CO2 on GACI level (semi-log)", cs(gy("A","semilog").b, gy("A","semilog").se, gy("A","semilog").p), f"{gy('A','semilog').b*mg:.2f} (x mean GACI {mg:.2f})", f"{gy('A','semilog').kpf:.1f}"],
        ["CO2 (Mt) on GACI level", cs(gy("A","levellevel").b, gy("A","levellevel").se, gy("A","levellevel").p), f"{gy('A','levellevel').b*mg/mc:.2f} (x mean GACI / mean CO2 {mc:.2f} Mt)", f"{gy('A','levellevel').kpf:.1f}"]]
for t, lbl in [(1, "Semi-log, low baseline-connectivity tercile"), (2, "Semi-log, middle"), (3, "Semi-log, high")]:
    s_ = gy("A", f"semilog_con{t}"); m_ = gy("A", f"meangaci_con{t}").b
    rows.append([lbl, cs(s_.b, s_.se, s_.p), f"{s_.b*m_:.2f} (x mean GACI {m_:.2f})", f"{s_.kpf:.1f}"])
T(["Specification", "2SLS (s.e.)", "Implied elasticity at means", "KP F"], rows, note="GACI (capacity-weighted mean) has a sample mean of 1.09 and a 1996 range of 0.56 to 2.57, so one unit is roughly a doubling for a typical country. " + SE)
B("Reading:", "The gradient survives in absolute terms: the same addition to the network raises emissions about ten times more in the bottom tercile than in the top. Level specifications are dominated by large emitters and imprecise; the semi-log results are the informative complement.")
B("In the manuscript:", "Layer 3 subsection 'What the elasticity represents: the take-off stage', ED Table 8.")

H("2. Weak post-2010 first stage, the take-off interpretation, the accident instrument")
B("What was done:", "The headline is now interpreted as the emissions response during the take-off stage of network development (abstract, Results, Discussion). The Aviation Safety Network accident-instrument pilot is reported in SI Table 3.")
tmp = R("_temporal_co2_cl.csv"); t8 = R("_temporal_pre2008_cl.csv"); txc = R("_temporal_excovid_cl.csv")
def tt(df, col, val):
    return df[(df[col] == val) & (df.outc == "ln_co2_tot")].iloc[0]
rows = [["Full sample 1996-2023", cs(tt(tmp,"period","full").b, tt(tmp,"period","full").se, tt(tmp,"period","full").p), f"{tt(tmp,'period','full').kpf:.1f}"],
        ["Network-expansion era 1996-2007", cs(tt(t8,"samp","1996-2007").b, tt(t8,"samp","1996-2007").se, tt(t8,"samp","1996-2007").p), f"{tt(t8,'samp','1996-2007').kpf:.1f}"],
        ["Mature era 2010-2023 excl. 2020-2021", cs(tt(txc,"samp","2010-2023_exCOVID").b, tt(txc,"samp","2010-2023_exCOVID").se, tt(txc,"samp","2010-2023_exCOVID").p), f"{tt(txc,'samp','2010-2023_exCOVID').kpf:.1f}"],
        ["Unrestricted 2010-2023 (COVID years included)", cs(tt(tmp,"period","2010-2023").b, tt(tmp,"period","2010-2023").se, tt(tmp,"period","2010-2023").p), f"{tt(tmp,'period','2010-2023').kpf:.1f}"]]
T(["Sample (log bunker CO2)", "2SLS (s.e.)", "KP F"], rows, note=SE)
asn = R("_asn_iv_results.csv")
def ga2(stage, z, outc):
    return asn[(asn.stage == stage) & (asn.z == z) & (asn.outc == outc)].iloc[0]
rows = []
for stage, z, lbl in [("IV_asn", "L.n_fatal", "Lagged fatal accidents, alone"), ("IV_asn", "L.ash_deaths", "Lagged accident deaths (asinh), alone"),
                      ("OVERID", "fey+L.n_fatal", "Feyrer + lagged fatal accidents (overidentified)"), ("OVERID", "fey+L.ash_deaths", "Feyrer + lagged accident deaths (overidentified)")]:
    r = ga2(stage, z, "ln_co2_tot")
    rows.append([lbl, cs(r.b, r.se, r.p), f"{r.kpf:.1f}", (f"{r.jp:.2f}" if stage == "OVERID" else "--")])
T(["Accident instrument (log bunker CO2)", "2SLS (s.e.)", "KP F", "Hansen J p-value"], rows, note="Weak first stages: partial F 4.4 (lagged fatal accidents), 2.0 (lagged deaths), 0.2 to 1.2 (accident indicators); the KP F column for the overidentified rows reflects the Feyrer instrument. Heteroskedasticity-robust standard errors (pilot not yet re-run with clustering). Accident variables: ASN country-year counts of fatal accidents and deaths, lagged one year.")
B("Reading:", "The identifying variation is concentrated in the expansion era; the elasticity halves in the mature era. The accident instrument has a weak first stage (partial F of 4.4 for lagged fatal accidents, 2.0 for lagged deaths, 0.2 to 1.2 for the accident indicators) and cannot stand alone; its point estimates are nonetheless close to the Feyrer estimate and the overidentification test does not reject, so it is reported as a consistency check rather than as an identification strategy.")
B("In the manuscript:", "Layer 3 subsection, ED Table 1, Methods 'Instruments and validity', SI Table 3.")

H("3. Sensitivity of the attribution to a common elasticity")
B("What was done:", "The 2023 attribution, share = 1 - exp(-b x change in log GACI 1996-2023) applied to 2023 emissions, is recomputed with country-specific elasticities: by baseline-connectivity tercile, by income tercile, from the continuous interaction (linear and quadratic), and with the semi-log slope applied to the absolute change in GACI.")
sens = R("_attribution_sensitivity.csv")
T(["Elasticity rule", "Attributed 2023 CO2 (Mt)", "Share of 2023 emissions (%)"], [[r.rule, f"{r.attributed_mt:.0f}", f"{r.share_pct:.1f}"] for r in sens.itertuples()],
  note="World 2023 bunker emissions in the 184-country sample: 832 Mt (838 Mt including countries outside the estimation sample, which gives the 42.5 percent headline).")
a = R("_attribution_hetero.csv")
a["att_c"] = (1 - (-a.b_common * a.dln_cwm).apply(math.exp)) * a.e_tot_t / 1e6
a["att_t"] = (1 - (-a.b_con3 * a.dln_cwm).apply(math.exp)) * a.e_tot_t / 1e6
a["att_q"] = (1 - (-a.b_quad * a.dln_cwm).apply(math.exp)) * a.e_tot_t / 1e6
top = a.sort_values("att_c", ascending=False).head(8)
rows = [[r.c, f"{r.gaci0:.2f}", f"{r.att_c:.1f}", f"{r.att_t:.1f}", f"{r.att_q:.1f}"] for r in top.itertuples()]
deu = a[a.c == "DEU"].iloc[0]; rows.append(["DEU", f"{deu.gaci0:.2f}", f"{deu.att_c:.1f}", f"{deu.att_t:.1f}", f"{deu.att_q:.1f}"])
T(["Country", "Baseline GACI 1996", "Common 5.67 (Mt)", "Connectivity-tercile rule", "Continuous (quadratic) rule"], rows)
B("Reading:", "The headline is the upper end of a 35 to 43 percent range. Downward revisions are largest for mature networks (United States -10 Mt, China -8, Japan -3, Spain -3) and Germany's negative attribution shrinks from -16 to -11; country ranks are unchanged.")
B("In the manuscript:", "Layer 5 attribution paragraph (range sentence), ED Table 10, Discussion.")

H("4. Continuous heterogeneity by baseline connectivity")
B("What was done:", "Quintile splits on 1996 GACI; pooled 2SLS with log GACI interacted with centred log baseline GACI (linear and quadratic), instruments interacted identically; fitted elasticities at baseline percentiles; a gradient figure.")
rows = []
for k in range(1, 6):
    r = gy("B", f"q{k}_loglog"); g0 = gy("B", f"q{k}_gaci0")
    rows.append([f"Q{k}", f"{g0.b:.2f} ({g0.v12:.2f} to {g0.v33:.2f})", cs(r.b, r.se, r.p) if r.kpf > 1 else "unidentified", f"{r.kpf:.1f}"])
T(["Quintile of 1996 GACI", "Baseline GACI (range)", "Log-log elasticity (s.e.)", "KP F"], rows, note=SE)
b1 = gy("B", "int_lin", "ln_gaci_cwm"); b2 = gy("B", "int_lin", "x_dev"); m0 = gy("C", "centre_lg0").b; lg0 = gy("C", "moment_lg0")
def blin(x): return b1.b + b2.b * (x - m0)
rows = [["ln GACI", cs(b1.b, b1.se, b1.p)], ["ln GACI x (ln baseline GACI - mean)", cs(b2.b, b2.se, b2.p)], ["Joint first-stage F", f"{b1.kpf:.1f}"],
        ["Fitted elasticity at p10 / p50 / p90 of baseline connectivity", f"{blin(lg0.v12):.2f} / {blin(lg0.v13):.2f} / {blin(lg0.v33):.2f}"]]
T(["Pooled linear interaction", "Estimate"], rows, note="The quadratic version adds a positive squared term (p = 0.08), indicating a floor near 4 rather than a continued decline. " + SE)
ht = R("_feyrer_hetero_tot_cl.csv")
def gh(panel, grp): return ht[(ht.panel == panel) & (ht.grp == grp)].iloc[0]
rows = [[lbl, cs(gh(p_, g).b, gh(p_, g).se, gh(p_, g).p), f"{gh(p_, g).KPF:.1f}"] for p_, g, lbl in
        [("A_income", "inc_low", "Income tercile: low"), ("A_income", "inc_mid", "middle"), ("A_income", "inc_high", "high"), ("A_income", "interact_hi", "  pooled interaction, top tercile x ln GACI"),
         ("B_conn", "con_low", "Connectivity tercile: low"), ("B_conn", "con_mid", "middle"), ("B_conn", "con_high", "high"), ("B_conn", "interact_hi", "  pooled interaction, top tercile x ln GACI"),
         ("C_region", "Europe", "Europe"), ("C_region", "AsiaPacific", "Asia-Pacific"), ("C_region", "Africa", "Africa"), ("C_region", "LatAm", "Latin America")]]
T(["Total CO2 elasticity by group (corrected; August draft used international CO2)", "2SLS (s.e.)", "KP F"], rows, fs=8.5, note=SE)
B("Reading:", "A step rather than a smooth slope: the two lowest quintiles (40 percent of countries) share the high elasticity, the decline occurs between the second and third quintiles, and the fourth quintile has no instrument variation. Split-sample first stages are weak under clustering, so the pooled interactions carry the gradient.")
B("In the manuscript:", "ED Table 9, gradient figure, ED Table 2, subsection text.")

H("III. Remaining weak points", 14)
P("(1) The inverse-distance spillover is not robust and the contiguity joint specification has a weak first stage, defended by Anderson-Rubin sets; (2) sectoral recomposition and cycle collinearity in the exclusion diagnostics; (3) identification concentrated in the expansion era and in low-baseline countries (top tercile F 0.2, fourth quintile unidentified); (4) the August heterogeneity mislabel must be communicated; (5) no causal airport estimate; (6) level specifications imprecise; (7) domestic-traffic results imprecise; (8) inference convention changed; (9) length above Nature limits and no local compilation (Overleaf required).")
P("Related files: Response_LZ_results_20260903.docx, Response_Yifu_results_20260903.docx, CO2_results_summary_20260903_LZ.xlsx, co2_overleaf_20260903_cl.zip.")
out = os.path.join(HERE, "Comment_response_memo_20260903.docx")
try:
    doc.save(out)
except PermissionError:
    out = out.replace(".docx", "_v2.docx")
    try:
        doc.save(out)
    except PermissionError:
        out = out.replace("_v2.docx", "_v3.docx"); doc.save(out)
    print("target locked; saved as", out)
txt = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
print("saved", out, "| tables:", len(doc.tables), "| em-dash:", txt.count(chr(8212)), "| words:", len(txt.split()))
