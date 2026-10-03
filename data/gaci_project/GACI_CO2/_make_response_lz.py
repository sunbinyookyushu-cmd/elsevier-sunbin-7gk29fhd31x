# -*- coding: utf-8 -*-
"""Build point-by-point response docx to LZ comments (2026-09-02)."""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Response_plan_LZ_comments_20260902.docx")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.0)

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcol)
    tcPr.append(shd)

def para(text, bold=False, italic=False, size=None, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Calibri"; r.font.color.rgb = RGBColor(0x23, 0x2A, 0x55)
    return h

def comment_box(text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, "EEF1F7")
    c.paragraphs[0].text = ""
    lab = c.paragraphs[0].add_run("Comment (LZ): ")
    lab.bold = True; lab.font.color.rgb = RGBColor(0x23, 0x2A, 0x55)
    first = True
    for line in text.strip().split("\n"):
        if first:
            c.paragraphs[0].add_run(line).italic = True
            first = False
        else:
            p = c.add_paragraph(); p.add_run(line).italic = True
            p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def response(items, label=True):
    if label:
        p = doc.add_paragraph()
        r = p.add_run("Response and plan: "); r.bold = True; r.font.color.rgb = RGBColor(0xB0, 0x6A, 0x1E)
        p.paragraph_format.space_after = Pt(2)
    for kind, txt in items:
        if kind == "p":
            q = doc.add_paragraph(txt); q.paragraph_format.space_after = Pt(4)
        elif kind == "b":
            q = doc.add_paragraph(txt, style="List Bullet"); q.paragraph_format.space_after = Pt(2)
        elif kind == "n":
            q = doc.add_paragraph(txt, style="List Number"); q.paragraph_format.space_after = Pt(2)
        elif kind == "s":
            q = doc.add_paragraph(); rr = q.add_run(txt); rr.bold = True; q.paragraph_format.space_after = Pt(2)

def table(headers, rows, widths=None, fs=9):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(fs)
        shade(c, "D9DEEA")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(fs)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------------------------------------------------------- title
para("Response to comments on the empirical analyses (GACI and aviation CO2)", bold=True, size=15, color="232A55")
para("Point-by-point response and implementation plan. Prepared 2 September 2026 by Sunbin Yoo. "
     "Comments are quoted from Comments_GAC&CO2_LZ.docx. Current manuscript: main_co2_nature_20260826.tex; "
     "current results workbook: CO2_results_summary_20260826_final.xlsx.", size=9.5, color="555555")

heading("Overall comment", 1)
comment_box("Overall, I think the current results are already strong and provide a clear and potentially important story about the carbon consequences of global air connectivity. However, if we target Nature-series journals, I suggest further strengthening the empirical analysis in four directions. The purpose is not to add more conventional robustness checks, but to make the causal mechanisms, spatial scope, and identification strategy sufficiently comprehensive for a broad climate and sustainability audience.")
response([
    ("p", "We agree with the direction and with the five-layer organisation proposed at the end of the comments. "
          "Two of the four requests (decomposition, spillovers) are largely covered by estimates that already exist in the results folder but are not yet displayed in the form requested; "
          "we will re-position and extend them. The other two (airport-level evidence, exclusion restriction) require new estimation. "
          "The table below summarises the status of each request; details follow comment by comment."),
])
table(["Request", "Already available", "To be added", "New displays"],
      [["1 Decompose the 5.67 elasticity", "Exact additive log decomposition on the accounting identity (flights, seats per flight, km per flight, CO2 per seat-km); IV mediation", "Identity written in the requested form; scale vs efficiency subtotals; decomposition by income group and by domestic/international; load-factor statement", "1 table + 1 waterfall figure (main text); 1 ED table"],
       ["2 Airport-level analysis", "5,354-airport panel 1996-2023 (97,165 airport-years); descriptive FE regressions; efficiency curve; hub mismatch table", "Heterogeneity by hub status, connectivity, orientation, baseline size, region and income; concentration of the attributed effect (top 1/5/10 percent); airport-level IV pilot", "1 table + 1 concentration figure; Methods paragraph on identification at airport level"],
       ["3 Formal spillover analysis", "Inverse-distance neighbour connectivity with matched Feyrer instrument (F = 134): own CO2 +7.45; neighbour mechanism decomposition", "Own-connectivity control; distance bands and continuous decay; within-region vs out-of-region exposure with region-by-year FE; placebo weights; spatial HAC", "1 extended spillover table + 1 decay figure; own-vs-neighbour decomposition table"],
       ["4 Exclusion restriction", "Size-by-cycle stress test (F 120 to 106); plausibly-exogenous bounds (85 / 42 percent); failed alternative instruments disclosed", "Non-aviation placebo outcomes; world-GDP-by-geography horse race; sea-geography placebo instrument; zero-first-stage test; development controls; alternative technology series", "1 diagnostics table (5 panels) + 1 ED figure"],
       ["5-layer structure", "Layers 1, 3, 5 exist in current form", "Re-order Results into Layers 1-5; move mediation to Extended Data", "Revised Results and Methods"]],
      widths=[3.2, 4.6, 5.2, 3.4])

# ---------------------------------------------------------------- comment 1
heading("1. Decompose the CO2 elasticity to explain where the 5.67 percent effect comes from", 1)
comment_box("""The most important result is that a 1% increase in GACI increases aviation CO2 by 5.67%. However, the reader will naturally ask: where does this very large elasticity come from? We should therefore conduct a systematic CO2 elasticity decomposition.
A useful starting point is to express total aviation CO2 as: CO2 = Flights x Seats/Flight x LoadFactor x CO2/Seat-km x Distance (1)
Alternatively, depending on the exact definitions and data availability, we can formulate the decomposition through flight activity, seat capacity, passenger activity, distance, and carbon intensity. The key point is that the decomposition should allow us to quantify how much of the 5.67% elasticity is associated with: the number of flights; seats per flight / aircraft size; passenger volume or load factor (maybe the data is unavailable); passenger-km/traffic expansion; average stage length or traffic composition; and emissions intensity, such as CO2 per seat-km or passenger-km.
We can then estimate the elasticity of each component with respect to GACI and show how these components add up to the overall CO2 elasticity: bCO2 = bTraffic + bAircraftSize + bLoadFactor + bDistance + bCarbonIntensity (2)
The exact decomposition should be consistent with the accounting identity used to construct airport/country-level emissions. This analysis would substantially improve the interpretation of the main result. In particular, it would allow us to distinguish between scale effects and efficiency effects. Rather than simply showing that connectivity increases CO2 while reducing emissions intensity, we can demonstrate quantitatively that the large expansion in aviation activity dominates the relatively small efficiency gains.
This is important because the 5.67% elasticity is sufficiently large that reviewers are likely to ask why the effect is so large. A transparent decomposition would provide a direct answer.""")
response([
    ("p", "This decomposition has already been estimated (co2_feyrer_mechanism.do, file _feyrer_mechanism.csv) but the current draft presents the mediation analysis instead. We will make the identity-based decomposition the main mechanism display, exactly in the form of equations (1) and (2)."),
    ("s", "What exists"),
    ("p", "The emissions data are built from scheduled flights, so the identity consistent with the data is CO2 = Flights x (Seats per flight) x (Km per flight) x (CO2 per seat-km), where km per flight is the average stage length and corresponds to the Distance term of equation (1). "
          "Each component is regressed on log GACI with the same Feyrer instrument, controls and fixed effects as Table 1, on the identical sample (N = 4,634, first-stage F = 154). Because the identity is multiplicative, the log elasticities add exactly:"),
])
table(["Component", "Elasticity", "s.e.", "Share of total"],
      [["Flights (frequency)", "+6.98***", "0.57", "123 percent"],
       ["Seats per flight (aircraft gauge)", "-0.30", "0.23", "-5 percent"],
       ["Km per flight (average stage length)", "-0.60***", "0.21", "-11 percent"],
       ["Subtotal: scale (= seat-km elasticity)", "+6.07***", "0.44", "107 percent"],
       ["CO2 per seat-km (carbon intensity)", "-0.40***", "0.12", "-7 percent"],
       ["Total CO2 (sum = Table 1 estimate)", "+5.67***", "0.41", "100 percent"],
       ["Memo: LTO share of CO2", "+0.014", "0.025", "composition"],
       ["Memo: international share of seat-km", "+0.096**", "0.048", "composition"]],
      widths=[6.5, 2.5, 2.0, 3.5])
response(label=False, items=[
    ("p", "The answer to the question of where 5.67 comes from is therefore that the scale response is a pure frequency response: connectivity adds flights, not larger aircraft or longer stages, and the efficiency gain (-0.40) offsets only about seven percent of it. Scale beats efficiency by a factor of roughly fifteen."),
    ("s", "What we will add"),
    ("n", "Main-text table in the form of equation (2), with scale and efficiency subtotals and the verification row showing that the components sum to the Table 1 estimate. A waterfall figure (0 to +6.98 to -0.30 to -0.60 to -0.40 = 5.67) as a main display."),
    ("n", "Decomposition by baseline income tercile (new do-file). The existing heterogeneity shows total elasticities of 11.4 (low income), 5.9 (middle) and 0.35 (high, not significant); the decomposition will show which component drives the gradient."),
    ("n", "Separate decompositions for domestic and international traffic (the delivery data carry a domestic/international split), to test whether stage length shortens on both segments."),
    ("n", "Load factor. The emissions are computed from scheduled capacity, so a load-factor term is constant by construction and cancels in the identity. We will state this in Methods and add a sensitivity paragraph that scales emissions by regional ICAO/IATA load factors. Passenger-km cannot be decomposed from these data."),
    ("n", "Carbon intensity sub-decomposition. Since gauge does not respond, the intensity gain must come through fleet fuel efficiency rather than up-gauging. We will ask Fangyu whether a seat-km-weighted fuel-efficiency-by-aircraft-type aggregate can be produced; if so, intensity will be split into within-type efficiency and type mix. Otherwise the fleet channel is stated as an interpretation with the LTO and stage-length shares as supporting evidence."),
    ("n", "A Methods paragraph on why the elasticity is large: OLS gives 3.59 and IV 5.67; the Feyrer instrument identifies a local effect among countries whose connectivity growth is driven by geography-exposed global aviation growth, which are predominantly network entrants (low-income, low-baseline-connectivity), for whom the frequency response is largest."),
    ("n", "The IV mediation analysis (product-of-coefficients and Imai-Keele-Yamamoto) moves to Extended Data as a complementary view."),
])

# ---------------------------------------------------------------- comment 2
heading("2. Add airport-level analysis to examine whether the global effect is concentrated in major hubs", 1)
comment_box("""Since our underlying dataset covers more than 6,000 airports, I think we should make better use of this particularly valuable feature of the data. In addition to the country-level causal estimates, we should conduct an airport-level analysis to examine whether the aggregate global effect is broadly distributed across airports or primarily driven by a relatively small number of major hubs.
The key question is: Is the global effect driven by a small number of hub airports? We could estimate airport-level effects of connectivity on CO2 and examine heterogeneity across: major international hubs vs. smaller airports; highly connected vs. weakly connected airports; domestic-oriented vs. internationally oriented airports; airports with different baseline traffic levels; and airports located in different regions or income groups.
It would also be useful to examine the contribution of airports to the aggregate effect. For example, we could rank airports by their estimated connectivity-induced emissions and calculate how much of the global effect is accounted for by the top 1%, 5%, 10%, etc. of airports.
If the results show strong concentration, this would provide an important policy insight: the climate consequences of connectivity expansion may be disproportionately generated by a relatively small number of globally important aviation nodes. If the effect is instead widely distributed, that is also an important finding because it would indicate that the carbon consequences of connectivity expansion are systemic rather than driven only by hubs. Therefore, the airport-level evidence would complement the country-level analysis and demonstrate that the main result is not merely an artifact of national aggregation.""")
response([
    ("s", "What exists"),
    ("p", "An airport-year panel (airport_co2_panel.csv: 5,354 airports, 1996-2023, 97,165 airport-years). The emissions delivery covers 6,356 airports; the panel keeps the 5,354 that are nodes of the GACI network, and the 1,002 airports without a GACI value account for 0.1 percent of departure CO2 over 1996-2023 (0.05 percent in 2023), so the restriction is immaterial for aggregate results and will be stated in Methods with descriptive regressions using airport fixed effects and country-by-year fixed effects (standard errors clustered by airport): log CO2 +3.72, log seat-km +4.18, log intensity -0.49, international share +0.33, all significant. "
          "Also an efficiency curve (median CO2 per 1,000 seat-km falls from 284 kg in the lowest connectivity ventile to 77 kg in the highest) and a 2023 hub table. "
          "An airport-level heritage-proximity instrument was piloted and abandoned because it failed the size-by-cycle test (F from 59.5 to 2.3); this is disclosed in Methods. Airport-level results are therefore currently descriptive, with the country-by-year fixed effects absorbing all national shocks."),
    ("s", "What we will add"),
    ("n", "Heterogeneity table (new do-file co2_airport_hetero.do), same specification, split samples and interactions along the five dimensions requested: (i) hub status (each country's largest airport; top 5 and top 1 percent by 1996 capacity) versus other airports; (ii) baseline connectivity terciles (1996 GACI); (iii) orientation (international share of seat-km: zero, below half, above half); (iv) baseline traffic terciles (1996 seat-km); (v) region and income group of the host country. Outcomes: log CO2, log seat-km, log intensity, international share."),
    ("n", "Contribution and concentration (new script). For each airport, connectivity-induced emissions are computed as CO2 in 2023 times [1 - exp(-b x change in log GACI 1996-2023)], with b taken from the country IV estimate (baseline) and from the airport-level and group-specific estimates as sensitivity. Airports are ranked and the shares of the global attributed effect accounted for by the top 1, 5, 10 and 25 percent are reported with a Lorenz curve and Gini coefficient. The benchmark is the concentration of physical emissions themselves in 2023, so the finding is whether connectivity-induced emissions are more or less concentrated than emissions. We will also check that airport-level attributed emissions aggregate to the country-level 356 Mt."),
    ("n", "Mechanical-correlation safeguard. Airport GACI embeds a capacity component, so we will add a specification using topology-only components (eigenvector, betweenness, closeness) as the regressor; intensity and international share are the informative outcomes in either case."),
    ("n", "Airport-level identification pilot, time-boxed to two days. Two candidates: (a) an airport-level Feyrer shifter, world aviation technology interacted with the airport's 1996 air market access, which varies within country under country-by-year fixed effects; (b) a shift-share exposure to the connectivity growth of the airport's 1996 foreign partner airports (leave-own-country-out). Both must pass the size-by-cycle test that eliminated the heritage instrument. If they pass, airport-level effects become causal; if not, the airport-level layer is presented as within-country descriptive evidence and the failed pilots are disclosed, as now."),
    ("n", "Manuscript placement: an 'Airport-level evidence' subsection in Layer 1 with one table and the concentration figure, plus one Methods paragraph."),
])

# ---------------------------------------------------------------- comment 3
heading("3. Expand neighbouring-country effects into a formal spillover analysis", 1)
comment_box("""The preliminary finding that a country's aviation CO2 responds to neighbouring countries' connectivity is potentially one of the most interesting aspects of the paper. I suggest developing this into a formal analysis of cross-border connectivity spillovers, rather than presenting it only as an additional result. We could then extend this analysis in several directions.
(a) First-order neighbour effects. Estimate the effect of neighbouring countries' GACI on country (i)'s aviation CO2 while controlling for its own GACI. This would directly establish whether connectivity expansion in one country generates additional aviation emissions in neighbouring countries.
(b) Distance-decay spillovers. Construct spatial exposure measures based on geographical distance and examine whether the spillover effect declines with distance. For example, we could distinguish: contiguous neighbours; countries within 500 km; 500-1,000 km; 1,000-2,000 km; and so forth. or use a continuous distance-decay function. This would help demonstrate whether the spillover is genuinely spatial rather than simply capturing broader regional trends.
(c) Regional spillovers. We could also examine spillovers within broader geographic regions or aviation markets. For example, connectivity expansion in one country may affect emissions in other countries within the same regional aviation network. This could be especially relevant for regions with highly integrated air transport systems.
Importantly, we should connect the spillover effects to the mechanisms/ decomposition already identified in the paper. This could become an important conceptual contribution of the paper: connectivity policies are not only domestic interventions; their carbon consequences can propagate through the international aviation network.""")
response([
    ("s", "What exists"),
    ("p", "Neighbour connectivity is the inverse-distance-weighted average of all other countries' log GACI (distances between aviation activity centroids, 100 km floor), instrumented by the identically weighted average of their Feyrer shifters (first-stage F = 134). A one percent rise in neighbour connectivity raises own total CO2 by 7.45 percent, international CO2 by 10.24 percent and intensity by 1.31 percent, with own shifter controlled in reduced form. "
          "The neighbour mechanism decomposition mirrors the own one: aircraft gauge +5.44 (p < 0.001), international share +1.05 (p < 0.001), stage length -2.03 (p = 0.045), frequency +2.73 not significant. Neighbours' connectivity therefore feeds larger, more international aircraft on somewhat shorter stages in the home country, the signature of feeder traffic into a growing regional hub. "
          "Jointly instrumenting own and neighbour connectivity is not possible with the current weights: the two shifters are nearly collinear within country and the joint first stage collapses (F about 5)."),
    ("s", "What we will add"),
    ("n", "(a) First-order effects with own GACI controlled. The headline specification instruments own and neighbour connectivity jointly. Because own GACI is itself endogenous, entering it as an exogenous control would contaminate the neighbour coefficient as well, so that specification is reported only as a transparently labelled reference column. To make joint instrumentation feasible, the weight matrix is re-defined (contiguity, five nearest neighbours, 500 km band) so that the neighbour shifter is no longer a smooth function of the same geography as the own shifter; Sanderson-Windmeijer conditional F statistics are reported for both endogenous variables. The current hybrid (neighbour instrumented, own shifter in reduced form) is kept as a second reference column. If no weight definition restores a joint first stage above conventional thresholds, this is reported as a limitation and the reduced-form hybrid remains the main estimate."),
    ("n", "(b) Distance decay. Leave-out averages by band: contiguous (Natural Earth borders), under 500 km, 500-1,000, 1,000-2,000, 2,000-5,000, over 5,000 km, each with its own band-matched instrument, entered jointly; plus a continuous kernel exp(-d/lambda) over a grid of lambda (250 to 5,000 km) with first-stage strength and fit reported. A coefficient profile figure will show the decay. Two placebos: random permutation of the weight rows (500 draws) and the most distant band alone, which should be close to zero if the effect is spatial rather than a common global trend."),
    ("n", "(c) Regional spillovers. Within-region and out-of-region leave-out exposures entered jointly, with regions defined as UN sub-regions and as aviation market blocs (EU single aviation market, ASEAN, GCC, Mercosur, North America, ECOWAS). The key robustness specification adds region-by-year fixed effects, which absorb any regional trend; a surviving distance-band effect is then genuinely spatial. Heterogeneity by region and by international orientation of the receiving country."),
    ("n", "Mechanism link. The own decomposition (Comment 1) and the neighbour decomposition will be shown side by side in one table: own connectivity works through frequency; neighbour connectivity works through gauge, internationalisation and shorter stages. This supports the conceptual point that hub growth in one country is fed by larger, more international aircraft on feeder-length stages from its neighbours, so the carbon consequences of connectivity policy cross borders. A domestic-only CO2 outcome will be added; spillovers should load on international rather than domestic emissions (consistent with 10.24 versus 7.45 now)."),
    ("n", "Inference: Conley spatial HAC standard errors at 500, 1,000 and 2,000 km cutoffs alongside region clustering."),
    ("n", "Manuscript placement: Layer 4 with an extended spillover table (panels: own-controlled, distance bands, regional with region-by-year FE), the decay figure, and the own-versus-neighbour decomposition table."),
])

# ---------------------------------------------------------------- comment 4
heading("4. Strengthen the exclusion restriction for the IV strategy", 1)
comment_box("""I think this is particularly important for a Nature-series submission. Our identification strategy uses an IV-2SLS regressions. However, Nature-series reviewers are likely to focus heavily on the exclusion restriction of IV, particularly because air connectivity is related to many dimensions of economic development. We therefore need more tests to support that the instrument affects aviation CO2 primarily through air connectivity.""")
response([
    ("s", "What exists"),
    ("p", "The instrument interacts world aviation technology with the country's 1996 air market access, with the sea market-access channel controlled directly. Three validity results are already in Methods: (i) a size-by-cycle stress test adding interactions of the global cycle with baseline population, GDP and airport capacity, which the instrument survives (F 120 to 106) and which eliminates the tourism-heritage and airport-heritage instruments; (ii) plausibly-exogenous bounds showing that a direct effect would need to account for 85 percent of the reduced form (42 percent for intensity) to overturn the results; (iii) reduced-form quintile plots."),
    ("s", "What we will add (in priority order; items 1-5 form the five panels of the diagnostics table, items 6-8 go to Extended Data)"),
    ("n", "Non-aviation placebo outcomes. The same reduced form and 2SLS applied to total CO2 excluding aviation (the OWID country series already in the folder minus our aviation series, so this panel can start immediately), road transport CO2 and power-sector CO2 (EDGAR sectoral), and where available maritime bunkers. If the instrument shifts aviation CO2 (reduced form 0.73) but not these, the general-development channel is rejected. Displayed as a reduced-form comparison bar chart (Extended Data)."),
    ("n", "Horse race on the content of the global cycle. Alongside aviation technology x geography, add world real GDP x geography, oil price x geography and world trade x geography. If the aviation-technology interaction retains the first stage and the second-stage coefficient, the identifying variation is aviation-specific rather than global growth."),
    ("n", "Placebo instrument. Aviation technology x 1996 sea market access should produce a weak first stage for GACI and a zero reduced form for aviation CO2; likewise shipping technology x air market access. This checks that the instrument works through aviation geography and not through geography in general."),
    ("n", "Zero-first-stage test (van Kippersluis and Rietveld). In the top tercile of baseline connectivity the shifter does not move GACI (first-stage F = 1.8). If the reduced form on CO2 is also zero in that group, a direct effect of the instrument is ruled out where it cannot operate through connectivity."),
    ("n", "Development-channel controls as bounds. Sequentially adding log GDP, trade openness, urbanisation, tourist arrivals and FDI (acknowledged as potentially bad controls) and reporting the coefficient path."),
    ("n", "Alternative technology series and distance decay: world flights, world seat-km, world GACI sum and a jet-fuel efficiency index as the technology term; market-access decay exponents of 0.5, 1 and 1.5. A narrow band of estimates shows the result does not depend on one functional form."),
    ("n", "Leads of the instrument (t+3, t+5) entered with the current value, with the caveat that the technology term is a global trend so this test is informative only about differential pre-trends by geography."),
    ("n", "Overidentification with the tourism instrument (Hansen J), not yet computed; it will be reported for reference only, given that instrument's weakness in the stress test. Conley spatial HAC plus two-way clustered standard errors throughout."),
])

# ---------------------------------------------------------------- layers
heading("5. Proposed organisation of the empirical analysis (five layers)", 1)
comment_box("""Taken together, I think the empirical analysis should be organized into the following layers. This would give the paper a much more complete structure and make the main contribution easier to communicate to a broad audience.
Layer 1: Causal effect. Country-level causal effect; Airport-level evidence; Establish the magnitude and robustness of the main 5.67% elasticity.
Layer 2: Mechanism/decomposition analysis. Traffic expansion VS. efficiency improvement; Decompose the total CO2 response into different parts.
Layer 3: Heterogeneity. When and where is the carbon effect largest?
Layer 4: Spillover effects. first-order neighbouring countries; distance-decay effects; regional spillovers.
Layer 5: Policy counterfactual""")
response([
    ("p", "We will restructure Results accordingly. Mapping of displays:"),
])
table(["Layer", "Content", "Main-text displays", "Extended Data"],
      [["1 Causal effect", "Country 2SLS (current Table 1); airport-level heterogeneity and concentration; identification pilot outcome", "Table 1; concentration figure", "Airport heterogeneity table; measures robustness"],
       ["2 Mechanism", "Identity decomposition with scale vs efficiency subtotals; by income and by domestic/international", "Decomposition table; waterfall figure", "Mediation (IV product and IKY); sub-decompositions"],
       ["3 Heterogeneity", "Income, baseline connectivity, region, intensity; temporal split with crisis years set aside (current)", "Heterogeneity coefficient plot", "Temporal plot; cutoff sweep"],
       ["4 Spillovers", "Own-controlled first-order effect; distance bands and decay; regional with region-by-year FE; own vs neighbour decomposition", "Spillover table; decay figure", "Placebo weights; kernel grid; Conley SEs"],
       ["5 Policy", "Attributed 356 Mt and its geography; social cost; SAF pathway; hub-coverage counterfactual based on Layer 1 concentration", "Attribution map; SAF figure", "SCC by country; carbon-price map"],
       ["Methods", "Identity in the form of equations (1)-(2); exclusion-restriction diagnostics", "", "Diagnostics table (5 panels); RF comparison figure"]],
      widths=[2.6, 6.4, 3.8, 3.6])
para("Nature-series display limits will be respected by keeping three tables and four to five figures in the main text and moving the rest to Extended Data and Supplementary Information.", size=10)

heading("6. Implementation schedule", 1)
table(["Step", "Work", "New inputs", "Estimate"],
      [["1", "Decomposition table, waterfall, income-group and domestic/international decompositions", "None (existing estimates plus one do-file)", "0.5 day"],
       ["2", "Exclusion-restriction diagnostics, panels 1-4", "EDGAR sectoral CO2; world GDP, oil price, world trade series (OWID total CO2 already in folder)", "1.5 days"],
       ["3", "Distance-band, kernel and regional spillovers with placebos", "Natural Earth borders; regional bloc mapping", "1.5 days"],
       ["4", "Airport heterogeneity and concentration analysis", "Topology components from the GACI construction files", "1 day"],
       ["5", "Airport-level identification pilot (time-boxed)", "Country population and coordinates for airport market access", "1-2 days"],
       ["6", "Own-controlled spillover with re-defined weights; remaining diagnostics panels", "None", "1 day"],
       ["7", "Manuscript restructuring into five layers; results workbook update; Overleaf synchronisation", "", "1 day"]],
      widths=[1.2, 7.6, 5.4, 2.2])
para("Total: seven to nine working days. Open items requiring a decision: whether to request the aircraft-type fuel-efficiency aggregate from Fangyu, and confirmation that EDGAR sectoral coverage (to 2022) is acceptable for the placebo outcomes.", size=10)

doc.save(OUT)
print("saved", OUT)
