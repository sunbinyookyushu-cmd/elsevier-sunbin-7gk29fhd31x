# -*- coding: utf-8 -*-
"""build_tables_docx.py -- parse gaci_tables_export.log -> publication Word tables.
Table 1: main OLS vs IV (robust + clustered SE), 3 outcomes x 2 treatments.
Table 2: heterogeneity (split-sample, robust SE), 3 outcomes x subgroups.
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = ['trade_share', 'trade_intensity', 'ln_tradevol']
OUTLAB = {'trade_share': 'Trade openness\n(trade % GDP)',
          'trade_intensity': 'Log trade/GDP', 'ln_tradevol': 'Log trade volume'}
DEC = {'trade_share': 2, 'trade_intensity': 3, 'ln_tradevol': 3}
GLAB = {'adv_LOW': 'Low aviation advantage', 'adv_HIGH': 'High aviation advantage',
        'coastal': 'Coastal', 'landlocked': 'Landlocked',
        'inc_LOW': 'Low income (baseline)', 'inc_HIGH': 'High income (baseline)',
        'size_SMALL': 'Small (baseline pop.)', 'size_LARGE': 'Large (baseline pop.)',
        'cont_AF': 'Africa', 'cont_AS': 'Asia', 'cont_EU': 'Europe', 'cont_LA': 'Latin America',
        'cont_ME': 'Middle East', 'cont_NA': 'North America', 'cont_SW': 'Pacific/SW'}
GORDER = ['coastal', 'landlocked', 'adv_LOW', 'adv_HIGH', 'inc_LOW', 'inc_HIGH',
          'size_SMALL', 'size_LARGE', 'cont_EU', 'cont_AS', 'cont_AF', 'cont_LA',
          'cont_ME', 'cont_NA', 'cont_SW']


def star(p):
    return '***' if p < .01 else '**' if p < .05 else '*' if p < .1 else ''


def fnum(x, d):
    return f'{x:,.{d}f}'


tab1, tab2 = {}, {}
for ln in open('gaci_tables_export.log', encoding='utf-8', errors='replace'):
    ln = ln.strip()
    if ln.startswith('TAB1|'):
        _, yv, tr, b, sr, pr, sc, pc, F, ob, op, N = ln.split('|')
        tab1[(yv, tr)] = dict(b=float(b), sr=float(sr), pr=float(pr), sc=float(sc),
                              pc=float(pc), F=float(F), ob=float(ob), op=float(op), N=int(float(N)))
    elif ln.startswith('TAB2|'):
        _, g, yv, b, se, p, F, N = ln.split('|')
        tab2[(g, yv)] = dict(b=float(b), se=float(se), p=float(p), F=float(F), N=int(float(N)))

doc = Document()
doc.styles['Normal'].font.name = 'Times New Roman'
doc.styles['Normal'].font.size = Pt(10)


def H(t):
    h = doc.add_paragraph(); r = h.add_run(t); r.bold = True; r.font.size = Pt(11)
    return h


def note(t):
    p = doc.add_paragraph(); r = p.add_run(t); r.font.size = Pt(8); r.italic = True


doc.add_heading('Air Connectivity and Trade: IV Estimates', level=1)
note('Instrument: Feyrer (2019)-style air/sea time-varying market access, '
     'feyrer_int = a_t x ln(air market access, 1996), with sea-distance market access '
     '(CERDI) as control. All models country and year fixed effects. '
     'Robust SE in (parentheses), country-clustered SE in [brackets].')

# ---------------- TABLE 1 ----------------
H('Table 1. Air connectivity and trade openness: OLS vs. IV (2SLS)')
t = doc.add_table(rows=0, cols=4); t.style = 'Table Grid'
hdr = t.add_row().cells
hdr[0].text = ''
for j, o in enumerate(OUT):
    hdr[j + 1].text = OUTLAB[o]
for cells in [hdr]:
    for c in cells:
        for para in c.paragraphs:
            for r in para.runs: r.bold = True


def panel(treat, title):
    row = t.add_row().cells
    row[0].merge(row[3]); row[0].text = title
    for r in row[0].paragraphs[0].runs: r.bold = True
    # IV beta (robust SE) [clustered SE]
    specs = [('Air connectivity (IV, 2SLS)', 'iv'),
             ('   Robust SE', 'sr'), ('   Clustered SE', 'sc'),
             ('OLS benchmark', 'ols'),
             ('First-stage KP F', 'F'), ('Observations', 'N')]
    for label, kind in specs:
        r = t.add_row().cells
        r[0].text = label
        for j, o in enumerate(OUT):
            d = tab1.get((o, treat)); dec = DEC[o]
            if d is None:
                r[j + 1].text = ''; continue
            if kind == 'iv':
                r[j + 1].text = fnum(d['b'], dec) + star(d['pr'])
            elif kind == 'sr':
                r[j + 1].text = f"({fnum(d['sr'], dec)})"
            elif kind == 'sc':
                r[j + 1].text = f"[{fnum(d['sc'], dec)}]" + ('+' if d['pc'] < .05 else '')
            elif kind == 'ols':
                r[j + 1].text = fnum(d['ob'], dec) + star(d['op'])
            elif kind == 'F':
                r[j + 1].text = fnum(d['F'], 1)
            elif kind == 'N':
                r[j + 1].text = f"{d['N']:,}"


panel('lng', 'Panel A. Treatment = total connectivity (log GACI sum)')
panel('ln_gaci_cwm', 'Panel B. Treatment = hub quality (log capacity-weighted-mean GACI)')
note('* p<0.10, ** p<0.05, *** p<0.01 (robust). "+" marks coefficients that remain '
     'significant at 5% under country-clustered SE. Hub-quality units are not comparable '
     'in scale to the connectivity-sum treatment. First-stage Kleibergen-Paap F is the '
     'weak-identification statistic.')

# ---------------- TABLE 2 ----------------
doc.add_paragraph('')
H('Table 2. Heterogeneity of the IV effect across country groups (split-sample)')
t2 = doc.add_table(rows=0, cols=5); t2.style = 'Table Grid'
h2 = t2.add_row().cells
h2[0].text = 'Subgroup'
for j, o in enumerate(OUT):
    h2[j + 1].text = OUTLAB[o]
h2[4].text = 'First-stage\nKP F'
for c in h2:
    for para in c.paragraphs:
        for r in para.runs: r.bold = True
for g in GORDER:
    if not any((g, o) in tab2 for o in OUT):
        continue
    r = t2.add_row().cells
    r[0].text = GLAB.get(g, g)
    Fval = None
    for j, o in enumerate(OUT):
        d = tab2.get((g, o))
        if d is None:
            r[j + 1].text = ''; continue
        r[j + 1].text = f"{fnum(d['b'], DEC[o])}{star(d['p'])}\n({fnum(d['se'], DEC[o])})"
        if o == 'trade_share':
            Fval = d['F']
    r[4].text = (fnum(Fval, 1) if Fval is not None else '') + (' (weak)' if (Fval is not None and Fval < 10) else '')
note('Each cell: 2SLS coefficient on log connectivity with robust SE in parentheses. '
     '* p<0.10, ** p<0.05, *** p<0.01. Moderators are time-invariant (baseline values), '
     'so the within-group estimator retains country and year fixed effects. Groups with '
     'first-stage F < 10 are weakly identified and reported for completeness only.')

doc.save('GACI_results_tables.docx')
print('wrote GACI_results_tables.docx  (TAB1 cells=%d, TAB2 cells=%d)' % (len(tab1), len(tab2)))
