# -*- coding: utf-8 -*-
"""41_table9_compact.py : one compact mechanism table (six links), numbers pulled from the result CSVs.
   Sheet Table9_compact in Tables_for_paper_20260928.xlsx."""
import os, pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
RD = dict(keep_default_na=False, na_values=["", "."])
mech = pd.read_csv("_mech_results.csv", **RD); comp = pd.read_csv("_components_results.csv", **RD); gic = pd.read_csv("_gic_dose_results.csv", **RD)
alt = pd.read_csv("_alt_measures_results.csv"); med = pd.read_csv("_median_splits_results.csv", **RD); stg = pd.read_csv("_gic_bysample_results.csv", **RD)
sp = pd.read_csv("_spill_results.csv", **RD); macro = pd.read_csv("_macro_channels_results.csv", **RD)
def st(p): return "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
def c(b, se, p, d=2): return f"{b:.{d}f}{st(p)} ({se:.{d}f})"
def g(df, **kw):
    x = df
    for k, v in kw.items(): x = x[x[k] == v]
    return x.iloc[0]
# ---- pull numbers ----
hub_sum = g(mech, block="A", item="hub_ctrl_sum", outcome="ln_gini_mkt", term="ln_gaci_max"); sum_ctrl = g(mech, block="A", item="hub_ctrl_sum", outcome="ln_gini_mkt", term="ln_gaci_sum")
hhi = g(mech, block="A", item="ols_hhi", outcome="ln_gini_mkt", term="ln_hhi")
eig = g(comp, block="C", item="topo_iv_cap_ctrl_eig", outcome="ln_gini_mkt", term="ln_hub_eig"); capn = g(comp, block="C", item="topo_iv_cap_ctrl_eig", outcome="ln_gini_mkt", term="ln_hub_cap")
mean_inc = g(gic, block="gic", outcome="ln_apt_all", spec="IV"); d10 = g(gic, block="gic", outcome="ln_apt_d10", spec="IV"); d7 = g(gic, block="gic", outcome="ln_apt_d7", spec="IV")
b50 = g(gic, block="gic", outcome="ln_apt_b50", spec="IV"); b50s = g(gic, block="gic_share", outcome="ln_spt_b50", spec="IV")
p5010 = alt[alt.outcome == "ln_p50_p10"].iloc[0]; p9050 = alt[alt.outcome == "ln_p90_p50"].iloc[0]
tour_b = g(mech, block="C", item="ln_tour_arr", outcome="ln_gini_mkt", term="ln_tour_arr"); ind_b = g(mech, block="C", item="emp_ind", outcome="ln_gini_mkt", term="emp_ind")
diff_b = g(macro, block="B", item="ln_tr_diff", outcome="ln_gini_mkt", term="ln_tr_diff"); gdp_b = g(mech, block="C", item="ln_gdppc", outcome="ln_gini_mkt", term="ln_gdppc"); cprime = g(mech, block="C", item="ln_gdppc", outcome="ln_gini_mkt", term="ln_gaci_max")
urb2 = g(med, var="urban", group="above", outcome="ln_gini_mkt"); urb1 = g(med, var="urban", group="below", outcome="ln_gini_mkt")
tour1 = g(mech, block="E", item="tour_rcpt_exp_1", outcome="ln_gini_mkt"); thin = g(stg, sample="con_low", outcome="ln_gini_mkt"); mid = g(stg, sample="con_mid", outcome="ln_gini_mkt")
contig = g(sp, panel="A", item="single_contig", var="nbr_g_contig", outcome="ln_gini_mkt")
wedge = g(mech, block="D", item="redistribution", outcome="wedge_pts"); tax = g(mech, block="D", item="redistribution", outcome="tax_gdp")
ROWS = [
 ("1", "The effect is carried by the hub airport's integration with the global core, not by the breadth of the network",
  "ln GACI_max instrumented, ln GACI_sum as control; eigenvector centrality instrumented, seat capacity as control",
  f"hub {c(hub_sum.b, hub_sum.se, hub_sum.p)} vs network {c(sum_ctrl.b, sum_ctrl.se, sum_ctrl.p)}; core integration {c(eig.b, eig.se, eig.p, 3)} vs capacity {c(capn.b, capn.se, capn.p, 3)}; airport HHI (OLS) {c(hhi.b, hhi.se, hhi.p, 3)}",
  "Table 9A; A10"),
 ("2", "The income it generates accrues to the upper half of the distribution",
  "Group-specific 2SLS: ln average pretax income of WID groups",
  f"mean income {c(mean_inc.b, mean_inc.se, mean_inc.p)}; p60-70 {c(d7.b, d7.se, d7.p)}; p90-100 {c(d10.b, d10.se, d10.p)}",
  "Table 4"),
 ("3", "The bottom half is bypassed: its income does not respond and it falls behind the median",
  "Same regressions for the bottom 50%; decile-mean ratios",
  f"bottom-50 income {c(b50.b, b50.se, b50.p)}; bottom-50 share {c(b50s.b, b50s.se, b50s.p)}; p50/p10 {c(p5010.iv_b, p5010.iv_se, p5010.iv_p)}; p90/p50 {c(p9050.iv_b, p9050.iv_se, p9050.iv_p)}",
  "Table 4; A14"),
 ("4", "The structural changes the hub induces all work toward equality and do not transmit the effect",
  "IV decomposition: mediator as control, hub instrumented (b = mediator coefficient on ln Gini)",
  f"b: tourist arrivals {c(tour_b.b, tour_b.se, tour_b.p, 3)}; industrial employment {c(ind_b.b, ind_b.se, ind_b.p, 4)}; differentiated exports {c(diff_b.b, diff_b.se, diff_b.p, 3)}; GDP per capita {c(gdp_b.b, gdp_b.se, gdp_b.p, 3)}; direct effect with GDP held fixed {c(cprime.b, cprime.se, cprime.p)}",
  "Table 9C; A12"),
 ("5", "The effect exists only where an urban, business-oriented hub economy exists, and it stays within the country",
  "Split samples on baseline characteristics; neighbours' connectivity instrumented",
  f"urbanisation above / below median {c(urb2.b, urb2.se, urb2.p)} / {c(urb1.b, urb1.se, urb1.p)}; non-tourism economies {c(tour1.b, tour1.se, tour1.p)}; thin networks {c(thin.b, thin.se, thin.p)} vs take-off hubs {c(mid.b, mid.se, mid.p)}; contiguous neighbours' hubs {c(contig.b, contig.se, contig.p)}",
  "Table 5b; 9E; A7; A9"),
 ("6", "Redistribution does not offset it",
  "Market-minus-disposable Gini wedge and tax revenue as outcomes",
  f"wedge {c(wedge.b, wedge.se, wedge.p, 1)} points; tax revenue / GDP {c(tax.b, tax.se, tax.p, 1)}",
  "Table 9D")]
wb = load_workbook("Tables_for_paper_20260928.xlsx")
if "Table9_compact" in wb.sheetnames: wb.remove(wb["Table9_compact"])
ws = wb.create_sheet("Table9_compact", wb.sheetnames.index("Table9_mechanisms"))
FN = "Times New Roman"; F_T = Font(name=FN, bold=True, size=11); F_H = Font(name=FN, bold=True, size=10); F_B = Font(name=FN, size=10); F_N = Font(name=FN, size=9)
TOP = Side(style="medium"); THIN = Side(style="thin")
ws["A1"] = "Table 9. Mechanism: a premium on access that bypasses the bottom half"; ws["A1"].font = F_T
r = 3
for j, h in enumerate(["", "Link", "Test", "Estimate (s.e.)", "Source"], 1):
    x = ws.cell(row=r, column=j, value=h); x.font = F_H; x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); x.border = Border(top=TOP, bottom=THIN)
for n, claim, test, est, src in ROWS:
    r += 1
    for j, v in enumerate([n, claim, test, est, src], 1):
        x = ws.cell(row=r, column=j, value=v); x.font = F_B; x.alignment = Alignment(horizontal="center" if j in (1, 5) else "left", vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 62
for j in range(1, 6): ws.cell(row=r, column=j).border = Border(bottom=TOP)
r += 2
for n in ["Notes: All estimates are 2SLS with the Feyrer instrument (except the airport-HHI row, OLS), with ln population, country and year fixed effects; standard errors clustered by country in parentheses. *** p<0.01, ** p<0.05, * p<0.1. Elasticities are with respect to ln GACI_max; the wedge is in Gini points. Full results in the source tables.",
          "Reading: the hub connects one city to the global core; the income it generates goes to households that use that connection, in the upper half of the distribution, while the bottom half does not respond. The structural changes the hub sets off are equalising and are outweighed. The effect requires an urban, business-oriented hub economy, does not cross borders, and is not offset by redistribution."]:
    ws.cell(row=r, column=1, value=n).font = F_N; r += 1
for j, w in enumerate([5, 44, 40, 62, 16], 1): ws.column_dimensions[get_column_letter(j)].width = w
wb.save("Tables_for_paper_20260928.xlsx"); print("saved Table9_compact")
for n, claim, test, est, src in ROWS: print(f"| {n} | {claim} | {est} | {src} |")
