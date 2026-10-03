# -*- coding: utf-8 -*-
"""39_add_macro_xlsx.py : Table A12 (macro channels: spatial concentration, export composition, factor shares, finance)."""
import os, pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
d = pd.read_csv("_macro_channels_results.csv", keep_default_na=False, na_values=["", "."])
FN = "Times New Roman"; F_T = Font(name=FN, bold=True, size=11); F_H = Font(name=FN, bold=True, size=10); F_B = Font(name=FN, size=10); F_I = Font(name=FN, italic=True, size=10); F_N = Font(name=FN, size=9)
TOP = Side(style="medium"); THIN = Side(style="thin")
def st(p): return "" if pd.isna(p) else "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
wb = load_workbook("Tables_for_paper_20260928.xlsx")
if "A12_macro_channels" in wb.sheetnames: wb.remove(wb["A12_macro_channels"])
ws = wb.create_sheet("A12_macro_channels", wb.sheetnames.index("A11_wage_premia") + 1)
ws["A1"] = "Table A12. Macro channels: spatial concentration, export composition, factor shares and finance"; ws["A1"].font = F_T
r = 3
def put(c, v, font=F_B, align="center", fmt=None):
    x = ws.cell(row=r, column=c, value=v); x.font = font; x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(c == 1))
    if fmt: x.number_format = fmt
def put_at(rr, c, v):
    x = ws.cell(row=rr, column=c, value=v); x.font = F_B; x.alignment = Alignment(horizontal="center")
GROUPS = [("Spatial concentration", [("ln_share_topreg", "ln share of top region in national GRP (DOSE)"), ("ln_theil_reg", "ln Theil index of regional GRP per capita (DOSE)"), ("ln_largest_city", "ln largest-city share of urban population"), ("primacy", "Urban primacy"), ("urban", "Urban population, %")]),
          ("Export composition (BACI, Rauch and value-to-weight)", [("sh_diff", "Differentiated-goods share of trade"), ("sh_homog", "Homogeneous-goods share of trade"), ("sh_hivw", "High value-to-weight share of trade"), ("sh_capital", "Capital-goods share of trade"), ("sh_interm", "Intermediate-goods share of trade"), ("ln_tr_diff", "ln differentiated-goods trade"), ("ln_tr_hivw", "ln high value-to-weight trade"), ("ln_nflow", "ln number of export partners"), ("ln_nprod", "ln number of export products"), ("ln_hitech", "ln high-tech exports, % of manufactured exports"), ("ln_ict_serv", "ln ICT service exports, % of service exports")]),
          ("Factor shares and finance", [("labsh", "Labour share of income (PWT)"), ("ln_priv_credit", "ln private credit, % of GDP"), ("gfcf_gdp", "Gross fixed capital formation, % of GDP"), ("resource_rents", "Natural resource rents, % of GDP"), ("remit_gdp", "Remittances, % of GDP"), ("ind_va", "Industry value added, % of GDP")])]
put(1, "Channel variable", F_H, "left")
for j, l in enumerate(["a-path: 2SLS", "a-path: OLS", "KP F", "N", "c' (hub) with channel as control", "b (channel) on ln Market Gini", "b (channel) on bottom-50 share"], 2): put(j, l, F_H)
for c in range(1, 9): ws.cell(row=r, column=c).border = Border(top=TOP, bottom=THIN)
r += 1
for glab, items in GROUPS:
    put(1, glab, F_I, "left"); r += 1
    for mv, lab in items:
        iv = d[(d.block == "A") & (d["item"] == "apath") & (d.outcome == mv)]; ol = d[(d.block == "A") & (d["item"] == "apath_ols") & (d.outcome == mv)]
        if not len(iv): continue
        put(1, lab, F_B, "left")
        put(2, f"{iv.b.iloc[0]:.3f}{st(iv.p.iloc[0])}"); put_at(r + 1, 2, f"({iv.se.iloc[0]:.3f})")
        put(3, f"{ol.b.iloc[0]:.3f}{st(ol.p.iloc[0])}"); put_at(r + 1, 3, f"({ol.se.iloc[0]:.3f})")
        put(4, float(iv.kpf.iloc[0]), fmt="0.0"); put(5, int(iv.N.iloc[0]), fmt="#,##0")
        cp = d[(d.block == "B") & (d["item"] == mv) & (d.outcome == "ln_gini_mkt") & (d.term == "ln_gaci_max")]; bg = d[(d.block == "B") & (d["item"] == mv) & (d.outcome == "ln_gini_mkt") & (d.term == mv)]; bb = d[(d.block == "B") & (d["item"] == mv) & (d.outcome == "ln_spt_b50") & (d.term == mv)]
        if len(cp): put(6, f"{cp.b.iloc[0]:.3f}{st(cp.p.iloc[0])}"); put_at(r + 1, 6, f"({cp.se.iloc[0]:.3f})")
        if len(bg): put(7, f"{bg.b.iloc[0]:.4f}{st(bg.p.iloc[0])}"); put_at(r + 1, 7, f"({bg.se.iloc[0]:.4f})")
        if len(bb): put(8, f"{bb.b.iloc[0]:.4f}{st(bb.p.iloc[0])}"); put_at(r + 1, 8, f"({bb.se.iloc[0]:.4f})")
        r += 2
r += 1; put(1, "Heterogeneity by baseline labour share (median split)", F_I, "left"); r += 1
for g, gl in [("1", "Labour share below median"), ("2", "Labour share above median")]:
    put(1, gl, F_B, "left")
    for j, o in enumerate(["ln_gini_mkt", "ln_spt_b50", "ln_apt_top_bot"], 2):
        x = d[(d.block == "C") & (d["item"] == f"labsh_{g}") & (d.outcome == o)]
        if len(x): put(j, f"{['Gini','b50 share','top-bot'][j-2]}: {x.b.iloc[0]:.3f}{st(x.p.iloc[0])} ({x.se.iloc[0]:.3f})")
    x = d[(d.block == "C") & (d["item"] == f"labsh_{g}") & (d.outcome == "ln_gini_mkt")]
    if len(x): put(5, float(x.kpf.iloc[0]), fmt="0.0"); put(6, int(x.N.iloc[0]), fmt="#,##0")
    r += 1
for c in range(1, 9): ws.cell(row=r - 1, column=c).border = Border(bottom=TOP)
r += 1
for n in ["Notes: a-path = regression of the channel variable on ln GACI_max (2SLS with the Feyrer instrument and OLS), ln population, country and year FE, SE clustered by country. c' and b from the IV decomposition (channel as control, hub instrumented). *** p<0.01, ** p<0.05, * p<0.1.",
          "Reading: hub connectivity shifts trade toward differentiated and high value-to-weight goods and more partners, but these shifts lower inequality (negative b), so they do not transmit the effect. There is no evidence that the hub raises spatial concentration of population; regional GRP concentration (DOSE) is unidentified (KP F 2). The labour share does not respond; the effect is concentrated in economies with a low baseline labour share.",
          "Sources: _macro_channels_results.csv (38_macro_channels.do); BACI mechanism panel of the companion trade paper; DOSE v2.9; WDI API; PWT 10.0."]:
    put(1, n, F_N, "left"); r += 1
ws.column_dimensions["A"].width = 52
for j in range(2, 9): ws.column_dimensions[get_column_letter(j)].width = 20
wb.save("Tables_for_paper_20260928.xlsx"); print("saved A12")
