# -*- coding: utf-8 -*-
"""22_extend_paper_xlsx.py : append heterogeneity (T4), spillover (T5) and aggregate incidence (T6) sheets to
   GIC_paper_tables_20260928.xlsx. Everything 2SLS Feyrer, ln pop, country + year FE, country-clustered SE.
   Sources: _hetero_co2form_results.csv, _spill_results.csv (20_hetero_spill.do), _aggregate_curve_*.csv (21)."""
import os, numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
RD = dict(keep_default_na=False, na_values=["", "."])
het = pd.read_csv("_hetero_co2form_results.csv", **RD)
sp = pd.read_csv("_spill_results.csv", **RD)
agc = pd.read_csv("_aggregate_curve_bycontinent.csv", index_col=0)
FA = "Arial"
H = Font(name=FA, bold=True, size=10); B = Font(name=FA, size=10); T = Font(name=FA, bold=True, size=12); N = Font(name=FA, italic=True, size=9, color="555555")
FILL = PatternFill("solid", fgColor="EDEDED"); THIN = Side(style="thin", color="999999")
def st(p):
    if pd.isna(p): return ""
    return "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
def cell(b, se, p): return f"{b:.3f}{st(p)} ({se:.3f})"
def put(ws, r, c, v, font=B, fmt=None, align=None):
    if isinstance(v, np.integer): v = int(v)
    if isinstance(v, np.floating): v = float(v)
    if isinstance(v, float) and np.isnan(v): v = None
    x = ws.cell(row=r, column=c, value=v); x.font = font
    if fmt: x.number_format = fmt
    if align: x.alignment = Alignment(horizontal=align)
    return x
def header(ws, r, cols):
    for j, c in enumerate(cols, 1):
        x = put(ws, r, j, c, H); x.fill = FILL; x.border = Border(top=THIN, bottom=THIN); x.alignment = Alignment(horizontal="center", wrap_text=True)
def widths(ws, w):
    for j, v in enumerate(w, 1): ws.column_dimensions[get_column_letter(j)].width = v

wb = load_workbook("GIC_paper_tables_20260928.xlsx")
for n in ["T4_heterogeneity", "T5_spillover", "T6_aggregate"]:
    if n in wb.sheetnames: wb.remove(wb[n])

OUTS = [("ln_gini_mkt", "ln market Gini"), ("ln_gini_disp", "ln disposable Gini"), ("ln_apt_top_bot", "Top half minus bottom half (income)"), ("ln_apt_d10", "Income p90-100"), ("ln_spt_b50", "Share bottom 50%")]

# ---------------- T4 heterogeneity ----------------
ws = wb.create_sheet("T4_heterogeneity")
put(ws, 1, 1, "Table 4. Heterogeneity of the connectivity effect (split-sample 2SLS, Feyrer IV)", T)
put(ws, 2, 1, "Each cell: 2SLS coefficient on ln GACI_max (se), country-clustered; ln pop, country + year FE. Terciles formed on the earliest observed value per country (income per capita, GACI_max, market Gini).", N)
put(ws, 3, 1, "'Interaction' rows: pooled regression with top tercile x ln GACI (instrumented by top tercile x Z); the cell reports the interaction. Temporal split follows the CO2 paper (1996-2007 expansion era; 2010-2023 mature era excl. 2020-21).", N)
put(ws, 4, 1, "Source: _hetero_co2form_results.csv (20_hetero_spill.do). Cells with N < 60 are skipped.", N)
r = 6; header(ws, r, ["Sample"] + [l for _, l in OUTS] + ["KP F (market Gini)", "N", "Countries"])
PAN = [("A_income", "Panel A. Baseline income per capita tercile", [("inc_low", "Low"), ("inc_mid", "Middle"), ("inc_high", "High"), ("interact_top_x_gaci", "Interaction: top tercile x ln GACI")]),
       ("B_conn", "Panel B. Baseline connectivity tercile", [("con_low", "Low"), ("con_mid", "Middle"), ("con_high", "High"), ("interact_top_x_gaci", "Interaction: top tercile x ln GACI")]),
       ("C_gini", "Panel C. Baseline market-Gini tercile", [("gini_low", "Low"), ("gini_mid", "Middle"), ("gini_high", "High"), ("interact_top_x_gaci", "Interaction: top tercile x ln GACI")]),
       ("D_time", "Panel D. Period", [("full", "Full sample 1996-2023"), ("1996_2007", "1996-2007"), ("2010_2023_ex2020_21", "2010-2023 excl. 2020-21"), ("2010_2023", "2010-2023"), ("excl_2020_2023", "1996-2019")]),
       ("E_region", "Panel E. Macro region", [("Europe", "Europe"), ("Asia", "Asia-Pacific"), ("Africa", "Africa"), ("LatAm", "Latin America"), ("MiddleEast", "Middle East"), ("excl_Europe", "Excluding Europe"), ("excl_LatAm", "Excluding Latin America")])]
for pan, plab, items in PAN:
    r += 1; put(ws, r, 1, plab, H)
    for key, lab in items:
        r += 1; put(ws, r, 1, lab)
        for j, (o, _) in enumerate(OUTS, 2):
            x = het[(het.panel == pan) & (het.group == key) & (het.outcome == o)]
            put(ws, r, j, cell(x.b.iloc[0], x.se.iloc[0], x.p.iloc[0]) if len(x) else "--", align="center")
        x = het[(het.panel == pan) & (het.group == key) & (het.outcome == "ln_gini_mkt")]
        if len(x):
            put(ws, r, 2 + len(OUTS), float(x.kpf.iloc[0]), fmt="0.0", align="center"); put(ws, r, 3 + len(OUTS), int(x.N.iloc[0]), fmt="#,##0", align="center")
            if not pd.isna(x.Nc.iloc[0]): put(ws, r, 4 + len(OUTS), int(x.Nc.iloc[0]), fmt="#,##0", align="center")
widths(ws, [40] + [22] * len(OUTS) + [12, 9, 10])

# ---------------- T5 spillover ----------------
ws = wb.create_sheet("T5_spillover")
put(ws, 1, 1, "Table 5. Spillovers from neighbours' connectivity (design of the CO2 paper)", T)
put(ws, 2, 1, "Neighbour exposure = leave-out weighted mean of other countries' ln GACI (capacity-weighted mean), instrumented by the identically weighted mean of their Feyrer shifters. 'Hybrid' rows control the own Feyrer shifter in reduced form; 'joint' rows instrument own ln GACI_max by the own shifter and neighbour exposure by neighbours' shifters.", N)
put(ws, 3, 1, "Exposures from ../GACI_CO2/spillover_bands.csv: inverse distance (inv), land contiguity (contig), five nearest (knn5), distance bands b1 <500 km, b2 500-1000, b3 1000-2000, b4 2000-5000, b5 >5000; kernel exp(-d/lambda). Presence flags (has_*) included as controls. Country-clustered SE.", N)
put(ws, 4, 1, "Source: _spill_results.csv (20_hetero_spill.do).", N)
SO = [("ln_gini_mkt", "ln market Gini"), ("ln_gini_disp", "ln disposable Gini"), ("ln_apt_top_bot", "Top half minus bottom half")]
r = 6; header(ws, r, ["Specification", "Coefficient on"] + [l for _, l in SO] + ["KP F (market Gini)", "N"])
ROWS = [("Panel A. First-order neighbour effect, own connectivity controlled", [
            ("A", "own_exog_control", "nbr_g_inv", "Own ln GACI as exogenous control; neighbour (inverse distance) instrumented", "neighbour"),
            ("A", "own_exog_control", "ln_gaci_max", "   same regression", "own ln GACI (exogenous)"),
            ("A", "hybrid_rf_control", "nbr_g_inv", "Own shifter in reduced form; neighbour (inverse distance) instrumented", "neighbour"),
            ("A", "hybrid_rf_control", "feyrer_int", "   same regression", "own shifter (reduced form)"),
            ("A", "joint_inv", "ln_gaci_max", "Joint 2SLS, inverse distance", "own ln GACI"),
            ("A", "joint_inv", "nbr_g_inv", "   same regression", "neighbour"),
            ("A", "joint_contig", "ln_gaci_max", "Joint 2SLS, land contiguity", "own ln GACI"),
            ("A", "joint_contig", "nbr_g_contig", "   same regression", "neighbour"),
            ("A", "joint_knn5", "ln_gaci_max", "Joint 2SLS, five nearest countries", "own ln GACI"),
            ("A", "joint_knn5", "nbr_g_knn5", "   same regression", "neighbour"),
            ("A", "joint_b1", "ln_gaci_max", "Joint 2SLS, within 500 km", "own ln GACI"),
            ("A", "joint_b1", "nbr_g_b1", "   same regression", "neighbour"),
            ("A", "joint_k500", "ln_gaci_max", "Joint 2SLS, kernel 500 km", "own ln GACI"),
            ("A", "joint_k500", "nbr_g_k500", "   same regression", "neighbour"),
            ("A", "joint_k1000", "ln_gaci_max", "Joint 2SLS, kernel 1000 km", "own ln GACI"),
            ("A", "joint_k1000", "nbr_g_k1000", "   same regression", "neighbour"),
            ("A", "single_contig", "nbr_g_contig", "Neighbour only (hybrid), contiguity", "neighbour"),
            ("A", "single_knn5", "nbr_g_knn5", "Neighbour only (hybrid), five nearest", "neighbour"),
            ("A", "single_k500", "nbr_g_k500", "Neighbour only (hybrid), kernel 500 km", "neighbour"),
            ("A", "single_k1000", "nbr_g_k1000", "Neighbour only (hybrid), kernel 1000 km", "neighbour")]),
        ("Panel B. Distance decay (own shifter in reduced form)", [
            ("B", "bands_joint", "nbr_g_b1", "Bands entered jointly: <500 km", "neighbour"),
            ("B", "bands_joint", "nbr_g_b2", "   500-1,000 km", "neighbour"),
            ("B", "bands_joint", "nbr_g_b3", "   1,000-2,000 km", "neighbour"),
            ("B", "bands_joint", "nbr_g_b4", "   2,000-5,000 km", "neighbour"),
            ("B", "bands_joint", "nbr_g_b5", "   beyond 5,000 km", "neighbour"),
            ("B", "band_single", "nbr_g_contig", "Band alone: contiguous", "neighbour"),
            ("B", "band_single", "nbr_g_b1", "   <500 km", "neighbour"),
            ("B", "band_single", "nbr_g_b2", "   500-1,000 km", "neighbour"),
            ("B", "band_single", "nbr_g_b3", "   1,000-2,000 km", "neighbour"),
            ("B", "band_single", "nbr_g_b4", "   2,000-5,000 km", "neighbour"),
            ("B", "band_single", "nbr_g_b5", "   beyond 5,000 km", "neighbour"),
            ("B", "kernel", "nbr_g_k250", "Kernel exp(-d/lambda), lambda = 250 km", "neighbour"),
            ("B", "kernel", "nbr_g_k500", "   500 km", "neighbour"),
            ("B", "kernel", "nbr_g_k1000", "   1,000 km", "neighbour"),
            ("B", "kernel", "nbr_g_k2000", "   2,000 km", "neighbour"),
            ("B", "kernel", "nbr_g_k5000", "   5,000 km", "neighbour")]),
        ("Panel C. Regional exposure and region-by-year fixed effects", [
            ("C", "subregion_inout", "nbr_g_inreg", "Within UN sub-region (leave-out mean)", "neighbour"),
            ("C", "subregion_inout", "nbr_g_outreg", "   outside sub-region, same regression", "neighbour"),
            ("C", "bloc_inout", "nbr_g_inbloc", "Within aviation bloc (leave-out mean)", "neighbour"),
            ("C", "bloc_inout", "nbr_g_outbloc", "   outside bloc, same regression", "neighbour"),
            ("C", "contig_joint_subregXyear", "ln_gaci_max", "Contiguous, joint 2SLS, sub-region x year FE", "own ln GACI"),
            ("C", "contig_joint_subregXyear", "nbr_g_contig", "   same regression", "neighbour"),
            ("C", "contig_hybrid_subregXyear", "nbr_g_contig", "Contiguous, own shifter in reduced form, sub-region x year FE", "neighbour")])]
for plab, items in ROWS:
    r += 1; put(ws, r, 1, plab, H)
    for pan, item, var, lab, on in items:
        r += 1; put(ws, r, 1, lab); put(ws, r, 2, on)
        for j, (o, _) in enumerate(SO, 3):
            x = sp[(sp.panel == pan) & (sp["item"] == item) & (sp["var"] == var) & (sp.outcome == o)]
            put(ws, r, j, cell(x.b.iloc[0], x.se.iloc[0], x.p.iloc[0]) if len(x) else "--", align="center")
        x = sp[(sp.panel == pan) & (sp["item"] == item) & (sp["var"] == var) & (sp.outcome == "ln_gini_mkt")]
        if len(x):
            put(ws, r, 3 + len(SO), float(x.kpf.iloc[0]), fmt="0.0", align="center"); put(ws, r, 4 + len(SO), int(x.N.iloc[0]), fmt="#,##0", align="center")
widths(ws, [62, 26] + [22] * len(SO) + [12, 9])

# ---------------- T6 aggregate ----------------
ws = wb.create_sheet("T6_aggregate")
put(ws, 1, 1, "Table 6. Aggregate incidence of observed hub-connectivity growth, 1996 to latest year (back-of-envelope)", T)
put(ws, 2, 1, "Implied = country's change in ln GACI_max over its window x 2SLS elasticity (Table 2), in percent of the group's average income; share changes in points of national income using the share elasticity and the initial share. Actual = observed change over the same window.", N)
put(ws, 3, 1, "Countries with >= 15 years of data. Continent rows are unweighted country means (as in the Gini contribution table T7 of the main workbook); WORLD is weighted by latest population. Implied changes exceed actual ones where other forces moved the distribution the other way (e.g. Latin America, Africa).", N)
put(ws, 4, 1, "Source: _aggregate_curve_bycontinent.csv, _aggregate_curve_bycountry.csv (21_aggregate_curve.py).", N)
r = 6
cols = ["Region", "Countries", "Mean dln GACI", "Bottom 50% income: implied %", "actual %", "p90-100 income: implied %", "actual %", "Mean income: implied %", "actual %",
        "Bottom 50% share: implied pts", "actual pts", "Top 10% share: implied pts", "actual pts", "Top 1% share: implied pts", "actual pts", "Market Gini: implied pts", "actual pts"]
header(ws, r, cols)
keys = ["n", "mean_dln_gaci", "implied_pct_inc_b50", "actual_pct_inc_b50", "implied_pct_inc_d10", "actual_pct_inc_d10", "implied_pct_inc_all", "actual_pct_inc_all",
        "implied_dshare_pts_b50", "actual_dshare_pts_b50", "implied_dshare_pts_d10", "actual_dshare_pts_d10", "implied_dshare_pts_t1", "actual_dshare_pts_t1", "implied_dgini_pts", "actual_dgini_pts"]
NAME = {"AF": "Africa", "AS": "Asia-Pacific", "EU": "Europe", "LA": "Latin America", "ME": "Middle East", "NA": "North America", "SW": "Oceania"}
for idx, row in agc.iterrows():
    r += 1; put(ws, r, 1, NAME.get(idx, idx))
    for j, k in enumerate(keys, 2):
        put(ws, r, j, row[k], fmt="#,##0" if k == "n" else "0.00", align="center")
widths(ws, [22, 10] + [14] * 15)
wb.save("GIC_paper_tables_20260928.xlsx"); print("saved", wb.sheetnames)
