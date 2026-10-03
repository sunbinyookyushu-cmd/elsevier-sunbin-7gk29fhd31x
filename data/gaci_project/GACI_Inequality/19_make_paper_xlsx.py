# -*- coding: utf-8 -*-
"""19_make_paper_xlsx.py : paper-ready workbook. Headline table (OLS + 2SLS Feyrer) and the incidence curves
   for GACI_max and GACI_cwm, 2SLS Feyrer with country-clustered SE only.
   Sources: _main_results_cl.csv (01_main_cl.do), _gic_dose_results.csv (12_gic_dose.do), _gic_cwm_results.csv (12c_gic_cwm.do).
   Output: GIC_paper_tables_20260928.xlsx"""
import os, pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
RD = dict(keep_default_na=False, na_values=["", "."])
main = pd.read_csv("_main_results_cl.csv", **RD)
gmax = pd.read_csv("_gic_dose_results.csv", **RD)
gcwm = pd.read_csv("_gic_cwm_results.csv", **RD)

G = ["d1", "d2", "d3", "d4", "d5", "d6", "d7", "d8", "d9", "d10", "t1", "t01", "b50", "m40", "all"]
LAB = {"d1": "p0-10", "d2": "p10-20", "d3": "p20-30", "d4": "p30-40", "d5": "p40-50", "d6": "p50-60", "d7": "p60-70",
       "d8": "p70-80", "d9": "p80-90", "d10": "p90-100", "t1": "Top 1%", "t01": "Top 0.1%", "b50": "Bottom 50%",
       "m40": "Middle 40%", "all": "All (mean income)"}
def st(p):
    if pd.isna(p): return ""
    return "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
def cell(b, se, p, d=3):
    return f"{b:.{d}f}{st(p)} ({se:.{d}f})"

FA = "Arial"
H = Font(name=FA, bold=True, size=10); B = Font(name=FA, size=10); T = Font(name=FA, bold=True, size=12); N = Font(name=FA, italic=True, size=9, color="555555")
FILL = PatternFill("solid", fgColor="EDEDED"); THIN = Side(style="thin", color="999999")
wb = Workbook(); wb.remove(wb.active)

def put(ws, r, c, v, font=B, fmt=None, align=None):
    if isinstance(v, np.integer): v = int(v)
    if isinstance(v, np.floating): v = float(v)
    if isinstance(v, float) and np.isnan(v): v = None
    x = ws.cell(row=r, column=c, value=v); x.font = font
    if fmt: x.number_format = fmt
    if align: x.alignment = Alignment(horizontal=align)
    return x

def header(ws, r, cols):
    for j, c in enumerate(cols, 1):
        x = put(ws, r, j, c, H); x.fill = FILL; x.border = Border(top=THIN, bottom=THIN); x.alignment = Alignment(horizontal="center", wrap_text=True)

def widths(ws, w):
    for j, v in enumerate(w, 1): ws.column_dimensions[get_column_letter(j)].width = v

# ---------------- T1 headline ----------------
ws = wb.create_sheet("T1_headline")
put(ws, 1, 1, "Table 1. Hub connectivity and income inequality, 1996-2023", T)
put(ws, 2, 1, "Dependent variable in logs. Country and year fixed effects, log population. 2SLS instruments ln GACI with the Feyrer interaction a_t x ln MA^air_1996. SE clustered by country. *** p<.01, ** p<.05, * p<.1.", N)
put(ws, 3, 1, "Source: _main_results_cl.csv (01_main_cl.do).", N)
r = 5
header(ws, r, ["", "Market-income Gini", "Disposable-income Gini", "N"])
panels = [("ln_gaci_max", "Panel A. Hub connectivity, ln GACI_max (headline)"), ("ln_gaci_cwm", "Panel B. Hub quality, ln GACI_cwm"), ("ln_gaci_sum", "Panel C. Total connectivity, ln GACI_sum")]
m = main[main.iv == "feyrer_int"]
for tr, lab in panels:
    r += 1; put(ws, r, 1, lab, H)
    x = m[m.treat == tr]
    fs = x[x.model == "FS"].iloc[0]
    for mod, mlab in [("OLS", "OLS"), ("IVcl", "2SLS, Feyrer IV")]:
        r += 1; put(ws, r, 1, mlab)
        for j, o in enumerate(["ln_gini_mkt", "ln_gini_disp"], 2):
            y = x[(x.outcome == o) & (x.model == mod)].iloc[0]; put(ws, r, j, cell(y.b, y.se, y.p), align="center")
        put(ws, r, 4, int(y.N), fmt="#,##0", align="center")
    r += 1; put(ws, r, 1, "First stage: coefficient on Z (se)"); put(ws, r, 2, cell(fs.b, fs.se, fs.p), align="center")
    r += 1; put(ws, r, 1, "Kleibergen-Paap F"); put(ws, r, 2, float(fs.kpf), fmt="0.0", align="center")
widths(ws, [44, 24, 24, 10])

# ---------------- curves ----------------
def curve_sheet(name, title, src, treat_lab, gapblock="gic_gap", extra_gap=None):
    ws = wb.create_sheet(name)
    put(ws, 1, 1, title, T)
    put(ws, 2, 1, f"Each row is one 2SLS regression of the outcome on {treat_lab} (Feyrer IV), log population, country and year FE; SE clustered by country. Income = ln average pretax national income of the group (WID, equal-split adults).", N)
    put(ws, 3, 1, "Share = ln income share of the group, computed from the identity ln share = ln group income - ln mean income (same sample as the income column). Stars: *** p<.01, ** p<.05, * p<.1.", N)
    put(ws, 4, 1, "Source: " + src + ".", N)
    r = 6
    header(ws, r, ["Group", "Income: elasticity", "se", "Share: elasticity", "se", "KP F", "N"])
    df = gcwm if "cwm" in name else gmax
    for g in G:
        a = df[(df.block == "gic") & (df.outcome == "ln_apt_" + g) & (df.spec == "IV")]
        s = df[(df.block == "gic_share") & (df.outcome == "ln_spt_" + g) & (df.spec == "IV")]
        if not len(a): continue
        a = a.iloc[0]; r += 1
        put(ws, r, 1, LAB[g]); put(ws, r, 2, f"{a.b:.3f}{st(a.p)}", align="right"); put(ws, r, 3, f"({a.se:.3f})", align="left")
        if len(s):
            s = s.iloc[0]; put(ws, r, 4, f"{s.b:.3f}{st(s.p)}", align="right"); put(ws, r, 5, f"({s.se:.3f})", align="left")
        put(ws, r, 6, float(a.kpf), fmt="0.0", align="center"); put(ws, r, 7, int(a.N), fmt="#,##0", align="center")
    r += 2; put(ws, r, 1, "Differences between groups (difference outcome, same regression)", H)
    r += 1; header(ws, r, ["Contrast", "Estimate", "se", "", "", "KP F", "N"])
    gl = {"ln_apt_d10_d1": "p90-100 minus p0-10", "ln_apt_d10_b50": "p90-100 minus Bottom 50%", "ln_apt_t1_d10": "Top 1% minus p90-100", "ln_apt_top_bot": "Top half (p50-100) minus bottom half (p0-50)"}
    gp = df[(df.block == gapblock) & (df.spec == "IV")]
    if extra_gap is not None: gp = pd.concat([gp, extra_gap])
    for k, l in gl.items():
        x = gp[gp.outcome == k]
        if not len(x): continue
        x = x.iloc[0]; r += 1
        put(ws, r, 1, l); put(ws, r, 2, f"{x.b:.3f}{st(x.p)}", align="right"); put(ws, r, 3, f"({x.se:.3f})", align="left")
        put(ws, r, 6, float(x.kpf), fmt="0.0", align="center"); put(ws, r, 7, int(x.N), fmt="#,##0", align="center")
    widths(ws, [44, 18, 10, 18, 10, 8, 9])
    return ws

extra = gcwm[(gcwm.block == "gic_gap_max") & (gcwm.spec == "IV")]
curve_sheet("T2_curve_GACI_max", "Table 2. Incidence of hub connectivity (ln GACI_max) across the income distribution", "_gic_dose_results.csv (12_gic_dose.do); top-half contrast from 12c_gic_cwm.do", "ln GACI_max", extra_gap=extra)
curve_sheet("T3_curve_GACI_cwm", "Table 3. Incidence of hub quality (ln GACI_cwm, capacity-weighted mean) across the income distribution", "_gic_cwm_results.csv (12c_gic_cwm.do)", "ln GACI_cwm")

# ---------------- long-format sheet for plotting ----------------
ws = wb.create_sheet("data_long")
put(ws, 1, 1, "Long format of Tables 2-3 (2SLS Feyrer, country-clustered) for figures.", N)
header(ws, 2, ["treatment", "outcome_type", "group", "b", "se", "p", "kpf", "N"])
r = 2
for tl, df in [("GACI_max", gmax), ("GACI_cwm", gcwm)]:
    for blk, ot, pre in [("gic", "income", "ln_apt_"), ("gic_share", "share", "ln_spt_")]:
        for g in G:
            x = df[(df.block == blk) & (df.outcome == pre + g) & (df.spec == "IV")]
            if not len(x): continue
            x = x.iloc[0]; r += 1
            for j, v in enumerate([tl, ot, LAB[g], x.b, x.se, x.p, x.kpf, int(x.N)], 1):
                put(ws, r, j, v, fmt="0.000" if isinstance(v, float) else None)
ws.freeze_panes = "A3"; widths(ws, [12, 12, 18, 10, 10, 10, 8, 8])

out = "GIC_paper_tables_20260928.xlsx"; wb.save(out); print("saved", out, wb.sheetnames)
