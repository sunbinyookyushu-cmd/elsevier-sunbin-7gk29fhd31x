# -*- coding: utf-8 -*-
"""37_add_components_premia_xlsx.py : append A10 (hub components) and A11 (ILOSTAT wage premia) to the paper workbook."""
import os, pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
FN = "Times New Roman"; F_T = Font(name=FN, bold=True, size=11); F_H = Font(name=FN, bold=True, size=10); F_B = Font(name=FN, size=10); F_I = Font(name=FN, italic=True, size=10); F_N = Font(name=FN, size=9)
TOP = Side(style="medium"); THIN = Side(style="thin")
def st(p): return "" if pd.isna(p) else "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
wb = load_workbook("Tables_for_paper_20260928.xlsx")
for n in ["A10_hub_components", "A11_wage_premia"]:
    if n in wb.sheetnames: wb.remove(wb[n])

class Sheet:
    def __init__(self, name, after, title):
        self.ws = wb.create_sheet(name, wb.sheetnames.index(after) + 1); self.r = 1
        self.put(1, title, F_T, "left"); self.r = 3
    def put(self, c, v, font=F_B, align="center", fmt=None):
        x = self.ws.cell(row=self.r, column=c, value=v); x.font = font; x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(c == 1))
        if fmt: x.number_format = fmt
    def put_at(self, rr, c, v):
        x = self.ws.cell(row=rr, column=c, value=v); x.font = F_B; x.alignment = Alignment(horizontal="center")
    def rule(self, ncol, top=False):
        for c in range(1, ncol + 1): self.ws.cell(row=self.r, column=c).border = Border(top=TOP, bottom=THIN) if top else Border(bottom=TOP)
    def notes(self, lines):
        self.r += 1
        for n in lines: self.put(1, n, F_N, "left"); self.r += 1
    def widths(self, w):
        for j, v in enumerate(w, 1): self.ws.column_dimensions[get_column_letter(j)].width = v

# ---------------- A10 ----------------
d = pd.read_csv("_components_results.csv", keep_default_na=False, na_values=["", "."])
S = Sheet("A10_hub_components", "A9_median_splits", "Table A10. Which component of hub connectivity carries the effect?")
COMP = [("gaci", "GACI (composite)"), ("cap", "Seat capacity"), ("deg", "Number of connections (degree)"), ("eig", "Eigenvector centrality (integration with the global core)"), ("close", "Closeness"), ("betw", "Flow betweenness (transit role)"), ("regimp", "Regional importance")]
OUTS = [("ln_gini_mkt", "ln Market Gini"), ("ln_gini_disp", "ln Disposable Gini"), ("ln_spt_b50", "Bottom 50% share"), ("ln_apt_top_bot", "Top half minus bottom half")]
S.put(1, "Hub component (ln), instrumented by the Feyrer shifter", F_H, "left")
for j, (_, l) in enumerate(OUTS, 2): S.put(j, l, F_H)
S.put(6, "First stage", F_H); S.put(7, "KP F", F_H); S.put(8, "N", F_H); S.rule(8, top=True); S.r += 1
S.put(1, "Panel A. One component at a time (2SLS)", F_I, "left"); S.r += 1
for k, lab in COMP:
    S.put(1, lab, F_B, "left")
    for j, (o, _) in enumerate(OUTS, 2):
        x = d[(d.block == "B") & (d["item"] == "iv_single") & (d.outcome == o) & (d.term == "ln_hub_" + k)]
        if len(x): S.put(j, f"{x.b.iloc[0]:.3f}{st(x.p.iloc[0])}"); S.put_at(S.r + 1, j, f"({x.se.iloc[0]:.3f})")
    fs = d[(d.block == "A") & (d.outcome == "ln_hub_" + k)]
    if len(fs): S.put(6, f"{fs.b.iloc[0]:.3f}{st(fs.p.iloc[0])}"); S.put_at(S.r + 1, 6, f"({fs.se.iloc[0]:.3f})")
    x = d[(d.block == "B") & (d["item"] == "iv_single") & (d.outcome == "ln_gini_mkt") & (d.term == "ln_hub_" + k)]
    if len(x): S.put(7, float(x.kpf.iloc[0]), fmt="0.0"); S.put(8, int(x.N.iloc[0]), fmt="#,##0")
    S.r += 2
S.r += 1; S.put(1, "Panel B. Horse race on ln Market Gini: one component instrumented, another as control", F_I, "left"); S.r += 1
for item, lab in [("topo_iv_cap_ctrl_eig", "Eigenvector instrumented, seat capacity as control"), ("cap_iv_topo_ctrl_betw", "Seat capacity instrumented, betweenness as control"), ("cap_iv_topo_ctrl_close", "Seat capacity instrumented, closeness as control"), ("cap_iv_topo_ctrl_deg", "Seat capacity instrumented, degree as control"), ("topo_iv_cap_ctrl_deg", "Degree instrumented, seat capacity as control"), ("topo_iv_cap_ctrl_betw", "Betweenness instrumented, seat capacity as control")]:
    x = d[(d.block == "C") & (d["item"] == item) & (d.outcome == "ln_gini_mkt")]
    if not len(x): continue
    S.put(1, lab, F_B, "left")
    for j, (_, rr) in enumerate(x.iterrows(), 2):
        S.put(j, "%s: %.3f%s (%.3f)" % (rr.term.replace("ln_hub_", ""), rr.b, st(rr.p), rr.se))
    S.put(7, float(x.kpf.iloc[0]), fmt="0.0"); S.put(8, int(x.N.iloc[0]), fmt="#,##0"); S.r += 1
S.r += 1; S.put(1, "Panel C. OLS, all components jointly (ln Market Gini)", F_I, "left"); S.r += 1
for _, rr in d[(d.block == "D") & (d.outcome == "ln_gini_mkt")].iterrows():
    S.put(1, dict(COMP).get(rr.term.replace("ln_hub_", ""), rr.term), F_B, "left"); S.put(2, "%.3f%s (%.3f)" % (rr.b, st(rr.p), rr.se)); S.r += 1
S.r -= 1; S.rule(8); S.r += 1
S.notes(["Notes: Hub = the airport with the highest GACI in the country-year; components from the airport-level GACI panel. 2SLS with the Feyrer instrument, ln population, country and year FE, SE clustered by country. *** p<0.01, ** p<0.05, * p<0.1.",
         "Reading: a single instrument moves every component (first-stage column), so one-at-a-time estimates are not separately identified. In the horse races that keep a first stage (KP F above 10), eigenvector centrality net of capacity raises inequality (0.133**) while capacity net of eigenvector does not; capacity net of betweenness raises it (0.089***) while betweenness is negative; joint OLS gives degree positive and betweenness negative. Integration with the global core network, rather than seat volume or a transit role, is the component associated with the effect. Suggestive only.",
         "Source: _components_results.csv (34_components.do); _hub_components.csv (33_build_hub_components.py)."])
S.widths([60] + [22] * 7)

# ---------------- A11 ----------------
w = pd.read_csv("_wage_premia_results.csv", keep_default_na=False, na_values=["", "."])
S = Sheet("A11_wage_premia", "A10_hub_components", "Table A11. Wage premia from ILOSTAT earnings: the skill channel cannot be identified")
S.put(1, "Outcome (log ratio of average monthly earnings, employees)", F_H, "left")
for j, l in enumerate(["2SLS", "OLS", "KP F", "N", "Countries", "Headline Gini on the same sample (2SLS)"], 2): S.put(j, l, F_H)
S.rule(7, top=True)
LAB = {"prem_adv_bas": "Advanced vs basic education", "prem_adv_int": "Advanced vs intermediate education", "prem_ser_agr": "Services vs agriculture", "prem_ser_ind": "Services vs industry", "prem_mkt_man": "Market services vs manufacturing", "prem_mgr_elem": "Managers vs elementary occupations", "prem_prof_elem": "Professionals vs elementary occupations", "prem_prof_craft": "Professionals vs craft workers", "gini_earn_total": "Gini of monthly earnings (level)", "ln_earn_total": "ln average monthly earnings"}
for k, lab in LAB.items():
    iv = w[(w.outcome == k) & (w.model == "IV")]; ol = w[(w.outcome == k) & (w.model == "OLS")]; hg = w[(w.outcome == k + "_sample_gini")]
    if not len(iv): continue
    S.r += 1; S.put(1, lab, F_B, "left")
    S.put(2, "%.2f%s (%.2f)" % (iv.b.iloc[0], st(iv.p.iloc[0]), iv.se.iloc[0])); S.put(3, "%.3f%s (%.3f)" % (ol.b.iloc[0], st(ol.p.iloc[0]), ol.se.iloc[0]))
    S.put(4, float(iv.kpf.iloc[0]), fmt="0.00"); S.put(5, int(iv.N.iloc[0]), fmt="#,##0"); S.put(6, int(iv.Nc.iloc[0]), fmt="#,##0")
    if len(hg): S.put(7, "%.2f%s (%.2f)" % (hg.b.iloc[0], st(hg.p.iloc[0]), hg.se.iloc[0]))
S.rule(7)
S.notes(["Notes: ILOSTAT EAR_EMTA (average monthly earnings of employees, both sexes) by education, economic activity and occupation; one source per country. 2SLS with the Feyrer instrument, ln population, country and year FE, SE clustered by country.",
         "Reading: earnings data exist mainly after 2010 and for about 100 countries; on that subsample the instrument has no first stage (KP F below 0.2) and even the headline Gini is unidentified, so the 2SLS columns are uninformative. OLS associations are all insignificant. The skill-premium channel cannot be tested causally with available cross-country wage data.",
         "Source: _wage_premia_results.csv (36_wage_premia.do); _ilo_premia.csv."])
S.widths([48] + [18] * 6)
wb.save("Tables_for_paper_20260928.xlsx"); print("saved A10, A11")
