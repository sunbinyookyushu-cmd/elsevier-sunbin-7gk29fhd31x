# -*- coding: utf-8 -*-
"""28_add_stage_defs_xlsx.py : append Table A8 (robustness of the development-stage split to its definition)
   to Tables_for_paper_20260928.xlsx. Source: _stage_defs_results.csv (27_stage_defs.do)."""
import os, numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
d = pd.read_csv("_stage_defs_results.csv", keep_default_na=False, na_values=["", "."])
FN = "Times New Roman"; F_T = Font(name=FN, bold=True, size=11); F_H = Font(name=FN, bold=True, size=10); F_B = Font(name=FN, size=10); F_I = Font(name=FN, italic=True, size=10); F_N = Font(name=FN, size=9)
TOP = Side(style="medium"); THIN = Side(style="thin")
def st(p): return "" if pd.isna(p) else "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
wb = load_workbook("Tables_for_paper_20260928.xlsx")
if "A8_stage_definitions" in wb.sheetnames: wb.remove(wb["A8_stage_definitions"])
ws = wb.create_sheet("A8_stage_definitions", wb.sheetnames.index("A7_spillovers_full") + 1)
ws["A1"] = "Table A8. Development-stage split under alternative definitions (2SLS, Feyrer instrument)"; ws["A1"].font = F_T
DEFS = [("g_max96_t3", "GACI_max in 1996, terciles (baseline)", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"}),
        ("g_max96_q4", "GACI_max in 1996, quartiles", {"1_q1": "Q1 (lowest)", "2_q2": "Q2", "3_q3": "Q3", "4_q4": "Q4 (highest)"}),
        ("g_max96_med", "GACI_max in 1996, median split", {"1_below": "Below median", "2_above": "Above median"}),
        ("g_max9600_t3", "GACI_max, mean 1996-2000, terciles", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"}),
        ("g_maxall_t3", "GACI_max, full-period country mean, terciles", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"}),
        ("g_cwm96_t3", "GACI_cwm (capacity-weighted mean) in 1996, terciles", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"}),
        ("g_sum96_t3", "GACI_sum (total connectivity) in 1996, terciles", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"}),
        ("g_share96_t3", "Share of the top airport in national connectivity, 1996, terciles", {"1_low": "Low (dispersed network)", "2_mid": "Middle", "3_high": "High (single dominant hub)"}),
        ("g_max96_wcont", "GACI_max in 1996, terciles within continent", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"}),
        ("g_airma96_t3", "Air market access in 1996 (instrument exposure), terciles", {"1_low": "Low", "2_mid": "Middle", "3_high": "High"})]
OUTS = [("ln_gini_mkt", "ln Market Gini"), ("ln_apt_all", "Mean income"), ("ln_apt_top_bot", "Top half minus bottom half"), ("ln_spt_b50", "Bottom 50% share"), ("ln_spt_d10", "Top 10% share")]
r = 3
ws.cell(row=r, column=1, value="Definition / group").font = F_H
for j, (_, l) in enumerate(OUTS, 2): x = ws.cell(row=r, column=j, value=l); x.font = F_H; x.alignment = Alignment(horizontal="center", wrap_text=True)
ws.cell(row=r, column=7, value="KP F").font = F_H; ws.cell(row=r, column=8, value="N").font = F_H; ws.cell(row=r, column=9, value="Countries").font = F_H
for c in range(1, 10): ws.cell(row=r, column=c).border = Border(top=TOP, bottom=THIN); ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
for dv, lab, groups in DEFS:
    r += 1; ws.cell(row=r, column=1, value=lab).font = F_I
    for g, gl in groups.items():
        x = d[(d.defn == dv) & (d.group == g)]
        if not len(x): continue
        r += 1; ws.cell(row=r, column=1, value="   " + gl).font = F_B
        for j, (o, _) in enumerate(OUTS, 2):
            y = x[x.outcome == o]
            if len(y):
                c1 = ws.cell(row=r, column=j, value=f"{y.b.iloc[0]:.3f}{st(y.p.iloc[0])}"); c1.font = F_B; c1.alignment = Alignment(horizontal="center")
                c2 = ws.cell(row=r + 1, column=j, value=f"({y.se.iloc[0]:.3f})"); c2.font = F_B; c2.alignment = Alignment(horizontal="center")
        y = x[x.outcome == "ln_gini_mkt"].iloc[0]
        for j, v, f in [(7, float(y.kpf), "0.0"), (8, int(y.N), "#,##0"), (9, int(y.Nc), "#,##0")]:
            c = ws.cell(row=r, column=j, value=v); c.number_format = f; c.font = F_B; c.alignment = Alignment(horizontal="center")
        r += 1
for c in range(1, 10): ws.cell(row=r, column=c).border = Border(bottom=TOP)
r += 2
for n in ["Notes: Split-sample 2SLS with the Feyrer instrument, ln population, country and year fixed effects; standard errors clustered by country in parentheses. *** p<0.01, ** p<0.05, * p<0.1. Groups are formed over countries on the stated baseline variable.",
          "Reading: the instrument has power (KP F above 7) only in groups with small-to-moderate baseline hubs; wherever it has power, the distributional pattern is the same (bottom-50 share down, top-10 share up). In the thinnest networks by GACI_max or GACI_cwm in 1996 (lowest tercile or quartile) there is no distributional tilt and the Gini elasticity is small; that near-zero is not robust to averaging the base over 1996-2000 (0.26**), but the absence of a tilt is. Splitting on the instrument's own exposure (air market access) removes all identifying variation, as expected.",
          "Source: _stage_defs_results.csv (27_stage_defs.do)."]:
    ws.cell(row=r, column=1, value=n).font = F_N; r += 1
ws.column_dimensions["A"].width = 52
for j in range(2, 10): ws.column_dimensions[get_column_letter(j)].width = 15
wb.save("Tables_for_paper_20260928.xlsx"); print("saved A8", wb.sheetnames[-4:])
