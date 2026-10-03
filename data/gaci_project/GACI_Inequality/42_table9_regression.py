# -*- coding: utf-8 -*-
"""42_table9_regression.py : Table 9 as a proper regression table (coefficient row, s.e. row), five panels.
   Sheet Table9_regression in Tables_for_paper_20260928.xlsx; also prints a markdown version."""
import os, pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
RD = dict(keep_default_na=False, na_values=["", "."])
mech = pd.read_csv("_mech_results.csv", **RD); comp = pd.read_csv("_components_results.csv", **RD); gic = pd.read_csv("_gic_dose_results.csv", **RD)
alt = pd.read_csv("_alt_measures_results.csv"); med = pd.read_csv("_median_splits_results.csv", **RD); stg = pd.read_csv("_gic_bysample_results.csv", **RD)
sp = pd.read_csv("_spill_results.csv", **RD); macro = pd.read_csv("_macro_channels_results.csv", **RD)
def st(p): return "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
def g(df, **kw):
    x = df
    for k, v in kw.items(): x = x[x[k] == v]
    return x.iloc[0] if len(x) else None
def cell(r, d=3): return (f"{r.b:.{d}f}{st(r.p)}", f"({r.se:.{d}f})") if r is not None else ("", "")
def kpn(r): return (float(r.kpf) if r is not None and not pd.isna(r.kpf) else None, int(r.N) if r is not None else None)

# ---------- assemble panels: each panel = (title, dep-var line, column headers, rows[(label, [cells per column])], stats rows) ----------
P = []
# Panel A
base = g(mech, block="A", item="baseline", outcome="ln_gini_mkt", term="ln_gaci_max")
hs_max = g(mech, block="A", item="hub_ctrl_sum", outcome="ln_gini_mkt", term="ln_gaci_max"); hs_sum = g(mech, block="A", item="hub_ctrl_sum", outcome="ln_gini_mkt", term="ln_gaci_sum")
hr_max = g(mech, block="A", item="hub_ctrl_rest", outcome="ln_gini_mkt", term="ln_gaci_max"); hr_rest = g(mech, block="A", item="hub_ctrl_rest", outcome="ln_gini_mkt", term="ln_gaci_rest")
eig = g(comp, block="C", item="topo_iv_cap_ctrl_eig", outcome="ln_gini_mkt", term="ln_hub_eig"); eig_cap = g(comp, block="C", item="topo_iv_cap_ctrl_eig", outcome="ln_gini_mkt", term="ln_hub_cap")
cap = g(comp, block="C", item="cap_iv_topo_ctrl_betw", outcome="ln_gini_mkt", term="ln_hub_cap"); cap_betw = g(comp, block="C", item="cap_iv_topo_ctrl_betw", outcome="ln_gini_mkt", term="ln_hub_betw")
hhi = g(mech, block="A", item="ols_hhi", outcome="ln_gini_mkt", term="ln_hhi"); shm = g(mech, block="A", item="ols_share_max", outcome="ln_gini_mkt", term="ln_share_max"); shm_sum = g(mech, block="A", item="ols_share_max", outcome="ln_gini_mkt", term="ln_gaci_sum")
P.append(("Panel A. Which connectivity? Dependent variable: ln market Gini",
          ["(1) Baseline", "(2) Hub + network", "(3) Hub + secondary airports", "(4) Core integration + capacity", "(5) Capacity + transit", "(6) Concentration (OLS)"],
          [("ln GACI_max (hub), instrumented", [base, hs_max, hr_max, None, None, None]),
           ("ln GACI_sum (whole network), control", [None, hs_sum, None, None, None, shm_sum]),
           ("ln GACI of secondary airports, control", [None, None, hr_rest, None, None, None]),
           ("ln hub eigenvector centrality, instrumented", [None, None, None, eig, None, None]),
           ("ln hub seat capacity", [None, None, None, eig_cap, cap, None]),
           ("ln hub flow betweenness, control", [None, None, None, None, cap_betw, None]),
           ("ln share of top airport in national GACI", [None, None, None, None, None, shm])],
          [base, hs_max, hr_max, eig, cap, shm], "In (4) the eigenvector term is instrumented and capacity is a control; in (5) capacity is instrumented and betweenness is a control."))
# Panel B
cols = [("ln_apt_all", "Mean income"), ("ln_apt_d7", "p60-70 income"), ("ln_apt_d10", "p90-100 income"), ("ln_apt_b50", "Bottom-50 income"), ("ln_spt_b50", "Bottom-50 share")]
rowsB = [g(gic, block="gic" if not o.startswith("ln_spt") else "gic_share", outcome=o, spec="IV") for o, _ in cols]
p5010 = alt[alt.outcome == "ln_p50_p10"].iloc[0]; p9050 = alt[alt.outcome == "ln_p90_p50"].iloc[0]
class R: pass
def mk(row): r = R(); r.b = row.iv_b; r.se = row.iv_se; r.p = row.iv_p; r.kpf = row.kpf; r.N = row.N; return r
rowsB += [mk(p5010), mk(p9050)]
P.append(("Panel B. Who gains? Dependent variables: WID group incomes and ratios (logs), 2SLS",
          [l for _, l in cols] + ["p50 / p10", "p90 / p50"],
          [("ln GACI_max, instrumented", rowsB)], rowsB, "Group incomes are average pretax national income per equal-split adult; share = ln group income minus ln mean income; ratios use decile means."))
# Panel C
meds = [("ln_gdppc", "GDP per capita (ln)", mech, "C")]
cprime = [g(df, block=blk, item=m, outcome="ln_gini_mkt", term="ln_gaci_max") for m, _, df, blk in meds]
bmed = [g(df, block=blk, item=m, outcome="ln_gini_mkt", term=m) for m, _, df, blk in meds]
apath = [g(mech if df is mech else macro, block="B" if df is mech else "A", item="apath", outcome=m) for m, _, df, blk in meds]
P.append(("Panel C. Not a growth effect. Dependent variable: ln market Gini, hub instrumented, GDP per capita as control",
          [l for _, l, _, _ in meds],
          [("ln GACI_max, instrumented (direct effect c')", cprime), ("Mediator (b)", bmed), ("Memo: hub -> mediator, 2SLS (a-path)", apath)],
          cprime, "GDP per capita is itself an outcome of connectivity (a-path), so it is not a valid control in the baseline; holding it fixed isolates the distributional effect from the growth effect. Other mediators (trade, tourism, sectoral structure) are reported in Appendix Table A12."))
# Panel D
urb2 = g(med, var="urban", group="above", outcome="ln_gini_mkt"); urb1 = g(med, var="urban", group="below", outcome="ln_gini_mkt")
tour1 = g(mech, block="E", item="tour_rcpt_exp_1", outcome="ln_gini_mkt"); thin = g(stg, sample="con_low", outcome="ln_gini_mkt"); mid = g(stg, sample="con_mid", outcome="ln_gini_mkt")
open2 = g(med, var="trade_gdp", group="above", outcome="ln_gini_mkt")
nb_own = g(sp, panel="A", item="joint_contig", var="ln_gaci_max", outcome="ln_gini_mkt"); nb_nbr = g(sp, panel="A", item="joint_contig", var="nbr_g_contig", outcome="ln_gini_mkt")
P.append(("Panel D. Where? Dependent variable: ln market Gini, 2SLS on subsamples; last column joint 2SLS with neighbours",
          ["Urbanisation above median", "Urbanisation below median", "Non-tourism economies", "Thin networks (low tercile)", "Take-off hubs (middle tercile)", "Open economies (trade above median)", "Own + contiguous neighbours"],
          [("ln GACI_max, instrumented", [urb2, urb1, tour1, thin, mid, open2, nb_own]), ("Neighbours' ln GACI (contiguity, leave-out mean), instrumented", [None] * 6 + [nb_nbr])],
          [urb2, urb1, tour1, thin, mid, open2, nb_own], "Splits at the median (or terciles) of the earliest observed value; tourism = receipts below median of exports share; connectivity terciles on GACI_max in 1996."))
# Panel E
wedge = g(mech, block="D", item="redistribution", outcome="wedge_pts"); wr_ = g(mech, block="D", item="redistribution", outcome="ln_wedge_ratio"); tax = g(mech, block="D", item="redistribution", outcome="tax_gdp")
disp = g(mech, block="A", item="baseline", outcome="ln_gini_disp", term="ln_gaci_max")
P.append(("Panel E. Is it offset? Dependent variables: redistribution measures, 2SLS",
          ["ln disposable Gini", "Market minus disposable Gini (points)", "ln (market / disposable Gini)", "Tax revenue, % GDP"],
          [("ln GACI_max, instrumented", [disp, wedge, wr_, tax])], [disp, wedge, wr_, tax], "The wedge is the redistribution achieved by taxes and transfers; a negative coefficient means less redistribution."))

# ---------- write ----------
wb = load_workbook("Tables_for_paper_20260928.xlsx")
if "Table9_regression" in wb.sheetnames: wb.remove(wb["Table9_regression"])
ws = wb.create_sheet("Table9_regression", wb.sheetnames.index("Table9_compact"))
FN = "Times New Roman"; F_T = Font(name=FN, bold=True, size=11); F_H = Font(name=FN, bold=True, size=10); F_B = Font(name=FN, size=10); F_I = Font(name=FN, italic=True, size=10); F_N = Font(name=FN, size=9)
TOP = Side(style="medium"); THIN = Side(style="thin")
ws["A1"] = "Table 9. Mechanism: a premium on access that bypasses the bottom half"; ws["A1"].font = F_T
r = 2; md = []
def put(c, v, font=F_B, align="center"):
    x = ws.cell(row=r, column=c, value=v); x.font = font; x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
for title, heads, rows, statrows, note in P:
    r += 1; put(1, title, F_I, "left"); r += 1
    put(1, "", F_H, "left")
    for j, h in enumerate(heads, 2): put(j, h, F_H)
    for c in range(1, len(heads) + 2): ws.cell(row=r, column=c).border = Border(top=TOP, bottom=THIN)
    md.append(f"\n**{title}**\n\n| | " + " | ".join(heads) + " |\n|---|" + "---|" * len(heads))
    for label, cells in rows:
        r += 1; put(1, label, F_B, "left"); line1 = []; line2 = []
        for j, cc in enumerate(cells, 2):
            dd = 4 if (label.startswith("Mediator") and cc is not None and abs(cc.b) < 0.05) else 3
            a, b = cell(cc, dd); put(j, a); line1.append(a); line2.append(b)
            if b: ws.cell(row=r + 1, column=j, value=b).font = F_B; ws.cell(row=r + 1, column=j).alignment = Alignment(horizontal="center")
        r += 1
        md.append(f"| {label} | " + " | ".join(a if a else "" for a in line1) + " |"); md.append("| | " + " | ".join(line2) + " |")
    r += 1; put(1, "KP F", F_B, "left"); kf = []
    for j, s in enumerate(statrows, 2):
        k, n = kpn(s)
        if k is not None: x = ws.cell(row=r, column=j, value=k); x.number_format = "0.0"; x.font = F_B; x.alignment = Alignment(horizontal="center"); kf.append(f"{k:.1f}")
        else: kf.append("")
    md.append("| KP F | " + " | ".join(kf) + " |")
    r += 1; put(1, "Observations", F_B, "left"); nn = []
    for j, s in enumerate(statrows, 2):
        k, n = kpn(s)
        if n is not None: x = ws.cell(row=r, column=j, value=n); x.number_format = "#,##0"; x.font = F_B; x.alignment = Alignment(horizontal="center"); nn.append(f"{n:,}")
        else: nn.append("")
    md.append("| N | " + " | ".join(nn) + " |")
    for c in range(1, len(heads) + 2): ws.cell(row=r, column=c).border = Border(bottom=THIN)
    r += 1; put(1, note, F_N, "left"); r += 1
for c in range(1, 9): ws.cell(row=r - 2, column=c).border = Border(bottom=TOP)
r += 1
ws.cell(row=r, column=1, value="Notes: 2SLS with the Feyrer instrument Z = a_t x ln MA_1996 unless marked OLS; all regressions include ln population, country and year fixed effects; standard errors clustered by country in parentheses. *** p<0.01, ** p<0.05, * p<0.1. Outcomes in logs unless stated.").font = F_N
ws.column_dimensions["A"].width = 52
for j in range(2, 9): ws.column_dimensions[get_column_letter(j)].width = 19
wb.save("Tables_for_paper_20260928.xlsx"); print("saved Table9_regression"); print("\n".join(md))
