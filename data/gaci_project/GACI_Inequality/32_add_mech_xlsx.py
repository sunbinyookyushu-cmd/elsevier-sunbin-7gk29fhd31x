# -*- coding: utf-8 -*-
"""32_add_mech_xlsx.py : Table 9 (testable mechanisms) appended to Tables_for_paper_20260928.xlsx.
   Panels: A hub vs rest of network; B a-paths (does connectivity move the mediator?); C IV decomposition
   (direct c' and mediator b for Gini, bottom-50 share, top-half minus bottom-half); D redistribution wedge;
   E heterogeneity by channel proxies. Source: _mech_results.csv (31_mechanisms.do)."""
import os, numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
d = pd.read_csv("_mech_results.csv", keep_default_na=False, na_values=["", "."])
FN = "Times New Roman"; F_T = Font(name=FN, bold=True, size=11); F_H = Font(name=FN, bold=True, size=10); F_B = Font(name=FN, size=10); F_I = Font(name=FN, italic=True, size=10); F_N = Font(name=FN, size=9)
TOP = Side(style="medium"); THIN = Side(style="thin")
st = lambda p: "" if pd.isna(p) else "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""
def cs(r, dd=3): return f"{r.b:.{dd}f}{st(r.p)}"
def ss(r, dd=3): return f"({r.se:.{dd}f})"
wb = load_workbook("Tables_for_paper_20260928.xlsx")
if "Table9_mechanisms" in wb.sheetnames: wb.remove(wb["Table9_mechanisms"])
ws = wb.create_sheet("Table9_mechanisms", wb.sheetnames.index("Table8_aggregate") + 1)
ws["A1"] = "Table 9. Testable mechanisms: which connectivity, what it moves, and what does not transmit the effect"; ws["A1"].font = F_T
r = [3]
def put(c, v, font=F_B, align="center", fmt=None):
    x = ws.cell(row=r[0], column=c, value=v); x.font = font; x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(c == 1))
    if fmt: x.number_format = fmt
def header(cols, first):
    put(1, first, F_H, "left")
    for j, c in enumerate(cols, 2): put(j, c, F_H)
    for j in range(1, len(cols) + 2): ws.cell(row=r[0], column=j).border = Border(top=TOP, bottom=THIN)
    r[0] += 1
def panel(lab): put(1, lab, F_I, "left"); r[0] += 1
def row2(label, cells, extra=None):
    put(1, label, F_B, "left")
    for j, c in enumerate(cells, 2): put(j, cs(c) if c is not None else "--")
    if extra:
        for j, v in extra.items(): put(j, v, F_B, "center", "#,##0" if isinstance(v, int) else "0.0")
    r[0] += 1
    for j, c in enumerate(cells, 2):
        if c is not None: put(j, ss(c))
    r[0] += 1
def get(**kw):
    x = d
    for k, v in kw.items(): x = x[x[k] == v]
    return x.iloc[0] if len(x) else None
OUTS = [("ln_gini_mkt", "ln Market Gini"), ("ln_gini_disp", "ln Disposable Gini"), ("ln_spt_b50", "Bottom 50% share"), ("ln_apt_top_bot", "Top half minus bottom half")]

# ---- Panel A ----
header([l for _, l in OUTS] + ["KP F", "N"], "")
panel("Panel A. Which connectivity? Hub (top airport) versus the rest of the network")
for item, term, lab in [("baseline", "ln_gaci_max", "ln GACI_max, baseline (2SLS)"),
                        ("hub_ctrl_sum", "ln_gaci_max", "ln GACI_max (2SLS), controlling ln GACI_sum"), ("hub_ctrl_sum", "ln_gaci_sum", "   ln GACI_sum (total connectivity, control)"),
                        ("hub_ctrl_rest", "ln_gaci_max", "ln GACI_max (2SLS), controlling secondary airports"), ("hub_ctrl_rest", "ln_gaci_rest", "   ln GACI of secondary airports (sum minus max, control)"),
                        ("rest_iv_ctrl_hub", "ln_gaci_rest", "Secondary airports instrumented, ln GACI_max as control"),
                        ("ols_share_max", "ln_share_max", "OLS: ln share of top airport in national GACI (with ln GACI_sum)"), ("ols_hhi", "ln_hhi", "OLS: ln Herfindahl of airport shares (with ln GACI_sum)")]:
    cells = [get(block="A", item=item, outcome=o, term=term) for o, _ in OUTS]
    x = get(block="A", item=item, outcome="ln_gini_mkt", term=term)
    row2(lab, cells, {6: float(x.kpf), 7: int(x.N)} if x is not None and not pd.isna(x.kpf) else ({7: int(x.N)} if x is not None else None))
# ---- Panel B ----
r[0] += 1
header(["2SLS", "OLS", "KP F", "N"], "Mediator (outcome of the a-path regression)")
panel("Panel B. Does connectivity move the candidate mediator? (coefficient on ln GACI_max)")
MED = [("ln_gdppc", "ln GDP per capita"), ("ln_seatkm", "ln departing seat-km"), ("ln_tour_arr", "ln international tourist arrivals"), ("tour_rcpt_exp", "Tourism receipts, % of exports"),
       ("intl_seat_share", "International share of seat-km"), ("trade_gdp", "Trade, % of GDP"), ("fdi_gdp", "FDI inflows, % of GDP"), ("urban", "Urban population, %"),
       ("emp_agr", "Employment in agriculture, %"), ("emp_ind", "Employment in industry, %"), ("emp_srv", "Employment in services, %"), ("manf_va", "Manufacturing value added, % of GDP"), ("srv_va", "Services value added, % of GDP"),
       ("wage_emp", "Wage and salaried workers, % of employment"), ("unemp", "Unemployment rate, %"), ("ter_enr2", "Tertiary enrolment, % gross"), ("lf_advanced", "Labour force with advanced education, %"), ("tax_gdp", "Tax revenue, % of GDP")]
for mv, lab in MED:
    iv = get(block="B", item="apath", outcome=mv); ol = get(block="B", item="apath_ols", outcome=mv)
    if iv is None: continue
    row2(lab, [iv, ol], {4: float(iv.kpf), 5: int(iv.N)})
# ---- Panel C ----
r[0] += 1
header(["c': ln GACI_max", "b: mediator", "c': ln GACI_max", "b: mediator", "c': ln GACI_max", "b: mediator", "KP F", "N"], "Mediator")
put(2, "ln Market Gini", F_H); put(4, "Bottom 50% share", F_H); put(6, "Top half minus bottom half", F_H); r[0] += 1
panel("Panel C. IV decomposition: direct effect of the hub (c') with the mediator held fixed, and the mediator's own coefficient (b)")
for mv, lab in MED:
    cells = []
    for o in ["ln_gini_mkt", "ln_spt_b50", "ln_apt_top_bot"]:
        cells += [get(block="C", item=mv, outcome=o, term="ln_gaci_max"), get(block="C", item=mv, outcome=o, term=mv)]
    if all(c is None for c in cells): continue
    x = get(block="C", item=mv, outcome="ln_gini_mkt", term="ln_gaci_max")
    put(1, lab, F_B, "left")
    for j, c in enumerate(cells, 2):
        if c is None: put(j, "--"); continue
        dd = 3 if j % 2 == 0 else 4
        put(j, f"{c.b:.{dd}f}{st(c.p)}")
    if x is not None: put(8, float(x.kpf), fmt="0.0"); put(9, int(x.N), fmt="#,##0")
    r[0] += 1
    for j, c in enumerate(cells, 2):
        if c is not None: put(j, f"({c.se:.{3 if j % 2 == 0 else 4}f})")
    r[0] += 1
# ---- Panel D ----
r[0] += 1
header(["2SLS", "OLS", "KP F", "N"], "Outcome")
panel("Panel D. Redistribution response")
for o, lab in [("wedge_pts", "Market minus disposable Gini, points"), ("ln_wedge_ratio", "ln(market Gini / disposable Gini)"), ("tax_gdp", "Tax revenue, % of GDP")]:
    iv = get(block="D", item="redistribution", outcome=o); ol = get(block="D", item="redistribution_ols", outcome=o)
    row2(lab, [iv, ol], {4: float(iv.kpf), 5: int(iv.N)})
# ---- Panel E ----
r[0] += 1
header(["ln Market Gini", "Bottom 50% share", "Top half minus bottom half", "Top 10% share", "KP F", "N"], "Sample (baseline median split)")
panel("Panel E. Heterogeneity by channel proxies (2SLS on ln GACI_max)")
for v, lab in [("intl_seat_share", "International share of seat-km"), ("tour_rcpt_exp", "Tourism receipts, % of exports"), ("srv_va", "Services value added, % of GDP"), ("wage_emp", "Wage and salaried workers, % of employment")]:
    for g, gl in [("1", "below median"), ("2", "above median")]:
        cells = [get(block="E", item=f"{v}_{g}", outcome=o) for o in ["ln_gini_mkt", "ln_spt_b50", "ln_apt_top_bot", "ln_spt_d10"]]
        x = cells[0]
        row2(f"{lab}: {gl}", cells, {6: float(x.kpf), 7: int(x.N)} if x is not None else None)
for j in range(1, 10): ws.cell(row=r[0] - 1, column=j).border = Border(bottom=TOP)
r[0] += 1
for n in ["Notes: 2SLS with the Feyrer instrument, ln population, country and year fixed effects; standard errors clustered by country in parentheses. *** p<0.01, ** p<0.05, * p<0.1. Mediators from WDI (API, 2026-09-28) and the companion CO2 panel (seat-km, international share).",
          "Panel C follows the IV decomposition of the companion paper: the mediator enters as a control while ln GACI_max is instrumented; c' is the direct effect and b the mediator's coefficient; the indirect effect is a x b with a from Panel B. Mediators with KP F below 5 (advanced-education labour force, tertiary enrolment, tax revenue) are shown for completeness.",
          "Reading: (A) the effect comes from the hub airport, not from the breadth of the network; (B) connectivity raises income, seat-km, tourist arrivals, industry employment and trade, and does not move urbanisation, FDI, services, formality or skill composition; (C) no measured mediator transmits the effect and the ones that respond (tourism, industrialisation, income) compress the distribution, so c' is at or above the total effect; (D) the market-disposable wedge narrows; (E) the effect is concentrated in non-tourism economies and, for shares, in internationally oriented networks.",
          "Source: _mech_results.csv (31_mechanisms.do)."]:
    put(1, n, F_N, "left"); r[0] += 1
ws.column_dimensions["A"].width = 58
for j in range(2, 10): ws.column_dimensions[get_column_letter(j)].width = 15
wb.save("Tables_for_paper_20260928.xlsx"); print("saved Table9", wb.sheetnames)
