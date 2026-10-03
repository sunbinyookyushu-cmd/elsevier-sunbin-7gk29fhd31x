# -*- coding: utf-8 -*-
"""18_make_gic_xlsx.py : stand-alone workbook for the incidence-curve regressions (WID income groups x GACI).
   Sources: _gic_dose_results.csv (12_gic_dose.do, Stata, country-clustered), _gic_stacked_results.csv (12b),
            _gic_spec_grid.csv (17, inference/spec grid), _acreg_results.csv (16), _gic_descriptive.csv (14).
   Output: GIC_results_20260928.xlsx"""
import os, pandas as pd, numpy as np
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

D = os.path.dirname(os.path.abspath(__file__)); os.chdir(D)
RD = dict(keep_default_na=False, na_values=["", "."])
gic = pd.read_csv("_gic_dose_results.csv", **RD)
stk = pd.read_csv("_gic_stacked_results.csv", **RD)
grid = pd.read_csv("_gic_spec_grid.csv", **RD)
acr = pd.read_csv("_acreg_results.csv", **RD)
des = pd.read_csv("_gic_descriptive.csv", **RD)

G = ["d1", "d2", "d3", "d4", "d5", "d6", "d7", "d8", "d9", "d10", "t1", "t01", "b50", "m40", "all"]
LAB = {"d1": "p0-10", "d2": "p10-20", "d3": "p20-30", "d4": "p30-40", "d5": "p40-50", "d6": "p50-60", "d7": "p60-70",
       "d8": "p70-80", "d9": "p80-90", "d10": "p90-100", "t1": "Top 1%", "t01": "Top 0.1%", "b50": "Bottom 50%",
       "m40": "Middle 40%", "all": "All (mean income)"}
def st(p):
    if pd.isna(p): return ""
    return "***" if p < .01 else "**" if p < .05 else "*" if p < .1 else ""

F_ARIAL = "Arial"
H_FONT = Font(name=F_ARIAL, bold=True, size=10); B_FONT = Font(name=F_ARIAL, size=10)
T_FONT = Font(name=F_ARIAL, bold=True, size=12); N_FONT = Font(name=F_ARIAL, italic=True, size=9, color="555555")
FILL = PatternFill("solid", fgColor="EDEDED"); THIN = Side(style="thin", color="999999")

wb = Workbook(); wb.remove(wb.active)

def sheet(name, title, note_lines, df, widths=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = T_FONT
    r = 2
    for ln in note_lines:
        ws.cell(row=r, column=1, value=ln).font = N_FONT; r += 1
    r += 1
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(row=r, column=j, value=c); x.font = H_FONT; x.fill = FILL
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        x.border = Border(bottom=THIN, top=THIN)
    hdr = r
    for _, row in df.iterrows():
        r += 1
        for j, c in enumerate(df.columns, 1):
            v = row[c]
            if isinstance(v, (float, np.floating)) and np.isnan(v): v = None
            if isinstance(v, np.integer): v = int(v)
            if isinstance(v, np.floating): v = float(v)
            x = ws.cell(row=r, column=j, value=v); x.font = B_FONT
            if isinstance(v, float): x.number_format = "0.000"
            if isinstance(v, int): x.number_format = "#,##0"
    for j in range(1, len(df.columns) + 1): ws.cell(row=r, column=j).border = Border(bottom=THIN)
    ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, max(12, min(28, len(str(c)) + 2)))
    return ws

# ---------------- README ----------------
ws = wb.create_sheet("README")
lines = [
    ("Incidence of hub connectivity across the income distribution", T_FONT),
    ("GACI x inequality project, stand-alone results for the group-specific regressions. Built 2026-09-28 by 18_make_gic_xlsx.py.", B_FONT),
    ("", B_FONT),
    ("Specification (headline of the paper, unchanged):", H_FONT),
    ("  y_ct = b ln GACI_max_ct + g ln pop_ct + country FE + year FE + e_ct;  2SLS with the Feyrer instrument Z_ct = a_t x ln MA^air_c,1996.", B_FONT),
    ("  Outcomes: ln average pretax national income of each WID group (aptinc, equal-split adults, 992j) and its log income share.", B_FONT),
    ("  Share regressions use the identity ln share_g = ln income_g - ln mean income + const, so every group is estimated on the same sample", B_FONT),
    ("  (WID rounds the bottom-decile share to zero in 271 country-years, which would otherwise drop them). Standard errors clustered by country unless stated.", B_FONT),
    ("  Panel 1996-2023, ~165 countries, N ~ 3,714 per regression. First-stage Kleibergen-Paap F = 10.5 (Feyrer), 2.3 with continent x year effects.", B_FONT),
    ("", B_FONT),
    ("Sheets:", H_FONT),
    ("  T1_group_elasticities  Elasticity of each group's income (and share) to ln GACI_max: OLS, 2SLS, 2SLS with continent x year FE.", B_FONT),
    ("  T2_contrasts_tests     Differences between groups (difference outcomes, exact) and the stacked-system tests (joint equality, linear trend).", B_FONT),
    ("  T3_inference           Headline and group 2SLS under country cluster, two-way cluster (country, year), Conley 1000/2000 km (uniform lags), Driscoll-Kraay (lag 5); acreg cross-check.", B_FONT),
    ("  T4_spec_grid           Scorecard over treatment (max / capacity-weighted mean) x sea-MA design control x FE x inference. No behavioural controls anywhere.", B_FONT),
    ("  T5_descriptive         Piketty-Saez-Zucman-style growth by group, countries split into terciles of GACI growth 1996-2023 (one obs per country).", B_FONT),
    ("  raw_*                  Untouched result CSVs.", B_FONT),
    ("", B_FONT),
    ("Reading:", H_FONT),
    ("  Upper-half incomes rise with hub connectivity at roughly the mean elasticity (1.2-1.4); lower-half incomes rise by less, so their shares fall.", B_FONT),
    ("  Group elasticities are not statistically distinguishable under the Feyrer IV (joint equality p = .74); the top-half vs bottom-half contrast is significant at 10%.", B_FONT),
    ("  Everything is insignificant with continent x year effects (weak first stage). Driscoll-Kraay SEs are reported for completeness but not adopted (lag truncation understates persistence).", B_FONT),
    ("  Stars: *** p<.01, ** p<.05, * p<.1.", B_FONT),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions["A"].width = 160

# ---------------- T1 ----------------
rows = []
for g in G:
    r = {"Group": LAB[g]}
    a = gic[(gic.block == "gic") & (gic.outcome == "ln_apt_" + g)]
    s = gic[(gic.block == "gic_share") & (gic.outcome == "ln_spt_" + g)]
    for lab, sub, spec in [("Income OLS", a, "OLS"), ("Income 2SLS", a, "IV"), ("Income 2SLS cont x year", a, "IV_contyear"),
                           ("Share OLS", s, "OLS"), ("Share 2SLS", s, "IV")]:
        x = sub[sub.spec == spec]
        if len(x):
            x = x.iloc[0]; r[lab + " b"] = x.b; r[lab + " se"] = x.se; r[lab + " sig"] = st(x.p)
            if "2SLS" in lab: r[lab + " KP F"] = x.kpf
            r[lab + " N"] = int(x.N)
        else:
            for k in [" b", " se", " sig"]: r[lab + k] = None
    rows.append(r)
T1 = pd.DataFrame(rows)
sheet("T1_group_elasticities", "T1. Elasticity of group income and income share to ln GACI_max",
      ["Each cell block is one regression: outcome = ln average pretax income of the group (columns 'Income') or its log share via the identity (columns 'Share').",
       "Country + year FE, ln population, SE clustered by country. 2SLS uses the Feyrer instrument. 'cont x year' replaces year FE by continent x year FE.",
       "Source: _gic_dose_results.csv (12_gic_dose.do)."], T1, widths={"Group": 18})

# ---------------- T2 ----------------
gap = gic[gic.block == "gic_gap"].copy()
gap["Contrast"] = gap.outcome.map({"ln_apt_d10_d1": "p90-100 minus p0-10", "ln_apt_d10_b50": "p90-100 minus Bottom 50%", "ln_apt_t1_d10": "Top 1% minus p90-100"})
gap["Method"] = "difference outcome, separate regression"
gap = gap.rename(columns={"spec": "Estimator", "b": "Estimate", "se": "SE", "p": "p-value", "kpf": "KP F"})
gap["sig"] = gap["p-value"].map(st)
T2a = gap[["Contrast", "Method", "Estimator", "Estimate", "SE", "sig", "p-value", "KP F", "N"]]
s2 = stk[stk.block != "stacked"].copy()
s2["Contrast"] = s2.param.map({"eq_d1_d10": "Joint test: d1 = ... = d10 (stat, p)", "eq_all12": "Joint test: all 12 groups equal (stat, p)",
                               "d10_minus_d1": "p90-100 minus p0-10", "top50_minus_bot50": "Top half (d6-d10) minus bottom half (d1-d5)",
                               "t1_minus_d10": "Top 1% minus p90-100", "level_at_median": "Linear-trend spec: elasticity at median decile",
                               "slope_per_decile": "Linear-trend spec: change in elasticity per decile"})
s2["Method"] = "stacked system (country x group x year), country#group and year#group FE"
s2 = s2.rename(columns={"spec": "Estimator", "b": "Estimate", "se": "SE", "p": "p-value", "kpf": "KP F"})
s2["sig"] = s2["p-value"].map(st)
T2 = pd.concat([T2a, s2[["Contrast", "Method", "Estimator", "Estimate", "SE", "sig", "p-value", "KP F", "N"]]], ignore_index=True)
sheet("T2_contrasts_tests", "T2. Are the group elasticities different from each other?",
      ["Difference outcomes give the exact difference of two group coefficients with a correct SE (same sample and regressors).",
       "Stacked rows: all groups in one 2SLS system with group x instrument; 'Joint test' rows report the Wald statistic in 'Estimate' and its p-value.",
       "The stacked 12-endogenous system's cluster-robust KP statistic is degenerate (0); use the per-equation KP F = 10.5 (Feyrer). Trend spec KP F = 5.2.",
       "Source: _gic_dose_results.csv (gic_gap block), _gic_stacked_results.csv (12b_gic_stacked.do)."], T2, widths={"Contrast": 46, "Method": 58})

# ---------------- T3 ----------------
base = grid[(grid.treat == "ln_gaci_max") & (grid.ctrl == "pop") & (grid.fe == "c+y")].copy()
SEL = {"cluster_c": "Cluster: country", "twoway_cy": "Two-way cluster: country, year", "conley_1000": "Conley 1000 km (uniform lags)",
       "conley_2000": "Conley 2000 km (uniform lags)", "dk5": "Driscoll-Kraay, lag 5 (not adopted)"}
OUTL = {"ln_gini_mkt": "Headline: ln market Gini", "ln_gini_disp": "Headline: ln disposable Gini", "gap_d10_b50": "Contrast: p90-100 minus Bottom 50%",
        "gap_top_bot": "Contrast: top half minus bottom half"}
for g in G: OUTL["ln_apt_" + g] = "Income: " + LAB[g]
for g in G[:-1]: OUTL["rel_" + g] = "Share: " + LAB[g]
rows = []
for o in OUTL:
    x = base[base.outcome == o]
    if not len(x): continue
    r = {"Outcome": OUTL[o], "Estimate (2SLS)": x.b.iloc[0], "KP F": x.F.iloc[0], "N": int(x.N.iloc[0])}
    for k, l in SEL.items():
        y = x[x.se == k].iloc[0]; r[l + " se"] = y.se_v; r[l + " sig"] = st(y.p)
    rows.append(r)
T3 = pd.DataFrame(rows)
sheet("T3_inference", "T3. Same 2SLS estimates under alternative standard errors (headline spec: GACI_max, ln pop, country + year FE)",
      ["Own implementation (15_conley_se.py / 17_gic_spec_grid.py): two-way FE removed by alternating projections; sandwich with product kernel. Cluster SE reproduces Stata to the dof convention (0.198 vs 0.194).",
       "Conley: Bartlett in great-circle distance between country centroids (Natural Earth; airport-mean fallback for 27 islands/city states); within-country pairs weighted 1 at every lag (as clustering).",
       "acreg cross-check (Colella et al. 2019; Bartlett in distance AND in time): market Gini 2000 km lag 30 -> se 0.161; lag 10 -> 0.125; 1000 km lag 30 -> 0.160. Disposable: 0.201 / 0.155 / 0.200.",
       "Lag-truncated kernels (acreg, Driscoll-Kraay) give smaller SEs than clustering because Gini residuals are highly persistent; they are reported, not adopted.",
       "Source: _gic_spec_grid.csv, _acreg_results.csv."], T3, widths={"Outcome": 40})

# ---------------- T4 ----------------
rows = []
for (xv, cl_, fe, se), x in grid.groupby(["treat", "ctrl", "fe", "se"], sort=False):
    x = x.set_index("outcome"); h = x.loc["ln_gini_mkt"]
    top = int(sum(x.loc["ln_apt_" + g].p < .1 for g in ["d7", "d8", "d9", "d10", "t1"]))
    bot = int(sum(x.loc["rel_" + g].p < .1 for g in ["d1", "d2", "d3", "d4", "d5", "b50"]))
    g1 = x.loc["gap_d10_b50"]; g2 = x.loc["gap_top_bot"]
    rows.append({"Treatment": {"ln_gaci_max": "GACI_max (hub)", "ln_gaci_cwm": "GACI_cwm (capacity-weighted mean)"}[xv],
                 "Design control": {"pop": "ln pop", "pop+seaMA": "ln pop + ln sea MA (CO2-paper design)"}[cl_],
                 "Fixed effects": {"c+y": "country + year", "c+cont#y": "country + continent x year"}[fe], "SE": SEL[se],
                 "Headline b": h.b, "Headline se": h.se_v, "Headline sig": st(h.p), "KP F": h.F,
                 "Top groups sig (of 5: p60-100, top1%)": top, "Bottom shares sig (of 6: d1-d5, b50)": bot,
                 "p90-100 - b50": g1.b, "se": g1.se_v, "sig": st(g1.p), "Top half - bottom half": g2.b, "se ": g2.se_v, "sig ": st(g2.p), "N": int(h.N)})
T4 = pd.DataFrame(rows)
sheet("T4_spec_grid", "T4. Specification scorecard for the incidence curve (2SLS Feyrer; no behavioural controls)",
      ["Rows vary only the treatment measure, the CO2-paper design control (ln sea market access), the fixed effects, and the inference. 'Top groups sig' counts income elasticities with p<.1 among p60-70 ... p90-100 and top 1%.",
       "The headline specification (first row) is the cleanest legitimate one; two-way clustering and Conley 1000 km leave it unchanged. Continent x year effects remove everything (weak first stage).",
       "Source: _gic_spec_grid.csv (17_gic_spec_grid.py)."], T4, widths={"Treatment": 30, "Design control": 36, "Fixed effects": 26, "SE": 32})

# ---------------- T5 ----------------
T5 = des.rename(columns={"label": "Group", "low_mean": "Low GACI-growth tercile: mean growth (%/yr)", "mid_mean": "Middle tercile (%/yr)",
                         "high_mean": "High GACI-growth tercile (%/yr)", "diff_high_low": "High minus low (pp/yr)", "se_diff": "SE of difference",
                         "n_low": "N low", "n_mid": "N mid", "n_high": "N high"}).drop(columns=["group"])
T5["sig"] = pd.Series(2 * (1 - stats.norm.cdf(np.abs(T5["High minus low (pp/yr)"] / T5["SE of difference"]))), index=T5.index).map(st)
sheet("T5_descriptive", "T5. Annualised real growth of group income, countries split by GACI growth (descriptive, one observation per country)",
      ["First to last year with complete WID and GACI data (span >= 15 years; 138 countries; median span 26 years). Terciles by annualised change in ln GACI_max (low mean -0.17 %/yr, mid 0.58, high 1.44).",
       "The level gap (~1.5 pp at every group) is catch-up growth of faster-globalising countries; only the slope across groups is distributional. Figure: fig_gic_descriptive.",
       "Source: _gic_descriptive.csv (14_fig_gic_descriptive.py)."], T5, widths={"Group": 14})

# ---------------- raw ----------------
for nm, df in [("raw_gic_dose", gic), ("raw_stacked", stk), ("raw_spec_grid", grid), ("raw_acreg", acr), ("raw_descriptive", des)]:
    ws = wb.create_sheet(nm)
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(row=1, column=j, value=c); x.font = H_FONT; x.fill = FILL
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, (float, np.floating)) and np.isnan(v): v = None
            if isinstance(v, np.integer): v = int(v)
            if isinstance(v, np.floating): v = float(v)
            ws.cell(row=i, column=j, value=v).font = B_FONT
    ws.freeze_panes = "A2"
    for j in range(1, len(df.columns) + 1): ws.column_dimensions[get_column_letter(j)].width = 14

out = "GIC_results_20260928.xlsx"; wb.save(out); print("saved", out, wb.sheetnames)
